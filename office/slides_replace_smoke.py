"""Verify versioned, run-preserving PowerPoint text replacement."""
import copy
import tempfile
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches
from lxml import etree

import session
import slides

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.pptx'; draft = root / 'draft'
    image = root / 'image.png'; Image.new('RGB', (10, 10), 'red').save(image)
    presentation = Presentation(); first = presentation.slides.add_slide(presentation.slide_layouts[6]); first.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = 'Wrong slide'
    target = presentation.slides.add_slide(presentation.slide_layouts[6]); target.shapes.add_picture(str(image), Inches(1), Inches(1))
    group = target.shapes.add_group_shape(); group_id = group.shape_id
    inner = group.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)); inner.text = 'Grouped'; inner_id = inner.shape_id
    grouped_table = target.shapes.add_table(1, 1, Inches(1), Inches(2), Inches(2), Inches(1)); grouped_table.table.cell(0, 0).text = 'Group table'; grouped_table_id = grouped_table.shape_id; group._element.append(grouped_table._element)
    box = target.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1)); paragraph = box.text_frame.paragraphs[0]
    first_run = paragraph.add_run(); first_run.text = 'Hello '; first_run.font.bold = True
    second_run = paragraph.add_run(); second_run.text = 'world!'; second_run.font.italic = True
    multi = target.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1)); multi.text = 'first\nsecond'
    broken = target.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)); broken.text = 'before'; broken.text_frame.paragraphs[0]._p.append(OxmlElement('a:br'))
    linked = target.shapes.add_textbox(Inches(1), Inches(6), Inches(4), Inches(1)); link_para=linked.text_frame.paragraphs[0]
    link_para.add_run().text='Visit '; link_run=link_para.add_run(); link_run.text='Link'; link_run.hyperlink.address='https://example.com';link_run.font.bold=True
    same_link=link_para.add_run();same_link.text='ed';same_link.hyperlink.address='https://example.com';same_link.font.italic=True
    other_link=link_para.add_run();other_link.text='other';other_link.hyperlink.address='https://example.org'
    link_para.add_run().text=' now'
    chart_data = ChartData(); chart_data.categories = ['A', 'B']; chart_data.add_series('Sales', [1, 2])
    target.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1), Inches(3), Inches(2), chart_data)
    table_shape = target.shapes.add_table(2, 3, Inches(5), Inches(3), Inches(4), Inches(2)); table_id = table_shape.shape_id
    table = table_shape.table; table.cell(0, 0).merge(table.cell(0, 1)); cell_para = table.cell(0, 0).text_frame.paragraphs[0]
    run = cell_para.add_run(); run.text = 'Sales '; run.font.bold = True
    run = cell_para.add_run(); run.text = 'report'; run.font.italic = True
    table.cell(0, 2).text = 'first\nsecond paragraph'
    base = root / 'base.pptx'; presentation.save(base)
    with zipfile.ZipFile(base) as original_zip, zipfile.ZipFile(source, 'w') as reordered:
        for entry in original_zip.infolist():
            data = original_zip.read(entry)
            if entry.filename == 'ppt/presentation.xml':
                tree = etree.fromstring(data); ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
                ids = tree.find('p:sldIdLst', ns); ids.insert(0, ids[1]); data = etree.tostring(tree, encoding='UTF-8', xml_declaration=True, standalone=True)
            reordered.writestr(entry, data)
    original = source.read_bytes()
    part = slides.slide_part(source, 0)
    assert part.endswith('slide2.xml')
    session.start(draft, source=source)
    target_info = session.edit_target(draft, 0, '/slide[1]/shape[1]')
    assert target_info['replace_text_path_supported'] and target_info['value'] == 'Hello world!'
    for path, find in [('/slide[1]/shape[2]', 'first'), ('/slide[1]/shape[3]/paragraph[1]', 'before'), ('/slide[1]/shape[4]/paragraph[1]', 'Visit Link'), ('/slide[1]/shape[4]/paragraph[1]', 'edother'), ('/slide[1]/shape[1]', 'missing')]:
        try:
            session.replace_text(draft, 0, path, find, 'bad')
            raise AssertionError('Unsafe replacement succeeded')
        except ValueError:
            pass
    assert session.info(draft)['latest_revision'] == 0
    link_draft=root/'link-draft';session.start(link_draft,source=source)
    link_result=session.replace_text(link_draft,0,'/slide[1]/shape[4]/paragraph[1]','nked',' reference')
    link_shape=next(shape for shape in Presentation(link_result['file']).slides[0].shapes if shape.shape_id==linked.shape_id)
    link_runs=link_shape.text_frame.paragraphs[0].runs
    assert link_shape.text=='Visit Li referenceother now' and link_runs[1].font.bold and link_runs[2].font.italic
    assert link_runs[1].hyperlink.address=='https://example.com' and link_runs[2].hyperlink.address=='https://example.com' and link_runs[3].hyperlink.address=='https://example.org'
    with zipfile.ZipFile(source) as before,zipfile.ZipFile(link_result['file']) as after:
        assert all(before.read(name)==after.read(name) for name in before.namelist() if name!=part)
    result = session.replace_text(draft, 0, '/slide[1]/shape[1]/paragraph[1]', 'llo wor', 'i earth')
    assert result['revision'] == 1 and source.read_bytes() == original
    updated = Presentation(result['file']); shape = next(shape for shape in updated.slides[0].shapes if shape.shape_id == box.shape_id)
    runs = shape.text_frame.paragraphs[0].runs
    assert shape.text == 'Hei earthld!' and runs[0].font.bold and runs[1].font.italic
    assert updated.slides[1].shapes[0].text == 'Wrong slide'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(result['file']) as after:
        assert set(before.namelist()) == set(after.namelist())
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name != part)
        old = etree.fromstring(before.read(part)); new = etree.fromstring(after.read(part))
        ns = {'p':'http://schemas.openxmlformats.org/presentationml/2006/main', 'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        for tree in (old, new):
            for node in tree.xpath(f'.//p:sp[p:nvSpPr/p:cNvPr[@id="{box.shape_id}"]]/p:txBody/a:p/a:r/a:t', namespaces=ns): node.text = 'MATCH'
        assert etree.tostring(old, method='c14n') == etree.tostring(new, method='c14n')
    grouped = session.replace_text(draft, 1, f'/slide[1]/group[@id={group_id}]/shape[@id={inner_id}]/paragraph[1]', 'Grouped', 'Inside')
    updated_group = next(shape for shape in Presentation(grouped['file']).slides[0].shapes if shape.shape_id == group_id)
    assert updated_group.shapes[0].text == 'Inside'
    try:
        session.replace_text(draft, 0, '/slide[1]/shape[1]', 'Hello', 'stale')
        raise AssertionError('Stale replacement succeeded')
    except ValueError as error:
        assert 'conflict' in str(error)
    table_path = f'/slide[1]/table[@id={table_id}]/tr[1]/tc[1]'
    target_info = session.edit_target(draft, 2, table_path)
    assert target_info['type'] == 'tc' and target_info['replace_text_path_supported'] and target_info['value'] == 'Sales report'
    try:
        session.edit(draft, 2, table_path, 'text', 'bad')
        raise AssertionError('Whole-cell replacement was accepted')
    except ValueError as error:
        assert 'text replacement' in str(error)
    try:
        session.replace_text(draft, 2, '/slide[1]/table[1]/tr[1]/tc[2]', 'x', 'bad')
        raise AssertionError('Merged placeholder accepted')
    except ValueError as error:
        assert 'placeholder' in str(error)
    table_result = session.replace_text(draft, 2, '/slide[1]/table[1]/tr[1]/tc[1]', 'les rep', 'LES REP')
    result_table = next(shape for shape in Presentation(table_result['file']).slides[0].shapes if shape.shape_id == table_id).table
    cell_para = result_table.cell(0, 0).text_frame.paragraphs[0]
    assert cell_para.text == 'SaLES REPort' and cell_para.runs[0].font.bold and cell_para.runs[1].font.italic
    assert result_table.cell(0, 0).is_merge_origin and result_table.cell(0, 1).is_spanned
    with zipfile.ZipFile(grouped['file']) as before, zipfile.ZipFile(table_result['file']) as after:
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name != part)
        old = etree.fromstring(before.read(part)); new = etree.fromstring(after.read(part))
        ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        selector = f'.//p:graphicFrame[p:nvGraphicFramePr/p:cNvPr[@id="{table_id}"]]/a:graphic/a:graphicData/a:tbl/a:tr[1]/a:tc[1]/a:txBody/a:p/a:r/a:t'
        for tree in (old, new):
            for node in tree.xpath(selector, namespaces=ns): node.text = 'MATCH'
        assert etree.tostring(old, method='c14n') == etree.tostring(new, method='c14n')
    try:
        session.replace_text(draft, 3, '/slide[1]/table[1]/tr[1]/tc[3]', 'second', 'bad')
        raise AssertionError('Ambiguous paragraph accepted')
    except ValueError as error:
        assert 'multiple paragraphs' in str(error)
    paragraph_result = session.replace_text(draft, 3, '/slide[1]/table[1]/tr[1]/tc[3]/paragraph[2]', 'second', 'updated')
    table_after = next(shape for shape in Presentation(paragraph_result['file']).slides[0].shapes if shape.shape_id == table_id).table
    assert table_after.cell(0, 2).text == 'first\nupdated paragraph'
    group_cell = session.replace_text(draft, 4, f'/slide[1]/group[@id={group_id}]/table[@id={grouped_table_id}]/tr[1]/tc[1]', 'Group', 'Nested')
    result_group = next(shape for shape in Presentation(group_cell['file']).slides[0].shapes if shape.shape_id == group_id)
    assert next(shape for shape in result_group.shapes if shape.shape_id == grouped_table_id).table.cell(0, 0).text == 'Nested table'
    assert session.info(draft)['latest_revision'] == 5
    duplicate = root / 'duplicate.pptx'
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(duplicate, 'w') as dst:
        for entry in src.infolist():
            data = src.read(entry)
            if entry.filename == part:
                tree = etree.fromstring(data); ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
                shape_tree = tree.find('p:cSld/p:spTree', ns)
                matching = shape_tree.xpath(f'p:sp[p:nvSpPr/p:cNvPr[@id="{box.shape_id}"]]', namespaces=ns)
                shape_tree.append(copy.deepcopy(matching[0])); data = etree.tostring(tree, encoding='UTF-8', xml_declaration=True, standalone=True)
            dst.writestr(entry, data)
    duplicate_draft = root / 'duplicate-draft'; session.start(duplicate_draft, source=duplicate)
    try:
        session.replace_text(duplicate_draft, 0, f'/slide[1]/shape[@id={box.shape_id}]/paragraph[1]', 'Hello', 'bad')
        raise AssertionError('Duplicate shape ID accepted')
    except ValueError as error:
        assert 'exactly one' in str(error)
    assert session.info(duplicate_draft)['latest_revision'] == 0
    shared = root / 'shared.pptx'
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(shared, 'w') as dst:
        for entry in src.infolist():
            data = src.read(entry)
            if entry.filename == 'ppt/presentation.xml':
                tree = etree.fromstring(data); ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
                items = tree.find('p:sldIdLst', ns); rid = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
                items[1].set(rid, items[0].get(rid)); data = etree.tostring(tree, encoding='UTF-8', xml_declaration=True, standalone=True)
            dst.writestr(entry, data)
    shared_draft = root / 'shared-draft'; session.start(shared_draft, source=shared)
    try:
        session.replace_text(shared_draft, 0, '/slide[1]/shape[1]/paragraph[1]', 'Hello', 'bad')
        raise AssertionError('Shared slide part accepted')
    except ValueError as error:
        assert 'multiple slides' in str(error)
    assert session.info(shared_draft)['latest_revision'] == 0
    alternate = root / 'alternate.pptx'
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(alternate, 'w') as dst:
        for entry in src.infolist():
            data = src.read(entry)
            if entry.filename == part:
                tree = etree.fromstring(data); ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
                shape_tree = tree.find('p:cSld/p:spTree', ns)
                shape_tree.append(etree.Element('{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent'))
                data = etree.tostring(tree, encoding='UTF-8', xml_declaration=True, standalone=True)
            dst.writestr(entry, data)
    alternate_draft = root / 'alternate-draft'; session.start(alternate_draft, source=alternate)
    try:
        session.replace_text(alternate_draft, 0, '/slide[1]/shape[1]/paragraph[1]', 'Hello', 'bad')
        raise AssertionError('Alternate content accepted')
    except ValueError as error:
        assert 'alternate content' in str(error)
    assert session.info(alternate_draft)['latest_revision'] == 0
print('Slides replacement smoke passed: reordered slides, groups, tables, run formatting and protected content')

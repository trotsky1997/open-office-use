"""Check the viewer's speaker notes routes and revision protection."""
import json
import subprocess
import sys
import tempfile
import urllib.error
import urllib.request
from pathlib import Path
from urllib.parse import urlsplit
from pptx import Presentation
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'deck.pptx'; draft = root / 'draft'
    deck = Presentation(); deck.slides.add_slide(deck.slide_layouts[1]); deck.save(source)
    session.start(draft, source=source)
    cli = Path(__file__).with_name('open_office_use.py')
    server = subprocess.Popen([sys.executable, str(cli), 'watch', str(draft)], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    try:
        url = json.loads(server.stdout.readline())['url']; origin = 'http://' + urlsplit(url).netloc
        def post(route, data, origin=origin):
            request = urllib.request.Request(url + route, data=json.dumps(data).encode(), headers={'Origin': origin, 'Content-Type': 'application/json'})
            with urllib.request.urlopen(request, timeout=30) as response: return json.load(response)
        assert not post('notes-read', {'revision': 0, 'slide': 1})['has_notes_slide']
        assert post('notes-replace', {'base_revision': 0, 'slide': 1, 'text': 'Notes'})['revision'] == 1
        assert post('notes-read', {'revision': 1, 'slide': 1})['text'] == 'Notes'
        assert post('notes-replace-text', {'base_revision': 1, 'slide': 1, 'paragraph': 1, 'find': 'Notes', 'replacement': 'Edited'})['revision'] == 2
        assert post('notes-read', {'revision': 2, 'slide': 1})['text'] == 'Edited'
        for route, data, request_origin, code in [
            ('notes-replace', {'base_revision': 0, 'slide': 1, 'text': 'Stale'}, origin, 400),
            ('notes-replace', {'base_revision': 2, 'slide': 1, 'text': 'Forbidden'}, 'http://evil.test', 403),
            ('notes-read', {'revision': 1, 'slide': 0}, origin, 400),
        ]:
            try: post(route, data, request_origin)
            except urllib.error.HTTPError as exc: assert exc.code == code, exc
            else: raise AssertionError('Invalid request accepted')
        assert session.info(draft)['latest_revision'] == 2
    finally:
        server.terminate(); server.wait(timeout=10)
print('Speaker notes viewer routes passed: read, save, stale revision and origin checks')

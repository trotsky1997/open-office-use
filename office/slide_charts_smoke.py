"""Inspect PowerPoint chart caches, revisions and bounded failure modes."""
import hashlib
import io
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
import open_office_use
import session
import slide_charts
import review_bundle

with tempfile.TemporaryDirectory() as tmp:
    root=Path(tmp); source=root/'chart.pptx'
    spec={'slides':[{'layout':'blank','elements':[{'type':'chart','chart_type':'column','x':1,'y':1,'width':6,'height':3,'categories':['Q1','Q2'],'series':[{'name':'Revenue','values':[100,120]}]}]}]}
    open_office_use.author(spec,source); result=slide_charts.inspect(source); chart=result['charts'][0]
    assert result['complete'] and result['snapshot_sha256']==hashlib.sha256(source.read_bytes()).hexdigest()
    assert chart['categories']==['Q1','Q2'] and chart['series'][0]['name']=='Revenue' and chart['series'][0]['values']==[100,120]
    assert chart['series'][0]['references']['val']['formula']=='Sheet1!$B$2:$B$3' and chart['workbook']['status']=='embedded'
    assert chart['reconciliation']['status']=='match' and chart['reconciliation']['complete']
    assert chart['reconciliation']['formula_freshness']=='unknown'
    mismatch=root/'mismatch.pptx'
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(mismatch,'w') as dst:
        for entry in src.infolist():
            data=src.read(entry.filename)
            if entry.filename.startswith('ppt/embeddings/') and entry.filename.endswith('.xlsx'):
                workbook=load_workbook(io.BytesIO(data)); workbook.active['B2']=999
                buffer=io.BytesIO(); workbook.save(buffer); data=buffer.getvalue()
            dst.writestr(entry,data)
    mismatch_report=slide_charts.inspect(mismatch)['charts'][0]['reconciliation']
    assert mismatch_report['status']=='mismatch' and mismatch_report['mismatch_count']>=1
    changed_reference=next(item['index'] for item in mismatch_report['references'] if item['status']=='mismatch')
    sync_draft=root/'sync-draft';session.start(sync_draft,source=mismatch)
    synced=session.sync_slide_chart_cache(sync_draft,0,chart['chart_part'],changed_reference)
    assert synced['updated_points']==1 and synced['revision']==1
    assert session.inspect_slide_charts(sync_draft,1)['charts'][0]['reconciliation']['status']=='match'
    with zipfile.ZipFile(mismatch) as before, zipfile.ZipFile(synced['output']) as after:
        assert {name for name in before.namelist() if before.read(name)!=after.read(name)} == {chart['chart_part'].lstrip('/')}
    cli_draft=root/'sync-cli'; session.start(cli_draft,source=mismatch)
    cli=Path(__file__).with_name('open_office_use.py'); call=subprocess.run([sys.executable,str(cli),'session','sync-slide-chart-cache',str(cli_draft),'--base','0','--part',chart['chart_part'],'--reference',str(changed_reference)],capture_output=True,text=True,check=True)
    assert json.loads(call.stdout)['updated_points']==1
    try: session.sync_slide_chart_cache(sync_draft,0,chart['chart_part'],changed_reference)
    except ValueError as exc: assert 'Revision conflict' in str(exc)
    else: raise AssertionError('Stale revision accepted')
    def failed_render(_snapshot, _out): raise RuntimeError('render unavailable')
    review=review_bundle.create(mismatch,root/'chart-review',failed_render)
    assert 'slide_chart_cache_mismatch' in review['flags']
    draft=root/'draft';session.start(draft,source=source);assert session.inspect_slide_charts(draft,0)['snapshot_sha256']==result['snapshot_sha256']
    cli=Path(__file__).with_name('open_office_use.py');call=subprocess.run([sys.executable,str(cli),'session','inspect-slide-charts',str(draft),'--revision','0'],capture_output=True,text=True,check=True)
    assert json.loads(call.stdout)['charts'][0]['series'][0]['values']==[100,120]
    external=root/'external.pptx'
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(external,'w') as dst:
        for entry in src.infolist():
            data=src.read(entry.filename)
            if entry.filename=='ppt/charts/_rels/chart1.xml.rels':
                tree=etree.fromstring(data); rel=tree[0]; rel.set('TargetMode','External'); rel.set('Target','https://example.invalid/workbook.xlsx'); data=etree.tostring(tree)
            dst.writestr(entry,data)
    external_report=slide_charts.inspect(external)
    assert not external_report['complete'] and external_report['charts'][0]['workbook']['status']=='external'
    missing=root/'missing.pptx'
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(missing,'w') as dst:
        for entry in src.infolist():
            data=src.read(entry.filename)
            if entry.filename=='ppt/charts/chart1.xml':
                tree=etree.fromstring(data); ref=tree.find('.//'+slide_charts.C+'val/'+slide_charts.C+'numRef'); ref.remove(ref.find(slide_charts.C+'numCache'));data=etree.tostring(tree)
            dst.writestr(entry,data)
    absent=slide_charts.inspect(missing);assert not absent['complete'] and any(item['reason']=='unsupported_or_missing_cache' for item in absent['issues'])
    assert not absent['charts'][0]['reconciliation']['complete']
    missing_draft=root/'missing-draft'; session.start(missing_draft,source=missing)
    try: session.sync_slide_chart_cache(missing_draft,0,chart['chart_part'],changed_reference)
    except ValueError: pass
    else: raise AssertionError('Missing cache accepted')
    assert session.info(missing_draft)['latest_revision']==0
    huge=root/'huge.pptx'
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(huge,'w') as dst:
        for entry in src.infolist():
            data=src.read(entry.filename)
            if entry.filename=='ppt/charts/chart1.xml':
                tree=etree.fromstring(data); tree.find('.//'+slide_charts.C+'val/'+slide_charts.C+'numRef/'+slide_charts.C+'numCache/'+slide_charts.C+'ptCount').set('val','99999999'); data=etree.tostring(tree)
            dst.writestr(entry,data)
    capped=slide_charts.inspect(huge);assert not capped['complete'] and capped['charts'][0]['series'][0]['values']==[] and capped['charts'][0]['series'][0]['values_truncated']
    many=root/'many.pptx'; deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);data=CategoryChartData();data.categories=[str(i) for i in range(201)];data.add_series('Values',list(range(201)));slide.shapes.add_chart(XL_CHART_TYPE.LINE,Inches(1),Inches(1),Inches(6),Inches(4),data);deck.save(many)
    bounded=slide_charts.inspect(many); assert not bounded['complete'] and bounded['charts'][0]['series'][0]['values_total']==201 and not bounded['charts'][0]['series'][0]['values']
print('Slide charts smoke passed: cached data, revision, missing and bounded caches')

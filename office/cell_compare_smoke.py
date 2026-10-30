"""Check formula and cached-value comparison on real workbooks."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

cli = Path(__file__).with_name('open_office_use.py')
def run(*args, ok=True):
    result = subprocess.run([sys.executable, str(cli), *map(str, args)], capture_output=True, text=True)
    assert (result.returncode == 0) == ok, (result.stdout, result.stderr)
    return json.loads(result.stdout) if ok else result.stderr

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); before = root / 'before.xlsx'; after = root / 'after.xlsx'
    workbook = Workbook(); ws = workbook.active; ws.title = 'Forecast'; ws['A1'] = 4; ws['A2'] = '=A1*3'; workbook.save(before)
    ws['A1'] = 5; ws['A2'] = '=A1*4'; ws['B2'] = 'New'; workbook.save(after)
    compared = run('compare-cells', before, after, '--sheet', 'Forecast', '--range', 'A1:B3')
    changes = {item['cell']: item for item in compared['changes']}
    assert compared['changed_count'] == 3 and set(changes) == {'A1', 'A2', 'B2'}, compared
    assert changes['A1']['before']['value'] == 4 and changes['A1']['after']['value'] == 5
    assert changes['A2']['before']['formula'] == '=A1*3' and changes['A2']['after']['formula'] == '=A1*4'
    assert 'formula' in changes['A2']['changed_fields'] and compared['cache_freshness'] == 'unknown'
    assert not run('compare-cells', before, before, '--sheet', 'Forecast', '--range', 'A1:B3')['changes']
    calculated = root / 'calculated.xlsx'; run('recalculate', before, '--out', calculated)
    cached = run('compare-cells', before, calculated, '--sheet', 'Forecast', '--range', 'A2')
    assert cached['changes'][0]['before']['value'] is None and cached['changes'][0]['after']['value'] == 12, cached
    assert 'formula' not in cached['changes'][0]['changed_fields'], cached
    assert 'not found' in run('compare-cells', before, after, '--sheet', 'Absent', '--range', 'A1', ok=False)
    for row in range(1, 503): ws.cell(row, 3, row)
    workbook.save(after)
    capped = run('compare-cells', before, after, '--sheet', 'Forecast', '--range', 'C1:C502')
    assert capped['changed_count'] == 502 and len(capped['changes']) == 500 and capped['truncated'], capped
print('Cell comparison smoke passed: formulas, cache, values, missing sheet and truncation')

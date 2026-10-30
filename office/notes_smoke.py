"""Exercise speaker notes creation, editing and version safety."""
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch
from pptx import Presentation
import notes
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'deck.pptx'; draft = root / 'draft'
    presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text = 'Title'; presentation.save(source)
    assert not notes.read(source, 1)['has_notes_slide']
    session.start(draft, source=source)
    original = session.read_notes(draft, 0, 1)
    try: session.replace_notes(draft, 0, 1, 'Opening', expected_sha256='0'*64)
    except ValueError as exc: assert 'snapshot changed' in str(exc)
    else: raise AssertionError('Changed notes source accepted')
    assert session.info(draft)['latest_revision'] == 0
    one = session.replace_notes(draft, 0, 1, 'Opening\nSecond line', expected_sha256=original['snapshot_sha256'])
    assert session.read_notes(draft, 0, 1)['text'] == ''
    assert session.read_notes(draft, 1, 1)['text'] == 'Opening\nSecond line'
    read_before=session.read_notes(draft,1,1); revision_file=draft/'revision-000001.pptx'; original_bytes=revision_file.read_bytes(); original_slide=notes._slide
    def mutate_after_snapshot(frozen, number):
        revision_file.write_bytes(source.read_bytes())
        return original_slide(frozen, number)
    with patch.object(notes, '_slide', side_effect=mutate_after_snapshot): captured=session.read_notes(draft,1,1)
    assert captured['text']=='Opening\nSecond line' and captured['snapshot_sha256']==read_before['snapshot_sha256']
    try: session.replace_notes(draft,1,1,'Should fail',expected_sha256=captured['snapshot_sha256'])
    except ValueError as exc: assert 'snapshot changed' in str(exc)
    else: raise AssertionError('Changed notes snapshot accepted')
    revision_file.write_bytes(original_bytes)
    two = session.replace_notes(draft, 1, 1, 'Updated speaker notes')
    assert session.read_notes(draft, 2, 1)['text'] == 'Updated speaker notes'
    assert Presentation(two['file']).slides[0].shapes.title.text == 'Title'
    with zipfile.ZipFile(one['file']) as a, zipfile.ZipFile(two['file']) as b:
        changed = {name for name in a.namelist() if a.read(name) != b.read(name)}
        assert len(changed) == 1 and next(iter(changed)).startswith('ppt/notesSlides/'), changed
    try: session.replace_notes(draft, 0, 1, 'stale')
    except ValueError as exc: assert 'Revision conflict' in str(exc)
    else: raise AssertionError('Stale edit succeeded')
    assert session.info(draft)['latest_revision'] == 2
    cli = Path(__file__).with_name('open_office_use.py')
    read = subprocess.run([sys.executable, str(cli), 'session', 'notes', str(draft), '--revision', '1', '--slide', '1'], capture_output=True, text=True, check=True)
    assert json.loads(read.stdout)['text'] == 'Opening\nSecond line'
    payload = root / 'notes.txt'; payload.write_text('CLI edit', encoding='utf-8')
    subprocess.run([sys.executable, str(cli), 'session', 'replace-notes', str(draft), '--base', '2', '--slide', '1', '--text-file', str(payload)], capture_output=True, text=True, check=True)
    assert session.read_notes(draft, 3, 1)['text'] == 'CLI edit'
    formatted = root / 'formatted.pptx'; formatted_draft = root / 'formatted-draft'
    presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[1]); frame = slide.notes_slide.notes_text_frame
    paragraph = frame.paragraphs[0]; first = paragraph.add_run(); first.text = 'Hello '; first.font.bold = True
    second = paragraph.add_run(); second.text = 'world'; second.font.italic = True
    frame.add_paragraph().text = 'Second paragraph'; presentation.save(formatted)
    session.start(formatted_draft, source=formatted)
    result = session.replace_notes_text(formatted_draft, 0, 1, 1, 'lo wo', 'LO WO')
    edited = Presentation(result['file']).slides[0].notes_slide.notes_text_frame
    assert edited.text == 'HelLO WOrld\nSecond paragraph', edited.text
    assert edited.paragraphs[0].runs[0].font.bold and edited.paragraphs[0].runs[1].font.italic
    with zipfile.ZipFile(formatted) as a, zipfile.ZipFile(result['file']) as b:
        changed = {name for name in a.namelist() if a.read(name) != b.read(name)}
        assert len(changed) == 1 and next(iter(changed)).startswith('ppt/notesSlides/'), changed
    for base, find in [(1, 'missing'), (1, 'l'), (0, 'Hello')]:
        try: session.replace_notes_text(formatted_draft, base, 1, 1, find, 'x')
        except ValueError: pass
        else: raise AssertionError('Invalid edit succeeded')
    assert session.info(formatted_draft)['latest_revision'] == 1
    subprocess.run([sys.executable, str(cli), 'session', 'replace-notes-text', str(formatted_draft), '--base', '1', '--slide', '1', '--paragraph', '2', '--find', 'Second', '--replace', 'Next'], capture_output=True, text=True, check=True)
    assert session.read_notes(formatted_draft, 2, 1)['text'].endswith('Next paragraph')
    hyperlink = root / 'hyperlink.pptx'; link_deck = Presentation(); link_slide = link_deck.slides.add_slide(link_deck.slide_layouts[1])
    link_paragraph=link_slide.notes_slide.notes_text_frame.paragraphs[0];link_paragraph.add_run().text='See '
    link_run=link_paragraph.add_run();link_run.text='Linked';link_run.hyperlink.address='https://example.com';link_run.font.bold=True
    next_run=link_paragraph.add_run();next_run.text=' notes';next_run.hyperlink.address='https://example.com';next_run.font.italic=True
    link_paragraph.add_run().text=' today';link_deck.save(hyperlink)
    rejected=root/'rejected.pptx'
    try: notes.replace_text(hyperlink,rejected,1,1,'See Linked','Other')
    except ValueError as exc: assert 'boundary' in str(exc)
    else: raise AssertionError('Cross-boundary hyperlink edit succeeded')
    assert not rejected.exists()
    link_draft=root/'notes-link-draft';session.start(link_draft,source=hyperlink)
    linked_before=session.read_notes(link_draft,0,1)
    try: session.replace_notes_text(link_draft,0,1,1,'ked notes','k references',expected_sha256='0'*64)
    except ValueError as exc: assert 'snapshot changed' in str(exc)
    else: raise AssertionError('Changed notes paragraph source accepted')
    linked_result=session.replace_notes_text(link_draft,0,1,1,'ked notes','k references',expected_sha256=linked_before['snapshot_sha256'])
    linked_frame=Presentation(linked_result['file']).slides[0].notes_slide.notes_text_frame;linked_runs=linked_frame.paragraphs[0].runs
    assert linked_frame.text=='See Link references today' and linked_runs[1].font.bold and linked_runs[2].font.italic
    assert linked_runs[1].hyperlink.address=='https://example.com' and linked_runs[2].hyperlink.address=='https://example.com'
    with zipfile.ZipFile(hyperlink) as before,zipfile.ZipFile(linked_result['file']) as after:
        changed={name for name in before.namelist() if before.read(name)!=after.read(name)}
        assert len(changed)==1 and next(iter(changed)).startswith('ppt/notesSlides/'),changed
print('Speaker notes smoke passed: create, replace, CLI and revision safety')

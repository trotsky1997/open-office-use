"""Verify text candidates are tied to revisions and duplicates stay ambiguous."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from session import set_selection

cli = Path(__file__).with_name('open_office_use.py')
def run(*args):
    result = subprocess.run([sys.executable, str(cli), *map(str, args)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    return json.loads(result.stdout)

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    for suffix in ('.docx', '.xlsx', '.pptx'):
        source = root / ('source' + suffix); directory = root / suffix[1:]
        if suffix == '.docx':
            obj = Document(); obj.add_paragraph('Shared needle'); obj.add_paragraph('Shared needle'); obj.save(source)
        elif suffix == '.xlsx':
            obj = Workbook(); obj.active['A1'] = 'Shared needle'; obj.active['A2'] = 'Shared needle'; obj.save(source)
        else:
            obj = Presentation()
            for _ in range(2):
                slide = obj.slides.add_slide(obj.slide_layouts[1]); slide.shapes.title.text = 'Shared needle'
            obj.save(source)
        run('session', 'start', directory, '--source', source)
        set_selection(directory, {'revision': 0, 'page': 1, 'box': [0, 0, 1, 1]})
        result = run('session', 'selection-candidates', directory)
        assert result['revision'] == 0 and not result['stale'] and result['requires_verification'], result
        assert result['returned_candidate_count'] >= 2, result
        assert len({candidate['path'] for candidate in result['candidates']}) >= 2, result
        assert all('Shared needle' in candidate['text'] for candidate in result['candidates']), result
        if suffix == '.docx':
            spec = root / 'ops.json'; spec.write_text(json.dumps([{'command': 'set', 'path': '/body/p[1]', 'props': {'text': 'Changed'}}]))
            run('session', 'apply', directory, '--base', '0', '--spec', spec)
            stale = run('session', 'selection-candidates', directory)
            assert stale['stale'] and stale['revision'] == 0 and stale['returned_candidate_count'] >= 2, stale
    empty = root / 'empty'; run('session', 'start', empty, '--format', 'docx')
    set_selection(empty, {'revision': 0, 'page': 1, 'box': [0, 0, 0.05, 0.05]})
    result = run('session', 'selection-candidates', empty)
    assert not result['candidates'] and result['requires_verification'], result
print('Candidate smoke passed: duplicates in three formats and empty selections')

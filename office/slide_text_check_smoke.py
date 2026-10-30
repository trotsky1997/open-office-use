"""Check rendered slide text against normal and overflowing text boxes."""
import tempfile
import json
import subprocess
import sys
import zipfile
import warnings
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE
import open_office_use
import slide_text_check
import review_bundle
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'deck.pptx'; deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    normal = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)); normal.text_frame.text = 'A clear normal heading'
    normal.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    overflow = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(1.5), Inches(.4)); overflow.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    overflow.text_frame.word_wrap = True; overflow.text_frame.text = 'A very long sentence that overflows this small text box and requires attention'
    overflow.text_frame.paragraphs[0].runs[0].font.size = Pt(24); deck.save(source)
    result = slide_text_check.scan(source, open_office_use.render)
    assert result['mapping_valid'] and result['text_shapes_checked'] == 2, result
    assert [item['path_hint'] for item in result['candidates']] == [f'/slide[1]/shape[@id={overflow.shape_id}]'], result
    bundle = review_bundle.create(source, root / 'review', open_office_use.render)
    report = json.loads(Path(bundle['report']).read_text())
    assert report['checks']['slide_text']['result']['candidate_count'] == 1 and 'slide_text_candidates' in bundle['flags']
    draft = root / 'draft'; session.start(draft, source=source)
    cli = Path(__file__).with_name('open_office_use.py')
    command = subprocess.run([sys.executable, str(cli), 'session', 'inspect-slide-text', str(draft), '--revision', '0'], capture_output=True, text=True, check=True)
    assert json.loads(command.stdout)['candidate_count'] == 1
    duplicate = root / 'duplicate.pptx'
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        with zipfile.ZipFile(duplicate, 'w') as archive:
            archive.writestr('same', 'one'); archive.writestr('same', 'two')
    try: slide_text_check.scan(duplicate, lambda *_: (_ for _ in ()).throw(AssertionError('render called')))
    except ValueError: pass
    else: raise AssertionError('Duplicate package accepted')
    hidden = root / 'hidden.pptx'; slide._element.set('show', '0'); deck.slides.add_slide(deck.slide_layouts[6]); deck.save(hidden)
    hidden_result = slide_text_check.scan(hidden, open_office_use.render)
    assert not hidden_result['mapping_valid'] and not hidden_result['candidates'] and not hidden_result['complete'], hidden_result
print('Rendered slide text check smoke passed')

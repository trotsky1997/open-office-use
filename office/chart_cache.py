"""Compare bounded chart caches with stored worksheet values in an XLSX snapshot."""
import math
import re
import shutil
import tempfile
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from lxml import etree

import ooxml

C = '{http://schemas.openxmlformats.org/drawingml/2006/chart}'
CT = '{http://schemas.openxmlformats.org/package/2006/content-types}'
CHART_TYPE = 'application/vnd.openxmlformats-officedocument.drawingml.chart+xml'
CHARTEX_TYPE = 'application/vnd.ms-office.chartex+xml'
REFERENCE = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_.]+))!(\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6})(?::(\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}))?\Z")


def parse_reference(formula):
    match = REFERENCE.fullmatch(formula or '')
    if not match: return None
    sheet = match.group(1).replace("''", "'") if match.group(1) is not None else match.group(2)
    if '[' in sheet or ']' in sheet: return None
    ref = match.group(3) + (':' + match.group(4) if match.group(4) else '')
    try: c1, r1, c2, r2 = range_boundaries(ref)
    except ValueError: return None
    if not (1 <= c1 <= c2 <= 16384 and 1 <= r1 <= r2 <= 1048576) or (c1 != c2 and r1 != r2): return None
    return sheet, c1, r1, c2, r2


def scan(file):
    source = Path(file).resolve(strict=True)
    if source.suffix.lower() != '.xlsx': raise ValueError('Chart cache inspection requires an XLSX file')
    if source.stat().st_size > 128 * 1024 * 1024: raise ValueError('Workbook exceeds the 128 MiB size limit')
    with tempfile.TemporaryDirectory(prefix='office-chart-cache-') as tmp:
        snapshot = Path(tmp) / 'workbook.xlsx'; shutil.copyfile(source, snapshot)
        if snapshot.stat().st_size > 128 * 1024 * 1024: raise ValueError('Workbook exceeds the 128 MiB size limit')
        try:
            with zipfile.ZipFile(snapshot) as archive:
                entries = archive.infolist(); names = [item.filename for item in entries]
                if len(entries) > 5000 or len(set(names)) != len(names) or sum(item.file_size for item in entries) > 128 * 1024 * 1024:
                    raise ValueError('Workbook exceeds package limits or contains duplicate parts')
                manifest = ooxml.parse_xml(archive.read('[Content_Types].xml')).getroot()
                if manifest.tag != CT + 'Types': raise ValueError('Invalid content type manifest')
                chart_parts = []
                for item in manifest.findall(CT + 'Override'):
                    kind = item.get('ContentType'); name = item.get('PartName', '')
                    if kind not in (CHART_TYPE, CHARTEX_TYPE): continue
                    if not name.startswith('/') or '..' in name.split('/') or name[1:] not in names:
                        raise ValueError('Invalid chart part declaration')
                    chart_parts.append((name[1:], kind))
                if len({name for name, _ in chart_parts}) != len(chart_parts): raise ValueError('Duplicate chart part declaration')
                workbook = load_workbook(snapshot, read_only=True, data_only=True, keep_links=False)
                try:
                    sheets = {sheet.title: sheet for sheet in workbook.worksheets}; results = []; complete = True; points_read = 0; limit_hit = False
                    # shortcut: inspect 50 chart parts, 100 references and 10,000 source cells; split or inspect larger workbooks separately.
                    for part, kind in chart_parts[:50]:
                        if kind != CHART_TYPE or archive.getinfo(part).file_size > 4 * 1024 * 1024:
                            results.append({'part': part, 'status': 'unsupported_chart', 'references': []}); complete = False; continue
                        chart = ooxml.parse_xml(archive.read(part)).getroot()
                        if chart.tag != C + 'chartSpace':
                            results.append({'part': part, 'status': 'unsupported_chart', 'references': []}); complete = False; continue
                        external_data = chart.find('.//' + C + 'externalData') is not None
                        extensions = chart.find('.//' + C + 'extLst') is not None
                        if external_data or extensions: complete = False
                        references = []
                        for node in chart.iter():
                            if node.tag not in (C+'numRef', C+'strRef', C+'multiLvlStrRef', C+'numLit', C+'strLit'): continue
                            if sum(len(item['references']) for item in results) + len(references) >= 100:
                                complete = False; limit_hit = True; break
                            parent = node.getparent(); role = parent.tag.rsplit('}',1)[-1] if parent is not None else None
                            formula_node = node.find(C+'f'); formula = formula_node.text if formula_node is not None else None
                            item = {'role': role, 'kind': node.tag.rsplit('}',1)[-1], 'formula': (formula or '')[:500], 'status': 'unverified', 'mismatches': []}
                            references.append(item)
                            if node.tag not in (C+'numRef', C+'strRef'):
                                item['reason'] = 'unsupported_source'; complete = False; continue
                            cache_nodes = node.findall(C+('numCache' if node.tag == C+'numRef' else 'strCache'))
                            if not cache_nodes:
                                item['status'] = 'missing_cache'; complete = False; continue
                            if len(cache_nodes) != 1 or len(node.findall(C+'f')) != 1:
                                item['reason'] = 'ambiguous_cache_or_formula'; complete = False; continue
                            cache = cache_nodes[0]
                            parsed = parse_reference(formula)
                            if parsed is None:
                                item['reason'] = 'unsupported_reference'; complete = False; continue
                            name,c1,r1,c2,r2 = parsed; count = (c2-c1+1)*(r2-r1+1)
                            if name not in sheets or count > 2000 or points_read + count > 10000:
                                item['reason'] = 'sheet_or_range_limit'; complete = False; continue
                            count_nodes = cache.findall(C+'ptCount'); count_node = count_nodes[0] if len(count_nodes) == 1 else None
                            count_text = count_node.get('val') if count_node is not None else None
                            points = {}; invalid = count_text is None or not count_text.isdecimal() or len(count_text)>8 or int(count_text)!=count
                            for point in cache.findall(C+'pt'):
                                index = point.get('idx'); values = point.findall(C+'v'); value = values[0] if len(values) == 1 else None
                                if index is None or not index.isdecimal() or len(index)>8 or value is None or len(value): invalid = True; break
                                index = int(index)
                                if index in points or index >= count: invalid = True; break
                                points[index] = value.text or ''
                            if invalid:
                                item['reason'] = 'invalid_or_nonmatching_cache_structure'; complete = False; continue
                            values = []
                            for row in sheets[name].iter_rows(min_row=r1,max_row=r2,min_col=c1,max_col=c2,values_only=False): values.extend(row)
                            values.extend([None] * (count-len(values))); points_read += count
                            unknown = 0; compared = 0; mismatches = 0
                            for index, source_cell in enumerate(values[:count]):
                                value = source_cell.value if source_cell is not None else None
                                cached = points.get(index); reason = None
                                if value is None or isinstance(value, bool) or getattr(source_cell, 'data_type', None) == 'e': unknown += 1; continue
                                if node.tag == C+'numRef':
                                    if not isinstance(value, (int,float)) or isinstance(value,float) and not math.isfinite(value): unknown += 1; continue
                                    if cached is None: reason = 'missing_point'
                                    elif len(cached) > 256: unknown += 1; continue
                                    else:
                                        try:
                                            left, right = Decimal(str(value)), Decimal(cached)
                                            if not left.is_finite() or not right.is_finite(): raise InvalidOperation
                                            if left != right: reason = 'value_differs'
                                        except InvalidOperation: unknown += 1; continue
                                else:
                                    if not isinstance(value, str): unknown += 1; continue
                                    if cached is None: reason = 'missing_point'
                                    elif cached != value: reason = 'value_differs'
                                compared += 1
                                if reason:
                                    mismatches += 1
                                    if len(item['mismatches']) < 20:
                                        address = f'{get_column_letter(c1 + (index if r1 == r2 else 0))}{r1 + (index if c1 == c2 else 0)}'
                                        item['mismatches'].append({'index': index, 'cell': address, 'reason': reason, 'stored_value': str(value)[:500], 'cached_value': cached[:500] if cached is not None else None})
                            item.update({'sheet': name, 'range': f'{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}', 'points': count, 'compared_points': compared,
                                         'unverified_points': unknown, 'mismatch_count': mismatches, 'mismatches_truncated': mismatches > len(item['mismatches']),
                                         'status': 'mismatch' if mismatches else 'unverified' if unknown else 'match'})
                            if unknown: complete = False
                        if not references: complete = False
                        results.append({'part': part, 'status': 'scanned' if references else 'no_references_found', 'external_data': external_data, 'extensions_present': extensions, 'references': references})
                        if limit_hit: break
                    if len(chart_parts) > 50: complete = False
                    refs = [ref for chart in results for ref in chart['references']]
                    return {'charts': results, 'chart_parts_total': len(chart_parts), 'complete': complete,
                            'mismatch_references': sum(ref['status']=='mismatch' for ref in refs),
                            'missing_caches': sum(ref['status']=='missing_cache' for ref in refs),
                            'unverified_references': sum(ref['status']=='unverified' for ref in refs), 'source_cells_read': points_read,
                            'notes': ['Comparison uses stored worksheet values; formula caches can be stale.', 'Chart parts may be unreferenced; unsupported chart types and data sources are not fully covered.', 'Blank values, hidden data, formatting and rendering may affect chart behavior.']}
                finally: workbook.close()
        except (zipfile.BadZipFile, KeyError, etree.XMLSyntaxError) as exc:
            raise ValueError(f'Invalid workbook package: {exc}') from exc

"""Bounded worksheet inspection and explicit LibreOffice recalculation."""
import os
import hashlib
import re
import json
import shutil
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from xml.etree.ElementTree import tostring

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


def inspect_workbook(file):
    """Return bounded sparse samples from a saved XLSX snapshot."""
    source = Path(file).resolve(strict=True)
    if source.suffix.lower() != '.xlsx': raise ValueError('Workbook inspection requires an XLSX file')
    if source.stat().st_size > 128 * 1024 * 1024: raise ValueError('Workbook exceeds the 128 MiB size limit')
    with tempfile.TemporaryDirectory(prefix='office-workbook-inspect-') as tmp:
        snapshot = Path(tmp) / 'workbook.xlsx'; shutil.copyfile(source, snapshot)
        if snapshot.stat().st_size > 128 * 1024 * 1024: raise ValueError('Workbook exceeds the 128 MiB size limit')
        try:
            with zipfile.ZipFile(snapshot) as archive:
                entries = archive.infolist(); names = [entry.filename for entry in entries]
                if len(entries) > 5000 or len(names) != len(set(names)) or sum(entry.file_size for entry in entries) > 128 * 1024 * 1024:
                    raise ValueError('Workbook exceeds package limits or contains duplicate parts')
                workbook = load_workbook(snapshot, read_only=True, data_only=False, keep_links=False)
                try:
                    worksheets = workbook.worksheets; sheet_names = workbook.sheetnames; worksheet_names = {sheet.title for sheet in worksheets}
                    chart_names = [name for name in sheet_names if name not in worksheet_names]
                    results = []
                    # shortcut: inspect only the first 50 worksheets and first 100 populated cells per sheet; use targeted reads for the rest.
                    for sheet in worksheets[:50]:
                        path = getattr(sheet, '_worksheet_path', None)
                        if not isinstance(path, str) or path not in names: raise ValueError(f'Cannot locate worksheet XML: {sheet.title}')
                        strings = sheet._shared_strings; sample = []; scanned = 0; complete = True
                        with archive.open(path) as stream:
                            stack, sheet_data = [], None
                            for event, element in ET.iterparse(stream, events=('start', 'end')):
                                ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                                tag = element.tag[len(ns):] if isinstance(element.tag, str) and element.tag.startswith(ns) else ''
                                if event == 'start':
                                    if not stack and element.tag != '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}worksheet': raise ValueError(f'Invalid worksheet XML: {sheet.title}')
                                    if tag == 'sheetData' and stack[-1:] == ['worksheet']:
                                        if sheet_data is not None: raise ValueError(f'Worksheet has duplicate sheet data: {sheet.title}')
                                        sheet_data = element
                                    stack.append(tag); continue
                                if stack[-3:] == ['sheetData', 'row', 'c']:
                                    if scanned >= 10000: complete = False; break
                                    scanned += 1; parts = {child.tag[len(ns):]: child for child in element if isinstance(child.tag, str) and child.tag.startswith(ns)}
                                    formula, value, inline = parts.get('f'), parts.get('v'), parts.get('is')
                                    if formula is not None or value is not None and value.text is not None or inline is not None:
                                        if len(sample) >= 100: complete = False; break
                                        cell_type = element.get('t', 'n'); raw = value.text if value is not None else None
                                        text = None; text_complete = True
                                        if formula is not None:
                                            text = '=' + formula.text if formula.text else None
                                            if not formula.text: text_complete = None
                                        elif cell_type == 's':
                                            if raw is None or not raw.isdecimal() or len(raw) > 10 or int(raw) >= len(strings):
                                                raise ValueError(f'Invalid shared string index in {sheet.title}')
                                            text = strings[int(raw)]
                                            text_complete = None
                                        elif cell_type == 'inlineStr' and inline is not None:
                                            x = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                                            text = ''.join(node.text or '' for child in inline for node in ([child] if child.tag == x + 't' else child.findall(x + 't') if child.tag == x + 'r' else []))
                                            text_complete = all(child.tag == x + 't' or child.tag == x + 'r' and all(item.tag in (x + 't', x + 'rPr') for item in child) for child in inline)
                                        else: text = raw
                                        if text is not None and not isinstance(text, str): text = str(text)
                                        if text is not None and re.search(r'_x[0-9a-fA-F]{4}_', text): text_complete = False
                                        sample.append({'cell': element.get('r'), 'type': cell_type, 'formula_kind': formula.get('t', 'normal') if formula is not None else None,
                                                       'value': text[:500] if text is not None else None, 'value_truncated': text is not None and len(text) > 500, 'text_complete': text_complete})
                                    element.clear()
                                elif stack[-2:] == ['sheetData', 'row']:
                                    element.clear()
                                    if sheet_data is not None: sheet_data.remove(element)
                                stack.pop()
                        if sheet_data is None: raise ValueError(f'Worksheet has no sheet data: {sheet.title}')
                        results.append({'name': sheet.title, 'state': sheet.sheet_state, 'declared_max_row': sheet.max_row,
                                        'declared_max_column': sheet.max_column, 'cells': sample, 'stored_cells_scanned': scanned, 'sample_complete': complete})
                    return {'type': 'xlsx', 'sheets': results, 'total_worksheets': len(worksheets), 'worksheets_truncated': len(worksheets) > 50,
                            'chartsheet_names': chart_names[:50], 'chartsheets_truncated': len(chart_names) > 50,
                            'limits': {'worksheets': 50, 'cells_per_sheet': 100, 'stored_cells_scanned_per_sheet': 10000, 'characters_per_value': 500},
                            'dimension_note': 'Declared dimensions may be inaccurate; samples cover stored cells only.', 'formula_cache_freshness': 'unknown'}
                finally:
                    workbook.close()
        except (zipfile.BadZipFile, ET.ParseError) as exc:
            raise ValueError(f'Invalid workbook package or worksheet XML: {exc}') from exc


def read_range(file, sheet, cell_range, include_style=False):
    source = Path(file).resolve(strict=True)
    if source.suffix.lower() != '.xlsx':
        raise ValueError('Worksheet reading requires an .xlsx file')
    if source.stat().st_size > 128 * 1024 * 1024: raise ValueError('Workbook exceeds the 128 MiB size limit')
    with tempfile.TemporaryDirectory(prefix='office-cell-read-') as tmp:
        snapshot = Path(tmp) / 'workbook.xlsx'; shutil.copyfile(source, snapshot)
        if snapshot.stat().st_size > 128 * 1024 * 1024: raise ValueError('Workbook exceeds the 128 MiB size limit')
        try:
            with zipfile.ZipFile(snapshot) as archive:
                entries = archive.infolist(); names = [entry.filename for entry in entries]
                if len(entries) > 5000 or len(names) != len(set(names)) or sum(entry.file_size for entry in entries) > 128 * 1024 * 1024:
                    raise ValueError('Workbook exceeds package limits or contains duplicate parts')
            return _read_range(snapshot, sheet, cell_range, include_style)
        except zipfile.BadZipFile as exc:
            raise ValueError('Invalid workbook package') from exc


def _read_range(file, sheet, cell_range, include_style=False):
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        if min_col is None or min_row is None or max_col is None or max_row is None:
            raise ValueError
        if (min_col < 1 or min_row < 1
                or max_col > 16384 or max_row > 1048576 or min_col > max_col or min_row > max_row
                or (max_col - min_col + 1) * (max_row - min_row + 1) > 10000):
            raise ValueError
    except (TypeError, ValueError):
        raise ValueError('Specify a valid rectangular cell range of at most 10,000 cells') from None
    formulas = load_workbook(file, read_only=True, data_only=False)
    values = None
    try:
        values = load_workbook(file, read_only=True, data_only=True)
        if sheet not in formulas.sheetnames:
            raise ValueError(f'Worksheet not found: {sheet}')
        f_rows = formulas[sheet].iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
        v_rows = values[sheet].iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
        rows = []; size = 0
        for _ in range(max_row - min_row + 1):
            f_row = next(f_rows, ())
            v_row = next(v_rows, ())
            row = []
            for col in range(max_col - min_col + 1):
                f = f_row[col] if col < len(f_row) else None
                v = v_row[col] if col < len(v_row) else None
                if f is not None and f.data_type == 'f':
                    formula = f.value
                    if not isinstance(formula, str):
                        formula = {'text': getattr(formula, 'text', None), 'ref': getattr(formula, 'ref', None)}
                    row.append({'formula': formula, 'value': v.value if v else None, 'type': v.data_type if v else 'n'})
                else:
                    row.append({'value': f.value if f else None, 'type': f.data_type if f else 'n'})
                size += len(json.dumps(row[-1], ensure_ascii=False, default=str))
                if size > 1000000: raise ValueError('Cell data exceeds the 1,000,000-character response limit; request a smaller range')
            if include_style:
                for index, cell in enumerate(row):
                    f = f_row[index] if index < len(f_row) else None
                    font, fill = getattr(f, 'font', None), getattr(f, 'fill', None)
                    color = getattr(fill, 'fgColor', None)
                    rgb = color.rgb if fill and fill.patternType == 'solid' and color and color.type == 'rgb' else None
                    if include_style == 'review':
                        # Style IDs are workbook-local, so compare resolved components instead.
                        style = formulas._cell_styles[getattr(f, '_style_id', 0)]
                        cell['format'] = {
                            name: tostring(collection[index].to_tree(), encoding='unicode')
                            for name, collection, index in (
                                ('font', formulas._fonts, style.fontId),
                                ('fill', formulas._fills, style.fillId),
                                ('border', formulas._borders, style.borderId),
                                ('alignment', formulas._alignments, style.alignmentId),
                                ('protection', formulas._protections, style.protectionId),
                            )
                        }
                        cell['format']['number_format'] = getattr(f, 'number_format', 'General')
                    else:
                        cell['style'] = {'bold': bool(getattr(font, 'bold', False)),
                                         'fill_rgb': rgb[-6:] if isinstance(rgb, str) else None,
                                         'number_format': getattr(f, 'number_format', 'General')}
                        cell['locked'] = getattr(getattr(f, 'protection', None), 'locked', True) is not False
                    size += len(json.dumps(cell.get('format', cell.get('style')), ensure_ascii=False, default=str))
                    if size > 1000000: raise ValueError('Cell data exceeds the 1,000,000-character response limit; request a smaller range')
            rows.append(row)
        return {'sheet': sheet, 'range': cell_range, 'rows': rows, 'cache_freshness': 'unknown'}
    finally:
        formulas.close()
        if values:
            values.close()


FORMULA_SCAN_MAX_CELLS = 1000000
FORMULA_SCAN_MAX_ISSUES = 200


def search_shared_strings(archive):
    """Read verifiable plain and rich strings from the package's shared string table."""
    types_ns = '{http://schemas.openxmlformats.org/package/2006/content-types}'
    main_ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    types = ET.fromstring(archive.read('[Content_Types].xml'))
    parts = [item.get('PartName', '').lstrip('/') for item in types.findall(types_ns + 'Override')
             if item.get('ContentType') == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml']
    if len(parts) != 1 or parts[0] not in archive.namelist(): return None
    strings = []
    with archive.open(parts[0]) as stream:
        root = None
        for event, node in ET.iterparse(stream, events=('start', 'end')):
            if root is None:
                root = node
                if root.tag != main_ns + 'sst': return None
            if event == 'end' and node.tag == main_ns + 'si':
                children = list(node); texts = []; valid = bool(children)
                plain = [child for child in children if child.tag == main_ns + 't']
                runs = [child for child in children if child.tag == main_ns + 'r']
                valid = valid and (len(plain) == 1 and len(children) == 1 or bool(runs) and len(runs) == len(children))
                for child in children:
                    if child.tag == main_ns + 't':
                        if len(child): valid = False
                        texts.append(child.text or '')
                    elif child.tag == main_ns + 'r':
                        run_text = child.findall(main_ns + 't')
                        if len(run_text) != 1 or len(child.findall(main_ns + 'rPr')) > 1 or any(len(item) for item in run_text) or any(item.tag not in (main_ns + 't', main_ns + 'rPr') for item in child): valid = False
                        texts.extend(item.text or '' for item in run_text)
                text = ''.join(texts)
                if re.search(r'_x[0-9a-fA-F]{4}_', text): valid = False
                strings.append((text, valid)); root.remove(node); node.clear()
    return strings


def search_cells(file, query, sheet=None, offset=0, snapshot_sha256=None):
    """Search stored worksheet text and formula text in a saved XLSX snapshot."""
    if not isinstance(query, str) or not query or len(query) > 200 or any(ord(c) < 32 for c in query):
        raise ValueError('Search text must contain 1 to 200 characters without control characters')
    if sheet is not None and (not isinstance(sheet, str) or not sheet): raise ValueError('Sheet name must be nonempty text')
    if isinstance(offset, bool) or not isinstance(offset, int) or offset < 0: raise ValueError('Offset must be a nonnegative integer')
    if offset and snapshot_sha256 is None: raise ValueError('Continuing a search requires the prior snapshot SHA-256')
    if snapshot_sha256 is not None and (not isinstance(snapshot_sha256, str) or not re.fullmatch(r'[0-9a-f]{64}', snapshot_sha256)):
        raise ValueError('Invalid snapshot SHA-256')
    source = Path(file).resolve(strict=True)
    if source.suffix.lower() != '.xlsx': raise ValueError('Cell search requires an XLSX file')
    if source.stat().st_size > 128 * 1024 * 1024: raise ValueError('Workbook exceeds the 128 MiB size limit')
    with tempfile.TemporaryDirectory(prefix='office-cell-search-') as tmp:
        snapshot = Path(tmp) / 'workbook.xlsx'; shutil.copyfile(source, snapshot)
        if snapshot.stat().st_size > 128 * 1024 * 1024: raise ValueError('Workbook exceeds the 128 MiB size limit')
        digest = hashlib.sha256()
        with snapshot.open('rb') as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b''): digest.update(chunk)
        digest = digest.hexdigest()
        if snapshot_sha256 is not None and digest != snapshot_sha256: raise ValueError('Workbook snapshot changed; restart the search')
        try:
            with zipfile.ZipFile(snapshot) as archive:
                entries = archive.infolist(); names = [entry.filename for entry in entries]
                if len(entries) > 5000 or len(names) != len(set(names)) or sum(entry.file_size for entry in entries) > 128 * 1024 * 1024:
                    raise ValueError('Workbook exceeds package limits or contains duplicate parts')
                workbook = load_workbook(snapshot, read_only=True, data_only=False, keep_links=False)
                try:
                    sheets = [(s.title, s.sheet_state, s._worksheet_path, s._shared_strings) for s in workbook.worksheets]
                    raw_strings = search_shared_strings(archive)
                    chartsheets = len(workbook.sheetnames) - len(sheets)
                    total_worksheets = len(sheets)
                    if sheet is not None:
                        sheets = [s for s in sheets if s[0] == sheet]
                        if not sheets: raise ValueError('Worksheet not found')
                    matches, scanned, unverified, scanned_sheets = [], 0, 0, []
                    complete = True; skipped = 0; stop_reason = None
                    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                    for name, state, path, strings in sheets:
                        if path not in names: raise ValueError(f'Cannot locate worksheet XML: {name}')
                        with archive.open(path) as stream:
                            stack, sheet_data = [], None
                            for event, element in ET.iterparse(stream, events=('start', 'end')):
                                tag = element.tag[len(ns):] if isinstance(element.tag, str) and element.tag.startswith(ns) else ''
                                if event == 'start':
                                    if not stack and element.tag != ns + 'worksheet': raise ValueError(f'Invalid worksheet XML: {name}')
                                    if tag == 'sheetData' and stack[-1:] == ['worksheet']:
                                        if sheet_data is not None: raise ValueError(f'Duplicate worksheet data: {name}')
                                        sheet_data = element
                                    stack.append(tag); continue
                                if stack[-3:] == ['sheetData', 'row', 'c']:
                                    if skipped < offset:
                                        skipped += 1; element.clear(); stack.pop(); continue
                                    if scanned >= 100000 or len(matches) >= 200:
                                        complete = False; stop_reason = 'match_limit' if len(matches) >= 200 else 'cell_limit'; break
                                    scanned += 1; children = {child.tag: child for child in element}
                                    formula, value, inline = children.get(ns + 'f'), children.get(ns + 'v'), children.get(ns + 'is')
                                    kind, text, uncertain = None, None, False
                                    if formula is not None:
                                        kind = 'formula'; text = '=' + formula.text if formula.text else None
                                        uncertain = text is None
                                    elif element.get('t') == 's':
                                        raw = value.text if value is not None else None
                                        if raw is None or not raw.isdecimal() or len(raw) > 10 or int(raw) >= len(strings): raise ValueError(f'Invalid shared string index in {name}')
                                        index = int(raw); kind, text, uncertain = 'text', strings[index], True
                                        if raw_strings is not None and len(raw_strings) == len(strings) and raw_strings[index][1]:
                                            text, uncertain = raw_strings[index][0], False
                                    elif element.get('t') == 'inlineStr' and inline is not None:
                                        kind = 'text'; text = ''.join(node.text or '' for child in inline for node in ([child] if child.tag == ns + 't' else child.findall(ns + 't') if child.tag == ns + 'r' else []))
                                        uncertain = any(child.tag not in (ns + 't', ns + 'r') or child.tag == ns + 'r' and any(item.tag not in (ns + 't', ns + 'rPr') for item in child) for child in inline)
                                    elif element.get('t') == 'inlineStr':
                                        uncertain = True
                                    elif element.get('t') == 'str' and value is not None:
                                        kind, text = 'text', value.text or ''
                                    if text is not None and (re.search(r'_x[0-9a-fA-F]{4}_', text) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', element.get('r', ''))): uncertain = True
                                    if uncertain: unverified += 1
                                    if text is not None and query in text:
                                        index = text.index(query); start = max(0, index - 100); end = min(len(text), index + len(query) + 100)
                                        matches.append({'sheet': name, 'sheet_state': state, 'cell': element.get('r'), 'kind': kind,
                                                        'excerpt': text[start:end], 'excerpt_start': start, 'text_length': len(text), 'text_complete': not uncertain})
                                    element.clear()
                                elif stack[-2:] == ['sheetData', 'row']:
                                    element.clear()
                                    if sheet_data is not None: sheet_data.remove(element)
                                stack.pop()
                            if sheet_data is None: raise ValueError(f'Worksheet has no sheet data: {name}')
                        scanned_sheets.append({'sheet': name, 'complete': complete})
                        if not complete: break
                    if skipped < offset: raise ValueError('Search offset exceeds the selected worksheets')
                    return {'query': query, 'sheet': sheet, 'case_sensitive': True, 'snapshot_sha256': digest,
                            'offset': offset, 'cells_skipped': skipped, 'next_offset': offset + scanned if not complete else None,
                            'stop_reason': stop_reason, 'matches': matches, 'match_count': len(matches),
                            'matches_truncated': stop_reason == 'match_limit', 'stored_cells_scanned': scanned,
                            'sheets': scanned_sheets, 'total_worksheets': total_worksheets, 'chartsheets_omitted': chartsheets,
                            'complete': complete and offset == 0 and unverified == 0, 'scan_complete': complete, 'unverified_text_cells': unverified,
                            'limits': {'stored_cells': 100000, 'matches': 200},
                            'notes': ['Counts and completeness apply to this page; continue with the same query and sheet.',
                                      'Search covers stored text and formula text, not calculated values or chart content.',
                                      'Complex shared strings, shared formula followers and Excel escape sequences may be unverified.']}
                finally: workbook.close()
        except (zipfile.BadZipFile, ET.ParseError) as exc:
            raise ValueError(f'Invalid workbook package or worksheet XML: {exc}') from exc


def scan_formulas(file):
    """Inspect stored formula caches and error cells across saved worksheets."""
    source = Path(file).resolve(strict=True)
    if source.suffix.lower() != '.xlsx':
        raise ValueError('Formula scanning requires an .xlsx file')
    if source.stat().st_size > 128 * 1024 * 1024:
        raise ValueError('Workbook exceeds the 128 MiB scan limit')
    with tempfile.TemporaryDirectory(prefix='office-formulas-') as tmp:
        snapshot = Path(tmp) / 'workbook.xlsx'
        shutil.copyfile(source, snapshot)
        if snapshot.stat().st_size > 128 * 1024 * 1024:
            raise ValueError('Workbook exceeds the 128 MiB scan limit')
        try:
            with zipfile.ZipFile(snapshot) as archive:
                entries = archive.infolist()
                names = [entry.filename for entry in entries]
                if len(entries) > 5000 or len(names) != len(set(names)) or sum(entry.file_size for entry in entries) > 128 * 1024 * 1024:
                    raise ValueError('Workbook exceeds package limits or contains duplicate parts')
                workbook = load_workbook(snapshot, read_only=True, data_only=False, keep_links=False)
                try:
                    sheets = [(sheet.title, getattr(sheet, '_worksheet_path', None)) for sheet in workbook.worksheets]
                finally:
                    workbook.close()
                counts = {'scanned_cells': 0, 'formulas': 0, 'formula_errors': 0,
                          'literal_errors': 0, 'empty_or_missing_cache': 0}
                errors, missing, scanned_sheets = [], [], []
                complete = True
                for name, path in sheets:
                    if not isinstance(path, str) or path not in names:
                        raise ValueError(f'Cannot locate worksheet XML: {name}')
                    sheet_complete = True
                    with archive.open(path) as stream:
                        stack, sheet_data = [], None
                        for event, element in ET.iterparse(stream, events=('start', 'end')):
                            tag = element.tag.rsplit('}', 1)[-1]
                            if event == 'start':
                                if tag == 'sheetData' and stack[-1:] == ['worksheet']:
                                    sheet_data = element
                                stack.append(tag)
                                continue
                            if stack[-3:] == ['sheetData', 'row', 'c']:
                                if counts['scanned_cells'] >= FORMULA_SCAN_MAX_CELLS:
                                    sheet_complete = complete = False
                                    break
                                counts['scanned_cells'] += 1
                                children = {child.tag.rsplit('}', 1)[-1]: child for child in element}
                                formula, value = children.get('f'), children.get('v')
                                is_formula = formula is not None
                                cached = value.text if value is not None else None
                                if is_formula:
                                    counts['formulas'] += 1
                                if element.get('t') == 'e':
                                    kind = 'formula_errors' if is_formula else 'literal_errors'
                                    counts[kind] += 1
                                    if len(errors) < FORMULA_SCAN_MAX_ISSUES:
                                        errors.append({'sheet': name, 'cell': element.get('r'), 'kind': kind,
                                                       'error': (cached or '')[:256], 'formula': (formula.text or '')[:500] if is_formula else None})
                                elif is_formula and (cached is None or cached == ''):
                                    counts['empty_or_missing_cache'] += 1
                                    if len(missing) < FORMULA_SCAN_MAX_ISSUES:
                                        missing.append({'sheet': name, 'cell': element.get('r'), 'formula': (formula.text or '')[:500]})
                                element.clear()
                            elif stack[-2:] == ['sheetData', 'row']:
                                element.clear()
                                if sheet_data is not None:
                                    sheet_data.remove(element)
                            stack.pop()
                    scanned_sheets.append({'sheet': name, 'complete': sheet_complete})
                    if not complete:
                        break
                return {**counts, 'sheets': scanned_sheets, 'total_sheets': len(sheets), 'complete': complete,
                        'errors': errors, 'errors_truncated': counts['formula_errors'] + counts['literal_errors'] > len(errors),
                        'empty_or_missing': missing, 'empty_or_missing_truncated': counts['empty_or_missing_cache'] > len(missing),
                        'cache_freshness': 'unknown'}
        except (zipfile.BadZipFile, ET.ParseError) as exc:
            raise ValueError(f'Invalid workbook package or worksheet XML: {exc}') from exc


def recalculate(source, output):
    source = Path(source).resolve(strict=True)
    output = Path(output).absolute()
    if source.suffix.lower() != '.xlsx' or output.suffix.lower() != '.xlsx':
        raise ValueError('Recalculation requires .xlsx input and output')
    if output.exists() or output.is_symlink():
        raise FileExistsError(f'Output already exists: {output}')
    executable = shutil.which('soffice') or shutil.which('libreoffice')
    if not executable:
        raise RuntimeError('LibreOffice (soffice) is required for recalculation')
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent, prefix='.office-calc-') as tmp:
        root = Path(tmp)
        inputs = root / 'input'; inputs.mkdir()
        results = root / 'result'; results.mkdir()
        copy = inputs / 'workbook.xlsx'
        shutil.copyfile(source, copy)
        profile = (root / 'profile').as_uri()
        result = subprocess.run([executable, '-env:UserInstallation=' + profile, '--headless',
                                 '--convert-to', 'xlsx', '--outdir', str(results), str(copy)],
                                capture_output=True, text=True, timeout=120, check=False)
        candidate = results / copy.name
        if result.returncode or not candidate.is_file():
            raise RuntimeError(f'LibreOffice recalculation failed: {result.stderr or result.stdout}')
        check = load_workbook(candidate, read_only=True, data_only=False)
        check.close()
        os.link(candidate, output)
    return str(output)


def worksheet_constraints(file, sheet):
    """Inspect multi-cell formulas, merged ranges and sheet protection without rewriting."""
    import xml.etree.ElementTree as ET
    import zipfile

    if Path(file).stat().st_size > 128 * 1024 * 1024: raise ValueError('Workbook exceeds the 128 MiB size limit')
    with zipfile.ZipFile(file) as archive:
        entries = archive.infolist(); names = [entry.filename for entry in entries]
        if len(entries) > 5000 or len(names) != len(set(names)) or sum(entry.file_size for entry in entries) > 128 * 1024 * 1024:
            raise ValueError('Workbook exceeds package limits or contains duplicate parts')
    workbook = load_workbook(file, read_only=True, data_only=False)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f'Worksheet not found: {sheet}')
        # openpyxl resolves workbook relationships to this package-relative path.
        path = getattr(workbook[sheet], '_worksheet_path', None)
        if not isinstance(path, str):
            raise RuntimeError('Cannot locate worksheet XML for formula protection')
        if sum(getattr(item, '_worksheet_path', None) == path for item in workbook.worksheets) != 1:
            raise ValueError('Worksheet part is referenced by multiple sheets')
    finally:
        workbook.close()
    ranges = []; merges = []; sheet_protected = False; protection_seen = False
    shared_ranges, shared_uses = {}, set()
    with zipfile.ZipFile(file) as archive, archive.open(path) as stream:
        for _, element in ET.iterparse(stream, events=('end',)):
            if element.tag == '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}mergeCell':
                try:
                    bounds = range_boundaries(element.get('ref'))
                    if any(v is None for v in bounds) or not (1 <= bounds[0] <= bounds[2] <= 16384 and 1 <= bounds[1] <= bounds[3] <= 1048576): raise ValueError
                except (TypeError, ValueError):
                    raise ValueError('Invalid merged-cell range') from None
                merges.append(bounds)
            elif element.tag == '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheetProtection':
                if protection_seen: raise ValueError('Worksheet has duplicate protection settings')
                protection_seen = True; flag = element.get('sheet', '0')
                if flag not in ('0', '1', 'true', 'false'): raise ValueError('Invalid worksheet protection setting')
                sheet_protected = flag in ('1', 'true')
            if element.tag.rsplit('}', 1)[-1] == 'f':
                kind, ref, shared_id = element.get('t'), element.get('ref'), element.get('si')
                if kind == 'shared':
                    if shared_id is None:
                        raise ValueError('Shared formula has no identifier')
                    shared_uses.add(shared_id)
                if kind in ('array', 'dataTable', 'shared') and ref:
                    try:
                        bounds = range_boundaries(ref)
                        if any(value is None for value in bounds) or not (1 <= bounds[0] <= bounds[2] <= 16384 and 1 <= bounds[1] <= bounds[3] <= 1048576):
                            raise ValueError
                    except (TypeError, ValueError):
                        raise ValueError('Invalid multi-cell formula range') from None
                    if kind == 'shared':
                        assert shared_id is not None
                        if shared_id in shared_ranges and shared_ranges[shared_id] != bounds:
                            raise ValueError('Conflicting shared formula ranges')
                        shared_ranges[shared_id] = bounds
                    ranges.append(bounds)
                elif kind in ('array', 'dataTable') and not ref:
                    raise ValueError('Multi-cell formula has no protected range')
            element.clear()
    if shared_uses - shared_ranges.keys():
        raise ValueError('Shared formula has no protected range')
    return {'formula_ranges': ranges, 'merged_ranges': merges, 'sheet_protected': sheet_protected}


def cell_constraint(constraints, col, row, locked, complex_formula=False):
    ranges = [bounds for bounds in constraints['merged_ranges'] if bounds[0] <= col <= bounds[2] and bounds[1] <= row <= bounds[3]]
    from openpyxl.utils.cell import get_column_letter
    anchor = f'{get_column_letter(ranges[0][0])}{ranges[0][1]}' if len(ranges) == 1 else None
    reason = ('overlapping_merges' if len(ranges) > 1 else
              'merged_placeholder' if ranges and (col, row) != ranges[0][:2] else
              'formula_range' if complex_formula or any(c1 <= col <= c2 and r1 <= row <= r2 for c1, r1, c2, r2 in constraints['formula_ranges']) else
              'locked_cell' if constraints['sheet_protected'] and locked else None)
    return reason, anchor

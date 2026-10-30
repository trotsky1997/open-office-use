"""Exercise slide geometry checks, transforms and explicit scan limits."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

import layout
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.pptx'; deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    def text(name, x, y, w, h):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); shape.name = name; shape.text = name; return shape
    text('Alpha', 1, 1, 2, 1); text('Beta', 2, 1.5, 2, 1); text('Overflow', 9.5, 7, 2, 1)
    group = slide.shapes.add_group_shape(); child = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)); child.name = 'Rotated child'; child.text = 'Group'
    group.left = Inches(5); group.top = Inches(3); group.width = Inches(2); group.height = Inches(1); group.rotation = 90
    text('Group target', 5.5, 2.5, 1, 2)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.1), Inches(1.1), Inches(0.5), Inches(0.5)); bg.name = 'Overlay'
    hidden = text('Hidden', 20, 20, 1, 1); hidden._element.find(layout.P + 'nvSpPr/' + layout.P + 'cNvPr').set('hidden', '1')
    deck.save(source); original = source.read_bytes(); result = layout.inspect(source)
    assert result['complete'] and result['slides_total'] == 1 and result['off_slide_count'] == 1, result
    assert result['off_slide'][0]['name'] == 'Overflow' and result['off_slide'][0]['bounds_inches'] == [9.5, 7.0, 11.5, 8.0]
    visible = result['off_slide'][0]['visible_polygon_inches']
    assert min(x for x, y in visible) == 9.5 and max(x for x, y in visible) == 10 and min(y for x, y in visible) == 7 and max(y for x, y in visible) == 7.5
    assert result['visibility_known'] and not result['hidden_slides']
    pairs = {frozenset((item['first']['name'], item['second']['name'])): item for item in result['overlaps']}
    alpha_beta = pairs[frozenset(('Alpha', 'Beta'))]
    assert alpha_beta['area_square_inches'] == 0.5 and alpha_beta['fraction_of_smaller'] == 0.25
    assert len(alpha_beta['first']['polygon_inches']) == 4 and len(alpha_beta['intersection_polygon_inches']) == 4
    grouped = pairs[frozenset(('Rotated child', 'Group target'))]
    child_descriptor = next(item for item in (grouped['first'], grouped['second']) if item['name'] == 'Rotated child')
    assert child_descriptor['bounds_inches'] == [5.5, 2.5, 6.5, 4.5] and '/group[@id=' in child_descriptor['path_hint']
    assert frozenset(('Alpha', 'Overlay')) in pairs and source.read_bytes() == original
    flipped = slide.shapes.add_group_shape()
    flipped.shapes.add_textbox(Inches(1), Inches(1), Inches(1.5), Inches(0.8)).text = 'First'
    flipped.shapes.add_textbox(Inches(4), Inches(1), Inches(1.5), Inches(0.8)).text = 'Second'
    flipped.left = Inches(2); flipped.top = Inches(2); flipped.width = Inches(6); flipped.height = Inches(1.5); flipped.rotation = 90
    flipped._element.grpSpPr.xfrm.set('flipH', '1')
    matrix = layout.group_transform(flipped, (1, 0, 0, 1, 0, 0)); first = layout.polygon(flipped.shapes[0], matrix); second = layout.polygon(flipped.shapes[1], matrix)
    assert round(min(y for x, y in first) / layout.EMU, 2) == 3.75 and round(min(y for x, y in second) / layout.EMU, 2) == -0.25
    diamond = [(0, 1), (1, 0), (2, 1), (1, 2)]; corner = [(0, 0), (0.2, 0), (0.2, 0.2), (0, 0.2)]
    assert layout.area(layout.intersect(diamond, corner)) == 0
    outside = root / 'outside.pptx'; off_deck = Presentation(); off_slide = off_deck.slides.add_slide(off_deck.slide_layouts[6])
    off_slide.shapes.add_textbox(Inches(12), Inches(1), Inches(2), Inches(1)).text = 'Outside one'
    off_slide.shapes.add_textbox(Inches(12.5), Inches(1), Inches(2), Inches(1)).text = 'Outside two'
    picture = root / 'image.png'; Image.new('RGB', (8, 8), 'red').save(picture)
    off_slide.shapes.add_picture(str(picture), Inches(13), Inches(2))
    off_deck.save(outside); off_result = layout.inspect(outside)
    assert off_result['off_slide_count'] == 3 and off_result['overlap_count'] == 0, off_result
    assert any('/picture[@id=' in item['path_hint'] for item in off_result['off_slide'])
    hidden_deck = Presentation(); hidden_deck.slides.add_slide(hidden_deck.slide_layouts[6]); hidden = hidden_deck.slides.add_slide(hidden_deck.slide_layouts[6]); hidden._element.set('show', '0')
    hidden_source = root / 'hidden.pptx'; hidden_deck.save(hidden_source); hidden_result = layout.inspect(hidden_source)
    assert hidden_result['visibility_known'] and hidden_result['hidden_slides'] == [2]
    invalid = root / 'invalid.pptx'; invalid.write_bytes(b'not a presentation')
    try:
        layout.inspect(invalid)
        raise AssertionError('Invalid package was accepted')
    except ValueError:
        pass
    draft = root / 'draft'; session.start(draft, source=source)
    assert session.inspect_layout(draft, 0)['revision'] == 0
    cli = Path(__file__).with_name('open_office_use.py')
    output = subprocess.run([sys.executable, str(cli), 'inspect-layout', str(source)], check=True, capture_output=True, text=True)
    assert json.loads(output.stdout)['overlap_count'] == result['overlap_count']
    output = subprocess.run([sys.executable, str(cli), 'session', 'inspect-layout', str(draft), '--revision', '0'], check=True, capture_output=True, text=True)
    assert json.loads(output.stdout)['revision'] == 0
    saved_shapes, saved_pairs, saved_issues = layout.MAX_SHAPES, layout.MAX_PAIRS, layout.MAX_ISSUES
    try:
        layout.MAX_SHAPES = 3; limited = layout.inspect(source)
        assert not limited['complete'] and limited['skipped_count'] and limited['shapes_scanned'] == 3
        layout.MAX_SHAPES = saved_shapes; layout.MAX_PAIRS = 1; limited = layout.inspect(source)
        assert not limited['complete'] and not limited['overlap_complete'] and limited['pairs_checked'] == 1
        layout.MAX_PAIRS = saved_pairs; layout.MAX_ISSUES = 1; limited = layout.inspect(source)
        assert limited['overlaps_truncated'] and len(limited['overlaps']) == 1 and limited['overlap_count'] > 1
        limited_off = layout.inspect(outside)
        assert limited_off['off_slide_count'] == 3 and len(limited_off['off_slide']) == 1 and limited_off['off_slide_truncated']
    finally:
        layout.MAX_SHAPES, layout.MAX_PAIRS, layout.MAX_ISSUES = saved_shapes, saved_pairs, saved_issues
print('Layout smoke passed: overlap, rotated groups, off-slide shapes and scan limits')

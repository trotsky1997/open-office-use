"""Small local file-oriented office harness."""
import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

import officecli_backend
import spreadsheet
import review


def kind(path):
    suffix = Path(path).suffix.lower()
    if suffix not in ('.docx', '.xlsx', '.pptx'):
        raise ValueError('Supported formats: .docx, .xlsx, .pptx')
    return suffix


def inspect(path):
    suffix = kind(path)
    if suffix == '.docx':
        doc = Document(path)
        return {'type': 'docx', 'paragraphs': [p.text for p in doc.paragraphs],
                'tables': [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]}
    if suffix == '.xlsx':
        wb = load_workbook(path, data_only=False)
        return {'type': 'xlsx', 'sheets': [{'name': ws.title, 'rows': [[cell.value for cell in row] for row in ws.iter_rows()] } for ws in wb.worksheets]}
    prs = Presentation(path)
    return {'type': 'pptx', 'slides': [[{'shape': i, 'text': shape.text} for i, shape in enumerate(slide.shapes) if shape.has_text_frame] for slide in prs.slides]}


def author(spec, output, source=None):
    suffix = kind(output)
    if source and kind(source) != suffix:
        raise ValueError('Source and output formats must match')
    out = Path(output)
    if out.exists():
        raise FileExistsError(f'Output already exists: {out}')
    if source and out.resolve() == Path(source).resolve():
        raise ValueError('Output must differ from source')
    if suffix == '.docx':
        doc = Document(source) if source else Document()
        for index, text in spec.get('replace_paragraphs', {}).items():
            # Setting paragraph text discards run-level formatting in that paragraph.
            doc.paragraphs[int(index)].text = text
        for text in spec.get('paragraphs', []):
            doc.add_paragraph(text)
        obj = doc
    elif suffix == '.xlsx':
        obj = load_workbook(source) if source else Workbook()
        for name, cells in spec.get('sheets', {}).items():
            ws = obj[name] if name in obj.sheetnames else obj.create_sheet(name)
            for address, value in cells.items():
                ws[address] = value
        if not source and 'Sheet' not in spec.get('sheets', {}) and len(obj.sheetnames) > 1:
            del obj['Sheet']
    else:
        obj = Presentation(source) if source else Presentation()
        for item in spec.get('replace_text', []):
            shape = obj.slides[item['slide']].shapes[item['shape']]
            if not shape.has_text_frame:
                raise ValueError('Selected shape has no text frame')
            shape.text = item['text']
        for item in spec.get('slides', []):
            slide = obj.slides.add_slide(obj.slide_layouts[1])
            slide.shapes.title.text = item['title']
            slide.placeholders[1].text = item.get('body', '')
    out.parent.mkdir(parents=True, exist_ok=True)
    # Save into the target directory, then exclusively create the requested output.
    with tempfile.TemporaryDirectory(dir=out.parent) as tmp:
        candidate = Path(tmp) / out.name
        obj.save(candidate)
        with out.open('xb') as destination, candidate.open('rb') as original:
            shutil.copyfileobj(original, destination)
    return str(out.resolve())


def render(source, output_dir):
    kind(source)
    source = Path(source).resolve(strict=True)
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    target = destination / (source.stem + '.pdf')
    if target.exists():
        raise FileExistsError(f'Output already exists: {target}')
    executable = shutil.which('soffice') or shutil.which('libreoffice')
    if not executable:
        raise RuntimeError('LibreOffice (soffice) is required for rendering')
    with tempfile.TemporaryDirectory() as tmp:
        profile = (Path(tmp) / 'profile').as_uri()
        result = subprocess.run([executable, '-env:UserInstallation=' + profile, '--headless', '--convert-to', 'pdf', '--outdir', tmp, str(source)], capture_output=True, text=True, timeout=120, check=False)
        pdf = Path(tmp) / target.name
        if result.returncode or not pdf.is_file():
            raise RuntimeError(f'LibreOffice conversion failed: {result.stderr or result.stdout}')
        with target.open('xb') as out, pdf.open('rb') as inp:
            shutil.copyfileobj(inp, out)
    return str(target.resolve())


def preview(source, output_dir):
    executable = shutil.which('pdftoppm')
    if not executable:
        raise RuntimeError('Poppler pdftoppm is required for page previews')
    destination = Path(output_dir).absolute()
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=destination.parent, prefix='.office-preview-') as tmp:
        pdf = render(source, tmp)
        result = subprocess.run([executable, '-png', '-r', '120', pdf, str(Path(tmp) / 'page')],
                                capture_output=True, text=True, timeout=120, check=False)
        pages = sorted(Path(tmp).glob('page-*.png'))
        if result.returncode or not pages:
            raise RuntimeError(f'Page rendering failed: {result.stderr}')
        destination.mkdir()
        published = []
        try:
            for page in pages:
                target = destination / page.name
                os.link(page, target)
                published.append(target)
        except OSError:
            for target in published:
                target.unlink()
            destination.rmdir()
            raise
    return [str(page) for page in published]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    read = commands.add_parser('inspect'); read.add_argument('file')
    edit = commands.add_parser('author'); edit.add_argument('--spec', required=True); edit.add_argument('--out', required=True); edit.add_argument('--source')
    render_command = commands.add_parser('render'); render_command.add_argument('file'); render_command.add_argument('--out-dir', required=True)
    images = commands.add_parser('preview'); images.add_argument('file'); images.add_argument('--out-dir', required=True)
    structured = commands.add_parser('structured'); structured.add_argument('file'); structured.add_argument('mode', choices=['outline', 'get', 'query']); structured.add_argument('target', nargs='?')
    batch = commands.add_parser('batch'); batch.add_argument('--spec', required=True); batch.add_argument('--out', required=True); batch.add_argument('--source')
    cells = commands.add_parser('cells'); cells.add_argument('file'); cells.add_argument('--sheet', required=True); cells.add_argument('--range', required=True, dest='cell_range')
    calc = commands.add_parser('recalculate'); calc.add_argument('file'); calc.add_argument('--out', required=True)
    audit = commands.add_parser('audit'); audit.add_argument('file')
    comparison = commands.add_parser('compare'); comparison.add_argument('before'); comparison.add_argument('after')
    commands.add_parser('mcp', help='Start the local MCP stdio server (requires the mcp extra)')
    args = parser.parse_args()
    try:
        if args.command == 'mcp':
            try:
                import office_mcp
            except ModuleNotFoundError as exc:
                if exc.name == 'mcp':
                    raise RuntimeError('Install the MCP extra from this checkout: pip install ./office[mcp]') from exc
                raise
            office_mcp.main()
            return 0
        if args.command == 'inspect':
            result = inspect(args.file)
        elif args.command == 'author':
            with open(args.spec, encoding='utf-8') as f:
                spec = json.load(f)
            result = {'output': author(spec, args.out, args.source)}
        elif args.command == 'render':
            result = {'pdf': render(args.file, args.out_dir)}
        elif args.command == 'preview':
            result = {'pages': preview(args.file, args.out_dir)}
        elif args.command == 'audit':
            result = officecli_backend.audit(args.file)
        elif args.command == 'compare':
            result = review.compare(args.before, args.after)
        elif args.command == 'cells':
            result = spreadsheet.read_range(args.file, args.sheet, args.cell_range)
        elif args.command == 'recalculate':
            result = {'output': spreadsheet.recalculate(args.file, args.out)}
        elif args.command == 'structured':
            if (args.mode == 'outline') != (args.target is None):
                raise ValueError('get and query require a target; outline takes no target')
            result = officecli_backend.read(args.file, args.mode, args.target)
        else:
            with open(args.spec, encoding='utf-8') as f:
                operations = json.load(f)
            result = officecli_backend.batch(operations, args.out, args.source)
        print(json.dumps(result, ensure_ascii=False, default=str))
    except (OSError, ValueError, KeyError, IndexError, TypeError, RuntimeError, subprocess.SubprocessError) as exc:
        print(f'office-use: {exc}', file=sys.stderr)
        return 1
    return 0

if __name__ == '__main__':
    sys.exit(main())

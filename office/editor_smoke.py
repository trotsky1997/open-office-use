"""Verify version-bound single-element edits against real OfficeCLI."""
import tempfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from pptx import Presentation

import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    for suffix in ('.docx', '.xlsx', '.pptx'):
        source = root / ('source' + suffix); draft = root / suffix[1:]
        if suffix == '.docx':
            obj = Document(); obj.add_paragraph('Original'); obj.save(source)
            path, mode, value = '/body/p[1]', 'text', 'Changed'
        elif suffix == '.xlsx':
            obj = Workbook(); obj.active['A1'] = '=1+2'; obj.save(source)
            path, mode, value = '/Sheet/A1', 'formula', '=2+3'
        else:
            obj = Presentation(); slide = obj.slides.add_slide(obj.slide_layouts[1]); slide.shapes.title.text = 'Original'; obj.save(source)
            path, mode, value = '/slide[1]/shape[1]', 'text', 'Changed'
        original = source.read_bytes()
        session.start(draft, source=source)
        target = session.edit_target(draft, 0, path)
        assert target['revision'] == 0 and target['mode'] == mode
        if suffix == '.xlsx':
            assert target['value'] == '=1+2'
        else:
            assert target['value'] == 'Original'
        try: session.edit(draft, 0, target['path'], mode, value, expected_sha256='0'*64)
        except ValueError as exc: assert 'snapshot changed' in str(exc)
        else: raise AssertionError('Wrong target snapshot accepted')
        edited = session.edit(draft, 0, target['path'], mode, value, expected_sha256=target['snapshot_sha256'])
        assert edited['revision'] == 1 and source.read_bytes() == original
        after = session.edit_target(draft, 1, target['path'])
        assert after['value'] == value, after
        if suffix == '.xlsx':
            assert load_workbook(edited['file']).active['A1'].value == '=2+3'
        try:
            session.edit(draft, 0, path, mode, value)
            raise AssertionError('stale edit succeeded')
        except ValueError as error:
            assert 'conflict' in str(error)
        try:
            session.edit(draft, 1, path, 'invalid', value)
            raise AssertionError('invalid mode succeeded')
        except ValueError:
            pass
        if suffix == '.xlsx':
            try:
                session.edit(draft, 1, path, 'value', '=7+8')
                raise AssertionError('formula in value mode succeeded')
            except ValueError as error:
                assert 'formula mode' in str(error)
        assert session.info(draft)['latest_revision'] == 1
    race_source = root/'race.docx'; doc = Document(); doc.add_paragraph('Original'); doc.save(race_source)
    other = root/'other.docx'; doc = Document(); doc.add_paragraph('Original extra'); doc.save(other)
    race_draft = root/'race-draft'; session.start(race_draft, source=race_source)
    original_target = session.edit_target
    def swap_after_target(*args):
        target = original_target(*args)
        (race_draft/'revision-000000.docx').write_bytes(other.read_bytes())
        return target
    with patch.object(session, 'edit_target', side_effect=swap_after_target):
        try: session.edit(race_draft, 0, '/body/p[1]', 'text', 'Changed')
        except ValueError as exc: assert 'snapshot changed' in str(exc)
        else: raise AssertionError('Changed target source accepted')
    assert session.info(race_draft)['latest_revision'] == 0
    stale_hash = session.edit_target(race_draft, 0, '/body/p[1]')['snapshot_sha256']
    (race_draft/'revision-000000.docx').write_bytes(race_source.read_bytes())
    try: session.replace_text(race_draft, 0, '/body/p[1]', 'Original', 'Changed', expected_sha256=stale_hash)
    except ValueError as exc: assert 'snapshot changed' in str(exc)
    else: raise AssertionError('Changed replacement source accepted')
    assert session.info(race_draft)['latest_revision'] == 0
    for operation in ('track_replace', 'track_insert', 'comment'):
        draft = root/('review-' + operation); session.start(draft, source=race_source)
        target = session.edit_target(draft, 0, '/body/p[1]')
        def execute(expected):
            if operation == 'track_replace': return session.track_replace_text(draft, 0, target['path'], 'Original', 'Revised', 'Reviewer', expected_sha256=expected)
            if operation == 'track_insert': return session.track_insert_text(draft, 0, target['path'], 'Original', 'New ', 'before', 'Reviewer', expected_sha256=expected)
            return session.add_comment(draft, 0, target['path'], 'Review', 'Reviewer', 'Original', expected_sha256=expected)
        try: execute('0'*64)
        except ValueError as exc: assert 'snapshot changed' in str(exc)
        else: raise AssertionError(f'Wrong snapshot accepted for {operation}')
        assert session.info(draft)['latest_revision'] == 0
        result = execute(target['snapshot_sha256']); assert result['revision'] == 1
        if operation == 'comment': assert session.list_comments(draft, 1)['total_comments'] == 1
        else: assert session.tracked_changes(draft, 1)['records']
    field_source = root/'field.docx'; doc = Document()
    first = doc.add_paragraph(); marker = OxmlElement('w:fldChar'); marker.set(qn('w:fldCharType'), 'begin'); first.add_run()._r.append(marker)
    doc.add_paragraph('Field result')
    last = doc.add_paragraph(); marker = OxmlElement('w:fldChar'); marker.set(qn('w:fldCharType'), 'end'); last.add_run()._r.append(marker)
    doc.add_paragraph('Outside field'); doc.save(field_source)
    field_draft = root/'field-draft'; session.start(field_draft, source=field_source)
    target = session.edit_target(field_draft, 0, '/body/p[2]')
    assert target['read_only_reason'] == 'word_field' and not target['replace_text_path_supported'], target
    assert session.edit_target(field_draft, 0, '/body/p[2]/r[1]')['read_only_reason'] == 'word_field'
    for operation in (lambda: session.edit(field_draft, 0, '/body/p[2]', 'text', 'Corrupt'),
                      lambda: session.edit(field_draft, 0, '/body/p[2]/r[1]', 'text', 'Corrupt'),
                      lambda: session.replace_text(field_draft, 0, '/body/p[2]', 'Field', 'Corrupt'),
                      lambda: session.track_replace_text(field_draft, 0, '/body/p[2]', 'Field', 'Corrupt', 'Reviewer'),
                      lambda: session.add_comment(field_draft, 0, '/body/p[2]', 'Comment', 'Reviewer')):
        try: operation()
        except ValueError as exc: assert 'field' in str(exc).lower(), exc
        else: raise AssertionError('A field result was editable')
    assert session.info(field_draft)['latest_revision'] == 0
    assert session.edit_target(field_draft, 0, '/body/p[4]')['read_only_reason'] is None
    assert session.replace_text(field_draft, 0, '/body/p[4]', 'Outside', 'After')['revision'] == 1
    literal_source = root / 'literal.xlsx'; book = Workbook(); book.active['A1'] = '00123'; book.save(literal_source)
    literal_draft = root / 'literal'; session.start(literal_draft, source=literal_source)
    assert session.edit_target(literal_draft, 0, '/Sheet/A1')['mode'] == 'literal'
    literal = session.edit(literal_draft, 0, '/Sheet/A1', 'literal', '=ordinary text')
    saved = load_workbook(literal['file']).active['A1']; assert saved.value == '=ordinary text' and saved.data_type == 's'
    assert session.edit_target(literal_draft, 1, '/Sheet/A1')['mode'] == 'literal'
    cleared = session.edit(literal_draft, 1, '/Sheet/A1', 'literal', '')
    assert load_workbook(cleared['file']).active['A1'].value is None
print('Editor smoke passed: docx, xlsx formula, pptx, stale and invalid edits')

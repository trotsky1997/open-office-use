"""Verify bounded speaker-notes comparisons across decks and revisions."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path
from pptx import Presentation
import review
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'deck.pptx'; draft = root / 'draft'
    deck = Presentation(); first = deck.slides.add_slide(deck.slide_layouts[1]); second = deck.slides.add_slide(deck.slide_layouts[1])
    first.notes_slide.notes_text_frame.text = 'Opening\nMiddle'; second.notes_slide.notes_text_frame.text = 'Opening\nMiddle'; deck.save(source)
    session.start(draft, source=source); session.replace_notes(draft, 0, 1, 'Opening\nEnding')
    result = session.compare_notes(draft, 0, 1, 2, 1)
    assert result['text_changed'] and '-Middle' in result['diff'] and '+Ending' in result['diff'] and not result['truncated'], result
    assert session.compare_notes(draft, 0, 0, 1, 2)['text_changed'] is False
    cli = Path(__file__).with_name('open_office_use.py')
    call = subprocess.run([sys.executable, str(cli), 'session', 'compare-notes', str(draft), '--before', '0', '--after', '1', '--before-slide', '2', '--after-slide', '1'], capture_output=True, text=True, check=True)
    assert json.loads(call.stdout)['text_changed'] is True
    long_a = root / 'long-a.pptx'; long_b = root / 'long-b.pptx'
    first.notes_slide.notes_text_frame.text = 'a' * 100000 + 'X'; deck.save(long_a)
    first.notes_slide.notes_text_frame.text = 'a' * 100000 + 'Y'; deck.save(long_b)
    clipped = review.compare_notes(long_a, long_b, 1, 1)
    assert clipped['text_changed'] is None and clipped['truncated'] is True, clipped
    blank = root / 'blank.pptx'; empty_deck = Presentation(); empty_deck.slides.add_slide(empty_deck.slide_layouts[1]); empty_deck.save(blank)
    added = review.compare_notes(blank, source, 1, 1)
    assert not added['before_has_notes'] and added['after_has_notes'] and added['text_changed']
print('Speaker notes comparison passed: explicit slide mapping, history and truncation')

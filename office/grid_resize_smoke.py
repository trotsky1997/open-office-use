"""Exercise bounded versioned row and column sizing."""
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from lxml import etree
import session
import review

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.xlsx'; book = Workbook(); sheet = book.active
    sheet['A1'] = 'Height'; sheet['D2'] = '=2+2'; sheet['H2'] = 'Other column'
    sheet.column_dimensions.group('C', 'F', hidden=True); sheet.column_dimensions['C'].width = 20
    sheet.row_dimensions[2].height = 40; book.save(source); original = source.read_bytes()
    draft = root / 'draft'; session.start(draft, source=source)
    initial = session.read_grid(draft, 0, 'Sheet', 'C1:F3')['dimensions']
    assert initial['columns']['D']['width'] == '20' and initial['columns']['D']['min'] == '3' and initial['columns']['D']['max'] == '6' and initial['columns']['D']['hidden'] == '1' and initial['rows']['2']['ht'] == '40', initial
    result = session.resize_grid(draft, 0, 'Sheet', 'C1:F3', 'D2', 12.5, 27)
    workbook = load_workbook(result['file']); sheet = workbook.active
    assert sheet['D2'].value == '=2+2' and sheet['H2'].value == 'Other column' and sheet.row_dimensions[2].height == 27
    comparison = session.compare_cells(draft, 0, 1, 'Sheet', 'C1:F3')
    dimension_diff = comparison['dimensions']
    assert result['formula_cache_effect'] == 'may_change' and comparison['changed_count'] == 1 and comparison['changes'][0]['changed_fields'] == ['value'], comparison
    assert dimension_diff['complete'] and dimension_diff['changed_count'] == 2, comparison
    assert {(item['kind'], item['target']) for item in dimension_diff['changes']} == {('column', 'D'), ('row', '2')}
    assert 'width' in next(item for item in dimension_diff['changes'] if item['kind'] == 'column')['changed_fields']
    dims = session.read_grid(draft, 1, 'Sheet', 'C1:F3')['dimensions']
    assert dims['columns']['C']['width'] == '20' and dims['columns']['D']['width'] == '12.5' and dims['columns']['E']['width'] == '20', dims
    assert all(dims['columns'][c]['hidden'] == '1' for c in 'CDEF') and dims['rows']['2']['ht'] == '27' and source.read_bytes() == original
    for width, height, target in [(None, None, 'D2'), (256, None, 'D2'), (0, None, 'D2'), (None, 410, 'D2'), (True, None, 'D2'), (12, None, 'B2')]:
        try: session.resize_grid(draft, 1, 'Sheet', 'C1:F3', target, width, height)
        except ValueError: pass
        else: raise AssertionError('Invalid dimensions accepted')
    assert session.info(draft)['latest_revision'] == 1
    cli = Path(__file__).with_name('open_office_use.py')
    command = subprocess.run([sys.executable, str(cli), 'session', 'resize-grid', str(draft), '--base', '1', '--sheet', 'Sheet', '--range', 'C1:F3', '--target-range', 'E3', '--column-width', '18'], capture_output=True, text=True, check=True)
    assert json.loads(command.stdout)['revision'] == 2
    cli_diff = subprocess.run([sys.executable, str(cli), 'session', 'compare-cells', str(draft), '--before', '0', '--after', '2', '--sheet', 'Sheet', '--range', 'C1:F3'], capture_output=True, text=True, check=True)
    assert json.loads(cli_diff.stdout)['dimensions']['changed_count'] == 3
    assert session.read_grid(draft, 2, 'Sheet', 'E3')['dimensions']['columns']['E']['width'] == '18'
    try: session.resize_grid(draft, 0, 'Sheet', 'C1:F3', 'D2', 15, None)
    except ValueError: pass
    else: raise AssertionError('Stale sizing accepted')
    signed = root / 'signed.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(signed, 'w') as after:
        for entry in before.infolist(): after.writestr(entry, before.read(entry))
        after.writestr('_xmlsignatures/signature.xml', '<signature/>')
    signed_draft = root / 'signed-draft'; session.start(signed_draft, source=signed)
    try: session.resize_grid(signed_draft, 0, 'Sheet', 'D2', 'D2', 15, None)
    except ValueError: pass
    else: raise AssertionError('Signed workbook sizing accepted')
    assert session.info(signed_draft)['latest_revision'] == 0
    protected = root / 'protected.xlsx'; protected_book = Workbook(); protected_book.active.protection.sheet = True; protected_book.save(protected)
    protected_draft = root / 'protected-draft'; session.start(protected_draft, source=protected)
    try: session.resize_grid(protected_draft, 0, 'Sheet', 'A1', 'A1', 15, None)
    except ValueError: pass
    else: raise AssertionError('Protected worksheet sizing accepted')
    overlap = root / 'overlap.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(overlap, 'w') as after:
        for info in before.infolist():
            data = before.read(info)
            if info.filename == 'xl/worksheets/sheet1.xml':
                tree = etree.fromstring(data); ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'; cols = tree.find(ns + 'cols'); etree.SubElement(cols, ns + 'col', min='4', max='4', width='9', customWidth='1'); data = etree.tostring(tree)
            after.writestr(info, data)
    overlap_draft = root / 'overlap-draft'; session.start(overlap_draft, source=overlap)
    assert 'overlapping_column_ranges' in session.read_grid(overlap_draft, 0, 'Sheet', 'D2')['dimensions']['issues']
    assert not review.compare_cells(source, overlap, 'Sheet', 'D2')['dimensions']['complete']
    try: session.resize_grid(overlap_draft, 0, 'Sheet', 'D2', 'D2', 18, None)
    except ValueError: pass
    else: raise AssertionError('Ambiguous column range accepted')
    assert session.info(overlap_draft)['latest_revision'] == 0
print('Grid sizing smoke passed')

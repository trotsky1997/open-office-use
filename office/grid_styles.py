"""Apply OfficeCLI-produced styles while retaining source worksheet content."""
import copy
import posixpath
import re
import tempfile
import zipfile
from pathlib import Path
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
import officecli_backend
import ooxml
import snapshot as snapshots

X = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REL = '{http://schemas.openxmlformats.org/package/2006/relationships}'
STYLE_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles'


def sheet_part(file, sheet):
    workbook = load_workbook(file, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet not in workbook.sheetnames: raise ValueError('Worksheet not found')
        path = getattr(workbook[sheet], '_worksheet_path', None)
        if not isinstance(path, str) or sum(getattr(item, '_worksheet_path', None) == path for item in workbook.worksheets) != 1:
            raise ValueError('Worksheet part is unavailable or shared')
        return path
    finally: workbook.close()


def style_part(archive):
    tree = ooxml.parse_xml(archive.read('xl/_rels/workbook.xml.rels'))
    if tree.getroot().tag != REL + 'Relationships': raise ValueError('Invalid workbook relationships')
    relationships = tree.getroot().findall(REL + 'Relationship')
    ids = [rel.get('Id') for rel in relationships]
    if any(not value for value in ids) or len(ids) != len(set(ids)): raise ValueError('Invalid or duplicate workbook relationship ID')
    rels = [r for r in relationships if r.get('Type') == STYLE_REL]
    if len(rels) != 1 or rels[0].get('TargetMode', 'Internal') != 'Internal': raise ValueError('Styles relationship is missing or ambiguous')
    target = rels[0].get('Target', '')
    if not target or any(c in target for c in ('\\', '%', '?', '#')): raise ValueError('Unsupported styles relationship')
    path = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join('xl', target))
    if path != 'xl/styles.xml' or path not in archive.namelist(): raise ValueError('Unsupported styles part location')
    return path


def equivalent(node):
    text = node.text or ''
    if len(node) and not text.strip(): text = ''
    return node.tag, tuple(sorted(node.attrib.items())), text, tuple(equivalent(child) for child in node)


def merge_styles(original, candidate):
    before = ooxml.parse_xml(original); after = ooxml.parse_xml(candidate)
    a, b = before.getroot(), after.getroot()
    if a.tag != X + 'styleSheet' or b.tag != X + 'styleSheet': raise ValueError('Invalid stylesheet')
    if len({node.tag for node in a}) != len(a) or len({node.tag for node in b}) != len(b): raise ValueError('Duplicate stylesheet collections')
    appendable = {X + name for name in ('numFmts', 'fonts', 'fills', 'borders', 'cellXfs')}
    allowed_order = [X + name for name in ('numFmts', 'fonts', 'fills', 'borders', 'cellStyleXfs', 'cellXfs', 'cellStyles', 'dxfs', 'tableStyles', 'colors', 'extLst')]
    for node in a:
        other = b.find(node.tag)
        if other is None or len(other) < len(node): raise ValueError('OfficeCLI changed existing styles')
        if any(equivalent(old) != equivalent(new) for old, new in zip(node, other)): raise ValueError('OfficeCLI changed existing styles')
        if node.tag not in appendable and equivalent(node) != equivalent(other): raise ValueError('OfficeCLI changed unsupported stylesheet metadata')
    for node in b:
        old = a.find(node.tag)
        if old is None and node.tag not in appendable: raise ValueError('OfficeCLI added unsupported stylesheet metadata')
        if node.tag in appendable:
            if node.tag not in allowed_order: raise ValueError('Unsupported stylesheet collection')
            added = list(node)[len(old) if old is not None else 0:]
            for item in added:
                for nested in item.iter():
                    if not isinstance(nested.tag, str) or not nested.tag.startswith(X) or any(key.startswith('{') and not key.startswith('{http://www.w3.org/XML/1998/namespace}') for key in nested.attrib):
                        raise ValueError('New style contains unsupported extension markup')
            if old is None:
                old = etree.Element(node.tag)
                following = next((existing for existing in a if existing.tag in allowed_order and allowed_order.index(existing.tag) > allowed_order.index(node.tag)), None)
                if following is None: a.append(old)
                else: a.insert(a.index(following), old)
            for item in added: old.append(copy.deepcopy(item))
            old.set('count', str(len(old)))
    xfs = a.find(X + 'cellXfs'); fonts = a.find(X + 'fonts'); fills = a.find(X + 'fills'); borders = a.find(X + 'borders'); base = a.find(X + 'cellStyleXfs')
    if any(value is None or len(value) == 0 for value in (xfs, fonts, fills, borders, base)): raise ValueError('Stylesheet is missing required collections')
    for xf in xfs:
        for attr, collection in (('fontId', fonts), ('fillId', fills), ('borderId', borders), ('xfId', base)):
            value = xf.get(attr, '0')
            if not re.fullmatch(r'[0-9]{1,10}', value) or int(value) >= len(collection): raise ValueError('Invalid stylesheet reference')
    formats = a.find(X + 'numFmts')
    ids = [node.get('numFmtId', '') for node in formats] if formats is not None else []
    if any(not re.fullmatch(r'[0-9]{1,10}', value) for value in ids) or len(ids) != len({int(value) for value in ids}):
        raise ValueError('Invalid or duplicate number format ID')
    custom_formats = {int(value) for value in ids}
    for xf in xfs:
        value = xf.get('numFmtId', '0')
        if not re.fullmatch(r'[0-9]{1,10}', value) or (int(value) >= 164 and int(value) not in custom_formats):
            raise ValueError('Invalid number format reference')
    return ooxml.serialize_xml(before, original), len(xfs)


def prepare_cells(root, addresses):
    targets = set(addresses); found = {}
    data = root.findall(X + 'sheetData')
    if len(data) != 1 or len(list(root.iter(X + 'sheetData'))) != 1: raise ValueError('Ambiguous worksheet data')
    rows = {}; previous_row = 0
    for row in data[0]:
        if row.tag != X + 'row' or not re.fullmatch(r'[0-9]{1,7}', row.get('r', '')): raise ValueError('Unsupported worksheet row')
        index = int(row.get('r'))
        if not 1 <= index <= 1048576 or index <= previous_row: raise ValueError('Duplicate or unordered rows')
        previous_row = index; rows[index] = row; previous_col = 0
        for cell in row:
            if cell.tag != X + 'c': raise ValueError('Unsupported row content')
            match = re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]{0,6})', cell.get('r', ''))
            if not match or int(match.group(2)) != index: raise ValueError('Invalid cell reference')
            column = column_index_from_string(match.group(1))
            if column > 16384 or column <= previous_col: raise ValueError('Duplicate or unordered cells')
            previous_col = column
            if cell.get('r') in targets: found[cell.get('r')] = cell
    for address in addresses:
        match = re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]{0,6})', address)
        if not match: raise ValueError('Invalid target cell')
        col, row_number = column_index_from_string(match.group(1)), int(match.group(2))
        row = rows.get(row_number)
        if row is None:
            row = etree.Element(X + 'row', r=str(row_number)); rows[row_number] = row
        if not any(cell.get('r') == address for cell in row):
            cell = etree.Element(X + 'c', r=address); found[address] = cell
            row.append(cell); row[:] = sorted(row, key=lambda item: column_index_from_string(re.match(r'[A-Z]+', item.get('r')).group()))
    data[0][:] = [rows[number] for number in sorted(rows)]
    dimensions = root.findall(X + 'dimension')
    if len(dimensions) > 1 or len(list(root.iter(X + 'dimension'))) != len(dimensions): raise ValueError('Ambiguous worksheet dimension')
    dimension = dimensions[0] if dimensions else None
    if dimension is not None:
        try:
            c1, r1, c2, r2 = range_boundaries(dimension.get('ref'))
            if any(value is None for value in (c1, r1, c2, r2)) or not 1 <= c1 <= c2 <= 16384 or not 1 <= r1 <= r2 <= 1048576: raise ValueError
        except (ValueError, TypeError): raise ValueError('Invalid worksheet dimension') from None
        for address in addresses:
            col, row = range_boundaries(address)[:2]; c1 = min(c1, col); c2 = max(c2, col); r1 = min(r1, row); r2 = max(r2, row)
        dimension.set('ref', f'{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}')
    return found


def apply(source, output, sheet, addresses, operations, expected_sha256=None):
    with tempfile.TemporaryDirectory(prefix='office-grid-source-') as tmp:
        frozen, digest = snapshots.copy(source, tmp)
        if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Workbook snapshot changed; reload the grid')
        return _apply(frozen, output, sheet, addresses, operations)


def _apply(source, output, sheet, addresses, operations):
    source = Path(source); output = Path(output)
    with zipfile.ZipFile(source) as archive:
        if any(name.startswith('_xmlsignatures/') for name in archive.namelist()): raise ValueError('Digitally signed workbooks are not supported for formatting')
        style_name = style_part(archive)
    with tempfile.TemporaryDirectory(prefix='office-grid-style-') as tmp:
        candidate = Path(tmp) / 'candidate.xlsx'; batch = officecli_backend.batch(operations, candidate, source)
        source_part = sheet_part(source, sheet)
        with zipfile.ZipFile(candidate) as archive:
            entries = archive.infolist(); names = [entry.filename for entry in entries]
            if len(entries) > 5000 or len(names) != len(set(names)) or sum(entry.file_size for entry in entries) > 128 * 1024 * 1024:
                raise ValueError('OfficeCLI candidate exceeds package limits or contains duplicate parts')
            candidate_part = sheet_part(candidate, sheet)
            if style_part(archive) != style_name: raise ValueError('OfficeCLI changed styles relationship')
            candidate_styles = archive.read(style_name)
            candidate_sheet = ooxml.parse_xml(archive.read(candidate_part))
        targets = set(addresses); styles = {}
        candidate_root = candidate_sheet.getroot()
        if candidate_root.tag != X + 'worksheet': raise ValueError('Invalid candidate worksheet')
        candidate_data = candidate_root.findall(X + 'sheetData')
        if len(candidate_data) != 1 or len(list(candidate_root.iter(X + 'sheetData'))) != 1: raise ValueError('Ambiguous candidate worksheet data')
        for row in candidate_data[0].findall(X + 'row'):
            for cell in row.findall(X + 'c'):
                address = cell.get('r')
                if address in targets:
                    if address in styles: raise ValueError('Candidate contains duplicate target cells')
                    match = re.fullmatch(r'[A-Z]{1,3}([1-9][0-9]{0,6})', address)
                    if not match or row.get('r') != match.group(1): raise ValueError('Invalid candidate cell row')
                    styles[address] = cell.get('s', '0')
        if set(styles) != targets: raise ValueError('OfficeCLI did not create every target cell')
        def edit(tree, archive):
            if tree.getroot().tag != X + 'worksheet': raise ValueError('Invalid worksheet')
            merged_styles, count = merge_styles(archive.read(style_name), candidate_styles)
            if any(not re.fullmatch(r'[0-9]{1,10}', value) or int(value) >= count for value in styles.values()): raise ValueError('Invalid target style index')
            targets_by_address = prepare_cells(tree.getroot(), addresses)
            for address, cell in targets_by_address.items(): cell.set('s', styles[address])
            return {style_name: merged_styles}
        ooxml.edit_xml(source, output, lambda _: source_part, edit)
        return {'output': str(output), 'result': batch['result'], 'formula_cache_effect': 'preserved'}

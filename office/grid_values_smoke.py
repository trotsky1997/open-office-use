"""Exercise grid value edits with preserved unrelated formula caches."""
import concurrent.futures
import tempfile
import json
import subprocess
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import xlsxwriter
import officecli_backend
import grid_values
import ooxml
import session
import spreadsheet

X = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.xlsx'; book = Workbook(); sheet = book.active
    sheet['A1'] = 5; sheet['A1'].font = Font(bold=True); sheet['A1'].number_format = '0.00'; sheet['B1'] = '=A1*2'; sheet['C1'] = '=3+3'; book.create_sheet('Other')['A1'] = '=10+10'; book.save(source)
    cached = root / 'cached.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(cached, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry)
            if entry.filename in ('xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml'):
                tree = etree.fromstring(data)
                for cell in tree.iter(X + 'c'):
                    if cell.get('r') in ('B1', 'C1', 'A1') and cell.find(X + 'f') is not None:
                        cell.find(X + 'v').text = {'B1': '10', 'C1': '6', 'A1': '20'}[cell.get('r')]
                data = etree.tostring(tree)
            after.writestr(entry, data)
    cached.replace(source); original = source.read_bytes(); draft = root / 'draft'; session.start(draft, source=source)
    result = session.edit_grid(draft, 0, 'Sheet', 'A1:D2', [{'cell': 'A1', 'value': '8'}, {'cell': 'C1', 'value': '=A1+1'}, {'cell': 'D2', 'value': '2026-10-30'}])
    assert result['formula_cache_effect'] == 'may_be_stale' and result['recalculation_requested']
    values = load_workbook(result['file'], data_only=True); formulas = load_workbook(result['file'])
    assert values.active['A1'].value == 8 and values.active['B1'].value == 10 and values.active['C1'].value is None and values['Other']['A1'].value == 20
    assert formulas.active['C1'].value == '=A1+1' and formulas.active['D2'].value == datetime(2026, 10, 30)
    assert formulas.active['A1'].font.bold and formulas.active['A1'].number_format == '0.00'
    scan = spreadsheet.scan_formulas(result['file']); assert scan['recalculation_requests'] == 1 and any(item['cell'] == 'C1' for item in scan['empty_or_missing'])
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(result['file']) as after:
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name not in ('xl/worksheets/sheet1.xml', 'xl/styles.xml'))
    audit = officecli_backend.audit(result['file']); assert audit['validation']['exit_code'] == 0 and audit['validation']['result']['success'], audit
    recalc = session.recalculate(draft, 1); recalculated = load_workbook(recalc['file'], data_only=True)
    assert recalculated.active['B1'].value == 16 and recalculated.active['C1'].value == 9
    assert source.read_bytes() == original
    literal_draft = root / 'literal-draft'; session.start(literal_draft, source=source)
    literal = session.edit_grid(literal_draft, 0, 'Sheet', 'A1:D2', [
        {'cell': 'A1', 'value': '00123', 'type': 'text'},
        {'cell': 'C1', 'value': '=not a formula', 'type': 'text'},
        {'cell': 'D2', 'value': ' line 1\nline 2 ', 'type': 'text'},
    ])
    assert 'office-use-literal-text' not in json.dumps(literal['result'])
    book = load_workbook(literal['file']); assert book.active['A1'].value == '00123' and book.active['A1'].data_type == 's'
    assert book.active['C1'].value == '=not a formula' and book.active['C1'].data_type == 's'
    assert book.active['D2'].value == ' line 1\nline 2 ' and book.active['A1'].font.bold
    assert load_workbook(literal['file'], data_only=True).active['B1'].value == 10
    for invalid in ('', 'bad\x01text', '_x0041_'):
        try: session.edit_grid(literal_draft, 1, 'Sheet', 'A1', [{'cell': 'A1', 'value': invalid, 'type': 'text'}])
        except ValueError: pass
        else: raise AssertionError('Invalid explicit text accepted')
    assert session.info(literal_draft)['latest_revision'] == 1
    shared_source = root / 'shared.xlsx'; wb = xlsxwriter.Workbook(str(shared_source)); ws = wb.add_worksheet('Sheet'); ws.write('A1', 'same'); ws.write('A2', 'same'); ws.write('B1', 5); wb.close()
    shared_draft = root / 'shared-draft'; session.start(shared_draft, source=shared_source)
    changed = session.edit_grid(shared_draft, 0, 'Sheet', 'A1:B2', [{'cell': 'A1', 'value': 'new text'}, {'cell': 'B1', 'value': 'hello'}])
    result_book = load_workbook(changed['file']); assert result_book.active['A1'].value == 'new text' and result_book.active['A2'].value == 'same' and result_book.active['B1'].value == 'hello'
    with zipfile.ZipFile(shared_source) as before, zipfile.ZipFile(changed['file']) as after:
        old = etree.fromstring(before.read('xl/sharedStrings.xml')); new = etree.fromstring(after.read('xl/sharedStrings.xml'))
        assert old.get('count') is not None and new.get('count') is None and new.get('uniqueCount') is None
        assert [etree.tostring(item) for item in old] == [etree.tostring(item) for item in new]
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name not in ('xl/worksheets/sheet1.xml', 'xl/sharedStrings.xml', 'xl/styles.xml'))
    same_count = root / 'same-count'; session.start(same_count, source=shared_source)
    one_text = session.edit_grid(same_count, 0, 'Sheet', 'A1', [{'cell': 'A1', 'value': 'different'}])
    with zipfile.ZipFile(shared_source) as before, zipfile.ZipFile(one_text['file']) as after:
        assert before.read('xl/sharedStrings.xml') == after.read('xl/sharedStrings.xml')
    shared_audit = officecli_backend.audit(changed['file']); assert shared_audit['validation']['exit_code'] == 0 and shared_audit['validation']['result']['success'], shared_audit
    cli_draft = root / 'cli'; session.start(cli_draft, source=source)
    cli = Path(__file__).with_name('open_office_use.py')
    command = subprocess.run([sys.executable, str(cli), 'session', 'edit-grid', str(cli_draft), '--base', '0', '--sheet', 'Sheet', '--range', 'A1', '--changes', json.dumps([{'cell': 'A1', 'value': '17'}])], capture_output=True, text=True, check=True)
    assert json.loads(command.stdout)['recalculation_requested'] and load_workbook(cli_draft / 'revision-000001.xlsx', data_only=True).active['B1'].value == 10
    command = subprocess.run([sys.executable, str(cli), 'session', 'edit-grid', str(cli_draft), '--base', '1', '--sheet', 'Sheet', '--range', 'A1', '--changes', json.dumps([{'cell': 'A1', 'value': '=literal', 'type': 'text'}])], capture_output=True, text=True, check=True)
    assert json.loads(command.stdout)['revision'] == 2 and load_workbook(cli_draft / 'revision-000002.xlsx').active['A1'].value == '=literal' and load_workbook(cli_draft / 'revision-000002.xlsx').active['A1'].data_type == 's'
    chain_source = root / 'chain.xlsx'; rel = '{http://schemas.openxmlformats.org/package/2006/relationships}'; ct = '{http://schemas.openxmlformats.org/package/2006/content-types}'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(chain_source, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry)
            if entry.filename == 'xl/_rels/workbook.xml.rels':
                tree = etree.fromstring(data); etree.SubElement(tree, rel + 'Relationship', Id='rIdChain', Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain', Target='calcChain.xml'); data = etree.tostring(tree)
            if entry.filename == '[Content_Types].xml':
                tree = etree.fromstring(data); etree.SubElement(tree, ct + 'Override', PartName='/xl/calcChain.xml', ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml'); data = etree.tostring(tree)
            after.writestr(entry, data)
        chain = etree.Element(X + 'calcChain'); etree.SubElement(chain, X + 'c', r='C1', i='1'); after.writestr('xl/calcChain.xml', etree.tostring(chain))
    chain_draft = root / 'chain-draft'; session.start(chain_draft, source=chain_source)
    chain_edit = session.edit_grid(chain_draft, 0, 'Sheet', 'C1', [{'cell': 'C1', 'value': '=A1+3'}])
    assert chain_edit['calculation_chain'] == 'removed'
    with zipfile.ZipFile(chain_source) as before, zipfile.ZipFile(chain_edit['file']) as after:
        assert 'xl/calcChain.xml' not in after.namelist()
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name not in ('xl/calcChain.xml', 'xl/_rels/workbook.xml.rels', '[Content_Types].xml', 'xl/worksheets/sheet1.xml', 'xl/styles.xml'))
        assert b'calcChain' not in after.read('xl/_rels/workbook.xml.rels') and b'calcChain' not in after.read('[Content_Types].xml')
    assert load_workbook(chain_edit['file'], data_only=True).active['B1'].value == 10
    assert load_workbook(chain_edit['file'], data_only=True).active['C1'].value is None
    audit = officecli_backend.audit(chain_edit['file']); assert audit['validation']['exit_code'] == 0 and audit['validation']['result']['success'], audit
    recalculated = session.recalculate(chain_draft, 1); assert load_workbook(recalculated['file'], data_only=True).active['C1'].value == 8
    unchanged_draft = root / 'chain-unchanged'; session.start(unchanged_draft, source=chain_source)
    unchanged = session.edit_grid(unchanged_draft, 0, 'Sheet', 'A1', [{'cell': 'A1', 'value': '8'}])
    assert unchanged['calculation_chain'] == 'unchanged'
    with zipfile.ZipFile(chain_source) as before, zipfile.ZipFile(unchanged['file']) as after: assert before.read('xl/calcChain.xml') == after.read('xl/calcChain.xml')
    unsafe = root / 'chain-referenced.xlsx'
    with zipfile.ZipFile(chain_source) as before, zipfile.ZipFile(unsafe, 'w') as after:
        for entry in before.infolist(): after.writestr(entry, before.read(entry))
        extra = etree.Element(rel + 'Relationships'); etree.SubElement(extra, rel + 'Relationship', Id='rId1', Type='urn:other', Target='../calcChain.xml')
        after.writestr('xl/worksheets/_rels/sheet1.xml.rels', etree.tostring(extra))
    with zipfile.ZipFile(unsafe) as archive:
        try: grid_values.remove_calculation_chain(archive, ooxml.parse_xml(archive.read('[Content_Types].xml')))
        except ValueError as error: assert 'additional references' in str(error)
        else: raise AssertionError('Referenced calculation chain removed')
    signed = root / 'signed.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(signed, 'w') as after:
        for entry in before.infolist(): after.writestr(entry, before.read(entry))
        after.writestr('_xmlsignatures/signature.xml', '<signature/>')
    signed_draft = root / 'signed-draft'; session.start(signed_draft, source=signed)
    try: session.edit_grid(signed_draft, 0, 'Sheet', 'A1', [{'cell': 'A1', 'value': '7'}])
    except ValueError: pass
    else: raise AssertionError('Signed grid edit accepted')
    assert session.info(signed_draft)['latest_revision'] == 0
    race = root / 'race'; session.start(race, source=source)
    def compete(value):
        try: return session.edit_grid(race, 0, 'Sheet', 'A1', [{'cell': 'A1', 'value': value}])['revision']
        except ValueError: return None
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool: outcomes = list(pool.map(compete, ['7', '8']))
    assert sorted(outcomes, key=lambda value: value is None) == [1, None]
print('Grid values preservation smoke passed')

"""Exercise Word header/footer inventory and versioned replacement."""
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.shared import Inches
from lxml import etree
import session
import word

with tempfile.TemporaryDirectory() as tmp:
    root=Path(tmp); source=root/'sections.docx'; draft=root/'draft'; doc=Document()
    section=doc.sections[0]; header=section.header; p=header.paragraphs[0]; p.add_run('Project ').bold=True; p.add_run('Alpha').italic=True
    table=header.add_table(rows=1,cols=1,width=Inches(3)); table.cell(0,0).text='Header table'
    section.first_page_header.paragraphs[0].text='First only'; footer=section.footer.paragraphs[0]; footer.add_run('Footer label')
    doc.add_paragraph('Body remains'); doc.add_section(WD_SECTION.NEW_PAGE); doc.save(source)
    listing=word.list_stories(source); assert listing['complete'] and listing['sections']==2
    header_story=next(x for x in listing['stories'] if any(p['text']=='Project Alpha' for p in x['paragraphs']))
    assert header_story['references']==[{'section':1,'type':'default'}]
    assert header_story['linked_sections']==[{'section':1,'type':'default','inherited':False},{'section':2,'type':'default','inherited':True}]
    assert listing['reference_resolution_complete']
    first_story=next(s for s in listing['stories'] if any(p['text']=='First only' for p in s['paragraphs']))
    assert first_story['linked_sections']==[{'section':1,'type':'first','inherited':False},{'section':2,'type':'first','inherited':True}]
    header_path=next(p['path'] for p in header_story['paragraphs'] if p['text']=='Project Alpha')
    table_path=next(p['path'] for p in header_story['paragraphs'] if p['text']=='Header table')
    footer_path=next(p['path'] for s in listing['stories'] for p in s['paragraphs'] if p['text']=='Footer label')
    session.start(draft, source=source); target=session.edit_target(draft,0,header_path)
    assert target['replace_only'] and target['story_impact']['linked_sections']==[{'relationship_id':header_story['id'],**link} for link in header_story['linked_sections']]
    first=session.replace_text(draft,0,header_path,'ect Al','ECT AL')
    changed=Document(first['file']); p=changed.sections[0].header.paragraphs[0]
    assert p.text=='ProjECT ALpha' and p.runs[0].bold and p.runs[1].italic
    assert changed.sections[1].header.paragraphs[0].text=='ProjECT ALpha' and changed.paragraphs[0].text=='Body remains'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(first['file']) as after:
        modified={name for name in before.namelist() if before.read(name)!=after.read(name)}
        assert modified=={header_story['part']}, modified
    session.replace_text(draft,1,table_path,'table','cell'); session.replace_text(draft,2,footer_path,'label','note')
    assert Document(draft/'revision-000003.docx').sections[0].footer.paragraphs[0].text=='Footer note'
    try: session.add_comment(draft,3,header_path,'No','Author')
    except ValueError: pass
    else: raise AssertionError('Header comment unexpectedly accepted')
    try: session.replace_text(draft,1,header_path,'Project','Bad')
    except ValueError as exc: assert 'conflict' in str(exc)
    else: raise AssertionError('Stale edit succeeded')
    cli=Path(__file__).with_name('open_office_use.py'); result=subprocess.run([sys.executable,str(cli),'session','stories',str(draft),'--revision','3'],capture_output=True,text=True,check=True)
    assert json.loads(result.stdout)['revision']==3
    field_doc=Document(); para=field_doc.sections[0].header.paragraphs[0]; para.add_run('label'); run=para.add_run(); fld=OxmlElement('w:fldChar'); fld.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType','begin'); run._r.append(fld)
    field_source=root/'field.docx'; field_doc.save(field_source); field_path=word.list_stories(field_source)['stories'][0]['paragraphs'][0]['path']
    try: word.replace_text(field_source,root/'bad.docx',field_path,'label','other')
    except ValueError: pass
    else: raise AssertionError('Field paragraph accepted')
    assert not (root/'bad.docx').exists()
    alias=root/'alias.docx'; ns={'w':word.W[1:-1], 'r':word.R[1:-1], 'rel':word.REL[1:-1]}
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(alias,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='word/_rels/document.xml.rels':
                tree=etree.fromstring(data); item=tree.xpath(f"./rel:Relationship[@Id='{header_story['id']}']",namespaces=ns)[0]
                duplicate=etree.fromstring(etree.tostring(item)); duplicate.set('Id','rIdAlias'); tree.append(duplicate); data=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
            if entry.filename=='word/document.xml':
                tree=etree.fromstring(data); last=tree.xpath('./w:body/w:sectPr',namespaces=ns)[0]
                ref=etree.Element(word.W+'headerReference'); ref.set(word.W+'type','default'); ref.set(word.R+'id','rIdAlias'); last.insert(0,ref); data=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
            after.writestr(entry,data)
    aliased=word.list_stories(alias); original=next(s for s in aliased['stories'] if s['id']==header_story['id']); other=next(s for s in aliased['stories'] if s['id']=='rIdAlias')
    assert original['same_part_ids']==['rIdAlias'] and other['linked_sections']==[{'section':2,'type':'default','inherited':False}]
    assert original['linked_sections']==[{'section':1,'type':'default','inherited':False}]
    alias_draft=root/'alias-draft'; session.start(alias_draft,source=alias); impact=session.edit_target(alias_draft,0,header_path)['story_impact']
    assert {link['section'] for link in impact['linked_sections']}=={1,2} and impact['same_part_ids']==['rIdAlias']
    ambiguous=root/'wrapped-section.docx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(ambiguous,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='word/document.xml':
                tree=etree.fromstring(data); section=tree.xpath('./w:body/w:p/w:pPr/w:sectPr',namespaces=ns)[0]; parent=section.getparent(); parent.remove(section)
                wrapper=etree.SubElement(tree.find(word.W+'body'),word.W+'customXml'); wrapper.append(section); data=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
            after.writestr(entry,data)
    assert not word.list_stories(ambiguous)['reference_resolution_complete']
    many=root/'many-sections.docx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(many,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='word/document.xml':
                tree=etree.fromstring(data); body=tree.find(word.W+'body'); last=body.find(word.W+'sectPr'); index=body.index(last)
                for _ in range(2000):
                    paragraph=etree.Element(word.W+'p'); properties=etree.SubElement(paragraph,word.W+'pPr'); etree.SubElement(properties,word.W+'sectPr'); body.insert(index,paragraph); index+=1
                data=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
            after.writestr(entry,data)
    limited=word.list_stories(many)
    assert limited['sections']==2002 and not limited['reference_resolution_complete']
print('Word story smoke passed: shared sections, headers, footers, tables and field protection')

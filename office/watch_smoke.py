"""Verify local watch routes, preview updates and access restrictions."""
import json
import signal
import subprocess
import sys
import tempfile
import urllib.error
import urllib.request
from urllib.parse import urlsplit
from pathlib import Path

from docx import Document

cli = Path(__file__).with_name('open_office_use.py')
def run(*args):
    result = subprocess.run([sys.executable, str(cli), *map(str, args)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    return json.loads(result.stdout)

def get(url, headers=None):
    with urllib.request.urlopen(urllib.request.Request(url, headers=headers or {}), timeout=60) as response:
        return response.read()

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.docx'; directory = root / 'session'
    doc = Document(); doc.add_paragraph('First'); doc.save(source)
    run('session', 'start', directory, '--source', source)
    server = subprocess.Popen([sys.executable, str(cli), 'watch', str(directory)], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    try:
        url = json.loads(server.stdout.readline())['url']
        assert b'Office session preview' in get(url)
        first = json.loads(get(url + 'state'))
        assert first['revision'] == 0 and first['pages'] > 0
        assert get(url + 'page/0/0').startswith(b'\x89PNG\r\n\x1a\n')
        origin = f"http://{urlsplit(url).netloc}"
        selection = {'revision': 0, 'page': 1, 'box': [0, 0, 1, 1]}
        request = urllib.request.Request(url + 'selection', data=json.dumps(selection).encode(), headers={'Origin': origin, 'Content-Type': 'application/json'})
        assert json.loads(urllib.request.urlopen(request).read()) == selection
        assert json.loads((directory / 'selection.json').read_text()) == selection
        candidates_request = urllib.request.Request(url + 'candidates', data=json.dumps({'selection': selection}).encode(), headers={'Origin': origin, 'Content-Type': 'application/json'})
        candidates = json.loads(urllib.request.urlopen(candidates_request, timeout=60).read())
        assert candidates['revision'] == 0 and any(item['path'] == '/body/p[1]' for item in candidates['candidates']), candidates
        mismatch = dict(selection, box=[0, 0, 0.5, 0.5])
        wrong_selection = urllib.request.Request(url + 'candidates', data=json.dumps({'selection': mismatch}).encode(), headers={'Origin': origin, 'Content-Type': 'application/json'})
        try:
            urllib.request.urlopen(wrong_selection)
            raise AssertionError('Mismatched selection was accepted')
        except urllib.error.HTTPError as error:
            assert error.code == 400
        forbidden_candidates = urllib.request.Request(url + 'candidates', data=json.dumps({'selection': selection}).encode(), headers={'Origin': 'http://evil.test', 'Content-Type': 'application/json'})
        try:
            urllib.request.urlopen(forbidden_candidates)
            raise AssertionError('Cross-origin search was accepted')
        except urllib.error.HTTPError as error:
            assert error.code == 403
        bad_origin = urllib.request.Request(url + 'selection', data=b'{}', headers={'Origin': 'http://evil.test', 'Content-Type': 'application/json'})
        try:
            urllib.request.urlopen(bad_origin)
            raise AssertionError('Unexpected cross-origin success')
        except urllib.error.HTTPError as error:
            assert error.code == 403
        spec = root / 'ops.json'; spec.write_text(json.dumps([{'command': 'set', 'path': '/body/p[1]', 'props': {'text': 'Changed'}}]))
        run('session', 'apply', directory, '--base', '0', '--spec', spec)
        second = json.loads(get(url + 'state'))
        assert second['revision'] == 1 and second['pages'] > 0
        stale_candidates = urllib.request.Request(url + 'candidates', data=json.dumps({'selection': selection}).encode(), headers={'Origin': origin, 'Content-Type': 'application/json'})
        try:
            urllib.request.urlopen(stale_candidates)
            raise AssertionError('Stale candidate lookup was accepted')
        except urllib.error.HTTPError as error:
            assert error.code == 400
        history = json.loads(get(url + 'revision/0/state'))
        assert history['revision'] == 0 and history['pages'] > 0
        assert get(url + 'page/0/0').startswith(b'\x89PNG\r\n\x1a\n')
        try:
            get(url + 'revision/999/state')
            raise AssertionError('Missing revision was served')
        except urllib.error.HTTPError as error:
            assert error.code == 404
        stale = urllib.request.Request(url + 'selection', data=json.dumps(selection).encode(), headers={'Origin': origin, 'Content-Type': 'application/json'})
        try:
            urllib.request.urlopen(stale)
            raise AssertionError('Unexpected stale selection success')
        except urllib.error.HTTPError as error:
            assert error.code == 400
        assert get(url + 'page/1/0').startswith(b'\x89PNG\r\n\x1a\n')
        restore = urllib.request.Request(url + 'restore', data=json.dumps({'base_revision': 1, 'revision': 0}).encode(), headers={'Origin': origin, 'Content-Type': 'application/json'})
        forbidden_restore = urllib.request.Request(url + 'restore', data=json.dumps({'base_revision': 1, 'revision': 0}).encode(), headers={'Origin': 'http://evil.test', 'Content-Type': 'application/json'})
        try:
            urllib.request.urlopen(forbidden_restore)
            raise AssertionError('Unexpected cross-origin restore')
        except urllib.error.HTTPError as error:
            assert error.code == 403
        assert json.loads(urllib.request.urlopen(restore).read())['revision'] == 2
        third = json.loads(get(url + 'state'))
        assert third['revision'] == 2 and third['revisions'] == [0, 1, 2]
        stale_restore = urllib.request.Request(url + 'restore', data=json.dumps({'base_revision': 1, 'revision': 0}).encode(), headers={'Origin': origin, 'Content-Type': 'application/json'})
        try:
            urllib.request.urlopen(stale_restore)
            raise AssertionError('Unexpected stale restore')
        except urllib.error.HTTPError as error:
            assert error.code == 400
        assert (directory / 'revision-000002.docx').read_bytes() == source.read_bytes()
        element = urllib.request.Request(url + 'element', data=json.dumps({'base_revision': 2, 'path': '/body/p[1]'}).encode(), headers={'Origin': origin, 'Content-Type': 'application/json'})
        target = json.loads(urllib.request.urlopen(element).read())
        assert target['value'] == 'First' and target['revision'] == 2
        edit_data = {'base_revision': 2, 'path': target['path'], 'mode': 'text', 'value': 'User edit'}
        forbidden_edit = urllib.request.Request(url + 'edit', data=json.dumps(edit_data).encode(), headers={'Origin': 'http://evil.test', 'Content-Type': 'application/json'})
        try:
            urllib.request.urlopen(forbidden_edit)
            raise AssertionError('Unexpected cross-origin edit')
        except urllib.error.HTTPError as error:
            assert error.code == 403
        edit = urllib.request.Request(url + 'edit', data=json.dumps(edit_data).encode(), headers={'Origin': origin, 'Content-Type': 'application/json'})
        assert json.loads(urllib.request.urlopen(edit).read())['revision'] == 3
        assert Document(directory / 'revision-000003.docx').paragraphs[0].text == 'User edit'
        stale_edit = urllib.request.Request(url + 'edit', data=json.dumps(edit_data).encode(), headers={'Origin': origin, 'Content-Type': 'application/json'})
        try:
            urllib.request.urlopen(stale_edit)
            raise AssertionError('Unexpected stale edit')
        except urllib.error.HTTPError as error:
            assert error.code == 400


        for address, headers, status in [(url.replace('/state', '') + '../state', {}, 404), (url + 'state', {'Host': 'example.com'}, 403)]:
            try:
                get(address, headers)
                raise AssertionError('Unexpected successful request')
            except urllib.error.HTTPError as error:
                assert error.code == status, error.code
    finally:
        server.send_signal(signal.SIGINT)
        server.wait(timeout=10)
print('Watch smoke passed: revision updates, page PNG, token path and host restriction')

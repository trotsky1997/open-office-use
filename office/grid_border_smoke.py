"""Check versioned borders, colors and existing workbook content."""
import tempfile
import zipfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from lxml import etree
import officecli_backend
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.xlsx'; book = Workbook(); sheet = book.active
    sheet['A1'] = 'Title'; sheet['A1'].font = Font(bold=True); sheet['B1'] = '=1+1'; book.save(source)
    patched = root / 'patched.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(patched, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                tree = etree.fromstring(data); ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                tree.xpath('./x:sheetData/x:row/x:c[@r="B1"]/x:v', namespaces=ns)[0].text = '2'; data = etree.tostring(tree)
            after.writestr(entry, data)
    patched.replace(source); draft = root / 'draft'; session.start(draft, source=source)
    first = session.format_grid(draft, 0, 'Sheet', 'A1:B1', 'A1:B1', {'border':'thin', 'border_color':'336699'})
    book = load_workbook(first['file']); assert book.active['A1'].font.b
    for cell in ('A1', 'B1'):
        for side in ('top', 'bottom', 'left', 'right'):
            line = getattr(book.active[cell].border, side)
            assert line.style == 'thin' and line.color.rgb[-6:] == '336699'
    assert load_workbook(first['file'], data_only=True).active['B1'].value == 2
    style = session.read_grid(draft, 1, 'Sheet', 'A1')['rows'][0][0]['style']
    assert all(value == {'style':'thin', 'color_rgb':'336699'} for value in style['border'].values())
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(first['file']) as after:
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name not in ('xl/styles.xml', 'xl/worksheets/sheet1.xml'))
    second = session.format_grid(draft, 1, 'Sheet', 'A1:B1', 'A1', {'border_bottom':'double'})
    cell = load_workbook(second['file']).active['A1']
    assert cell.border.bottom.style == 'double' and cell.border.top.style == 'thin' and cell.font.b
    assert officecli_backend.audit(second['file'])['validation']['exit_code'] == 0
    for props in ({'border':'wat'}, {'border_color':'ff0000'}, {'border':'none','border_color':'ff0000'}, {'border':'thin','border_left':'thick'}, {'border':'thin','border_color':'red'}):
        try: session.format_grid(draft, 2, 'Sheet', 'A1', 'A1', props)
        except ValueError: pass
        else: raise AssertionError(f'Invalid border accepted: {props}')
    assert session.info(draft)['latest_revision'] == 2
    cleared = session.format_grid(draft, 2, 'Sheet', 'A1', 'A1', {'border':'none'})
    cell = load_workbook(cleared['file']).active['A1']
    assert all(getattr(cell.border, side).style is None for side in ('top', 'bottom', 'left', 'right'))
print('Grid border smoke passed: styles, color, preservation and validation')

"""Check atomic ordered text edits across Office formats."""
import concurrent.futures
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    doc = Document(); paragraph = doc.add_paragraph(); paragraph.add_run('Hello ').bold = True; paragraph.add_run('world').italic = True; doc.add_paragraph('Second item')
    doc.save(root / 'doc.docx'); session.start(root / 'doc', source=root / 'doc.docx')
    changes = [{'path': '/body/p[1]', 'find': 'Hello', 'replacement': 'Welcome'}, {'path': '/body/p[2]', 'find': 'Second', 'replacement': 'Final'}]
    failed = [changes[0], {'path': '/body/p[2]', 'find': 'missing', 'replacement': 'bad'}]
    try: session.replace_text_batch(root / 'doc', 0, failed)
    except ValueError as error: assert 'Replacement 2' in str(error)
    else: raise AssertionError('Invalid batch accepted')
    assert session.info(root / 'doc')['latest_revision'] == 0 and not list((root / 'doc').glob('.office-replace-batch-*'))
    cli = Path(__file__).with_name('open_office_use.py')
    output = subprocess.run([sys.executable, str(cli), 'session', 'replace-text-batch', str(root / 'doc'), '--base', '0', '--changes', json.dumps(changes)], capture_output=True, text=True, check=True)
    assert json.loads(output.stdout)['count'] == 2
    updated = Document(root / 'doc/revision-000001.docx'); assert updated.paragraphs[0].text == 'Welcome world' and updated.paragraphs[1].text == 'Final item'
    assert updated.paragraphs[0].runs[0].bold and updated.paragraphs[0].runs[1].italic
    wb = Workbook(); wb.active['A1'] = 'Alpha'; wb.active['A2'] = 'Beta'; wb.save(root / 'book.xlsx'); session.start(root / 'book', source=root / 'book.xlsx')
    result = session.replace_text_batch(root / 'book', 0, [{'path':'/Sheet/A1','find':'Alpha','replacement':'One'}, {'path':'/Sheet/A2','find':'Beta','replacement':'Two'}])
    sheet = load_workbook(result['file']).active; assert sheet['A1'].value == 'One' and sheet['A2'].value == 'Two' and result['recalculation_requested']
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6]); slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = 'Old title'; slide.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1)).text = 'Old body'
    deck.save(root / 'deck.pptx'); session.start(root / 'deck', source=root / 'deck.pptx')
    edits = [{'path':'/slide[1]/shape[1]','find':'Old','replacement':'New'}, {'path':'/slide[1]/shape[2]','find':'Old','replacement':'New'}]
    def race(_):
        try: return session.replace_text_batch(root / 'deck', 0, edits)['revision']
        except ValueError: return None
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool: outcomes = list(pool.map(race, range(2)))
    assert sorted(outcomes, key=lambda x: x is None) == [1, None]
    slide = Presentation(root / 'deck/revision-000001.pptx').slides[0]; assert [shape.text for shape in slide.shapes] == ['New title', 'New body']
print('Atomic ordered text batch smoke passed')

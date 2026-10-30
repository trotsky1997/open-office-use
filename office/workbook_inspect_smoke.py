"""Exercise bounded sparse workbook inspection and revision history."""
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from openpyxl import Workbook
import spreadsheet
import session

with tempfile.TemporaryDirectory() as tmp:
    root=Path(tmp); source=root/'sparse.xlsx'; draft=root/'draft'; book=Workbook(); sheet=book.active; sheet.title='Sparse'
    sheet['A1']='Beginning'; sheet['XFD1048576']='Distant'; sheet['B2']='=A1'; sheet['B3']='z'*600
    hidden=book.create_sheet('Hidden'); hidden.sheet_state='hidden'
    for row in range(1,103): hidden.cell(row,1,value=f'row {row}')
    book.save(source)
    # Make the declared dimension incorrect while retaining the distant stored cell.
    rewritten=root/'incorrect-dimension.xlsx'
    with zipfile.ZipFile(source) as before,zipfile.ZipFile(rewritten,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='xl/worksheets/sheet1.xml': data=data.replace(b'<dimension ref="A1:XFD1048576"/>',b'<dimension ref="A1"/>')
            after.writestr(entry,data)
    result=spreadsheet.inspect_workbook(rewritten); assert result['total_worksheets']==2 and not result['worksheets_truncated']
    sparse, second=result['sheets']; cells={c['cell']:c for c in sparse['cells']}
    assert sparse['declared_max_row']==1 and sparse['sample_complete'] and cells['XFD1048576']['value']=='Distant'
    assert cells['B2']['value']=='=A1' and cells['B3']['value_truncated'] and len(cells['B3']['value'])==500
    assert second['state']=='hidden' and len(second['cells'])==100 and not second['sample_complete']
    session.start(draft,source=rewritten); assert session.inspect_workbook(draft,0)['sheets'][0]['cells']==sparse['cells']
    cli=Path(__file__).with_name('open_office_use.py'); call=subprocess.run([sys.executable,str(cli),'session','inspect-workbook',str(draft),'--revision','0'],capture_output=True,text=True,check=True)
    assert json.loads(call.stdout)['revision']==0
    via_cli=subprocess.run([sys.executable,str(cli),'inspect',str(rewritten)],capture_output=True,text=True,check=True)
    assert len(via_cli.stdout)<100000 and json.loads(via_cli.stdout)['sheets'][0]['cells'][-1]['cell']=='XFD1048576'
    shared=root/'shared.xlsx'; import xlsxwriter
    writer=xlsxwriter.Workbook(str(shared)); tab=writer.add_worksheet(); tab.write('A1','共享😀'); tab.write_formula('B1','=1+2',None,3); writer.close()
    shared_result=spreadsheet.inspect_workbook(shared)['sheets'][0]['cells']; assert shared_result[0]['value']=='共享😀' and shared_result[1]['value']=='=1+2'
    bounded=root/'bounded.xlsx'; bound_book=Workbook(); bound_sheet=bound_book.active
    for row in range(1,10001): bound_sheet.cell(row,1).number_format='0.00'
    bound_sheet['A10001']='Beyond scan limit'; bound_book.save(bounded)
    bound_view=spreadsheet.inspect_workbook(bounded)['sheets'][0]
    assert bound_view['stored_cells_scanned']==10000 and not bound_view['sample_complete'] and not bound_view['cells']
    chart_file=root/'chart.xlsx'; writer=xlsxwriter.Workbook(str(chart_file)); tab=writer.add_worksheet('Data'); tab.write_column('A1',[1,2]); chart=writer.add_chart({'type':'column'}); chart.add_series({'values':'=Data!$A$1:$A$2'}); chart_sheet=writer.add_chartsheet('Chart'); chart_sheet.set_chart(chart); writer.close()
    chart_view=spreadsheet.inspect_workbook(chart_file); assert chart_view['chartsheet_names']==['Chart'] and chart_view['total_worksheets']==1
    many=root/'many.xlsx'; many_book=Workbook()
    for index in range(51): many_book.create_sheet(f'Sheet{index + 1}')
    many_book.save(many); many_view=spreadsheet.inspect_workbook(many)
    assert many_view['total_worksheets']==52 and len(many_view['sheets'])==50 and many_view['worksheets_truncated']
print('Workbook inspection passed: sparse distant cells, dimensions, hidden sheets, bounds and revisions')

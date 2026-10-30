"""Exercise common spreadsheet formatting through versioned grid edits."""
import json
import subprocess
import sys
import tempfile
import concurrent.futures
import zipfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from lxml import etree
import session
import grid_styles
import officecli_backend

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.xlsx'; draft = root / 'draft'; book = Workbook(); sheet = book.active
    sheet['A1'] = 'Report'; sheet['A1'].font = Font(bold=True, name='Calibri', size=11); sheet['A1'].fill = PatternFill('solid', fgColor='FFEECC')
    sheet['B1'] = '=1+1'; sheet.conditional_formatting.add('B1:B2', CellIsRule(operator='greaterThan', formula=['1'], font=Font(color='FF0000'))); book.create_sheet('Other')['A1'] = '=3+3'; book.save(source)
    patched = root / 'cached.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(patched, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                tree = etree.fromstring(data); ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                tree.xpath('./s:sheetData/s:row/s:c[@r="B1"]/s:v', namespaces=ns)[0].text = '99.5'; data = etree.tostring(tree)
            after.writestr(entry, data)
    patched.replace(source); original = source.read_bytes(); session.start(draft, source=source)
    props = {'italic': True, 'underline': 'single', 'font_name': 'Arial', 'font_size': 14.5, 'font_color': '336699', 'horizontal': 'center', 'vertical': 'top', 'wrap': True, 'numberformat': '0.0000%'}
    result = session.format_grid(draft, 0, 'Sheet', 'A1:B1', 'A1:B1', props)
    assert result['formula_cache_effect'] == 'preserved' and load_workbook(result['file'], data_only=True).active['B1'].value == 99.5
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(result['file']) as after:
        assert set(before.namelist()) == set(after.namelist())
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name not in ('xl/styles.xml', 'xl/worksheets/sheet1.xml'))
        old_sheet = etree.fromstring(before.read('xl/worksheets/sheet1.xml')); new_sheet = etree.fromstring(after.read('xl/worksheets/sheet1.xml'))
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        for tree in (old_sheet, new_sheet):
            for cell in tree.xpath('./s:sheetData/s:row/s:c[@r="A1" or @r="B1"]', namespaces=ns): cell.attrib.pop('s', None)
        assert grid_styles.equivalent(old_sheet) == grid_styles.equivalent(new_sheet)
        original_styles = before.read('xl/styles.xml'); modified_styles = etree.fromstring(original_styles); modified_styles.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}fonts')[0][0].set('val', 'Changed')
        try: grid_styles.merge_styles(original_styles, etree.tostring(modified_styles))
        except ValueError: pass
        else: raise AssertionError('Mutation of an existing font accepted')
    audit = officecli_backend.audit(result['file']); assert audit['validation']['exit_code'] == 0 and audit['validation']['result']['success'], audit
    output = load_workbook(result['file']); cell = output.active['A1']
    assert cell.value == 'Report' and cell.font.b and cell.font.i and cell.font.u == 'single', cell.font
    assert cell.font.name == 'Arial' and cell.font.sz == 14.5 and cell.font.color.rgb[-6:] == '336699' and cell.number_format == '0.0000%', cell.font
    assert cell.fill.fgColor.rgb[-6:] == 'FFEECC' and cell.alignment.horizontal == 'center' and cell.alignment.vertical == 'top' and cell.alignment.wrap_text, cell.alignment
    assert output.active['B1'].value == '=1+1' and source.read_bytes() == original
    style = session.read_grid(draft, 1, 'Sheet', 'A1')['rows'][0][0]['style']
    assert style['italic'] and style['underline'] == 'single' and style['font_name'] == 'Arial' and style['font_size'] == 14.5 and style['font_color_rgb'] == '336699'
    assert style['horizontal'] == 'center' and style['vertical'] == 'top' and style['wrap']
    signed = root / 'signed.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(signed, 'w') as after:
        for entry in before.infolist(): after.writestr(entry, before.read(entry))
        after.writestr('_xmlsignatures/signature.xml', '<signature/>')
    signed_draft = root / 'signed-draft'; session.start(signed_draft, source=signed)
    try: session.format_grid(signed_draft, 0, 'Sheet', 'A1', 'A1', {'bold': True})
    except ValueError: pass
    else: raise AssertionError('Signed workbook formatting accepted')
    assert session.info(signed_draft)['latest_revision'] == 0
    race = root / 'race'; session.start(race, source=source)
    def compete(color):
        try: return session.format_grid(race, 0, 'Sheet', 'A1', 'A1', {'font_color': color})['revision']
        except ValueError: return None
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool: outcomes = list(pool.map(compete, ['112233', '334455']))
    assert sorted(outcomes, key=lambda value: value is None) == [1, None] and session.info(race)['latest_revision'] == 1
    for props in ({'font_size': True}, {'font_size': 500}, {'font_size': float('nan')}, {'font_size': 10**400}, {'font_color': 'red'}, {'horizontal': 'bad'}, {'wrap': 'yes'}, {'font_name': '\n'}):
        try: session.format_grid(draft, 1, 'Sheet', 'A1:B1', 'A1', props)
        except ValueError: pass
        else: raise AssertionError(f'Invalid style accepted: {props}')
    assert session.info(draft)['latest_revision'] == 1
    cli = Path(__file__).with_name('open_office_use.py')
    command = subprocess.run([sys.executable, str(cli), 'session', 'format-grid', str(draft), '--base', '1', '--sheet', 'Sheet', '--range', 'A1:B1', '--style-range', 'A1', '--properties', json.dumps({'italic': False, 'wrap': False, 'underline': 'none'})], capture_output=True, text=True, check=True)
    assert json.loads(command.stdout)['revision'] == 2
    changed = load_workbook(draft / 'revision-000002.xlsx').active['A1']
    assert not changed.font.i and changed.font.u is None and not changed.alignment.wrap_text and changed.font.b and changed.font.name == 'Arial'
    blank = session.format_grid(draft, 2, 'Sheet', 'D5', 'D5', {'bold': True})
    blank_book = load_workbook(blank['file']); assert blank_book.active['D5'].font.b and load_workbook(blank['file'], data_only=True).active['B1'].value == 99.5
    with zipfile.ZipFile(blank['file']) as archive:
        tree = etree.fromstring(archive.read('xl/worksheets/sheet1.xml')); ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        assert tree.find(ns + 'dimension').get('ref') == 'A1:D5'
print('Spreadsheet grid formatting smoke passed')

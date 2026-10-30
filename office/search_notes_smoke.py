"""Search speaker notes and apply a targeted edit in a saved presentation."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path
from pptx import Presentation
import session
import slides

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.pptx'; deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1]); first.shapes.title.text = 'needle in title'
    frame = first.notes_slide.notes_text_frame; frame.text = 'First needle'; frame.add_paragraph().text = 'Second needle'
    hidden = deck.slides.add_slide(deck.slide_layouts[1]); hidden._element.set('show', '0'); hidden.notes_slide.notes_text_frame.text = 'Hidden needle'
    deck.slides.add_slide(deck.slide_layouts[1]); deck.save(source)
    result = slides.search_text(source, 'needle', scope='notes')
    assert result['complete'] and result['scope'] == 'notes' and [(m['slide'], m['notes_paragraph'], m['hidden']) for m in result['matches']] == [(1, 1, False), (1, 2, False), (2, 1, True)]
    assert slides.search_text(source, 'needle')['matches'][0]['path'].endswith('/paragraph[1]')
    assert len(slides.search_text(source, 'needle', slide=2, scope='notes')['matches']) == 1
    draft = root / 'draft'; session.start(draft, source=source)
    found = session.search_slides(draft, 0, 'needle', scope='notes'); match = found['matches'][0]
    result = session.replace_notes_text(draft, 0, match['slide'], match['notes_paragraph'], 'needle', 'target')
    assert session.read_notes(draft, 1, 1)['text'].startswith('First target')
    assert len(session.search_slides(draft, 1, 'needle', scope='notes')['matches']) == 2
    assert all(item['replace_text_path_supported'] for item in found['matches'])
    single_draft=root/'note-path-draft';session.start(single_draft,source=source)
    note_target=session.edit_target(single_draft,0,found['matches'][0]['path']);assert note_target['value']=='First needle' and note_target['replace_only'] and note_target['replace_text_path_supported']
    session.replace_text(single_draft,0,found['matches'][0]['path'],'needle','subject')
    assert session.read_notes(single_draft,1,1)['text'].startswith('First subject')
    batch_draft=root/'mixed-batch-draft';session.start(batch_draft,source=source)
    title_path=f'/slide[1]/shape[@id={first.shapes.title.shape_id}]/paragraph[1]'
    changes=[{'path':title_path,'find':'needle','replacement':'subject'}, {'path':'/slide[1]/notes/paragraph[1]','find':'needle','replacement':'subject'}, {'path':'/slide[1]/notes/paragraph[2]','find':'needle','replacement':'topic'}]
    try: session.replace_text_batch(batch_draft,0,changes+[{'path':'/slide[1]/notes/paragraph[2]','find':'missing','replacement':'x'}])
    except ValueError as exc: assert 'Replacement 4' in str(exc)
    else: raise AssertionError('Failed mixed batch was published')
    assert session.info(batch_draft)['latest_revision']==0
    batch=session.replace_text_batch(batch_draft,0,changes);assert batch['revision']==1 and batch['count']==3
    edited=Presentation(batch['file']);assert edited.slides[0].shapes.title.text=='subject in title'
    assert edited.slides[0].notes_slide.notes_text_frame.text=='First subject\nSecond topic' and edited.slides[1].notes_slide.notes_text_frame.text=='Hidden needle'
    assert Presentation(source).slides[0].notes_slide.notes_text_frame.text=='First needle\nSecond needle'
    cli = Path(__file__).with_name('open_office_use.py')
    output = subprocess.run([sys.executable, str(cli), 'search-slides', str(source), '--scope', 'notes', '--query', 'needle'], text=True, capture_output=True, check=True)
    assert json.loads(output.stdout)['scope'] == 'notes' and len(json.loads(output.stdout)['matches']) == 3
    many = root / 'many.pptx'; deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[1]); frame = slide.notes_slide.notes_text_frame; frame.text = 'needle'
    for _ in range(104): frame.add_paragraph().text = 'needle'
    deck.save(many); first_page = slides.search_text(many, 'needle', scope='notes')
    assert len(first_page['matches']) == 100 and first_page['next_offset'] is not None
    second_page = slides.search_text(many, 'needle', offset=first_page['next_offset'], snapshot_sha256=first_page['snapshot_sha256'], scope='notes')
    assert len(second_page['matches']) == 5 and second_page['scan_complete']
    try: slides.search_text(source, 'needle', offset=1, snapshot_sha256='0'*64, scope='notes')
    except ValueError: pass
    else: raise AssertionError('Changed snapshot accepted')
    try: slides.search_text(source, 'needle', scope='invalid')
    except ValueError: pass
    else: raise AssertionError('Invalid scope accepted')
print('Speaker notes search smoke passed: hidden slides, paragraphs, targeted editing and snapshot safety')

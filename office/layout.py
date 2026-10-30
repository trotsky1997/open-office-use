"""Bounded geometric review of PowerPoint slide objects."""
import math
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.exc import PackageNotFoundError
from pptx.enum.shapes import MSO_SHAPE_TYPE

P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
EMU = 914400
MAX_SHAPES = 10000
MAX_PAIRS = 200000
MAX_ISSUES = 200
MAX_DEPTH = 8


def compose(left, right):
    a, b, c, d, e, f = left; g, h, i, j, k, l = right
    return (a*g+c*h, b*g+d*h, a*i+c*j, b*i+d*j, a*k+c*l+e, b*k+d*l+f)


def translate(x, y):
    return (1, 0, 0, 1, x, y)


def scale(x, y):
    return (x, 0, 0, y, 0, 0)


def rotate(angle):
    angle = math.radians(angle); c, s = math.cos(angle), math.sin(angle)
    return (c, s, -s, c, 0, 0)


def point(matrix, x, y):
    a, b, c, d, e, f = matrix
    return (a*x+c*y+e, b*x+d*y+f)


def around(cx, cy, angle, flip_x=1, flip_y=1):
    return compose(translate(cx, cy), compose(rotate(angle), compose(scale(flip_x, flip_y), translate(-cx, -cy))))


def group_transform(shape, parent):
    xfrm = shape._element.find(P + 'grpSpPr/' + A + 'xfrm')
    if xfrm is None:
        raise ValueError('Group has no transform')
    parts = [xfrm.find(A + name) for name in ('off', 'ext', 'chOff', 'chExt')]
    if any(item is None for item in parts):
        raise ValueError('Group transform is incomplete')
    off, ext, child_off, child_ext = parts
    x, y = int(off.get('x')), int(off.get('y'))
    width, height = int(ext.get('cx')), int(ext.get('cy'))
    child_x, child_y = int(child_off.get('x')), int(child_off.get('y'))
    child_w, child_h = int(child_ext.get('cx')), int(child_ext.get('cy'))
    if width <= 0 or height <= 0 or child_w <= 0 or child_h <= 0:
        raise ValueError('Group transform has an empty extent')
    rotation = int(xfrm.get('rot', '0')) / 60000
    flip_x = -1 if xfrm.get('flipH') in ('1', 'true') else 1
    flip_y = -1 if xfrm.get('flipV') in ('1', 'true') else 1
    mapping = compose(translate(x, y), compose(scale(width / child_w, height / child_h), translate(-child_x, -child_y)))
    return compose(parent, compose(around(x + width/2, y + height/2, rotation, flip_x, flip_y), mapping))


def polygon(shape, parent):
    x, y, width, height = shape.left, shape.top, shape.width, shape.height
    if any(value is None for value in (x, y, width, height)) or width < 0 or height < 0:
        raise ValueError('Shape has incomplete geometry')
    local = around(x + width/2, y + height/2, shape.rotation or 0)
    matrix = compose(parent, local)
    points = [point(matrix, x, y), point(matrix, x+width, y), point(matrix, x+width, y+height), point(matrix, x, y+height)]
    if not all(math.isfinite(value) for pair in points for value in pair):
        raise ValueError('Shape has non-finite geometry')
    return points


def area(points):
    return abs(sum(x1*y2-x2*y1 for (x1, y1), (x2, y2) in zip(points, points[1:] + points[:1]))) / 2


def intersect(subject, clip):
    orientation = sum(x1*y2-x2*y1 for (x1, y1), (x2, y2) in zip(clip, clip[1:] + clip[:1]))
    sign = 1 if orientation >= 0 else -1
    for a, b in zip(clip, clip[1:] + clip[:1]):
        if not subject:
            break
        result = []
        def distance(p):
            return sign * ((b[0]-a[0])*(p[1]-a[1]) - (b[1]-a[1])*(p[0]-a[0]))
        previous = subject[-1]; prev_distance = distance(previous)
        for current in subject:
            current_distance = distance(current)
            if (current_distance >= 0) != (prev_distance >= 0):
                factor = prev_distance / (prev_distance - current_distance)
                result.append((previous[0] + factor*(current[0]-previous[0]), previous[1] + factor*(current[1]-previous[1])))
            if current_distance >= 0:
                result.append(current)
            previous, prev_distance = current, current_distance
        subject = result
    return subject


def inspect(file, selection=None):
    source = Path(file).resolve(strict=True)
    if source.suffix.lower() != '.pptx':
        raise ValueError('Slide layout inspection requires a PPTX file')
    if source.stat().st_size > 128 * 1024 * 1024:
        raise ValueError('Presentation exceeds the 128 MiB scan limit')
    with tempfile.TemporaryDirectory(prefix='office-layout-') as tmp:
        snapshot = Path(tmp) / 'presentation.pptx'; shutil.copyfile(source, snapshot)
        if snapshot.stat().st_size > 128 * 1024 * 1024:
            raise ValueError('Presentation exceeds the 128 MiB scan limit')
        try:
            with zipfile.ZipFile(snapshot) as archive:
                entries = archive.infolist(); names = [entry.filename for entry in entries]
                if len(entries) > 5000 or len(names) != len(set(names)) or sum(entry.file_size for entry in entries) > 128 * 1024 * 1024:
                    raise ValueError('Presentation exceeds package limits or contains duplicate parts')
            presentation = Presentation(snapshot)
            slide_width, slide_height = presentation.slide_width, presentation.slide_height
            slides = presentation.slides
            visibility = [slide._element.get('show') for slide in slides]
            hidden_slides = [index for index, value in enumerate(visibility, 1) if value in ('0', 'false')]
            visibility_known = all(value in (None, '0', '1', 'false', 'true') for value in visibility)
            if not slide_width or not slide_height or slide_width <= 0 or slide_height <= 0:
                raise ValueError('Presentation has no valid slide size')
            off_slide, overlaps, skipped, selected = [], [], [], []
            selected_count = 0
            totals = {'shapes_scanned': 0, 'text_shapes': 0, 'pairs_checked': 0, 'off_slide_count': 0, 'overlap_count': 0, 'skipped_count': 0}
            complete, overlap_complete, shape_limit = True, True, False
            fonts = {'threshold_pt': 14, 'runs_scanned': 0, 'paragraph_default_runs': 0, 'unresolved_runs': 0, 'candidate_count': 0, 'candidates': [], 'scan_limit_reached': False}
            # shortcut: run and paragraph defaults only; resolve theme and effective rendered sizes when supported.
            def scan_fonts(frame, descriptor, cell=None):
                if fonts['scan_limit_reached']: return
                for paragraph in frame.paragraphs:
                    runs = paragraph.runs
                    paragraph_defaults = paragraph._p.findall(A + 'pPr')
                    defaults = paragraph_defaults[0].findall(A + 'defRPr') if len(paragraph_defaults) == 1 else []
                    raw_size = defaults[0].get('sz') if len(defaults) == 1 else None
                    default_size = int(raw_size) / 100 if raw_size and re.fullmatch(r'[0-9]{1,6}', raw_size) and 100 <= int(raw_size) <= 400000 else None
                    if paragraph._p.findall(A + 'fld') or paragraph.text.strip() and not runs: fonts['unresolved_runs'] += 1
                    for run in runs:
                        if not run.text.strip(): continue
                        if fonts['runs_scanned'] >= 100000:
                            fonts['scan_limit_reached'] = True; return
                        fonts['runs_scanned'] += 1
                        size = run.font.size
                        size_source = 'run' if size is not None else 'paragraph'
                        size_pt = size.pt if size is not None else default_size
                        if size_pt is None:
                            fonts['unresolved_runs'] += 1; continue
                        if size_source == 'paragraph': fonts['paragraph_default_runs'] += 1
                        if size_pt < fonts['threshold_pt']:
                            fonts['candidate_count'] += 1
                            if len(fonts['candidates']) < MAX_ISSUES:
                                fonts['candidates'].append({'slide': descriptor['slide'], 'path_hint': descriptor['path_hint'], 'id': descriptor['id'], 'name': descriptor['name'], 'cell': cell, 'size_pt': size_pt, 'size_source': size_source, 'text_preview': run.text[:160], 'visible_polygon_inches': descriptor['visible_polygon_inches']})
            def issue(bucket, item):
                if len(bucket) < MAX_ISSUES:
                    bucket.append(item)
            def skip(slide, path, reason):
                nonlocal complete
                complete = False; totals['skipped_count'] += 1
                issue(skipped, {'slide': slide, 'path_hint': path, 'reason': str(reason)[:300]})
            for slide_number, slide in enumerate(slides, 1):
                objects = []
                def walk(shapes, matrix, prefix, depth):
                    nonlocal shape_limit, selected_count
                    for shape in shapes:
                        if totals['shapes_scanned'] >= MAX_SHAPES:
                            skip(slide_number, prefix, 'Shape scan limit reached'); shape_limit = True; return
                        totals['shapes_scanned'] += 1
                        try:
                            metadata = next(shape._element.iter(P + 'cNvPr'), None)
                            if metadata is not None and metadata.get('hidden') in ('1', 'true'):
                                continue
                            shape_id, shape_type = shape.shape_id, shape.shape_type
                        except (ValueError, TypeError, AttributeError, OverflowError) as exc:
                            skip(slide_number, prefix, exc); continue
                        if shape_type == MSO_SHAPE_TYPE.GROUP:
                            path = f'{prefix}/group[@id={shape_id}]'
                            if depth >= MAX_DEPTH:
                                skip(slide_number, path, 'Group nesting limit reached'); continue
                            try:
                                walk(shape.shapes, group_transform(shape, matrix), path, depth + 1)
                                if shape_limit: return
                            except (ValueError, TypeError, AttributeError, OverflowError) as exc:
                                skip(slide_number, path, exc)
                            continue
                        path = None
                        try:
                            local = shape._element.tag.rsplit('}', 1)[-1]
                            kind = {'sp': 'shape', 'pic': 'picture', 'cxnSp': 'connector'}.get(local)
                            if local == 'graphicFrame':
                                kind = 'table' if shape.has_table else ('chart' if shape.has_chart else None)
                            path = f'{prefix}/{kind}[@id={shape_id}]' if kind else None
                            points = polygon(shape, matrix)
                            left = min(p[0] for p in points); top = min(p[1] for p in points)
                            right = max(p[0] for p in points); bottom = max(p[1] for p in points)
                            descriptor = {'slide': slide_number, 'id': shape_id, 'name': str(shape.name or '')[:200],
                                          'path_hint': path, 'bounds_inches': [round(value / EMU, 4) for value in (left, top, right, bottom)],
                                          'polygon_inches': [[round(x / EMU, 4), round(y / EMU, 4)] for x, y in points]}
                            if left < -1 or top < -1 or right > slide_width + 1 or bottom > slide_height + 1:
                                totals['off_slide_count'] += 1; issue(off_slide, descriptor)
                            text = shape.text_frame.text if shape.has_text_frame else ''
                            if shape.has_table:
                                text = '\n'.join(cell.text for row in shape.table.rows for cell in row.cells)
                            has_text = bool(text.strip()); descriptor['has_text'] = has_text; descriptor['draw_order'] = totals['shapes_scanned']
                            if has_text:
                                descriptor['text_preview'] = text[:160]; descriptor['text_truncated'] = len(text) > 160
                            if has_text:
                                totals['text_shapes'] += 1
                            full_area = area(points)
                            visible = intersect(points[:], [(0, 0), (slide_width, 0), (slide_width, slide_height), (0, slide_height)])
                            descriptor['visible_polygon_inches'] = [[round(x / EMU, 4), round(y / EMU, 4)] for x, y in visible]
                            if shape.has_table:
                                for row_index, row in enumerate(shape.table.rows, 1):
                                    for column_index, cell in enumerate(row.cells, 1):
                                        scan_fonts(cell.text_frame, descriptor, [row_index, column_index])
                            elif shape.has_text_frame:
                                scan_fonts(shape.text_frame, descriptor)
                            if selection and selection['page'] == slide_number and path and area(visible) > 1:
                                x1, y1, x2, y2 = selection['box']
                                box = [(x1 * slide_width, y1 * slide_height), (x2 * slide_width, y1 * slide_height),
                                       (x2 * slide_width, y2 * slide_height), (x1 * slide_width, y2 * slide_height)]
                                overlap = area(intersect(visible[:], box))
                                if overlap > 1:
                                    selected_count += 1
                                    if len(selected) < 50:
                                        selected.append({**descriptor, 'selection_overlap_fraction': round(overlap / area(visible), 6)})
                            if full_area > 1 and area(visible) > 1:
                                visible_box = (min(p[0] for p in visible), min(p[1] for p in visible), max(p[0] for p in visible), max(p[1] for p in visible))
                                objects.append((descriptor, visible, visible_box, full_area))
                        except (ValueError, TypeError, AttributeError, OverflowError) as exc:
                            skip(slide_number, path, exc)
                walk(slide.shapes, (1, 0, 0, 1, 0, 0), f'/slide[{slide_number}]', 0)
                for index, (first, first_poly, first_box, first_area) in enumerate(objects):
                    if not overlap_complete:
                        break
                    if not first['has_text']:
                        continue
                    for other_index, (second, second_poly, second_box, second_area) in enumerate(objects):
                        if index == other_index or (second['has_text'] and other_index < index):
                            continue
                        if totals['pairs_checked'] >= MAX_PAIRS:
                            overlap_complete = False; break
                        totals['pairs_checked'] += 1
                        if first_box[0] >= second_box[2] or second_box[0] >= first_box[2] or first_box[1] >= second_box[3] or second_box[1] >= first_box[3]:
                            continue
                        overlap_polygon = intersect(first_poly[:], second_poly); intersection = area(overlap_polygon)
                        if intersection > 1:
                            totals['overlap_count'] += 1
                            bottom, top = (first, second) if first['draw_order'] < second['draw_order'] else (second, first)
                            issue(overlaps, {'slide': slide_number, 'first': bottom, 'second': top,
                                             'intersection_polygon_inches': [[round(x / EMU, 4), round(y / EMU, 4)] for x, y in overlap_polygon],
                                             'area_square_inches': round(intersection / (EMU*EMU), 12),
                                             'fraction_of_smaller': round(min(1, intersection / min(first_area, second_area)), 8)})
                if shape_limit:
                    break
            return {**totals, 'slides_total': len(slides), 'slide_size_inches': [round(slide_width / EMU, 6), round(slide_height / EMU, 6)],
                    'hidden_slides': hidden_slides, 'visibility_known': visibility_known, 'selection_matches': selected,
                    'selection_match_count': selected_count, 'complete': complete and overlap_complete,
                    'geometry_complete': complete, 'overlap_complete': overlap_complete,
                    'off_slide': off_slide, 'off_slide_truncated': totals['off_slide_count'] > len(off_slide),
                    'overlaps': overlaps, 'overlaps_truncated': totals['overlap_count'] > len(overlaps),
                    'skipped': skipped, 'skipped_truncated': totals['skipped_count'] > len(skipped),
                    'font_size': {**fonts, 'candidates_truncated': fonts['candidate_count'] > len(fonts['candidates']),
                                  'complete': complete and not fonts['scan_limit_reached'] and not fonts['unresolved_runs'],
                                  'scope': 'Run and paragraph-default sizes only; other inheritance, autofit, group scaling, masters and rendered legibility require review.'},
                    'units': 'inches', 'overlaps_are_candidates': True}
        except (zipfile.BadZipFile, PackageNotFoundError, etree.XMLSyntaxError, KeyError) as exc:
            raise ValueError(f'Invalid presentation package: {exc}') from exc

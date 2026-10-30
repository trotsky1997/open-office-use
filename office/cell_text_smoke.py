"""Verify XLSX rich-text replacement, shared references and package preservation."""
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import xlsxwriter
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

import session
import officecli_backend

X = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.xlsx'; draft = root / 'draft'
    book = xlsxwriter.Workbook(str(source)); sheet = book.add_worksheet('Sales Data')
    bold = book.add_format({'bold': True}); italic = book.add_format({'italic': True}); cell_style = book.add_format({'bg_color': '#FFCC00'})
    for address in ('A1', 'A2'):
        sheet.write_rich_string(address, bold, 'Hello ', italic, 'world!', cell_style)
    sheet.write_rich_string('A3', bold, 'Hei earth', italic, 'ld!', cell_style)
    sheet.write_comment('A1', 'Keep this comment'); sheet.write_formula('B1', '=LEN(A1)', None, 12)
    sheet.merge_range('C1:D1', 'Merged cell', cell_style)
    sheet.write_url('E1', 'https://example.com', string='linked text')
    sheet.write_array_formula('F1:F2', '{=ROW(F1:F2)}'); sheet.write_number('G1', 123)
    sheet.write_string('N1', 'hello'); sheet.write_string('N2', 'hello'); sheet.write_string('N3', 'good')
    chart = book.add_chart({'type': 'column'}); chart.add_series({'values': "='Sales Data'!$G$1"}); sheet.insert_chart('H1', chart)
    other_sheet = book.add_worksheet('Other'); other_sheet.write_string('A1', 'other sheet'); other_sheet.hide()
    book.close(); original = source.read_bytes(); session.start(draft, source=source)
    target = session.edit_target(draft, 0, '/Sales Data/A1')
    assert target['replace_text_path_supported'] and target['value'] == 'Hello world!'
    for path, find, message in [('/Sales Data/A1', 'l', 'exactly once'), ('/Sales Data/D1', 'x', 'placeholder'),
                                ('/Sales Data/F2', '0', 'formula'), ('/Sales Data/G1', '123', 'text cell'), ('/Sales Data/A1', '_x0041_', 'escaped')]:
        try:
            session.replace_text(draft, 0, path, find, 'bad')
            raise AssertionError('Unsupported replacement succeeded')
        except ValueError as error:
            assert message in str(error), (path, error)
    first = session.replace_text(draft, 0, '/Sales Data/A1', 'llo wor', 'i earth')
    result = load_workbook(first['file'], rich_text=True); rich = result['Sales Data']['A1'].value
    assert str(rich) == 'Hei earthld!' and rich[0].font.b and rich[1].font.i
    assert str(result['Sales Data']['A2'].value) == 'Hello world!' and result['Sales Data']['A1'].fill.fgColor.rgb.endswith('FFCC00')
    assert result['Sales Data']['A1'].comment.text == 'Keep this comment'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(first['file']) as after:
        old = ET.fromstring(before.read('xl/sharedStrings.xml')); new = ET.fromstring(after.read('xl/sharedStrings.xml'))
        assert old.get('count') == new.get('count') and int(new.get('uniqueCount')) == int(old.get('uniqueCount')) + 1
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name not in ('xl/sharedStrings.xml', 'xl/worksheets/sheet1.xml'))
    second = session.replace_text(draft, 1, '/Sales Data/A1', 'earth', 'planet')
    code, validation, _ = officecli_backend.invoke('validate', second['file'])
    assert code == 0 and validation['success'] is True, validation
    result = load_workbook(second['file'], rich_text=True); rich = result['Sales Data']['A1'].value
    assert str(rich) == 'Hei planetld!' and rich[0].font.b and rich[1].font.i
    assert str(result['Sales Data']['A2'].value) == 'Hello world!' and str(result['Sales Data']['A3'].value) == 'Hei earthld!'
    assert load_workbook(second['file'], data_only=True)['Sales Data']['B1'].value == 12
    with zipfile.ZipFile(first['file']) as before, zipfile.ZipFile(second['file']) as after:
        old = ET.fromstring(before.read('xl/sharedStrings.xml')); new = ET.fromstring(after.read('xl/sharedStrings.xml'))
        assert old.get('count') == new.get('count') and int(new.get('uniqueCount')) == int(old.get('uniqueCount')) + 1
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name not in ('xl/sharedStrings.xml', 'xl/worksheets/sheet1.xml'))
    linked = session.replace_text(draft, 2, '/Sales Data/E1', 'linked', 'updated')
    cell = load_workbook(linked['file'])['Sales Data']['E1']
    assert cell.value == 'updated text' and cell.hyperlink.target == 'https://example.com'
    other = session.replace_text(draft, 3, '/Other/A1', 'other', 'hidden')
    assert load_workbook(other['file'])['Other']['A1'].value == 'hidden sheet'
    with zipfile.ZipFile(linked['file']) as before, zipfile.ZipFile(other['file']) as after:
        assert before.read('xl/worksheets/sheet1.xml') == after.read('xl/worksheets/sheet1.xml')
    reused = session.replace_text(draft, 4, '/Sales Data/N1', 'hello', 'good')
    reused_sheet = load_workbook(reused['file'])['Sales Data']
    assert reused_sheet['N1'].value == 'good' and reused_sheet['N2'].value == 'hello'
    with zipfile.ZipFile(other['file']) as before, zipfile.ZipFile(reused['file']) as after:
        assert before.read('xl/sharedStrings.xml') == after.read('xl/sharedStrings.xml')
    merged = session.replace_text(draft, 5, '/Sales Data/C1', 'Merged', 'Updated')
    merged_sheet = load_workbook(merged['file'])['Sales Data']
    assert merged_sheet['C1'].value == 'Updated cell' and 'C1:D1' in {str(item) for item in merged_sheet.merged_cells.ranges}
    assert source.read_bytes() == original
    try:
        session.replace_text(draft, 0, '/Sales Data/A1', 'Hello', 'bad')
        raise AssertionError('Stale replacement succeeded')
    except ValueError as error:
        assert 'conflict' in str(error)

    inconsistent = root / 'inconsistent.xlsx'
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(inconsistent, 'w') as dst:
        for entry in src.infolist():
            data = src.read(entry)
            if entry.filename == 'xl/sharedStrings.xml':
                tree = ET.fromstring(data); tree.set('uniqueCount', '0'); data = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            dst.writestr(entry, data)
    invalid_draft = root / 'invalid-draft'; session.start(invalid_draft, source=inconsistent)
    try:
        session.replace_text(invalid_draft, 0, '/Sales Data/A1', 'world', 'planet')
        raise AssertionError('Inconsistent shared string count accepted')
    except ValueError as error:
        assert 'unique count' in str(error)
    assert session.info(invalid_draft)['latest_revision'] == 0

    shared_sheet = root / 'shared-sheet.xlsx'
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(shared_sheet, 'w') as dst:
        for entry in src.infolist():
            data = src.read(entry)
            if entry.filename == 'xl/_rels/workbook.xml.rels':
                relations = ET.fromstring(data); sheets = [item for item in relations if item.get('Type', '').endswith('/worksheet')]
                sheets[1].set('Target', sheets[0].get('Target')); data = ET.tostring(relations, encoding='utf-8', xml_declaration=True)
            dst.writestr(entry, data)
    shared_draft = root / 'shared-sheet-draft'; session.start(shared_draft, source=shared_sheet)
    try:
        session.replace_text(shared_draft, 0, '/Sales Data/A1', 'Hello', 'bad')
        raise AssertionError('Shared worksheet part accepted')
    except ValueError as error:
        assert 'multiple sheets' in str(error)
    assert session.info(shared_draft)['latest_revision'] == 0

    long_source = root / 'long.xlsx'; long_draft = root / 'long-draft'; long_book = Workbook()
    long_book.active['A1'] = 'a' * 32766 + 'Z'; long_book.save(long_source); session.start(long_draft, source=long_source)
    try:
        session.replace_text(long_draft, 0, '/Sheet/A1', 'Z', '😀')
        raise AssertionError('Oversized UTF-16 text accepted')
    except ValueError as error:
        assert '32,767' in str(error)
    assert session.info(long_draft)['latest_revision'] == 0

    escape_source = root / 'escape.xlsx'; escape_draft = root / 'escape-draft'; escape = Workbook()
    escape.active['A1'] = '_xX041_'; escape.save(escape_source); session.start(escape_draft, source=escape_source)
    try:
        session.replace_text(escape_draft, 0, '/Sheet/A1', 'X', '0')
        raise AssertionError('Replacement created an Excel escape sequence')
    except ValueError as error:
        assert 'escaped' in str(error)
    assert session.info(escape_draft)['latest_revision'] == 0

    inline_source = root / 'inline.xlsx'; inline_draft = root / 'inline-draft'; inline = Workbook()
    inline.active['A1'] = CellRichText([TextBlock(InlineFont(b=True), 'Hello '), TextBlock(InlineFont(i=True), 'world!')]); inline.save(inline_source)
    session.start(inline_draft, source=inline_source)
    changed = session.replace_text(inline_draft, 0, '/Sheet/A1', 'llo wor', 'i earth')
    text = load_workbook(changed['file'], rich_text=True).active['A1'].value
    assert str(text) == 'Hei earthld!' and text[0].font.b and text[1].font.i
    with zipfile.ZipFile(inline_source) as before, zipfile.ZipFile(changed['file']) as after:
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name != 'xl/worksheets/sheet1.xml')
print('Cell text smoke passed: inline and shared rich text, formulas, merges and links')

"""Exercise OfficeCLI validation failures and formula issue detection."""
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook

cli = Path(__file__).with_name('open_office_use.py')
def audit(file):
    result = subprocess.run([sys.executable, str(cli), 'audit', str(file)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    return json.loads(result.stdout)

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    doc = Document(); doc.add_paragraph('Hello'); good = root / 'good.docx'; doc.save(good)
    result = audit(good)
    assert result['validation']['exit_code'] == 0 and result['validation']['result']['success'] is True, result
    bad = root / 'invalid.docx'
    with zipfile.ZipFile(good) as source, zipfile.ZipFile(bad, 'w', zipfile.ZIP_DEFLATED) as dest:
        for part in source.namelist():
            data = source.read(part)
            if part == 'word/document.xml':
                data = data.replace(b'<w:body>', b'<w:body><w:bogus/>')
            dest.writestr(part, data)
    result = audit(bad)
    assert result['validation']['exit_code'] != 0 and result['validation']['result']['success'] is False, result
    assert 'bogus' in str(result['validation']['result']), result
    workbook = Workbook(); workbook.active['A1'] = '=1/0'; xlsx = root / 'error.xlsx'; workbook.save(xlsx)
    result = audit(xlsx)
    assert any(issue['subtype'] == 'formula_eval_error' for issue in result['issues']['issues']), result
print('Audit smoke passed: valid file, schema violation and spreadsheet formula issue')

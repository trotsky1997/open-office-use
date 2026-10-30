"""Inspect rendered PowerPoint table cell text, including merged cells."""
import tempfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE
import open_office_use
import slide_text_check
import session

with tempfile.TemporaryDirectory() as tmp:
    source = Path(tmp) / 'table.pptx'; deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(5), Inches(.75)); table = shape.table
    table.cell(0, 0).text = 'Normal'; table.cell(0, 1).text = 'Overflowing text in a small cell should be checked carefully because it may be clipped or extend beyond the available area of the slide'
    table.cell(0, 1).text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    second = deck.slides.add_slide(deck.slide_layouts[6]); merged = second.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(5), Inches(1)); merged.table.cell(0, 0).merge(merged.table.cell(0, 1)); merged.table.cell(0, 0).text = 'Merged cell content'; deck.save(source)
    result = slide_text_check.scan(source, open_office_use.render)
    assert result['mapping_valid'] and result['complete'] and result['table_cells_scanned'] == 4 and result['table_cells_checked'] == 3, result
    prefix = f'/slide[1]/table[@id={shape.shape_id}]'
    assert [item['path_hint'] for item in result['candidates']] == [prefix + '/tr[1]/tc[2]'], result
    draft = source.parent / 'draft'; session.start(draft, source=source)
    target = result['candidates'][0]
    edited = session.replace_text(draft, 0, target['path_hint'], target['source_text'], 'Fits')
    checked = slide_text_check.scan(edited['file'], open_office_use.render)
    assert checked['candidate_count'] == 0 and checked['complete'], checked
    merge_doc = Presentation(source); t = merge_doc.slides[1].shapes[0].table
    t.cell(0, 1).text = 'Hidden spanned text'; malformed = source.with_name('spanned.pptx'); merge_doc.save(malformed)
    bad = slide_text_check.scan(malformed, open_office_use.render)
    assert not bad['complete'] and bad['skipped'][0]['reason'] == 'unsupported_merged_cell', bad
print('Rendered slide table text check passed')

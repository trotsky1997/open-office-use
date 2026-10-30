"""Inspect saved external relationships without connecting to their targets."""
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
import review
import review_bundle
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); target = 'https://example.invalid/report?token=example'; files = []
    doc = Document(); doc.add_paragraph('Link'); doc.part.relate_to(target, RT.HYPERLINK, is_external=True); files.append(root/'a.docx'); doc.save(files[-1])
    book = Workbook(); book.active['A1'] = 'Link'; book.active['A1'].hyperlink = target; files.append(root/'a.xlsx'); book.save(files[-1])
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6]); box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)); run = box.text_frame.paragraphs[0].add_run(); run.text = 'Link'; run.hyperlink.address = target; files.append(root/'a.pptx'); deck.save(files[-1])
    cli = Path(__file__).with_name('open_office_use.py')
    for file in files:
        report = review.external_relationships(file)
        assert report['complete'] and report['external_count'] == 1 and report['relationships'][0]['target'] == target, report
        draft = root/(file.stem + file.suffix[1:]); session.start(draft, source=file)
        assert session.external_relationships(draft, 0)['snapshot_sha256'] == report['snapshot_sha256']
        out = subprocess.run([sys.executable, str(cli), 'session', 'external-links', str(draft), '--revision', '0'], check=True, capture_output=True, text=True)
        assert json.loads(out.stdout)['external_count'] == 1
        bundle = review_bundle.create(file, root/(draft.name+'-report'), lambda *_: (_ for _ in ()).throw(RuntimeError('skip render')))
        assert 'external_relationships_present' in bundle['flags']
    bad = root/'bad.docx'
    with zipfile.ZipFile(files[0]) as src, zipfile.ZipFile(bad, 'w') as dst:
        for entry in src.infolist():
            data = src.read(entry.filename)
            if entry.filename == 'word/_rels/document.xml.rels': data = b'<bad'
            dst.writestr(entry, data)
    broken = review.external_relationships(bad)
    assert not broken['complete'] and broken['issues'][0]['reason'] == 'invalid_relationship_xml'
    many = root/'many.docx'
    with zipfile.ZipFile(files[0]) as src, zipfile.ZipFile(many, 'w') as dst:
        for entry in src.infolist():
            data = src.read(entry.filename)
            if entry.filename == 'word/_rels/document.xml.rels':
                from lxml import etree
                tree = etree.fromstring(data); ns = '{http://schemas.openxmlformats.org/package/2006/relationships}'
                for i in range(201):
                    etree.SubElement(tree, ns+'Relationship', Id='external'+str(i), Type='urn:test', Target='https://example.invalid/'+'x'*1200, TargetMode='External')
                data = etree.tostring(tree)
            dst.writestr(entry, data)
    capped = review.external_relationships(many)
    assert capped['complete'] and capped['external_count'] == 202 and capped['truncated'] and capped['fields_truncated'] and len(capped['relationships']) == 200
print('External relationship smoke passed: three formats, revisions, CLI and malformed XML')

"""Verify bounded, consistent XLSX range reads and long-cell grid safety."""
import shutil
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
import xlsxwriter
import session
import spreadsheet

with tempfile.TemporaryDirectory() as tmp:
    root=Path(tmp); original=root/'original.xlsx'; later=root/'later.xlsx'
    for path,formula,cached in [(original,'=1+1',2),(later,'=3+4',7)]:
        book=xlsxwriter.Workbook(str(path)); sheet=book.add_worksheet('Sheet'); sheet.write_formula('A1',formula,None,cached); book.close()
    real_load=spreadsheet.load_workbook; calls=[]
    def moving_source(*args,**kwargs):
        workbook=real_load(*args,**kwargs); calls.append(args[0])
        if len(calls)==1: shutil.copyfile(later,original)
        return workbook
    with patch.object(spreadsheet,'load_workbook',moving_source):
        result=spreadsheet.read_range(original,'Sheet','A1')
    assert result['rows'][0][0]['formula']=='=1+1' and result['rows'][0][0]['value']==2 and calls[0]==calls[1],result
    sparse=root/'sparse.xlsx'; book=Workbook(); book.active['A1']='start'; book.active['XFD1048576']='far'; book.save(sparse)
    assert spreadsheet.read_range(sparse,'Sheet','XFD1048576')['rows'][0][0]['value']=='far'
    inaccurate=root/'inaccurate.xlsx'
    with zipfile.ZipFile(sparse) as before,zipfile.ZipFile(inaccurate,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='xl/worksheets/sheet1.xml': data=data.replace(b'<dimension ref="A1:XFD1048576"/>',b'<dimension ref="A1"/>')
            after.writestr(entry,data)
    assert spreadsheet.read_range(inaccurate,'Sheet','XFD1048576')['rows'][0][0]['value']=='far'
    longfile=root/'long.xlsx'; book=xlsxwriter.Workbook(str(longfile)); sheet=book.add_worksheet('Sheet'); long_text='x'*10000
    for row in range(150): sheet.write_string(row,0,long_text)
    book.close()
    assert spreadsheet.read_range(longfile,'Sheet','A1')['rows'][0][0]['value']==long_text
    try: spreadsheet.read_range(longfile,'Sheet','A1:A150')
    except ValueError as exc: assert 'response limit' in str(exc)
    else: raise AssertionError('Oversized response accepted')
    draft=root/'draft'; session.start(draft,source=longfile); cell=session.read_grid(draft,0,'Sheet','A1')['rows'][0][0]
    assert cell['value']==long_text and not cell['editable'] and cell['read_only_reason']=='value_too_long'
    try: session.edit_grid(draft,0,'Sheet','A1',[{'cell':'A1','value':'short'}])
    except ValueError as exc: assert 'value_too_long' in str(exc)
    else: raise AssertionError('Long cell edit accepted')
    assert session.info(draft)['latest_revision']==0
    emoji_file=root/'emoji.xlsx'; emoji=Workbook(); emoji.active['A1']='😀'*2000; emoji.active['B1']='😀'*2001; emoji.save(emoji_file)
    emoji_draft=root/'emoji-draft'; session.start(emoji_draft,source=emoji_file); emoji_grid=session.read_grid(emoji_draft,0,'Sheet','A1:B1')['rows'][0]
    assert emoji_grid[0]['editable'] and emoji_grid[1]['read_only_reason']=='value_too_long'
print('Range safety passed: snapshot consistency, sparse bounds, response limit and long cells')

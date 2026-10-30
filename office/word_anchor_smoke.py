"""Verify Word hyperlink and comment boundaries using a real DOCX."""
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

import session


def hyperlink(doc, paragraph, parts):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), doc.part.relate_to('https://example.com', RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    for index, value in enumerate(parts):
        run = OxmlElement('w:r'); props = OxmlElement('w:rPr'); props.append(OxmlElement('w:b' if index == 0 else 'w:i')); run.append(props)
        text = OxmlElement('w:t'); text.text = value; run.append(text); link.append(run)
    paragraph._p.append(link)


with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.docx'; draft = root / 'draft'; doc = Document()
    para = doc.add_paragraph('Before '); hyperlink(doc, para, ['linked', ' text']); para.add_run(' after')
    duplicate = doc.add_paragraph('same '); hyperlink(doc, duplicate, ['same'])
    comment_para = doc.add_paragraph('prefix '); commented = comment_para.add_run('commented'); comment_para.add_run(' suffix')
    doc.add_comment(commented, text='Keep this comment', author='Reviewer')
    tracked = doc.add_paragraph('tracked '); tracked._p.append(OxmlElement('w:ins'))
    doc.save(source); original = source.read_bytes(); session.start(draft, source=source)
    for path, find in [('/body/p[1]', 'Before linked'), ('/body/p[2]', 'same'), ('/body/p[3]', 'prefix commented'), ('/body/p[4]', 'tracked')]:
        try:
            session.replace_text(draft, 0, path, find, 'bad')
            raise AssertionError('Ambiguous or boundary-crossing replacement succeeded')
        except ValueError:
            pass
    assert session.info(draft)['latest_revision'] == 0
    edited = session.replace_text(draft, 0, '/body/p[1]', 'ked te', 'KED TE')
    paragraph = Document(edited['file']).paragraphs[0]; link = paragraph.hyperlinks[0]
    assert paragraph.text == 'Before linKED TExt after' and link.address == 'https://example.com'
    assert link.text == 'linKED TExt' and link.runs[0].bold and link.runs[1].italic
    commented_edit = session.replace_text(draft, 1, '/body/p[3]', 'commented', 'reviewed')
    result = Document(commented_edit['file'])
    assert result.paragraphs[2].text == 'prefix reviewed suffix' and next(iter(result.comments)).text == 'Keep this comment'
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(commented_edit['file']) as after:
        assert source.read_bytes() == original and set(before.namelist()) == set(after.namelist())
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name != 'word/document.xml')
        from lxml import etree
        old = etree.fromstring(before.read('word/document.xml')); new = etree.fromstring(after.read('word/document.xml'))
        for kind in ('commentRangeStart', 'commentRangeEnd', 'commentReference'):
            old_markers = old.xpath(f'.//w:{kind}', namespaces=ns); new_markers = new.xpath(f'.//w:{kind}', namespaces=ns)
            assert len(old_markers) == len(new_markers) == 1 and dict(old_markers[0].attrib) == dict(new_markers[0].attrib)
        paragraph = new.xpath('/w:document/w:body/w:p[3]', namespaces=ns)[0]
        children = list(paragraph); start = next(i for i, node in enumerate(children) if node.tag.endswith('}commentRangeStart'))
        end = next(i for i, node in enumerate(children) if node.tag.endswith('}commentRangeEnd'))
        assert ''.join(children[i].xpath('.//w:t/text()', namespaces=ns)[0] for i in range(start + 1, end) if children[i].xpath('.//w:t/text()', namespaces=ns)) == 'reviewed'
print('Word anchor smoke passed: hyperlinks, comments, boundaries and package preservation')

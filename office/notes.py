"""Read and replace slide speaker notes in saved PowerPoint files."""
import os
import re
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation

import slides
import ooxml


def _slide(source, number):
    if Path(source).suffix.lower() != '.pptx':
        raise ValueError('Speaker notes require a PPTX file')
    if Path(source).stat().st_size > 128 * 1024 * 1024: raise ValueError('Document exceeds the 128 MiB size limit')
    if not zipfile.is_zipfile(source): raise ValueError('Invalid PowerPoint package')
    with zipfile.ZipFile(source) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or len({e.filename for e in entries}) != len(entries) or sum(e.file_size for e in entries) > 128 * 1024 * 1024:
            raise ValueError('Document exceeds package limits or contains duplicate parts')
    if isinstance(number, bool) or not isinstance(number, int) or number < 1:
        raise ValueError('Slide number must be a positive integer')
    slides.slide_part(source, number - 1)  # Also rejects duplicate slide relationships.
    presentation = Presentation(source)
    return presentation, presentation.slides[number - 1]


def read(source, number):
    _, slide = _slide(source, number)
    if not slide.has_notes_slide:
        return {'slide': number, 'has_notes_slide': False, 'text': '', 'truncated': False}
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        raise ValueError('Notes slide has no speaker notes placeholder')
    text = frame.text
    return {'slide': number, 'has_notes_slide': True, 'text': text[:100000], 'truncated': len(text) > 100000}


def replace_text(source, output, number, paragraph, find, replacement):
    if isinstance(paragraph, bool) or not isinstance(paragraph, int) or paragraph < 1:
        raise ValueError('Paragraph number must be a positive integer')
    target = {}
    def part(file):
        presentation, slide = _slide(file, number)
        if not slide.has_notes_slide: raise ValueError('Slide has no speaker notes')
        note = slide.notes_slide; frame = note.notes_text_frame
        if frame is None: raise ValueError('Notes slide has no speaker notes placeholder')
        name = str(note.part.partname).lstrip('/')
        if sum(other.has_notes_slide and str(other.notes_slide.part.partname).lstrip('/') == name for other in presentation.slides) != 1:
            raise ValueError('Notes slide is referenced by multiple slides')
        shape = frame._element.getparent(); c_nv = shape.find(slides.P + 'nvSpPr/' + slides.P + 'cNvPr')
        if c_nv is None or not c_nv.get('id'): raise ValueError('Notes placeholder has no shape ID')
        target['id'] = c_nv.get('id')
        return name
    def edit(tree, archive):
        root = tree.getroot()
        if root.tag != slides.P + 'notes': raise ValueError('Invalid notes slide XML')
        parent = root.find(slides.P + 'cSld/' + slides.P + 'spTree')
        if parent is None or parent.find(slides.MC + 'AlternateContent') is not None:
            raise ValueError('Notes shape tree is missing or contains alternate content')
        shapes = [shape for shape in parent.findall(slides.P + 'sp') if (item := shape.find(slides.P + 'nvSpPr/' + slides.P + 'cNvPr')) is not None and item.get('id') == target['id']]
        if len(shapes) != 1: raise ValueError('Notes placeholder shape ID is ambiguous')
        body = shapes[0].find(slides.P + 'txBody')
        if body is None: raise ValueError('Notes placeholder has no text body')
        paragraphs = body.findall(slides.A + 'p')
        if paragraph > len(paragraphs): raise ValueError('Notes paragraph does not exist')
        selected = paragraphs[paragraph - 1]
        if any(child.tag not in (slides.A + 'pPr', slides.A + 'r', slides.A + 'endParaRPr') for child in selected):
            raise ValueError('Notes paragraph contains unsupported content such as fields or line breaks')
        nodes = []
        for run in selected.findall(slides.A + 'r'):
            if any(child.tag not in (slides.A + 'rPr', slides.A + 't') for child in run) or run.find('.//' + slides.A + 'hlinkClick') is not None or run.find('.//' + slides.A + 'hlinkMouseOver') is not None:
                raise ValueError('Notes paragraph contains unsupported run content or hyperlinks')
            nodes.extend(run.findall(slides.A + 't'))
        ooxml.replace_nodes(nodes, find, replacement, preserve_space=True)
    return ooxml.edit_xml(source, output, part, edit)


def replace(source, output, number, text):
    source = Path(source).resolve(strict=True); output = Path(output).absolute()
    if output.suffix.lower() != '.pptx': raise ValueError('Output must be a PPTX file')
    if not isinstance(text, str) or len(text) > 100000: raise ValueError('Notes must be text of at most 100,000 characters')
    if any(ord(c) < 32 and c not in '\n\t' for c in text): raise ValueError('Notes contain unsupported control characters')
    if source.stat().st_size > 128 * 1024 * 1024: raise ValueError('Document exceeds the 128 MiB size limit')
    if not zipfile.is_zipfile(source): raise ValueError('Invalid PowerPoint package')
    if output.exists() or output.is_symlink(): raise FileExistsError(f'Output already exists: {output}')
    with zipfile.ZipFile(source) as original:
        entries = original.infolist(); names = {entry.filename for entry in entries}
        if len(entries) > 5000 or len(names) != len(entries) or sum(e.file_size for e in entries) > 128 * 1024 * 1024:
            raise ValueError('Document exceeds package limits or contains duplicate parts')
        if any(name.startswith('_xmlsignatures/') for name in names): raise ValueError('Digitally signed documents are not supported')
        presentation, slide = _slide(source, number)
        slide_name = str(slide.part.partname).lstrip('/')
        had_notes = slide.has_notes_slide
        notes_slide = slide.notes_slide
        frame = notes_slide.notes_text_frame
        if frame is None: raise ValueError('Notes slide has no speaker notes placeholder')
        note_name = str(notes_slide.part.partname).lstrip('/')
        if sum(other.has_notes_slide and str(other.notes_slide.part.partname).lstrip('/') == note_name for other in presentation.slides) != 1:
            raise ValueError('Notes slide is referenced by multiple slides')
        frame.text = text
        with tempfile.TemporaryDirectory(dir=output.parent, prefix='.office-notes-') as tmp:
            generated = Path(tmp) / 'generated.pptx'; candidate = Path(tmp) / 'candidate.pptx'
            presentation.save(generated)
            with zipfile.ZipFile(generated) as edited:
                new_names = set(edited.namelist())
                if not names <= new_names: raise ValueError('Notes edit would remove package parts')
                allowed = {note_name}
                if not had_notes:
                    slide_rel = str(Path(slide_name).parent / '_rels' / (Path(slide_name).name + '.rels'))
                    allowed.update({'[Content_Types].xml', 'ppt/_rels/presentation.xml.rels', slide_rel})
                added = new_names - names
                if any(not re.fullmatch(r'ppt/(?:notesSlides|notesMasters)/(?:_rels/)?[^/]+|ppt/theme/theme[0-9]+\.xml', name) for name in added):
                    raise ValueError('Notes edit would add unexpected package parts')
                changed = {name for name in names if original.read(name) != edited.read(name)}
                if changed - allowed: raise ValueError('Notes edit would modify unrelated package parts')
                if len(new_names) > 5000 or sum(e.file_size for e in edited.infolist()) > 128 * 1024 * 1024:
                    raise ValueError('Edited package exceeds limits')
                with zipfile.ZipFile(candidate, 'w') as result:
                    result.comment = original.comment
                    for entry in entries: result.writestr(entry, edited.read(entry.filename) if entry.filename in changed else original.read(entry.filename))
                    for name in sorted(added): result.writestr(name, edited.read(name), compress_type=zipfile.ZIP_DEFLATED)
            os.link(candidate, output)
    return str(output)

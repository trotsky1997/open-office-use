"""Check reference-based validation choices and unsupported sources."""
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
import session

with TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.xlsx'; draft = root / 'draft'
    book = Workbook(); target = book.active; target.title = 'Input'
    options = book.create_sheet("Team's options"); options['A1'] = '001'; options['A2'] = '=text'; options['A2'].data_type = 's'
    options['B1'] = '=1+1'; options['C1'] = 3
    for cell, formula in [('A1', "'Team''s options'!$A$1:$A$2"), ('B1', "'Team''s options'!$B$1"), ('C1', "'Team''s options'!$C$1"), ('D1', 'Choices'), ('E1', '$A$1:$A$101'), ('F1', "'Missing'!$A$1")]:
        validation = DataValidation(type='list', formula1=formula); target.add_data_validation(validation); validation.add(cell)
    book.save(source); session.start(draft, source=source)
    grid = session.read_grid(draft, 0, 'Input', 'A1:F1')
    rules = {rule['cells'][0]: rule for rule in grid['data_validation']['rules']}
    assert rules['A1']['list_resolution'] == 'resolved' and rules['A1']['list_values'] == ['001', '=text']
    assert rules['A1']['list_source'] == {'sheet': "Team's options", 'range': '$A$1:$A$2'}
    assert [rules[cell]['list_resolution'] for cell in ['B1','C1','D1','E1','F1']] == ['unverified_values', 'unverified_values', 'unsupported_reference', 'range_limit', 'missing_sheet']
    saved = session.edit_grid(draft, 0, 'Input', 'A1:F1', [{'cell':'A1', 'value':'=text', 'type':'text'}])
    edited = load_workbook(saved['file']); assert edited['Input']['A1'].value == '=text' and edited['Input']['A1'].data_type == 's'
    assert edited['Input'].data_validations.dataValidation[0].formula1 == "'Team''s options'!$A$1:$A$2"
print('Reference validation smoke passed')

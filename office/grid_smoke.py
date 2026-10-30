"""Test bounded, atomic XLSX grid edits with real OfficeCLI."""
import tempfile
import json
import signal
import subprocess
import sys
import urllib.error
import urllib.request
import zipfile
from urllib.parse import urlsplit
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
import session
import review

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.xlsx'; draft = root / 'draft'
    workbook = Workbook(); sheet = workbook.active; sheet.title = 'Sales Data'
    sheet['A1'] = 'Name'; sheet['B1'] = 4; sheet['B2'] = '=B1*2'; sheet['C2'] = 'Remove'
    workbook.save(source); original = source.read_bytes()
    session.start(draft, source=source)
    assert session.sheet_names(draft, 0)['sheets'] == ['Sales Data']
    grid = session.read_grid(draft, 0, 'Sales Data', 'A1:C2')
    assert grid['rows'][1][1]['formula'] == '=B1*2' and grid['rows'][1][0]['value'] is None
    result = session.edit_grid(draft, 0, 'Sales Data', 'A1:C2', [
        {'cell': 'B1', 'value': '7'}, {'cell': 'B2', 'value': '=B1*3'}, {'cell': 'C2', 'value': ''}])
    changed = load_workbook(result['file']).active
    assert changed['A1'].value == 'Name' and changed['B1'].value == 7
    assert changed['B2'].value == '=B1*3' and changed['C2'].value is None
    assert source.read_bytes() == original and session.info(draft)['latest_revision'] == 1
    for changes in ([{'cell': 'D1', 'value': 'x'}], [{'cell': 'A1', 'value': 'x'}, {'cell': 'a1', 'value': 'y'}],
                    [{'cell': 'A1', 'value': '='}]):
        try:
            session.edit_grid(draft, 1, 'Sales Data', 'A1:C2', changes)
            raise AssertionError('Invalid change was accepted')
        except ValueError:
            pass
    try:
        session.edit_grid(draft, 0, 'Sales Data', 'A1:C2', [{'cell': 'A1', 'value': 'stale'}])
        raise AssertionError('Stale edit was accepted')
    except ValueError as error:
        assert 'conflict' in str(error)
    try:
        session.read_grid(draft, 1, 'Sales Data', 'A1:Z100')
        raise AssertionError('Oversized grid was accepted')
    except ValueError:
        pass
    assert session.info(draft)['latest_revision'] == 1
    array_source = root / 'array.xlsx'; array = Workbook(); array.active['A1'] = ArrayFormula(ref='A1:A2', text='=ROW(A1:A2)'); array.active['C1'] = DataTableFormula(ref='C1:C2'); array.save(array_source)
    array_draft = root / 'array-draft'; session.start(array_draft, source=array_source)
    assert not session.read_grid(array_draft, 0, 'Sheet', 'A1:B2')['rows'][0][0]['editable']
    assert not session.read_grid(array_draft, 0, 'Sheet', 'A2:B2')['rows'][0][0]['editable']
    assert not session.read_grid(array_draft, 0, 'Sheet', 'C2')['rows'][0][0]['editable']
    try:
        session.edit_grid(array_draft, 0, 'Sheet', 'A1:B2', [{'cell': 'A1', 'value': '7'}])
        raise AssertionError('Array formula edit was accepted')
    except ValueError as error:
        assert 'Array' in str(error)
    try:
        session.edit_grid(array_draft, 0, 'Sheet', 'A2:B2', [{'cell': 'A2', 'value': '7'}])
        raise AssertionError('Array range edit was accepted')
    except ValueError as error:
        assert 'Array' in str(error)
    shared_source = root / 'shared.xlsx'; shared = Workbook(); shared.active['B1'] = '=A1*2'; shared.active['B2'] = '=A2*2'; shared.save(shared_source)
    patched = root / 'shared-patched.xlsx'
    with zipfile.ZipFile(shared_source) as archive, zipfile.ZipFile(patched, 'w', zipfile.ZIP_DEFLATED) as output:
        for name in archive.namelist():
            data = archive.read(name)
            if name == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'<f>A1*2</f>', b'<f t="shared" si="0" ref="B1:B2">A1*2</f>')
                data = data.replace(b'<f>A2*2</f>', b'<f t="shared" si="0"/>')
            output.writestr(name, data)
    shared_draft = root / 'shared-draft'; session.start(shared_draft, source=patched)
    assert not session.read_grid(shared_draft, 0, 'Sheet', 'B2')['rows'][0][0]['editable']
    try:
        session.edit_grid(shared_draft, 0, 'Sheet', 'B2', [{'cell': 'B2', 'value': '9'}])
        raise AssertionError('Shared formula range edit was accepted')
    except ValueError as error:
        assert 'shared' in str(error)
    format_draft = root / 'format-draft'; session.start(format_draft, source=source)
    formatted = session.format_grid(format_draft, 0, 'Sales Data', 'A1:C2', 'B1:B2', {'bold': True, 'fill': 'FFCC00', 'numberformat': '0.00'})
    cells = load_workbook(formatted['file']).active
    assert cells['B1'].value == 4 and cells['B2'].value == '=B1*2'
    assert cells['B1'].font.bold and cells['B2'].fill.fgColor.rgb[-6:] == 'FFCC00' and cells['B1'].number_format == '0.00'
    diff = review.compare_cells(source, formatted['file'], 'Sales Data', 'A1:C2')
    assert {entry['cell'] for entry in diff['changes']} == {'B1', 'B2'}
    assert {'format.font', 'format.fill', 'format.number_format'} <= set(diff['changes'][0]['changed_fields'])
    style = session.read_grid(format_draft, 1, 'Sales Data', 'B1')['rows'][0][0]['style']
    assert style == {'bold': True, 'fill_rgb': 'FFCC00', 'number_format': '0.00'}
    for base, target, properties in [(0, 'B1', {'bold': True}), (1, 'D1', {'bold': True}), (1, 'B1', {'fill': 'evil'}), (1, 'B1', {'value': 'bad'})]:
        try:
            session.format_grid(format_draft, base, 'Sales Data', 'A1:C2', target, properties)
            raise AssertionError('Invalid formatting accepted')
        except ValueError:
            pass
    try:
        session.format_grid(array_draft, 0, 'Sheet', 'A2:B2', 'A2', {'bold': True})
        raise AssertionError('Protected formula formatting accepted')
    except ValueError:
        pass
    cli = Path(__file__).with_name('open_office_use.py')
    server = subprocess.Popen([sys.executable, str(cli), 'watch', str(draft)], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    try:
        url = json.loads(server.stdout.readline())['url']
        origin = f"http://{urlsplit(url).netloc}"
        def post(route, payload, request_origin=origin):
            request = urllib.request.Request(url + route, data=json.dumps(payload).encode(), headers={'Origin': request_origin, 'Content-Type': 'application/json'})
            with urllib.request.urlopen(request, timeout=60) as response:
                return json.loads(response.read())
        assert post('sheets', {'base_revision': 1})['sheets'] == ['Sales Data']
        assert post('grid', {'base_revision': 1, 'sheet': 'Sales Data', 'cell_range': 'A1:C2'})['rows'][1][1]['formula'] == '=B1*3'
        payload = {'base_revision': 1, 'sheet': 'Sales Data', 'cell_range': 'A1:C2', 'changes': [{'cell': 'A1', 'value': 'Updated'}]}
        try:
            post('grid-edit', payload, 'http://evil.test')
            raise AssertionError('Cross-origin edit was accepted')
        except urllib.error.HTTPError as error:
            assert error.code == 403
        assert post('grid-format', {'base_revision': 1, 'sheet': 'Sales Data', 'cell_range': 'A1:C2', 'style_range': 'B1', 'properties': {'bold': True}})['revision'] == 2
        payload['base_revision'] = 2
        assert post('grid-edit', payload)['revision'] == 3
        assert load_workbook(draft / 'revision-000003.xlsx').active['A1'].value == 'Updated'
        request = {'before': 0, 'after': 3, 'sheet': 'Sales Data', 'cell_range': 'A1:C2'}
        comparison = post('cell-comparison', request)
        assert comparison['changed_count'] == 4 and any('format.font' in change['changed_fields'] for change in comparison['changes']), comparison
        for bad, origin_value in [(dict(request, after=2), origin), (dict(request, before=True), origin), (request, 'http://evil.test')]:
            try:
                post('cell-comparison', bad, origin_value)
                raise AssertionError('Invalid cell comparison was accepted')
            except urllib.error.HTTPError as error:
                assert error.code == (403 if origin_value != origin else 400)

        try:
            post('grid-edit', payload)
            raise AssertionError('Stale grid edit was accepted')
        except urllib.error.HTTPError as error:
            assert error.code == 400
    finally:
        server.send_signal(signal.SIGINT); server.wait(timeout=10)
print('Grid smoke passed: multiple cells, formulas, clear, range and revision safety')

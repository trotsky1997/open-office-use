"""Check Word comment insertion and package preservation."""
import tempfile
import json
import subprocess
import sys
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import officecli_backend
import session
import word

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.docx'; draft = root / 'draft'
    doc = Document(); para = doc.add_paragraph(); para.add_run('Important ').bold = True; para.add_run('text').italic = True
    table = doc.add_table(rows=1, cols=1); table.cell(0, 0).text = 'Table value'
    tracked = doc.add_paragraph('Track '); inserted = OxmlElement('w:ins'); inserted.set(qn('w:id'), '2'); inserted.set(qn('w:author'), 'Author'); tracked._p.append(inserted)
    doc.save(source); original = source.read_bytes(); session.start(draft, source=source)
    first = session.add_comment(draft, 0, '/body/p[1]', 'Please review\nSecond line', 'Reviewer')
    assert first['comment_id'] == 0
    result = Document(first['file']); comments = list(result.comments)
    assert len(comments) == 1 and comments[0].author == 'Reviewer' and comments[0].text == 'Please review\nSecond line'
    assert result.paragraphs[0].text == 'Important text' and result.paragraphs[0].runs[0].bold and result.paragraphs[0].runs[1].italic
    assert source.read_bytes() == original
    validation = officecli_backend.audit(first['file'])['validation']; assert validation['exit_code'] == 0, validation
    with zipfile.ZipFile(source) as a, zipfile.ZipFile(first['file']) as b:
        changed = {'word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'}
        assert set(b.namelist()) == set(a.namelist()) | {'word/comments.xml'}
        assert all(a.read(name) == b.read(name) for name in a.namelist() if name not in changed)
    second = session.add_comment(draft, 1, '/body/table[1]/row[1]/cell[1]/p[1]', 'Review table', 'Editor')
    assert second['comment_id'] == 1
    listing = session.list_comments(draft, 2)
    assert listing['total_comments'] == 2 and listing['anchor_scan_complete'], listing
    assert listing['comments'][0]['text'] == 'Please review\nSecond line' and listing['comments'][0]['anchor_text'] == 'Important text', listing
    assert listing['comments'][0]['paragraph_paths'] == ['/body/p[1]'] and listing['comments'][0]['anchor_status'] == 'range', listing
    assert listing['comments'][1]['paragraph_paths'] == ['/body/tbl[1]/tr[1]/tc[1]/p[1]'] and listing['comments'][1]['anchor_text'] == 'Table value', listing
    assert word.list_comments(source)['total_comments'] == 0
    comments = list(Document(second['file']).comments)
    assert len(comments) == 2 and comments[1].text == 'Review table'
    for base, path in [(1, '/body/p[1]'), (2, '/body/p[1]'), (2, '/body/p[2]')]:
        try: session.add_comment(draft, base, path, 'rejected', 'Reviewer')
        except ValueError: pass
        else: raise AssertionError('Stale or unsupported paragraph accepted')
    multi_draft = root / 'multi-draft'; session.start(multi_draft, source=second['file'])
    session.resolve_comment(multi_draft, 0, '0', True)
    session.resolve_comment(multi_draft, 1, '1', True)
    multi = session.list_comments(multi_draft, 2)
    assert [comment['resolved'] for comment in multi['comments']] == [True, True], multi
    assert session.info(draft)['latest_revision'] == 2
    resolve_draft = root / 'resolve-draft'; session.start(resolve_draft, source=first['file'])
    loaded_comments = session.list_comments(resolve_draft,0)
    try: session.resolve_comment(resolve_draft,0,'0',True,expected_sha256='0'*64)
    except ValueError as exc: assert 'snapshot changed' in str(exc)
    else: raise AssertionError('Stale comment resolution accepted')
    resolved = session.resolve_comment(resolve_draft, 0, '0', True, expected_sha256=loaded_comments['snapshot_sha256'])
    resolved_data = session.list_comments(resolve_draft, 1)
    assert resolved_data['comments'][0]['resolved'] is True and resolved_data['comments'][0]['reply_parent_status'] == 'root', resolved_data
    assert Document(resolved['file']).paragraphs[0].text == 'Important text'
    reopened = session.resolve_comment(resolve_draft, 1, '0', False)
    assert session.list_comments(resolve_draft, 2)['comments'][0]['resolved'] is False
    with zipfile.ZipFile(resolved['file']) as a, zipfile.ZipFile(reopened['file']) as b:
        assert set(a.namelist()) == set(b.namelist())
        assert all(a.read(name) == b.read(name) for name in a.namelist() if name != 'word/commentsExtended.xml')
    validation = officecli_backend.audit(resolved['file'])['validation']; assert validation['exit_code'] == 0, validation
    try: session.resolve_comment(resolve_draft, 0, '0', False)
    except ValueError as exc: assert 'conflict' in str(exc).lower()
    else: raise AssertionError('Stale comment resolution succeeded')

    external = Document(); old_paragraph = external.add_paragraph(); old_run = old_paragraph.add_run('Already reviewed')
    external.add_comment(old_run, text='Existing comment', author='Other'); external.add_paragraph('Another paragraph')
    external_file = root / 'existing.docx'; external.save(external_file); external_draft = root / 'existing-draft'; session.start(external_draft, source=external_file)
    external_result = session.add_comment(external_draft, 0, '/body/p[2]', 'New comment', 'Reviewer')
    assert [item.text for item in Document(external_result['file']).comments] == ['Existing comment', 'New comment']
    external_listing = word.list_comments(external_result['file']); assert external_listing['comments'][0]['anchor_text'] == 'Already reviewed', external_listing
    spanning = Document(); p1 = spanning.add_paragraph('First anchor'); p2 = spanning.add_paragraph('Second anchor')
    spanning.add_comment([p1.runs[0], p2.runs[0]], text='Cross paragraph', author='Reviewer')
    span_file = root / 'span.docx'; spanning.save(span_file); item = word.list_comments(span_file)['comments'][0]
    assert item['anchor_text'] == 'First anchor\nSecond anchor' and item['paragraph_paths'] == ['/body/p[1]', '/body/p[2]'], item
    assert item['anchor_status'] == 'range'
    span_draft = root / 'span-draft'; session.start(span_draft, source=span_file)
    session.reply_comment(span_draft, 0, '0', 'Reply to both', 'Editor')
    span_reply = session.list_comments(span_draft, 1)['comments'][1]
    assert span_reply['anchor_text'] == 'First anchor\nSecond anchor' and span_reply['reply_to_comment_id'] == '0', span_reply

    many = Document()
    for number in range(201):
        paragraph = many.add_paragraph(str(number)); many.add_comment(paragraph.runs[0], text=f'Comment {number}', author='Reviewer')
    many_file = root / 'many.docx'; many.save(many_file); many_result = word.list_comments(many_file)
    assert many_result['total_comments'] == 201 and len(many_result['comments']) == 200 and many_result['truncated']
    cli = Path(__file__).with_name('open_office_use.py')
    listed_cli = subprocess.run([sys.executable, str(cli), 'session', 'comments', str(draft), '--revision', '2'], check=True, capture_output=True, text=True)
    assert json.loads(listed_cli.stdout)['total_comments'] == 2
    reply_draft = root / 'reply-draft'; session.start(reply_draft, source=first['file'])
    live = reply_draft/'revision-000000.docx'; saved_live = live.read_bytes(); list_impl = word._list_comments
    def mutate_live(frozen):
        live.write_bytes(source.read_bytes())
        return list_impl(frozen)
    with patch.object(word, '_list_comments', side_effect=mutate_live): loaded_comments = session.list_comments(reply_draft,0)
    assert loaded_comments['total_comments'] == 1
    try: session.reply_comment(reply_draft,0,'0','Reply','Editor',expected_sha256=loaded_comments['snapshot_sha256'])
    except ValueError as exc: assert 'snapshot changed' in str(exc)
    else: raise AssertionError('Changed source accepted after comment read')
    live.write_bytes(saved_live)
    try: session.reply_comment(reply_draft,0,'0','Reply','Editor',expected_sha256='0'*64)
    except ValueError as exc: assert 'snapshot changed' in str(exc)
    else: raise AssertionError('Stale comment reply accepted')
    reply = session.reply_comment(reply_draft, 0, '0', 'Thanks for reviewing\nWill revise', 'Editor', expected_sha256=loaded_comments['snapshot_sha256'])
    assert reply['comment_id'] == 1 and reply['parent_id'] == '0'
    reply_listing = session.list_comments(reply_draft, 1); root_item, reply_item = reply_listing['comments']
    assert reply_item['reply_to_comment_id'] == '0' and reply_item['anchor_text'] == root_item['anchor_text'] == 'Important text', reply_listing
    assert reply_item['text'] == 'Thanks for reviewing\nWill revise'
    assert Document(reply['file']).paragraphs[0].runs[0].bold and Document(reply['file']).paragraphs[0].runs[1].italic
    validation = officecli_backend.audit(reply['file'])['validation']; assert validation['exit_code'] == 0, validation
    repeated_draft = root / 'repeat-draft'; session.start(repeated_draft, source=reply['file'])
    session.reply_comment(repeated_draft, 0, '0', 'Another reply', 'Editor')
    repeated = session.list_comments(repeated_draft, 1)['comments']
    assert [item['reply_to_comment_id'] for item in repeated] == [None, '0', '0']
    assert all(item['anchor_text'] == 'Important text' for item in repeated)
    session.resolve_comment(reply_draft, 1, '0', True)
    for base, parent in [(0, '0'), (2, '0'), (2, '1')]:
        try: session.reply_comment(reply_draft, base, parent, 'No', 'Editor')
        except ValueError: pass
        else: raise AssertionError('Stale, resolved or nonroot reply accepted')
    assert session.info(reply_draft)['latest_revision'] == 2
    cli_reply_draft = root / 'cli-reply'; session.start(cli_reply_draft, source=first['file'])
    subprocess.run([sys.executable, str(Path(__file__).with_name('open_office_use.py')), 'session', 'reply-comment', str(cli_reply_draft), '--base', '0', '--parent-id', '0', '--text', 'CLI reply', '--author', 'Editor'], check=True, capture_output=True)
    assert session.list_comments(cli_reply_draft, 1)['comments'][1]['reply_to_comment_id'] == '0'
    cli_state_draft = root / 'cli-state'; session.start(cli_state_draft, source=first['file'])
    subprocess.run([sys.executable, str(Path(__file__).with_name('open_office_use.py')), 'session', 'resolve-comment', str(cli_state_draft), '--base', '0', '--id', '0', '--resolved', 'true'], check=True, capture_output=True)
    assert session.list_comments(cli_state_draft, 1)['comments'][0]['resolved'] is True
    cli_draft = root / 'cli-draft'; session.start(cli_draft, source=source)
    subprocess.run([sys.executable, str(Path(__file__).with_name('open_office_use.py')), 'session', 'add-comment', str(cli_draft), '--base', '0', '--path', '/body/p[1]', '--text', 'CLI review', '--author', 'Editor'], check=True, capture_output=True)
    assert next(iter(Document(cli_draft / 'revision-000001.docx').comments)).text == 'CLI review'
print('Comment smoke passed: new and existing comments, tables, formatting and conflicts')

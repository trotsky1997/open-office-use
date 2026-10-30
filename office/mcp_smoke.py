"""Check the stdio MCP lifecycle and real artifact operations."""
import asyncio
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from pptx import Presentation
from pptx.util import Inches
from session import set_selection
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client


async def main():
    cli = Path(__file__).with_name('open_office_use.py')
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        async with stdio_client(StdioServerParameters(command=sys.executable, args=[str(cli), 'mcp'])) as (read, write):
            async with ClientSession(read, write) as session:
                await session.initialize()
                tools = await session.list_tools()
                assert {t.name for t in tools.tools} == {'inspect_file', 'structured_read', 'apply_batch', 'render_pdf', 'preview_pages', 'view_page', 'read_cells', 'recalculate_workbook', 'compare_files', 'audit_file', 'start_session', 'session_info', 'session_apply', 'session_export', 'session_selection', 'session_selection_image', 'session_selection_text', 'session_selection_candidates', 'session_restore', 'session_edit_target', 'session_edit', 'session_replace_text', 'compare_cells', 'session_sheets', 'session_read_grid', 'session_edit_grid', 'session_format_grid', 'scan_formulas'}
                output = root / 'report.docx'
                operations = [{'command': 'add', 'parent': '/body', 'type': 'paragraph', 'props': {'text': 'MCP works'}}]
                created = await session.call_tool('apply_batch', {'output': str(output), 'operations': operations})
                assert not created.isError and output.exists(), created
                readback = await session.call_tool('structured_read', {'file': str(output), 'mode': 'get', 'target': '/body/p[1]'})
                assert not readback.isError and 'MCP works' in str(readback), readback
                revised = root / 'revised.docx'
                edit = await session.call_tool('apply_batch', {'source': str(output), 'output': str(revised), 'operations': [{'command': 'set', 'path': '/body/p[1]', 'props': {'text': 'Revised'}}]})
                assert not edit.isError, edit
                compared = await session.call_tool('compare_files', {'before': str(output), 'after': str(revised)})
                assert not compared.isError and 'Revised' in str(compared.structuredContent['text_diff']), compared
                audited = await session.call_tool('audit_file', {'file': str(revised)})
                assert not audited.isError and audited.structuredContent['validation']['result']['success'] is True, audited
                word_draft = root / 'word-draft'
                word_started = await session.call_tool('start_session', {'directory': str(word_draft), 'source': str(output)})
                assert not word_started.isError, word_started
                word_edit = await session.call_tool('session_replace_text', {'directory': str(word_draft), 'base_revision': 0, 'path': '/body/p[1]', 'find': 'MCP works', 'replacement': 'MCP changed'})
                assert not word_edit.isError and word_edit.structuredContent['revision'] == 1, word_edit
                table_source = root / 'table.docx'; table_doc = Document()
                table_doc.add_table(rows=1, cols=1).cell(0, 0).text = 'table content'; table_doc.save(table_source)
                table_draft = root / 'table-draft'
                table_started = await session.call_tool('start_session', {'directory': str(table_draft), 'source': str(table_source)})
                assert not table_started.isError, table_started
                table_edit = await session.call_tool('session_replace_text', {'directory': str(table_draft), 'base_revision': 0, 'path': '/body/tbl[1]/tr[1]/tc[1]/p[1]', 'find': 'content', 'replacement': 'result'})
                assert not table_edit.isError and Document(table_draft / 'revision-000001.docx').tables[0].cell(0, 0).text == 'table result', table_edit
                slide_source = root / 'slides.pptx'; deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
                slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = 'slide content'; deck.save(slide_source)
                slide_draft = root / 'slide-draft'
                slide_started = await session.call_tool('start_session', {'directory': str(slide_draft), 'source': str(slide_source)})
                assert not slide_started.isError, slide_started
                slide_edit = await session.call_tool('session_replace_text', {'directory': str(slide_draft), 'base_revision': 0, 'path': '/slide[1]/shape[1]/paragraph[1]', 'find': 'content', 'replacement': 'result'})
                assert not slide_edit.isError and Presentation(slide_draft / 'revision-000001.pptx').slides[0].shapes[0].text == 'slide result', slide_edit
                link_source = root / 'link.docx'; linked = Document(); paragraph = linked.add_paragraph('Read ')
                hyperlink = OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'), linked.part.relate_to('https://example.com', RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                run = OxmlElement('w:r'); content = OxmlElement('w:t'); content.text = 'linked text'; run.append(content); hyperlink.append(run); paragraph._p.append(hyperlink); linked.save(link_source)
                link_draft = root / 'link-draft'; link_started = await session.call_tool('start_session', {'directory': str(link_draft), 'source': str(link_source)})
                assert not link_started.isError, link_started
                link_edit = await session.call_tool('session_replace_text', {'directory': str(link_draft), 'base_revision': 0, 'path': '/body/p[1]', 'find': 'linked', 'replacement': 'updated'})
                assert not link_edit.isError, link_edit
                linked_result = Document(link_draft / 'revision-000001.docx'); link_node = linked_result.paragraphs[0]._p.find(qn('w:hyperlink'))
                assert linked_result.part.rels[link_node.get(qn('r:id'))].target_ref == 'https://example.com', link_edit
                directory = root / 'draft'
                started = await session.call_tool('start_session', {'directory': str(directory), 'source': str(output)})
                assert not started.isError and started.structuredContent['latest_revision'] == 0, started
                updated = await session.call_tool('session_apply', {'directory': str(directory), 'base_revision': 0, 'operations': [{'command': 'set', 'path': '/body/p[1]', 'props': {'text': 'Session edit'}}]})
                assert not updated.isError and updated.structuredContent['revision'] == 1, updated
                set_selection(directory, {'revision': 1, 'page': 1, 'box': [0, 0, 1, 1]})
                selected = await session.call_tool('session_selection', {'directory': str(directory)})
                assert not selected.isError and selected.structuredContent['stale'] is False, selected
                selected_text = await session.call_tool('session_selection_text', {'directory': str(directory)})
                assert not selected_text.isError and 'Session edit' in selected_text.structuredContent['text'], selected_text
                matches = await session.call_tool('session_selection_candidates', {'directory': str(directory)})
                assert not matches.isError and matches.structuredContent['requires_verification'] and matches.structuredContent['candidates'], matches
                cropped = await session.call_tool('session_selection_image', {'directory': str(directory)})
                assert not cropped.isError and cropped.content[0].type == 'image', cropped
                advanced = await session.call_tool('session_apply', {'directory': str(directory), 'base_revision': 1, 'operations': []})
                assert not advanced.isError, advanced
                stale_selection = await session.call_tool('session_selection', {'directory': str(directory)})
                assert not stale_selection.isError and stale_selection.structuredContent['stale'] is True, stale_selection
                restored = await session.call_tool('session_restore', {'directory': str(directory), 'base_revision': 2, 'revision': 0})
                assert not restored.isError and restored.structuredContent['revision'] == 3, restored
                outdated_restore = await session.call_tool('session_restore', {'directory': str(directory), 'base_revision': 2, 'revision': 0})
                assert outdated_restore.isError, outdated_restore
                target = await session.call_tool('session_edit_target', {'directory': str(directory), 'base_revision': 3, 'path': '/body/p[1]'})
                assert not target.isError and target.structuredContent['value'] == 'MCP works', target
                edited = await session.call_tool('session_edit', {'directory': str(directory), 'base_revision': 3, 'path': target.structuredContent['path'], 'mode': 'text', 'value': 'Edited via MCP'})
                assert not edited.isError and edited.structuredContent['revision'] == 4, edited

                conflict = await session.call_tool('session_apply', {'directory': str(directory), 'base_revision': 0, 'operations': []})
                assert conflict.isError and 'conflict' in str(conflict).lower(), conflict
                exported = await session.call_tool('session_export', {'directory': str(directory), 'revision': 1, 'output': str(root / 'session.docx')})
                assert not exported.isError and (root / 'session.docx').exists(), exported
                invalid = await session.call_tool('apply_batch', {'output': str(output), 'operations': operations})
                assert invalid.isError and 'already exists' in str(invalid), invalid
                images = await session.call_tool('preview_pages', {'file': str(output), 'output_dir': str(root / 'pages')})
                assert not images.isError, images
                paths = images.structuredContent['pages']
                assert paths and Path(paths[0]).read_bytes().startswith(b'\x89PNG\r\n\x1a\n')
                image = await session.call_tool('view_page', {'file': paths[0]})
                assert not image.isError and image.content[0].type == 'image' and image.content[0].mimeType == 'image/png', image
                workbook = Workbook(); workbook.active['A1'] = 4; workbook.active['A2'] = '=A1*3'
                original = root / 'numbers.xlsx'; workbook.save(original)
                calculated = root / 'calculated.xlsx'
                calc = await session.call_tool('recalculate_workbook', {'file': str(original), 'output': str(calculated)})
                assert not calc.isError, calc
                grid_draft = root / 'grid-draft'
                grid_started = await session.call_tool('start_session', {'directory': str(grid_draft), 'source': str(original)})
                assert not grid_started.isError, grid_started
                sheets = await session.call_tool('session_sheets', {'directory': str(grid_draft), 'base_revision': 0})
                assert not sheets.isError and sheets.structuredContent['sheets'] == ['Sheet'], sheets
                grid = await session.call_tool('session_read_grid', {'directory': str(grid_draft), 'base_revision': 0, 'sheet': 'Sheet', 'cell_range': 'A1:A2'})
                assert not grid.isError and grid.structuredContent['rows'][1][0]['formula'] == '=A1*3', grid
                saved_grid = await session.call_tool('session_edit_grid', {'directory': str(grid_draft), 'base_revision': 0, 'sheet': 'Sheet', 'cell_range': 'A1:A2', 'changes': [{'cell': 'A1', 'value': '5'}]})
                assert not saved_grid.isError and saved_grid.structuredContent['revision'] == 1, saved_grid
                styled = await session.call_tool('session_format_grid', {'directory': str(grid_draft), 'base_revision': 1, 'sheet': 'Sheet', 'cell_range': 'A1:A2', 'style_range': 'A1', 'properties': {'bold': True}})
                assert not styled.isError and styled.structuredContent['revision'] == 2, styled
                format_diff = await session.call_tool('compare_cells', {'before': str(grid_draft / 'revision-000001.xlsx'), 'after': str(grid_draft / 'revision-000002.xlsx'), 'sheet': 'Sheet', 'cell_range': 'A1:A2'})
                assert not format_diff.isError and 'format.font' in format_diff.structuredContent['changes'][0]['changed_fields'], format_diff
                scanned = await session.call_tool('scan_formulas', {'file': str(calculated)})
                assert not scanned.isError and scanned.structuredContent['formulas'] == 1 and scanned.structuredContent['complete'], scanned
                cells = await session.call_tool('read_cells', {'file': str(calculated), 'sheet': 'Sheet', 'cell_range': 'A1:A2'})
                assert not cells.isError and cells.structuredContent['rows'][1][0]['value'] == 12, cells
                diff = await session.call_tool('compare_cells', {'before': str(original), 'after': str(calculated), 'sheet': 'Sheet', 'cell_range': 'A1:A2'})
                assert not diff.isError and any(item['cell'] == 'A2' and item['after']['value'] == 12 for item in diff.structuredContent['changes']), diff

                wrong = await session.call_tool('structured_read', {'file': str(output), 'mode': 'get'})
                assert wrong.isError, wrong
    print('MCP stdio smoke passed: discovery, authoring, read, preview and errors')


if __name__ == '__main__':
    asyncio.run(main())

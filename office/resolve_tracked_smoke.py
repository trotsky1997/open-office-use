"""Exercise atomic acceptance and rejection of selected Word run revisions."""
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import officecli_backend
import session

with tempfile.TemporaryDirectory() as temp:
    root = Path(temp); source = root / 'source.docx'; draft = root / 'draft'; doc = Document()
    p = doc.add_paragraph('Start '); p.add_run('old').bold = True; p.add_run(' text').italic = True; p.add_run(' end'); doc.save(source)
    session.start(draft, source=source); changed = session.track_replace_text(draft, 0, '/body/p[1]', 'old text', 'new', 'Reviewer')
    changed_file = Path(changed['file']); records = session.tracked_changes(draft, 1)['records']; ids = [item['id'] for item in records]
    for accept, expected in ((True, 'Start new end'), (False, 'Start old text end')):
        isolated = root / ('accept' if accept else 'reject'); session.start(isolated, source=changed_file)
        for bad_ids in ([ids[0], '999999999999999999'], [ids[0], ids[0]]):
            try: session.resolve_tracked_runs(isolated, 0, 'word/document.xml', bad_ids, accept)
            except ValueError: pass
            else: raise AssertionError('invalid revision selection accepted')
            assert session.info(isolated)['latest_revision'] == 0
        result = session.resolve_tracked_runs(isolated, 0, 'word/document.xml', ids, accept)
        paragraph = Document(result['file']).paragraphs[0]
        assert paragraph.text == expected and paragraph.runs[1].bold is True
        if not accept: assert paragraph.runs[2].italic is True
        assert session.tracked_changes(isolated, 1)['records_scanned'] == 0
        with zipfile.ZipFile(changed_file) as before, zipfile.ZipFile(result['file']) as after:
            assert all(before.read(name) == after.read(name) for name in before.namelist() if name != 'word/document.xml')
        audit = officecli_backend.audit(result['file']); assert audit['validation']['result']['success'], audit
    cli_draft = root / 'cli'; session.start(cli_draft, source=changed_file); cli = Path(__file__).with_name('open_office_use.py')
    out = subprocess.run([sys.executable, str(cli), 'session', 'resolve-tracked-runs', str(cli_draft), '--base', '0', '--part', 'word/document.xml', '--id', ids[0], '--id', ids[1], '--decision', 'accept'], capture_output=True, text=True, check=True)
    assert json.loads(out.stdout)['revision'] == 1
    marked = root / 'paragraph-mark.docx'; mark_doc = Document(); paragraph = mark_doc.add_paragraph('text')
    props = paragraph._p.get_or_add_pPr(); run_props = OxmlElement('w:rPr'); marker = OxmlElement('w:del'); marker.set(qn('w:id'), '3'); run_props.append(marker); props.append(run_props); mark_doc.save(marked)
    marker_draft = root / 'marker'; session.start(marker_draft, source=marked)
    try: session.resolve_tracked_runs(marker_draft, 0, 'word/document.xml', ['3'], True)
    except ValueError: pass
    else: raise AssertionError('paragraph mark revision accepted as a run')
    assert session.info(marker_draft)['latest_revision'] == 0
    header_file = root / 'header.docx'; header_doc = Document(); header = header_doc.sections[0].header.paragraphs[0]
    insertion = OxmlElement('w:ins'); insertion.set(qn('w:id'), '10'); run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'header text'; run.append(text); insertion.append(run); header._p.append(insertion); header_doc.save(header_file)
    header_draft = root / 'header'; session.start(header_draft, source=header_file)
    header_record = session.tracked_changes(header_draft, 0)['records'][0]
    header_result = session.resolve_tracked_runs(header_draft, 0, header_record['part'], ['10'], True)
    assert Document(header_result['file']).sections[0].header.paragraphs[0].text == 'header text'
    move_file = root / 'move.docx'; move_doc = Document(); move_paragraph = move_doc.add_paragraph()
    marker = OxmlElement('w:moveFromRangeStart'); marker.set(qn('w:id'), '2'); move_paragraph._p.addprevious(marker)
    insertion = OxmlElement('w:ins'); insertion.set(qn('w:id'), '3'); run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'text'; run.append(text); insertion.append(run); move_paragraph._p.append(insertion); move_doc.save(move_file)
    move_draft = root / 'move'; session.start(move_draft, source=move_file)
    try: session.resolve_tracked_runs(move_draft, 0, 'word/document.xml', ['3'], True)
    except ValueError as exc: assert 'ranges' in str(exc)
    else: raise AssertionError('tracked move range was accepted')
    assert session.info(move_draft)['latest_revision'] == 0
    duplicate_file = root / 'duplicate.docx'; duplicate_doc = Document(); p = duplicate_doc.add_paragraph()
    for raw_id in ('3', '0003'):
        insertion = OxmlElement('w:ins'); insertion.set(qn('w:id'), raw_id); run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'x'; run.append(text); insertion.append(run); p._p.append(insertion)
    duplicate_doc.save(duplicate_file); duplicate_draft = root / 'duplicate'; session.start(duplicate_draft, source=duplicate_file)
    try: session.resolve_tracked_runs(duplicate_draft, 0, 'word/document.xml', ['3'], True)
    except ValueError as exc: assert 'ambiguous' in str(exc)
    else: raise AssertionError('ambiguous IDs were accepted')
    assert session.info(duplicate_draft)['latest_revision'] == 0
print('Resolve tracked runs smoke passed: accept, reject, formatting, atomic failures and paragraph marks')

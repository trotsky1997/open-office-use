"""Verify rendered page order, pagination and safe output handling."""
import io
import hashlib
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

import open_office_use
import visual

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    deck = Presentation()
    for color in ((255, 0, 0), (0, 255, 0), (0, 0, 255), (255, 255, 0)):
        slide = deck.slides.add_slide(deck.slide_layouts[6]); fill = slide.background.fill
        fill.solid(); fill.fore_color.rgb = RGBColor(*color)
    marker = deck.slides[1].shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(1.5), Inches(1), Inches(.75))
    marker.fill.solid(); marker.fill.fore_color.rgb = RGBColor(255, 0, 0); marker.line.fill.background()
    for _ in range(8): deck.slides.add_slide(deck.slide_layouts[6])
    source = root / 'slides.pptx'; deck.save(source)
    data, meta = visual.contact_sheet(source, 2, 2, open_office_use.render)
    assert meta == {'start_page': 2, 'end_page': 3, 'total_pages': 12, 'render_advisory': None, 'snapshot_sha256': hashlib.sha256(source.read_bytes()).hexdigest()}, meta
    sheet = Image.open(io.BytesIO(data)); assert sheet.size == (776, 360)
    assert sheet.getpixel((196, 170))[1] > 200 and sheet.getpixel((576, 170))[2] > 200
    assert visual.contact_sheet(source, 10, 12, open_office_use.render)[1]['end_page'] == 12
    page, page_meta = visual.page_image(source, 2, open_office_use.render, [0, 0, .5, .5])
    image = Image.open(io.BytesIO(page))
    assert page_meta['page'] == 2 and page_meta['total_pages'] == 12 and image.width == 1800 and page_meta['raster_scale'] == 3600
    assert image.getpixel((image.width * 4 // 5, image.height * 4 // 5))[1] > 200
    detail, detail_meta = visual.page_image(source, 2, open_office_use.render, [.2, .2, .3, .3])
    detail_image = Image.open(io.BytesIO(detail))
    assert detail_meta['raster_scale'] == 7200 and 715 <= detail_image.width <= 725
    for x, y in ((5, 5), (detail_image.width//2, detail_image.height//2), (detail_image.width-5, detail_image.height-5)):
        red, green, blue = detail_image.getpixel((x, y))[:3]; assert red > 240 and green < 20 and blue < 20, (x, y, red, green, blue)
    assert visual.page_image(source, 2, open_office_use.render, [0, 0, 5e-324, 5e-324])[1]['raster_scale'] == 7200
    for invalid in ([0, 0, 0, 1], [0, 0, 1.01, 1], [0, 0, float('nan'), 1], [False, 0, 1, 1]):
        try: visual.page_image(source, 1, open_office_use.render, invalid)
        except ValueError: pass
        else: raise AssertionError('Invalid crop accepted')
    cli = Path(__file__).with_name('open_office_use.py'); single = root / 'single.png'
    subprocess.run([sys.executable, str(cli), 'render-page', str(source), '--out', str(single), '--page', '3'], check=True, capture_output=True)
    with Image.open(single) as image:
        assert image.getpixel((image.width // 2, image.height // 2))[2] > 200
    _, identical = visual.compare_pages(source, source, 2, 2, open_office_use.render)
    assert identical['changed_pixels'] == 0, identical
    _, changed = visual.compare_pages(source, source, 2, 3, open_office_use.render)
    assert changed['changed_fraction'] > .9 and changed['channel_threshold'] == 16, changed
    comparison = root / 'compare.png'
    subprocess.run([sys.executable, str(cli), 'compare-pages', str(source), str(source), '--before-page', '2', '--after-page', '3', '--out', str(comparison)], check=True, capture_output=True)
    with Image.open(comparison) as image: assert image.width > image.height
    portrait = Presentation(); portrait.slide_width, portrait.slide_height = Inches(7.5), Inches(10); portrait.slides.add_slide(portrait.slide_layouts[6]); portrait_path = root / 'portrait.pptx'; portrait.save(portrait_path)
    try: visual.compare_pages(source, portrait_path, 1, 1, open_office_use.render)
    except ValueError as exc: assert 'different dimensions' in str(exc)
    else: raise AssertionError('Mismatched page sizes accepted')
    output = root / 'contact.png'
    subprocess.run([sys.executable, str(cli), 'contact-sheet', str(source), '--out', str(output), '--start', '12'], check=True, capture_output=True)
    original = output.read_bytes()
    try: visual.save_contact_sheet(source, output, 1, 2, open_office_use.render)
    except FileExistsError: pass
    else: raise AssertionError('Existing output overwritten')
    assert output.read_bytes() == original
    for start, count in ((0, 1), (13, 1), (1, 13)):
        try: visual.contact_sheet(source, start, count, open_office_use.render)
        except ValueError: pass
        else: raise AssertionError('Invalid range accepted')
    doc = Document(); doc.add_paragraph('Page one'); doc.add_page_break(); doc.add_paragraph('Page two'); doc_path = root / 'doc.docx'; doc.save(doc_path)
    assert visual.contact_sheet(doc_path, 2, 12, open_office_use.render)[1]['total_pages'] == 2
    assert visual.page_image(doc_path, 2, open_office_use.render)[1]['page'] == 2
    assert visual.page_image(doc_path, 2, open_office_use.render, [.99999, .99999, 1, 1])[1]['image_size'][0] > 0
    book = Workbook(); book.active['A1'] = 'Workbook'; book_path = root / 'book.xlsx'; book.save(book_path)
    assert visual.contact_sheet(book_path, 1, 12, open_office_use.render)[1]['total_pages'] == 1
    assert visual.page_image(book_path, 1, open_office_use.render)[1]['total_pages'] == 1
print('Visual smoke passed: page comparison, crops, formats, pagination and output protection')

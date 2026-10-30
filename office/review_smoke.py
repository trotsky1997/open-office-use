"""Compare real Office files including a formatting-only change."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

cli = Path(__file__).with_name('open_office_use.py')
def compare(left, right):
    result = subprocess.run([sys.executable, str(cli), 'compare', str(left), str(right)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    return json.loads(result.stdout)

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    for suffix in ('.docx', '.xlsx', '.pptx'):
        old = root / ('old' + suffix); new = root / ('new' + suffix)
        if suffix == '.docx':
            obj = Document(); obj.add_paragraph('Before'); obj.save(old)
            obj.paragraphs[0].text = 'After'; obj.save(new)
        elif suffix == '.xlsx':
            obj = Workbook(); obj.active['A1'] = 'Before'; obj.save(old)
            obj.active['A1'] = 'After'; obj.save(new)
        else:
            obj = Presentation(); slide = obj.slides.add_slide(obj.slide_layouts[1]); slide.shapes.title.text = 'Before'; obj.save(old)
            slide.shapes.title.text = 'After'; obj.save(new)
        result = compare(old, new)
        assert result['parts']['modified'] and 'Before' in str(result['text_diff']) and 'After' in str(result['text_diff']), result
        same = compare(old, old)
        assert not any(same['parts'].values()) and not same['text_diff'], same
    old = root / 'old.docx'; styled = root / 'styled.docx'
    doc = Document(old); doc.paragraphs[0].runs[0].bold = True; doc.save(styled)
    result = compare(old, styled)
    assert result['parts']['modified'] and not result['text_diff'], result
    assert result['text_view_may_be_truncated'] is True
print('Review smoke passed: three formats, unchanged file and formatting-only changes')

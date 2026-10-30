"""Exercise paginated Word header/footer search and targeted editing."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path
from docx import Document
from docx.shared import Inches
import session
import word

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'stories.docx'; draft = root / 'draft'; doc = Document()
    header = doc.sections[0].header; header.paragraphs[0].text = 'header needle'
    table = header.add_table(rows=1, cols=1, width=Inches(2)); table.cell(0, 0).text = 'table needle'
    for i in range(205): header.add_paragraph(f'needle {i}')
    doc.sections[0].footer.paragraphs[0].text = 'footer needle'; doc.save(source)
    first = word.search_stories(source, 'needle'); assert len(first['matches']) == 100 and first['next_offset'] is not None
    second = word.search_stories(source, 'needle', first['next_offset'], first['snapshot_sha256'])
    third = word.search_stories(source, 'needle', second['next_offset'], second['snapshot_sha256'])
    hits = first['matches'] + second['matches'] + third['matches']
    assert len(hits) == 208 and third['scan_complete']
    assert any('/tbl[1]/tr[1]/tc[1]/p[1]' in item['path'] for item in hits)
    assert hits[-1]['path'].startswith('/footer') and hits[-1]['part'].startswith('word/footer')
    assert not word.search_body(source, 'needle')['matches']
    session.start(draft, source=source)
    late = next(item for item in hits if item['excerpt'] == 'needle 204')
    assert session.edit_target(draft, 0, late['path'])['value'] == 'needle 204'
    edit = session.replace_text(draft, 0, late['path'], 'needle', 'changed')
    assert session.search_stories(draft, 1, 'changed')['matches'][0]['path'] == late['path']
    cli = Path(__file__).with_name('open_office_use.py')
    result = subprocess.run([sys.executable, str(cli), 'session', 'search-word', str(draft), '--revision', '1', '--scope', 'stories', '--query', 'changed'], capture_output=True, text=True, check=True)
    assert json.loads(result.stdout)['matches'][0]['path'] == late['path']
    try: word.search_stories(edit['file'], 'needle', first['next_offset'], first['snapshot_sha256'])
    except ValueError: pass
    else: raise AssertionError('Changed snapshot accepted')
print('Word header/footer search smoke passed')

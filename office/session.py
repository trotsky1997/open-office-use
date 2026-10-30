"""Persistent local draft sessions with immutable, optimistic revisions."""
import json
import hashlib
import math
import os
import posixpath
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook
from lxml import etree
from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
from pptx import Presentation
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries, quote_sheetname

import cell_text
import chart_cache
import grid_styles
import grid_values
import layout
import notes
import officecli_backend
import ooxml
import review
import spreadsheet
import snapshot as snapshots
import slides
import slide_text_check
import slide_charts
import word

FORMATS = ('.docx', '.xlsx', '.pptx')
MAIN_TYPES = {
    '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml',
    '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',
    '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml',
}


def validate_source(source, suffix):
    """Check the saved snapshot is an Office package of the requested type."""
    try:
        with zipfile.ZipFile(source) as archive:
            names = archive.namelist()
            if names.count('[Content_Types].xml') != 1 or '_rels/.rels' not in names:
                raise ValueError('Missing or ambiguous Office package metadata')
            metadata = archive.getinfo('[Content_Types].xml')
            if metadata.file_size > 1024 * 1024:
                raise ValueError('Office package content types exceed 1 MiB')
            root = ooxml.parse_xml(archive.read(metadata)).getroot()
            namespace = '{http://schemas.openxmlformats.org/package/2006/content-types}'
            if root.tag != namespace + 'Types': raise ValueError('Invalid Office package content types')
            overrides = {item.get('PartName', '').lstrip('/'): item.get('ContentType') for item in root.findall(namespace + 'Override')}
            defaults = {item.get('Extension', '').lower(): item.get('ContentType') for item in root.findall(namespace + 'Default')}
            if not any((overrides.get(name) or defaults.get(name.rsplit('.', 1)[-1].lower())) == MAIN_TYPES[suffix] for name in names):
                raise ValueError('Office package type does not match the file extension')
    except (zipfile.BadZipFile, zipfile.LargeZipFile, OSError, EOFError, ooxml.etree.XMLSyntaxError) as exc:
        raise ValueError('Invalid Office package') from exc


def start(directory, source=None, format=None):
    directory = Path(directory).absolute()
    if (source is None) == (format is None):
        raise ValueError('Provide exactly one of source or format')
    if source is not None:
        source = Path(source).resolve(strict=True)
        suffix = source.suffix.lower()
    else:
        if format is None:
            raise ValueError('Format is required for a blank session')
        suffix = '.' + format.lower().lstrip('.')
    if suffix not in FORMATS:
        raise ValueError('Supported formats: docx, xlsx, pptx')
    directory.mkdir(parents=False)
    try:
        initial = directory / ('revision-000000' + suffix)
        if source:
            with tempfile.TemporaryDirectory(dir=directory, prefix='.office-source-') as tmp:
                candidate, _ = snapshots.copy(source, tmp)
                validate_source(candidate, suffix)
                os.replace(candidate, initial)
        else:
            officecli_backend.batch([], initial)
        with (directory / 'session.json').open('x', encoding='utf-8') as f:
            json.dump({'version': 1, 'suffix': suffix}, f)
    except BaseException:
        shutil.rmtree(directory)
        raise
    return info(directory)


def info(directory):
    directory = Path(directory).resolve(strict=True)
    with (directory / 'session.json').open(encoding='utf-8') as f:
        manifest = json.load(f)
    suffix = manifest.get('suffix')
    if manifest.get('version') != 1 or suffix not in FORMATS:
        raise ValueError('Invalid session manifest')
    pattern = re.compile(r'revision-(\d{6})' + re.escape(suffix) + r'\Z')
    revisions = []
    for path in directory.iterdir():
        match = pattern.fullmatch(path.name)
        if match:
            if path.is_symlink() or not path.is_file():
                raise ValueError('Invalid session revision file')
            revisions.append(int(match.group(1)))
    revisions.sort()
    if not revisions or revisions[0] != 0 or revisions != list(range(len(revisions))):
        raise ValueError('Session revisions are missing or inconsistent')
    return {'directory': str(directory), 'format': suffix[1:], 'latest_revision': revisions[-1],
            'revisions': [{'revision': rev, 'file': str(directory / f'revision-{rev:06d}{suffix}')} for rev in revisions]}


def revision_file(state, revision):
    if isinstance(revision, bool) or not isinstance(revision, int) or revision < 0 or revision > state['latest_revision']:
        raise ValueError('Session revision does not exist')
    return Path(state['revisions'][revision]['file'])


def dump(directory, revision, path='/'):
    state = info(directory); source = revision_file(state, revision)
    return {'revision': revision, **officecli_backend.dump(source, path)}


def apply(directory, base_revision, operations, expected_sha256=None):
    state = info(directory)
    source = revision_file(state, base_revision)
    if base_revision != state['latest_revision']:
        raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999:
        raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f"revision-{base_revision + 1:06d}.{state['format']}"
    try:
        with tempfile.TemporaryDirectory(prefix='office-apply-source-') as tmp:
            frozen, digest = snapshots.copy(source, tmp)
            if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Document snapshot changed; reload the target')
            result = officecli_backend.batch(operations, output, frozen)
    except FileExistsError as exc:
        raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': result['output'], 'result': result['result'], **({'formula_cache_effect': result['formula_cache_effect']} if 'formula_cache_effect' in result else {})}


def restore(directory, base_revision, revision):
    state = info(directory)
    revision_file(state, base_revision)
    source = revision_file(state, revision)
    if base_revision != state['latest_revision']:
        raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999:
        raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f"revision-{base_revision + 1:06d}.{state['format']}"
    with tempfile.TemporaryDirectory(dir=state['directory'], prefix='.office-restore-') as tmp:
        candidate, _ = snapshots.copy(source, tmp)
        try:
            os.link(candidate, output)
        except FileExistsError as exc:
            raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'restored_revision': revision, 'file': str(output)}


def list_comments(directory, revision):
    state = info(directory)
    source = revision_file(state, revision)
    if state['format'] != 'docx':
        raise ValueError('Comment inspection requires a DOCX session')
    return {'revision': revision, **word.list_comments(source)}


def tracked_changes(directory, revision):
    state = info(directory); source = revision_file(state, revision)
    return {'revision': revision, **word.tracked_changes(source)}


def external_relationships(directory, revision):
    state = info(directory); source = revision_file(state, revision)
    return {'revision': revision, **review.external_relationships(source)}


def inspect_slide_navigation(directory, revision):
    state = info(directory); source = revision_file(state, revision)
    return {'revision': revision, **slides.inspect_navigation(source)}


def inspect_slide_charts(directory, revision):
    state = info(directory); source = revision_file(state, revision)
    if state['format'] != 'pptx': raise ValueError('Slide chart inspection requires a PPTX session')
    return {'revision': revision, **slide_charts.inspect(source)}


def set_slide_chart_text_color(directory, base_revision, part, color, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'pptx': raise ValueError('Slide chart color editing requires a PPTX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.pptx'
    try: result = slide_charts.set_text_color(source, output, part, color, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, **result}


def set_slide_chart_literal_text(directory, base_revision, part, kind, text, series=None, point=None, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'pptx': raise ValueError('Slide chart text editing requires a PPTX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.pptx'
    try: result = slide_charts.set_literal_text(source, output, part, kind, text, series, point, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, **result}


def set_slide_chart_literal_value(directory, base_revision, part, series, point, value, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'pptx': raise ValueError('Slide chart value editing requires a PPTX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.pptx'
    try: result = slide_charts.set_literal_value(source, output, part, series, point, value, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, **result}


def set_slide_chart_axis_bounds(directory, base_revision, part, axis_id, minimum, maximum, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'pptx': raise ValueError('Slide chart axis editing requires a PPTX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.pptx'
    try: result = slide_charts.set_axis_bounds(source, output, part, axis_id, minimum, maximum, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, **result}


def sync_slide_chart_cache(directory, base_revision, part, reference, allow_formula_caches=False, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'pptx': raise ValueError('Slide chart cache sync requires a PPTX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.pptx'
    try: result = slide_charts.sync(source, output, part, reference, allow_formula_caches, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, **result}


def slide_chart_workbook(source, part, temporary):
    if not isinstance(part, str) or not part.startswith('/'): raise ValueError('Provide a chart part from slide chart inspection')
    inventory = slide_charts.inspect(source)
    charts = [chart for chart in inventory['charts'] if chart.get('chart_part') == part]
    if not charts or any(chart['workbook']['status'] != 'embedded' for chart in charts): raise ValueError('Chart with embedded workbook was not available in the bounded inventory')
    if any(chart['workbook']['part'] != charts[0]['workbook']['part'] for chart in charts): raise ValueError('Ambiguous embedded workbook')
    workbook = charts[0]['workbook']
    if workbook['size'] > 16 * 1024 * 1024: raise ValueError('Embedded workbook exceeds 16 MiB')
    path = Path(temporary) / 'embedded.xlsx'
    with zipfile.ZipFile(source) as archive: blob = archive.read(workbook['part'].lstrip('/'))
    if len(blob) != workbook['size'] or hashlib.sha256(blob).hexdigest() != workbook['sha256']: raise ValueError('Embedded workbook changed during inspection')
    path.write_bytes(blob)
    return path, workbook


def read_slide_chart_grid(directory, revision, part, sheet, cell_range):
    state = info(directory); source = revision_file(state, revision)
    if state['format'] != 'pptx': raise ValueError('Slide chart grid requires a PPTX session')
    with tempfile.TemporaryDirectory(prefix='office-slide-chart-grid-') as tmp:
        snapshot, digest = snapshots.copy(source, tmp)
        path, workbook = slide_chart_workbook(snapshot, part, tmp)
        nested = Path(tmp) / 'draft'; start(nested, source=path)
        result = read_grid(nested, 0, sheet, cell_range)
    result['revision'] = revision; result['snapshot_sha256'] = digest; result['chart_part'] = part; result['workbook'] = workbook
    return result


def slide_chart_sync_targets(source, inventory, workbook_part, require_match=False):
    if not inventory['complete']: raise ValueError('Atomic chart synchronization requires a complete slide chart inventory')
    charts = {}
    for chart in inventory['charts']:
        if chart.get('workbook', {}).get('part') != workbook_part: continue
        reconciliation = chart.get('reconciliation', {})
        if not reconciliation.get('complete') or reconciliation.get('status') not in (('match',) if require_match else ('match', 'mismatch')):
            raise ValueError('Atomic chart synchronization requires fully comparable chart caches')
        charts[chart['chart_part']] = chart
    if not charts: raise ValueError('Embedded workbook has no inspected charts')
    package = Presentation(source).part.package
    if any(not rel.is_external and str(rel.target_part.partname) == workbook_part for rel in package._rels.values()):
        raise ValueError('Embedded workbook has a relationship outside the inspected charts')
    for owner in package.iter_parts():
        for relation in owner.rels.values():
            if not relation.is_external and str(relation.target_part.partname) == workbook_part and str(owner.partname) not in charts:
                raise ValueError('Embedded workbook has a relationship outside the inspected charts')
    return charts


def edit_slide_chart_grid(directory, base_revision, part, sheet, cell_range, changes, sync_caches=False):
    if not isinstance(sync_caches, bool): raise ValueError('sync_caches must be a boolean')
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'pptx': raise ValueError('Slide chart grid requires a PPTX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.pptx'
    with tempfile.TemporaryDirectory(dir=state['directory'], prefix='.office-slide-chart-grid-') as tmp:
        snapshot, digest = snapshots.copy(source, tmp)
        path, workbook = slide_chart_workbook(snapshot, part, tmp)
        if sync_caches:
            baseline = slide_charts.inspect(snapshot)
            related = slide_chart_sync_targets(snapshot, baseline, workbook['part'], require_match=True)
        nested = Path(tmp) / 'draft'; start(nested, source=path)
        result = edit_grid(nested, 0, sheet, cell_range, changes)
        updated = Path(result['file']).read_bytes()
        if len(updated) > 16 * 1024 * 1024: raise ValueError('Edited embedded workbook exceeds 16 MiB')
        def replace_workbook(_tree, archive):
            workbook_part = workbook['part'].lstrip('/')
            if sum(item.file_size for item in archive.infolist()) - archive.getinfo(workbook_part).file_size + len(updated) > 128 * 1024 * 1024:
                raise ValueError('Edited presentation exceeds package size limit')
            return {workbook_part: updated}
        candidate = Path(tmp) / 'edited.pptx' if sync_caches else output
        try:
            ooxml.edit_xml(snapshot, candidate, lambda _: part.lstrip('/'), replace_workbook)
            synchronized = []
            if sync_caches:
                after = slide_charts.inspect(candidate)
                current = slide_chart_sync_targets(candidate, after, workbook['part'])
                if set(current) != set(related): raise ValueError('Chart workbook relationships changed during edit')
                pending = [(chart_part, ref['index']) for chart_part, chart in current.items() for ref in chart['reconciliation']['references'] if ref['status'] == 'mismatch']
                if len(pending) > 30: raise ValueError('Too many chart references to synchronize atomically')
                for index, (chart_part, ref) in enumerate(pending):
                    destination = Path(tmp) / f'synced-{index}.pptx'
                    synced = slide_charts.sync(candidate, destination, chart_part, ref)
                    synchronized.append({'chart_part': chart_part, 'reference': ref, 'updated_points': synced['updated_points']})
                    candidate = destination
                final = slide_chart_sync_targets(candidate, slide_charts.inspect(candidate), workbook['part'], require_match=True)
                if set(final) != set(related): raise ValueError('Chart workbook relationships changed during synchronization')
                os.link(candidate, output)
        except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'chart_part': part, 'workbook_part': workbook['part'],
            'source_sha256': digest, 'chart_cache_effect': 'matched_saved_values' if sync_caches else 'may_be_stale',
            'synchronized_references': synchronized, 'formula_cache_effect': result['formula_cache_effect'],
            'data_validation': result['data_validation']}


def inspect_layout(directory, revision):
    state = info(directory)
    source = revision_file(state, revision)
    if state['format'] != 'pptx':
        raise ValueError('Slide layout inspection requires a PPTX session')
    return {'revision': revision, **layout.inspect(source)}



def inspect_slide_text(directory, revision, render):
    state = info(directory); source = revision_file(state, revision)
    if state['format'] != 'pptx': raise ValueError('Slide text inspection requires a PPTX session')
    return {'revision': revision, **slide_text_check.scan(source, render)}

def search_slides(directory, revision, query, slide=None, offset=0, snapshot_sha256=None, scope='slides'):
    state = info(directory); source = revision_file(state, revision)
    return {'revision': revision, **slides.search_text(source, query, slide, offset, snapshot_sha256, scope)}


def recalculate(directory, base_revision, allow_dynamic_arrays=False):
    state = info(directory)
    source = revision_file(state, base_revision)
    if state['format'] != 'xlsx':
        raise ValueError('Recalculation requires an XLSX session')
    if base_revision != state['latest_revision']:
        raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999:
        raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.xlsx'
    try:
        spreadsheet.recalculate(source, output, allow_dynamic_arrays=allow_dynamic_arrays)
    except FileExistsError as exc:
        raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'engine': 'LibreOffice', 'dynamic_array_override': allow_dynamic_arrays}


def scan_formulas(directory, base_revision):
    state = info(directory)
    source = revision_file(state, base_revision)
    if state['format'] != 'xlsx':
        raise ValueError('Formula scanning requires an XLSX session')
    if base_revision != state['latest_revision']:
        raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    result = spreadsheet.scan_formulas(source)
    latest = info(directory)['latest_revision']
    if base_revision != latest:
        raise ValueError(f'Revision conflict: latest revision is {latest}')
    return {'revision': base_revision, **result}


def inspect_workbook(directory, revision):
    state = info(directory); source = revision_file(state, revision)
    return {'revision': revision, **spreadsheet.inspect_workbook(source)}


def search_cells(directory, revision, query, sheet=None, offset=0, snapshot_sha256=None):
    state = info(directory); source = revision_file(state, revision)
    return {'revision': revision, **spreadsheet.search_cells(source, query, sheet, offset, snapshot_sha256)}


def scan_chart_caches(directory, revision):
    state = info(directory); source = revision_file(state, revision)
    return {'revision': revision, **chart_cache.scan(source)}


def set_chart_literal_series_name(directory, base_revision, part, series, text, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'xlsx': raise ValueError('Chart series editing requires an XLSX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.xlsx'
    try: result = chart_cache.set_literal_series_name(source, output, part, series, text, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, **result}


def set_chart_literal_category(directory, base_revision, part, reference, point, text, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'xlsx': raise ValueError('Chart category editing requires an XLSX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.xlsx'
    try: result = chart_cache.set_literal_category(source, output, part, reference, point, text, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, **result}


def set_chart_literal_value(directory, base_revision, part, reference, point, value, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'xlsx': raise ValueError('Chart literal editing requires an XLSX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.xlsx'
    try: result = chart_cache.set_literal_value(source, output, part, reference, point, value, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, **result}


def sync_chart_cache(directory, base_revision, part, reference, allow_formula_caches=False, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'xlsx': raise ValueError('Chart cache sync requires an XLSX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.xlsx'
    try: result = chart_cache.sync(source, output, part, reference, allow_formula_caches, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, **result}


def export(directory, revision, output):
    state = info(directory)
    source = revision_file(state, revision)
    output = Path(output).absolute()
    if output.suffix.lower() != source.suffix.lower():
        raise ValueError('Output format must match the session')
    if output.resolve().is_relative_to(Path(state['directory'])):
        raise ValueError('Export must be outside the session directory')
    if output.exists() or output.is_symlink():
        raise FileExistsError(f'Output already exists: {output}')
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent, prefix='.office-export-') as tmp:
        candidate, _ = snapshots.copy(source, tmp)
        os.link(candidate, output)
    return {'revision': revision, 'output': str(output)}


def validate_selection(selection):
    if not isinstance(selection, dict) or set(selection) != {'revision', 'page', 'box'}:
        raise ValueError('Selection requires revision, page and box')
    revision, page, box = selection['revision'], selection['page'], selection['box']
    if isinstance(revision, bool) or not isinstance(revision, int) or revision < 0:
        raise ValueError('Invalid selection revision')
    if isinstance(page, bool) or not isinstance(page, int) or page < 1:
        raise ValueError('Invalid selection page')
    if not isinstance(box, list) or len(box) != 4 or any(isinstance(v, bool) or not isinstance(v, (int, float)) for v in box):
        raise ValueError('Selection box requires four normalized coordinates')
    x1, y1, x2, y2 = box
    if not (0 <= x1 < x2 <= 1 and 0 <= y1 < y2 <= 1):
        raise ValueError('Invalid selection box')
    return {'revision': revision, 'page': page, 'box': box}


def set_selection(directory, selection):
    selection = validate_selection(selection)
    state = info(directory)
    if selection['revision'] != state['latest_revision']:
        raise ValueError('Selection revision is stale; select the latest page')
    directory = Path(state['directory'])
    with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=directory, prefix='.selection-', delete=False) as f:
        temp = Path(f.name)
        try:
            json.dump(selection, f, allow_nan=False)
        except BaseException:
            temp.unlink(missing_ok=True)
            raise
    try:
        os.replace(temp, directory / 'selection.json')
    finally:
        temp.unlink(missing_ok=True)
    return selection


def get_selection(directory):
    state = info(directory)
    with (Path(state['directory']) / 'selection.json').open(encoding='utf-8') as f:
        selection = validate_selection(json.load(f))
    revision_file(state, selection['revision'])
    return {**selection, 'latest_revision': state['latest_revision'],
            'stale': selection['revision'] != state['latest_revision']}


def edit_target(directory, base_revision, path):
    if not isinstance(path, str) or not path.startswith('/') or len(path) > 512 or '\x00' in path:
        raise ValueError('Provide a valid element path of at most 512 characters')
    state = info(directory); source = revision_file(state, base_revision)
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    with tempfile.TemporaryDirectory(prefix='office-edit-target-') as tmp:
        frozen, digest = snapshots.copy(source, tmp)
        return {**_edit_target(directory, base_revision, path, state, frozen), 'snapshot_sha256': digest}


def _edit_target(directory, base_revision, path, state, source):
    if state['format'] == 'docx' and word.NOTE_PATH.fullmatch(path):
        part = word.paragraph_part(source, path)
        with zipfile.ZipFile(source) as archive:
            tree = ooxml.parse_xml(archive.read(part)); target = word.find_paragraph(tree, path)
            text = ''.join(node.text or '' for node in target.iter(word.W + 't'))
        if len(text) > 4000: raise ValueError('Note paragraph exceeds the editor text limit')
        if info(directory)['latest_revision'] != base_revision: raise ValueError('Revision conflict: latest revision changed')
        match = word.NOTE_PATH.fullmatch(path)
        field = word.field_overlap(target)
        return {'read_only_reason': 'word_field' if field else None, 'revision': base_revision, 'path': path, 'type': 'paragraph', 'format': 'docx', 'mode': 'text', 'value': text, 'replace_text_path_supported': not field, 'replace_only': True, 'note_impact': {'kind': match.group(1), 'id': match.group(2), 'part': part}}
    if state['format'] == 'docx' and word.STORY_PATH.fullmatch(path):
        inventory = word.list_stories(source)
        match = word.STORY_PATH.fullmatch(path)
        stories = [item for item in inventory['stories'] if item['id'] == match.group(2) and item['type'] == match.group(1)]
        if len(stories) != 1: raise ValueError('Header or footer relationship does not exist')
        story = stories[0]
        with zipfile.ZipFile(source) as archive:
            tree = ooxml.parse_xml(archive.read(story['part']))
            target = word.find_paragraph(tree, path)
            text = ''.join(node.text or '' for node in target.iter(word.W + 't'))
        if len(text) > 4000: raise ValueError('Header or footer paragraph exceeds the editor text limit')
        paragraph = {'text': text}
        if info(directory)['latest_revision'] != base_revision: raise ValueError('Revision conflict: latest revision changed')
        field = word.field_overlap(target)
        links = [{'relationship_id': item['id'], **link} for item in inventory['stories'] if item['part'] == story['part'] for link in item['linked_sections']]
        return {'read_only_reason': 'word_field' if field else None, 'revision': base_revision, 'path': path, 'type': 'paragraph', 'format': 'docx', 'mode': 'text', 'value': paragraph['text'], 'replace_text_path_supported': not field, 'replace_only': True,
                'story_impact': {'part': story['part'], 'linked_sections': links, 'same_part_ids': story['same_part_ids'], 'complete': inventory['reference_resolution_complete']}}
    note_match = slides.NOTES_PATH.fullmatch(path) if state['format'] == 'pptx' else None
    if note_match:
        result = notes.read_paragraph(source, int(note_match.group(1)), int(note_match.group(2)))
        if len(result['text']) > 4000: raise ValueError('Notes paragraph exceeds the editor text limit')
        if info(directory)['latest_revision'] != base_revision: raise ValueError('Revision conflict: latest revision changed')
        return {'revision': base_revision, 'path': path, 'type': 'paragraph', 'format': 'pptx', 'mode': 'text', 'value': result['text'], 'read_only_reason': result['read_only_reason'], 'replace_text_path_supported': result['read_only_reason'] is None, 'replace_only': True}
    slide_match = slides.PATH.fullmatch(path) if state['format'] == 'pptx' and slides.paragraph_path(path) else None
    if slide_match and slide_match.group(2) and '/table[' in path:
        part = slides.slide_part(source, int(slide_match.group(1)) - 1)
        with zipfile.ZipFile(source) as archive:
            if archive.getinfo(part).file_size > 8 * 1024 * 1024: raise ValueError('Slide XML exceeds inspection limit')
            tree = ooxml.parse_xml(archive.read(part)); paragraph = slides.find_paragraph(tree, path)
            text = ''.join(node.text or '' for node in paragraph.iter(slides.A + 't'))
            try: slides.paragraph_nodes(paragraph); reason = None
            except ValueError as exc: reason = str(exc)
        if len(text) > 4000: raise ValueError('Slide paragraph exceeds the editor text limit')
        if info(directory)['latest_revision'] != base_revision: raise ValueError('Revision conflict: latest revision changed')
        return {'revision': base_revision, 'path': path, 'type': 'paragraph', 'format': 'pptx', 'mode': 'text', 'value': text, 'read_only_reason': reason, 'replace_text_path_supported': reason is None, 'replace_only': True}
    result = officecli_backend.read(source, 'get', path)
    if not isinstance(result, dict) or result.get('matches') != 1 or not isinstance(result.get('results'), list) or len(result['results']) != 1:
        raise ValueError('The path must identify exactly one editable element')
    item = result['results'][0]
    supported = {'docx': {'paragraph', 'run'}, 'xlsx': {'cell'},
                 'pptx': {'title', 'shape', 'textbox', 'placeholder', 'paragraph', 'run', 'tc'}}
    if not isinstance(item, dict) or item.get('type') not in supported[state['format']] or not isinstance(item.get('path'), str):
        raise ValueError('Element type is not supported by this editor')
    format_data = item.get('format')
    if not isinstance(format_data, dict):
        format_data = {}
    formula = format_data.get('formula') if state['format'] == 'xlsx' else None
    text = item.get('text')
    if text is not None and not isinstance(text, str):
        text = str(text)
    if isinstance(formula, str):
        mode = 'formula'
        value = formula if formula.startswith('=') else '=' + formula
    else:
        mode = ('literal' if format_data.get('type') in {'String', 'SharedString', 'InlineString'} else 'value') if state['format'] == 'xlsx' else 'text'
        value = text or ''
    if len(value) > 4000:
        raise ValueError('Element value exceeds the 4,000-character editor limit')
    read_only_reason = None; link_context = {}
    if state['format'] == 'docx' and item['type'] in ('paragraph', 'run'):
        paragraph_path = item['path'] if item['type'] == 'paragraph' else item['path'].rsplit('/', 1)[0]
        if not word.paragraph_path(paragraph_path): read_only_reason = 'unsupported_word_path'
        elif word.paragraph_has_field(source, paragraph_path): read_only_reason = 'word_field'
    if state['format'] == 'xlsx':
        match = cell_text.PATH.fullmatch(item['path'])
        if not match or not cell_text.cell_path(item['path']):
            read_only_reason = 'unsupported_cell_path'
        else:
            address = match.group(2).upper() + match.group(3)
            grid = _read_grid(directory, base_revision, match.group(1), address, source); cell = grid['rows'][0][0]
            read_only_reason = cell['read_only_reason']; link_context = {'hyperlink': cell['hyperlink'], 'hyperlink_inspection': grid['hyperlink_inspection']}
    latest = info(directory)['latest_revision']
    if base_revision != latest:
        raise ValueError(f'Revision conflict: latest revision is {latest}')
    return {'revision': base_revision, 'path': item['path'], 'type': item['type'],
            'format': state['format'], 'mode': mode, 'value': value, 'read_only_reason': read_only_reason, **link_context,
            'replace_text_path_supported': (state['format'] == 'docx' and read_only_reason is None and item['type'] == 'paragraph' and word.paragraph_path(item['path'])) or (state['format'] == 'pptx' and item['type'] in {'paragraph', 'title', 'shape', 'textbox', 'placeholder', 'tc'} and slides.paragraph_path(item['path'])) or (state['format'] == 'xlsx' and read_only_reason is None and mode in ('value', 'literal') and format_data.get('type') in {'String', 'SharedString', 'InlineString'} and cell_text.cell_path(item['path']))}


def edit(directory, base_revision, path, mode, value, expected_sha256=None):
    if not isinstance(value, str) or len(value) > 4000:
        raise ValueError('Value must be text of at most 4,000 characters')
    target = edit_target(directory, base_revision, path)
    if expected_sha256 is not None and target['snapshot_sha256'] != expected_sha256: raise ValueError('Document snapshot changed; reload the target')
    if target.get('read_only_reason'): raise ValueError(f"Element cannot be edited: {target['read_only_reason']}")
    if target.get('replace_only'): raise ValueError('Use text replacement for this paragraph')
    if target['format'] == 'pptx' and target['type'] == 'tc':
        raise ValueError('Use text replacement to edit PowerPoint table cells')
    allowed = {'xlsx': {'value', 'formula', 'literal'}, 'docx': {'text'}, 'pptx': {'text'}}
    if mode not in allowed[target['format']]:
        raise ValueError('Edit mode is not supported for this format')
    if mode == 'value' and value.lstrip().startswith('='):
        raise ValueError('Choose formula mode for a value starting with =')
    if mode == 'formula' and (not value.startswith('=') or len(value) < 2):
        raise ValueError('A formula must start with = and contain an expression')
    if target['format'] == 'xlsx':
        match = cell_text.PATH.fullmatch(target['path'])
        address = match.group(2).upper() + match.group(3)
        return edit_grid(directory, base_revision, match.group(1), address, [{'cell': address, 'value': value, **({'type': 'text'} if mode == 'literal' and value else {})}], expected_sha256=target['snapshot_sha256'])
    return apply(directory, base_revision, [{'command': 'set', 'path': target['path'], 'props': {mode: value}}], expected_sha256=target['snapshot_sha256'])


def replace_document_text(format, source, output, path, find, replacement):
    note = slides.NOTES_PATH.fullmatch(path) if format == 'pptx' and isinstance(path, str) and len(path) <= 512 else None
    if note:
        return notes.replace_text(source, output, int(note.group(1)), int(note.group(2)), find, replacement)
    return {'docx': word, 'pptx': slides, 'xlsx': cell_text}[format].replace_text(source, output, path, find, replacement)


def replace_text(directory, base_revision, path, find, replacement, expected_sha256=None):
    state = info(directory)
    source = revision_file(state, base_revision)
    if base_revision != state['latest_revision']:
        raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999:
        raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f"revision-{base_revision + 1:06d}.{state['format']}"
    try:
        with tempfile.TemporaryDirectory(prefix='office-replace-source-') as tmp:
            frozen, digest = snapshots.copy(source, tmp)
            if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Document snapshot changed; reload the target')
            replace_document_text(state['format'], frozen, output, path, find, replacement)
    except FileExistsError as exc:
        raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'path': path, **({'formula_cache_effect': 'may_be_stale', 'recalculation_requested': True} if state['format'] == 'xlsx' else {})}


def refresh_word_toc(directory, base_revision, path, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'docx': raise ValueError('TOC refresh requires a DOCX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f"revision-{base_revision + 1:06d}.docx"
    try: result = word.refresh_authored_toc(source, output, path, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'path': path, **result}


def replace_text_batch(directory, base_revision, changes):
    if not isinstance(changes, list) or not 1 <= len(changes) <= 25:
        raise ValueError('Provide 1 to 25 text replacements')
    for change in changes:
        if not isinstance(change, dict) or set(change) != {'path', 'find', 'replacement'} or not all(isinstance(change[key], str) for key in change):
            raise ValueError('Each replacement requires text path, find and replacement')
    state = info(directory); source = revision_file(state, base_revision)
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f"revision-{base_revision + 1:06d}.{state['format']}"
    with tempfile.TemporaryDirectory(dir=state['directory'], prefix='.office-replace-batch-') as tmp:
        previous = source
        for index, change in enumerate(changes):
            candidate = Path(tmp) / f"step-{index:02d}.{state['format']}"
            try: replace_document_text(state['format'], previous, candidate, change['path'], change['find'], change['replacement'])
            except ValueError as exc: raise ValueError(f'Replacement {index + 1}: {exc}') from exc
            if previous != source: previous.unlink()
            previous = candidate
        try: os.link(previous, output)
        except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'count': len(changes), **({'formula_cache_effect': 'may_be_stale', 'recalculation_requested': True} if state['format'] == 'xlsx' else {})}


def track_replace_text(directory, base_revision, path, find, replacement, author, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'docx': raise ValueError('Tracked replacement requires a DOCX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.docx'
    try:
        with tempfile.TemporaryDirectory(prefix='office-track-source-') as tmp:
            frozen, digest = snapshots.copy(source, tmp)
            if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Document snapshot changed; reload the target')
            word.track_replace_text(frozen, output, path, find, replacement, author)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'path': path}



def track_insert_text(directory, base_revision, path, anchor, text, position, author, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'docx': raise ValueError('Tracked insertion requires a DOCX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.docx'
    try:
        with tempfile.TemporaryDirectory(prefix='office-track-source-') as tmp:
            frozen, digest = snapshots.copy(source, tmp)
            if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Document snapshot changed; reload the target')
            result = word.track_insert_text(frozen, output, path, anchor, text, position, author)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'path': path, 'id': result['id'], 'part': result['part'], 'position': position}

def resolve_tracked_runs(directory, base_revision, part, ids, accept, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'docx': raise ValueError('Tracked revision resolution requires a DOCX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.docx'
    try: word.resolve_tracked_runs(source, output, part, ids, accept, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'part': part, 'ids': ids, 'accepted': accept}




def inventory_slides(directory, revision):
    state = info(directory); source = revision_file(state, revision)
    if state['format'] != 'pptx': raise ValueError('Slide inventory requires a PPTX session')
    return {'revision': revision, **slides.inventory(source)}

def reorder_slides(directory, base_revision, order, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'pptx': raise ValueError('Slide reorder requires a PPTX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.pptx'
    try: slides.reorder(source, output, order, expected_sha256=expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'order': order}

def read_notes(directory, revision, slide):
    state = info(directory); source = revision_file(state, revision)
    result = notes.read(source, slide); result['revision'] = revision
    return result


def list_stories(directory, revision):
    state = info(directory); source = revision_file(state, revision)
    return {'revision': revision, **word.list_stories(source)}


def search_body(directory, revision, query, offset=0, snapshot_sha256=None):
    state = info(directory); source = revision_file(state, revision)
    return {'revision': revision, **word.search_body(source, query, offset, snapshot_sha256)}



def search_notes(directory, revision, query, offset=0, snapshot_sha256=None):
    state = info(directory); source = revision_file(state, revision)
    return {'revision': revision, **word.search_notes(source, query, offset, snapshot_sha256)}


def search_stories(directory, revision, query, offset=0, snapshot_sha256=None):
    state = info(directory); source = revision_file(state, revision)
    return {'revision': revision, **word.search_stories(source, query, offset, snapshot_sha256)}


def compare_cells(directory, before_revision, after_revision, sheet, cell_range):
    state = info(directory)
    before = revision_file(state, before_revision); after = revision_file(state, after_revision)
    return {'before_revision': before_revision, 'after_revision': after_revision, **review.compare_cells(before, after, sheet, cell_range)}

def compare_notes(directory, before_revision, after_revision, before_slide, after_slide):
    state = info(directory)
    before = revision_file(state, before_revision); after = revision_file(state, after_revision)
    result = review.compare_notes(before, after, before_slide, after_slide)
    return {'before_revision': before_revision, 'after_revision': after_revision, **result}


def replace_notes(directory, base_revision, slide, text, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'pptx': raise ValueError('Speaker notes require a PPTX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.pptx'
    try:
        with tempfile.TemporaryDirectory(prefix='office-notes-source-') as tmp:
            frozen, digest = snapshots.copy(source, tmp)
            if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Document snapshot changed; reload notes')
            notes.replace(frozen, output, slide, text)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'slide': slide}


def replace_notes_text(directory, base_revision, slide, paragraph, find, replacement, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'pptx': raise ValueError('Speaker notes require a PPTX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.pptx'
    try:
        with tempfile.TemporaryDirectory(prefix='office-notes-source-') as tmp:
            frozen, digest = snapshots.copy(source, tmp)
            if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Document snapshot changed; reload notes')
            notes.replace_text(frozen, output, slide, paragraph, find, replacement)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'slide': slide, 'paragraph': paragraph}


def reply_comment(directory, base_revision, parent_id, text, author, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'docx': raise ValueError('Comment replies require a DOCX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.docx'
    try:
        with tempfile.TemporaryDirectory(prefix='office-comment-source-') as tmp:
            frozen, digest = snapshots.copy(source, tmp)
            if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Document snapshot changed; reload comments')
            result = word.reply_comment(frozen, output, parent_id, text, author)
    except FileExistsError as exc:
        raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, **result}


def resolve_comment(directory, base_revision, comment_id, resolved, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'docx': raise ValueError('Comment state requires a DOCX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.docx'
    try:
        with tempfile.TemporaryDirectory(prefix='office-comment-source-') as tmp:
            frozen, digest = snapshots.copy(source, tmp)
            if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Document snapshot changed; reload comments')
            word.set_comment_resolved(frozen, output, comment_id, resolved)
    except FileExistsError as exc:
        raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'comment_id': comment_id, 'resolved': resolved}


def add_comment(directory, base_revision, path, text, author, anchor=None, expected_sha256=None):
    state = info(directory)
    source = revision_file(state, base_revision)
    if state['format'] != 'docx':
        raise ValueError('Comments require a DOCX session')
    if base_revision != state['latest_revision']:
        raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999:
        raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.docx'
    try:
        with tempfile.TemporaryDirectory(prefix='office-comment-source-') as tmp:
            frozen, digest = snapshots.copy(source, tmp)
            if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Document snapshot changed; reload the target')
            result = word.add_comment(frozen, output, path, text, author, anchor)
    except FileExistsError as exc:
        raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'path': path, 'comment_id': result['comment_id']}



def read_picture_file(path):
    source = Path(path).resolve(strict=True)
    if source.stat().st_size > 10 * 1024 * 1024: raise ValueError('Image must be at most 10 MiB')
    with tempfile.TemporaryDirectory(prefix='office-picture-source-') as tmp:
        frozen, _ = snapshots.copy(source, tmp)
        if frozen.stat().st_size > 10 * 1024 * 1024: raise ValueError('Image must be at most 10 MiB')
        return frozen.read_bytes()

def set_slide_navigation(directory, base_revision, path, target_slide, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'pptx': raise ValueError('Slide navigation requires a PPTX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.pptx'
    try: slides.set_navigation(source, output, path, target_slide, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'path': path, 'target_slide': target_slide}


def set_slide_description(directory, base_revision, path, description, expected_sha256=None):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'pptx': raise ValueError('Slide description requires a PPTX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.pptx'
    try: slides.set_description(source, output, path, description, expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'path': path}


def replace_picture(directory, base_revision, path, image_data, expected_sha256=None):
    state = info(directory)
    source = revision_file(state, base_revision)
    if state['format'] != 'pptx':
        raise ValueError('Picture replacement requires a PPTX session')
    if base_revision != state['latest_revision']:
        raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999:
        raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.pptx'
    try:
        slides.replace_picture(source, output, path, image_data, expected_sha256)
    except FileExistsError as exc:
        raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'path': path}


def grid_bounds(cell_range):
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        if min_col is None or min_row is None or max_col is None or max_row is None:
            raise ValueError
        if (min_col < 1 or min_row < 1 or max_col > 16384 or max_row > 1048576
                or min_col > max_col or min_row > max_row
                or (max_col - min_col + 1) * (max_row - min_row + 1) > 200):
            raise ValueError
        return min_col, min_row, max_col, max_row
    except (TypeError, ValueError):
        raise ValueError('Specify a valid rectangular range of at most 200 cells') from None


def grid_validations(root, bounds):
    """Describe standard validation rules intersecting the bounded grid."""
    ns = grid_styles.X
    # shortcut: bounded standard rules only; add extension support when needed.
    issues, rules = set(), []
    containers = root.findall(ns + 'dataValidations')
    if len(containers) > 1: issues.add('multiple_validation_containers')
    if any(isinstance(node.tag, str) and node.tag.endswith('}dataValidations') and node.tag != ns + 'dataValidations' for node in root.iter()):
        issues.add('unsupported_validation_extension')
    min_col, min_row, max_col, max_row = bounds
    for container in containers:
        children = list(container)
        if len(children) > 1000: issues.add('rule_scan_limit')
        if container.get('count') is not None and container.get('count') != str(len(children)): issues.add('rule_count_mismatch')
        for item in children[:1000]:
            if item.tag != ns + 'dataValidation':
                issues.add('unsupported_rule'); continue
            refs = item.get('sqref', '')
            if len(refs) > 32000 or len(refs.split()) > 500:
                issues.add('reference_limit'); continue
            addresses = set()
            if not refs.strip(): issues.add('invalid_reference')
            for ref in refs.split():
                try:
                    a, b, c, d = range_boundaries(ref)
                    if not all(isinstance(v, int) for v in (a, b, c, d)) or not (1 <= a <= c <= 16384 and 1 <= b <= d <= 1048576): raise ValueError()
                except (ValueError, TypeError):
                    issues.add('invalid_reference'); continue
                for row in range(max(b, min_row), min(d, max_row) + 1):
                    for col in range(max(a, min_col), min(c, max_col) + 1): addresses.add(f'{get_column_letter(col)}{row}')
            if not addresses: continue
            if len(rules) >= 50:
                issues.add('rule_output_limit'); continue
            rule = {'cells': sorted(addresses), 'attributes': {}, 'formula1': None, 'formula2': None}
            for key, value in item.attrib.items():
                if key == 'sqref': continue
                if len(value) > 256: issues.add('field_truncated')
                rule['attributes'][key] = value[:256]
            for name in ('formula1', 'formula2'):
                formulas = item.findall(ns + name)
                if len(formulas) > 1: issues.add('ambiguous_formula')
                if formulas:
                    value = formulas[0].text or ''
                    if len(value) > 256: issues.add('field_truncated')
                    rule[name] = value[:256]
            if any(child.tag not in (ns + 'formula1', ns + 'formula2') for child in item): issues.add('unsupported_rule_content')
            rules.append(rule)
    return {'rules': rules, 'complete': not issues, 'issues': sorted(issues), 'evaluated': False}


def grid_conditional_formats(root, bounds):
    """Summarize standard conditional-format rules intersecting the grid."""
    ns = grid_styles.X; issues, rules = set(), []; scanned = 0
    # shortcut: summarize standard rules without evaluating styles; inspect extension rules in the target application.
    if any(isinstance(node.tag, str) and node.tag.endswith('}conditionalFormatting') and (node.tag != ns+'conditionalFormatting' or node.getparent() is not root) for node in root.iter()): issues.add('unsupported_conditional_format_extension')
    containers = root.findall(ns+'conditionalFormatting')
    if len(containers) > 1000: issues.add('container_scan_limit')
    min_col, min_row, max_col, max_row = bounds
    for container in containers[:1000]:
        refs = container.get('sqref', '')
        if len(refs) > 32000 or len(refs.split()) > 500:
            issues.add('reference_limit'); continue
        addresses = set()
        if not refs.strip(): issues.add('invalid_reference')
        for ref in refs.split():
            try:
                a, b, c, d = range_boundaries(ref)
                if not all(isinstance(v, int) for v in (a,b,c,d)) or not (1 <= a <= c <= 16384 and 1 <= b <= d <= 1048576): raise ValueError()
            except (ValueError, TypeError):
                issues.add('invalid_reference'); continue
            for row in range(max(b,min_row), min(d,max_row)+1):
                for col in range(max(a,min_col), min(c,max_col)+1): addresses.add(f'{get_column_letter(col)}{row}')
        for item in container:
            scanned += 1
            if scanned > 1000: issues.add('rule_scan_limit'); break
            if item.tag != ns+'cfRule': issues.add('unsupported_rule'); continue
            priority = item.get('priority', '')
            if not item.get('type') or not priority.isascii() or not priority.isdigit() or len(priority) > 10 or int(priority) < 1: issues.add('invalid_rule_metadata')
            if not addresses: continue
            if len(rules) >= 50: issues.add('rule_output_limit'); continue
            attributes = {}
            for key,value in item.attrib.items():
                if len(value) > 256: issues.add('field_truncated')
                attributes[key] = value[:256]
            formulas = item.findall(ns+'formula')
            if len(formulas) > 3: issues.add('formula_limit')
            values = []
            for formula in formulas[:3]:
                value = formula.text or ''
                if len(value) > 256: issues.add('field_truncated')
                values.append(value[:256])
            rules.append({'cells': sorted(addresses), 'attributes': attributes, 'formulas': values})
        if scanned > 1000: break
    return {'rules': rules, 'complete': not issues, 'issues': sorted(issues), 'evaluated': False, 'style_details_included': False}


def resolve_grid_lists(source, sheet, validation):
    """Resolve bounded absolute list ranges containing verified plain text."""
    candidates = [rule for rule in validation['rules'] if rule['attributes'].get('type') == 'list' and rule.get('formula1') and not rule['formula1'].startswith('"')]
    if not candidates: return
    if not validation['complete']:
        for rule in candidates: rule['list_resolution'] = 'incomplete_metadata'
        return
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        with zipfile.ZipFile(source) as archive:
            spreadsheet.prepare_shared_strings(workbook, archive=archive)
            try:
                workbook_root = ooxml.parse_xml(archive.read('xl/workbook.xml')).getroot()
                containers = workbook_root.findall(grid_styles.X + 'definedNames')
                names = list(containers[0]) if len(containers) == 1 else []
                sheets = workbook_root.findall(grid_styles.X + 'sheets')
                names_complete = (workbook_root.tag == grid_styles.X + 'workbook' and len(containers) <= 1 and len(names) <= 10000
                                  and all(item.tag == grid_styles.X + 'definedName' for item in names)
                                  and len(sheets) == 1 and [item.get('name') for item in sheets[0]] == workbook.sheetnames)
            except (KeyError, ooxml.etree.XMLSyntaxError):
                names, names_complete = [], False
        cache = {}
        for rule in candidates:
            formula = rule['formula1']
            name_match = re.fullmatch(r'=?([A-Za-z_][A-Za-z0-9_.]*)', formula)
            if name_match:
                if not names_complete:
                    rule['list_resolution'] = 'incomplete_names'; continue
                name = name_match.group(1)
                matching = [item for item in names if item.get('name', '').casefold() == name.casefold()]
                index = str(workbook.sheetnames.index(sheet))
                if any(item.get('localSheetId') is not None and (not item.get('localSheetId').isdigit() or len(item.get('localSheetId')) > 6 or int(item.get('localSheetId')) >= len(workbook.sheetnames)) for item in matching):
                    rule['list_resolution'] = 'ambiguous_name'; continue
                local = [item for item in matching if item.get('localSheetId') is not None and int(item.get('localSheetId')) == int(index)]
                selected = local or [item for item in matching if item.get('localSheetId') is None]
                if len(selected) != 1:
                    rule['list_resolution'] = 'ambiguous_name' if selected else 'missing_name'; continue
                item = selected[0]
                if any(item.get(flag) not in (None, '0', 'false') for flag in ('function', 'vbProcedure', 'xlm')):
                    rule['list_resolution'] = 'unsupported_name'; continue
                formula = item.text or ''
                rule['list_name'] = {'name': item.get('name'), 'scope': sheet if local else 'workbook'}
            # shortcut: names must resolve to absolute one-dimensional ranges of verified text; add formula evaluation when supported.
            match = re.fullmatch(r"=?(?:(?:'((?:[^']|'')+)'|([A-Za-z0-9_.]+))!)?(\$[A-Za-z]{1,3}\$[1-9][0-9]{0,6})(?::(\$[A-Za-z]{1,3}\$[1-9][0-9]{0,6}))?", formula)
            if not match or name_match and not (match.group(1) or match.group(2)):
                rule['list_resolution'] = 'unsupported_reference'; continue
            target = (match.group(1).replace("''", "'") if match.group(1) else match.group(2)) or sheet
            ref = match.group(3) + (':' + match.group(4) if match.group(4) else '')
            key = (target, ref)
            if key not in cache:
                if len(cache) >= 20:
                    rule['list_resolution'] = 'reference_limit'; continue
                try:
                    a, b, c, d = range_boundaries(ref)
                    if target not in workbook.sheetnames: raise ValueError('missing_sheet')
                    if not (1 <= a <= c <= 16384 and 1 <= b <= d <= 1048576 and (a == c or b == d) and (c - a + 1) * (d - b + 1) <= 100): raise ValueError('range_limit')
                    cells = [cell for row in workbook[target].iter_rows(min_row=b, max_row=d, min_col=a, max_col=c) for cell in row]
                    values = [cell.value for cell in cells]
                    if any(cell.data_type not in ('s', 'inlineStr') or not isinstance(value, str) or isinstance(value, spreadsheet.UnverifiedSharedString) or not value or len(value) > 256 or re.search(r'_x[0-9a-fA-F]{4}_', value) for cell, value in zip(cells, values)):
                        raise ValueError('unverified_values')
                    cache[key] = ('resolved', values)
                except ValueError as exc:
                    cache[key] = (str(exc), None)
            status, values = cache[key]
            rule['list_resolution'] = status
            if values is not None:
                rule['list_values'] = values
                rule['list_source'] = {'sheet': target, 'range': ref}
    finally:
        workbook.close()


def grid_hyperlinks(root, archive, part, bounds):
    """Inspect bounded worksheet hyperlink metadata without following targets."""
    ns = grid_styles.X; rel_ns = grid_styles.REL; rid_ns = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    issues = set(); mapped = {}; links_scanned = 0
    containers = root.findall(ns + 'hyperlinks')
    if len(containers) > 1: issues.add('multiple_hyperlink_containers')
    if any(isinstance(node.tag, str) and node.tag.endswith('}hyperlink') and node.tag != ns + 'hyperlink' for node in root.iter()): issues.add('unsupported_hyperlink_extension')
    rel_path = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
    relationships = {}; duplicate_ids = set()
    if rel_path in archive.namelist():
        try:
            rel_root = ooxml.parse_xml(archive.read(rel_path)).getroot()
            if rel_root.tag != rel_ns + 'Relationships': raise ValueError('Invalid relationships')
            entries = list(rel_root)
            if len(entries) > 10000: issues.add('relationship_scan_limit')
            for entry in entries[:10000]:
                if entry.tag != rel_ns + 'Relationship' or not entry.get('Id'): issues.add('invalid_relationship'); continue
                identity = entry.get('Id')
                if identity in relationships: duplicate_ids.add(identity); issues.add('duplicate_relationship_id')
                else: relationships[identity] = entry
        except (ValueError, KeyError, etree.XMLSyntaxError): issues.add('invalid_relationship_part')
    for container in containers[:10]:
        items = list(container)
        if len(items) > 10000: issues.add('hyperlink_scan_limit')
        for link in items[:10000]:
            links_scanned += 1
            if link.tag != ns + 'hyperlink': issues.add('unsupported_hyperlink'); continue
            ref = link.get('ref', '')
            try:
                c1, r1, c2, r2 = range_boundaries(ref)
                if None in (c1, r1, c2, r2) or not (1 <= c1 <= c2 <= 16384 and 1 <= r1 <= r2 <= 1048576): raise ValueError
            except (TypeError, ValueError): issues.add('invalid_hyperlink_range'); continue
            left, top, right, bottom = max(c1,bounds[0]), max(r1,bounds[1]), min(c2,bounds[2]), min(r2,bounds[3])
            if left > right or top > bottom: continue
            identity = link.get(rid_ns + 'id'); location = link.get('location'); target = None; status = 'internal'
            if identity is not None:
                relation = relationships.get(identity)
                if identity in duplicate_ids or relation is None or relation.get('Type') != rid_ns[1:-1] + '/hyperlink' or relation.get('TargetMode') != 'External' or not relation.get('Target'):
                    status = 'unresolved'; issues.add('unresolved_hyperlink')
                else: status = 'external'; target = relation.get('Target')
            elif not location:
                status = 'unresolved'; issues.add('unresolved_hyperlink')
            truncated = any(value is not None and len(value) > 2048 for value in (target, location))
            if truncated: issues.add('hyperlink_target_truncated')
            details = {'ref': ref, 'status': status, 'target': target[:2048] if target is not None else None, 'location': location[:2048] if location is not None else None, 'truncated': truncated}
            for row in range(top,bottom+1):
                for col in range(left,right+1):
                    address = f'{get_column_letter(col)}{row}'
                    if address in mapped:
                        mapped[address] = {'ref': None, 'status': 'ambiguous', 'target': None, 'location': None, 'truncated': False}; issues.add('overlapping_hyperlinks')
                    else: mapped[address] = details.copy()
    return mapped, {'complete': not issues, 'issues': sorted(issues), 'links_scanned': links_scanned}


def normalize_hyperlink_location(location):
    if not isinstance(location, dict) or set(location) != {'sheet', 'cell'} or not isinstance(location['sheet'], str) or not location['sheet']: raise ValueError('Internal target requires a sheet and single cell')
    match = re.fullmatch(r'([A-Za-z]{1,3})([1-9][0-9]{0,6})', location['cell']) if isinstance(location['cell'], str) else None
    if match is None: raise ValueError('Internal target requires a single cell address')
    cell = match.group(1).upper() + match.group(2); grid_bounds(cell)
    return {'sheet': location['sheet'], 'cell': cell}


def set_grid_hyperlink(directory, base_revision, sheet, cell, url, expected_sha256=None, location=None):
    """Set or remove a single-cell hyperlink without rewriting cell data."""
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'xlsx': raise ValueError('Hyperlink editing requires an XLSX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    match = re.fullmatch(r'([A-Za-z]{1,3})([1-9][0-9]{0,6})', cell) if isinstance(cell, str) else None
    if match is None: raise ValueError('Provide a single cell address')
    cell = match.group(1).upper() + match.group(2); bounds = grid_bounds(cell)
    if location is not None:
        if url is not None: raise ValueError('Internal target cannot include a URL')
        location = normalize_hyperlink_location(location)
    if url is not None:
        if not isinstance(url, str) or not url or len(url) > 2048 or any(c.isspace() or ord(c) < 32 or ord(c) == 127 or 0xd800 <= ord(c) <= 0xdfff or ord(c) in (0xfffe, 0xffff) for c in url): raise ValueError('Invalid hyperlink URL')
        try: parsed = urlsplit(url); hostname = parsed.hostname
        except ValueError: raise ValueError('Invalid hyperlink URL') from None
        if parsed.scheme.lower() not in ('http', 'https', 'mailto') or (parsed.scheme.lower() in ('http', 'https') and not hostname) or (parsed.scheme.lower() == 'mailto' and not parsed.path): raise ValueError('Hyperlink must be an HTTP, HTTPS or mailto URL')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.xlsx'
    with tempfile.TemporaryDirectory(prefix='office-link-edit-') as tmp:
        frozen, digest = snapshots.copy(source, tmp)
        if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Workbook snapshot changed; reload the grid')
        grid = _read_grid(directory, base_revision, sheet, cell, frozen); target = grid['rows'][0][0]
        if not target['editable']: raise ValueError(f"Cell cannot be edited: {target['read_only_reason']}")
        if not grid['hyperlink_inspection']['complete']: raise ValueError('Hyperlink inspection is incomplete')
        existing = target['hyperlink']
        if existing and range_boundaries(existing['ref']) != bounds: raise ValueError('Multi-cell hyperlinks require a separate edit')
        if url is None and location is None and not existing: raise ValueError('Cell has no hyperlink')
        if (url is not None or location is not None) and target.get('value') is None and not target.get('formula'): raise ValueError('Enter a cell value before adding a hyperlink')
        if location is not None: grid_styles.sheet_part(frozen, location['sheet'])
        part = grid_styles.sheet_part(frozen, sheet); rel_path = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
        def edit(tree, archive):
            ns = grid_styles.X; rel_ns = grid_styles.REL; rid_ns = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
            root = tree.getroot(); containers = root.findall(ns + 'hyperlinks')
            if root.tag != ns + 'worksheet' or len(containers) > 1: raise ValueError('Unsupported worksheet hyperlink structure')
            container = containers[0] if containers else None
            node = None
            if existing:
                matches = [item for item in container if item.tag == ns + 'hyperlink' and range_boundaries(item.get('ref')) == bounds]
                if len(matches) != 1: raise ValueError('Hyperlink target is ambiguous')
                node = matches[0]
                if set(node.attrib) - {'ref', 'location', 'tooltip', 'display', rid_ns + 'id'} or len(node): raise ValueError('Unsupported hyperlink properties')
            raw = archive.read(rel_path) if rel_path in archive.namelist() else None
            if raw is None:
                if url is None: rel_tree = None
                else:
                    ct = ooxml.parse_xml(archive.read('[Content_Types].xml')).getroot(); ct_ns = '{http://schemas.openxmlformats.org/package/2006/content-types}'
                    defaults = [item for item in ct.findall(ct_ns + 'Default') if item.get('Extension', '').lower() == 'rels']
                    if len(defaults) != 1 or defaults[0].get('ContentType') != 'application/vnd.openxmlformats-package.relationships+xml': raise ValueError('Relationship content type is unavailable')
                    rel_tree = etree.ElementTree(etree.Element(rel_ns + 'Relationships', nsmap={None: rel_ns[1:-1]}))
            else:
                rel_tree = ooxml.parse_xml(raw)
                if rel_tree.getroot().tag != rel_ns + 'Relationships': raise ValueError('Invalid worksheet relationships')
            rel_root = rel_tree.getroot() if rel_tree is not None else None
            if rel_root is not None:
                entries = list(rel_root); ids = [item.get('Id') for item in entries]
                if len(entries) > 10000 or any(item.tag != rel_ns + 'Relationship' or not item.get('Id') for item in entries) or len(set(ids)) != len(ids): raise ValueError('Ambiguous worksheet relationships')
            old_id = node.get(rid_ns + 'id') if node is not None else None; rel_changed = False
            if url is None and location is None:
                container.remove(node)
                if not len(container): root.remove(container)
            else:
                if container is None:
                    before = {'sheetPr','dimension','sheetViews','sheetFormatPr','cols','sheetData','sheetCalcPr','sheetProtection','protectedRanges','scenarios','autoFilter','sortState','dataConsolidate','customSheetViews','mergeCells','phoneticPr','conditionalFormatting','dataValidations'}
                    after = {'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
                    children = list(root)
                    if any(child.tag not in {ns+name for name in before | after} for child in children): raise ValueError('Unsupported worksheet structure for hyperlink insertion')
                    container = etree.Element(ns + 'hyperlinks'); position = next((i for i,child in enumerate(children) if child.tag in {ns+name for name in after}),len(children)); root.insert(position,container)
                if node is None: node = etree.SubElement(container, ns + 'hyperlink', ref=cell)
                if location is not None:
                    node.attrib.pop(rid_ns + 'id', None)
                    node.set('location', quote_sheetname(location['sheet']) + '!' + location['cell'])
                else:
                    node.attrib.pop('location', None)
                    used = {item.get('Id') for item in rel_root}; number = 1
                    while f'rId{number}' in used: number += 1
                    identity = f'rId{number}'; node.set(rid_ns + 'id', identity)
                    etree.SubElement(rel_root, rel_ns + 'Relationship', Id=identity, Type=rid_ns[1:-1] + '/hyperlink', Target=url, TargetMode='External'); rel_changed = True
            if old_id and rel_root is not None and not any(old_id in element.attrib.values() for element in root.iter()):
                old = [item for item in rel_root if item.get('Id') == old_id]
                if len(old) != 1 or old[0].get('Type') != rid_ns[1:-1] + '/hyperlink' or old[0].get('TargetMode') != 'External': raise ValueError('Unsupported old hyperlink relationship')
                rel_root.remove(old[0]); rel_changed = True
            if not rel_changed: return {}
            serialized = ooxml.serialize_xml(rel_tree, raw) if raw is not None else etree.tostring(rel_tree, encoding='UTF-8', xml_declaration=True)
            return {rel_path: serialized}
        try: ooxml.edit_xml(frozen, output, lambda _: part, edit, allow_new=True, expected_sha256=digest)
        except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'sheet': sheet, 'cell': cell, 'action': 'remove' if url is None and location is None else 'set_internal' if location is not None else 'set'}


def read_grid(directory, base_revision, sheet, cell_range):
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'xlsx': raise ValueError('Grid editing requires an XLSX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    with tempfile.TemporaryDirectory(prefix='office-grid-read-') as tmp:
        frozen, digest = snapshots.copy(source, tmp)
        return {**_read_grid(directory, base_revision, sheet, cell_range, frozen), 'snapshot_sha256': digest}


def _read_grid(directory, base_revision, sheet, cell_range, source):
    if not isinstance(sheet, str) or not sheet or '/' in sheet or '\x00' in sheet:
        raise ValueError('Invalid sheet name')
    min_col, min_row, max_col, max_row = grid_bounds(cell_range)
    constraints = spreadsheet.worksheet_constraints(source, sheet, (min_col, min_row, max_col, max_row), True, True)
    result = spreadsheet.read_range(source, sheet, cell_range, include_style=True)
    complex_addresses = {f'{get_column_letter(min_col + col_index)}{min_row + row_index}' for row_index, row in enumerate(result['rows']) for col_index, cell in enumerate(row) if isinstance(cell.get('formula'), dict)}
    dynamic_anchors = grid_values.unexpanded_dynamic_anchors(source, sheet, complex_addresses)
    rows = []
    for row_index, row in enumerate(result['rows']):
        cells = []
        for col_index, cell in enumerate(row):
            col, row_number = min_col + col_index, min_row + row_index
            address = f'{get_column_letter(col)}{row_number}'
            reason, origin = spreadsheet.cell_constraint(constraints, col, row_number, cell['locked'], isinstance(cell.get('formula'), dict))
            overlaps = [bounds for bounds in constraints['formula_ranges'] if bounds[0] <= col <= bounds[2] and bounds[1] <= row_number <= bounds[3]]
            dynamic_anchor = address in dynamic_anchors and overlaps == [(col, row_number, col, row_number)] and not any(bounds[0] <= col <= bounds[2] and bounds[1] <= row_number <= bounds[3] for bounds in constraints['merged_ranges']) and not (constraints['sheet_protected'] and cell['locked'])
            if dynamic_anchor and reason == 'formula_range': reason = None
            formula = cell.get('formula')
            display = formula if isinstance(formula, str) else formula.get('text') if isinstance(formula, dict) else cell.get('value')
            if reason is None and 'formula' not in cell and cell.get('text_complete') is False:
                reason = 'unverified_text'
            if reason is None and isinstance(display, str) and len(display.encode('utf-16-le')) // 2 > 4000:
                reason = 'value_too_long'
            cells.append({'cell': address, **cell, 'dynamic_formula_anchor': dynamic_anchor,
                          'editable': reason is None, 'read_only_reason': reason, 'merged_anchor': origin})
        rows.append(cells)
    part = grid_styles.sheet_part(source, sheet)
    with zipfile.ZipFile(source) as archive:
        root = ooxml.parse_xml(archive.read(part)).getroot(); views = root.findall(grid_styles.X + 'sheetViews')
        validations = grid_validations(root, (min_col, min_row, max_col, max_row))
        conditional_formats = grid_conditional_formats(root, (min_col, min_row, max_col, max_row))
        hyperlinks, hyperlink_inspection = grid_hyperlinks(root, archive, part, (min_col, min_row, max_col, max_row))
        if len(views) != 1 or len(views[0]) != 1 or views[0][0].tag != grid_styles.X + 'sheetView':
            view_info = {'pane': None, 'issue': 'unsupported_worksheet_views'}
        else:
            panes = views[0][0].findall(grid_styles.X + 'pane')
            view_info = {'pane': dict(panes[0].attrib) if len(panes) == 1 else None, 'issue': 'ambiguous_panes' if len(panes) > 1 else None}
    for row in rows:
        for cell in row: cell['hyperlink'] = hyperlinks.get(cell['cell'])
    resolve_grid_lists(source, sheet, validations)
    latest = info(directory)['latest_revision']
    if base_revision != latest:
        raise ValueError(f'Revision conflict: latest revision is {latest}')
    return {'revision': base_revision, 'sheet': sheet, 'range': cell_range, 'sheet_protected': constraints['sheet_protected'], 'dimensions': constraints['dimensions'], 'worksheet_view': view_info,
            'rows': rows, 'data_validation': validations, 'conditional_formatting': conditional_formats, 'hyperlink_inspection': hyperlink_inspection, 'cache_freshness': 'unknown'}


def table_formula_node(part, node):
    ns = grid_styles.X
    tag = node.tag
    if part.startswith('xl/worksheets/') and part.endswith('.xml') and tag in (ns+'f', ns+'formula', ns+'formula1', ns+'formula2'): return True
    if part.startswith('xl/tables/') and part.endswith('.xml') and tag in (ns+'calculatedColumnFormula', ns+'totalsRowFormula'): return True
    if part == 'xl/workbook.xml' and tag == ns+'definedName': return True
    return part.startswith('xl/charts/') and tag == '{http://schemas.openxmlformats.org/drawingml/2006/chart}f'


def rewrite_table_formula(formula, table_name, old_name, new_name, implicit_table=None):
    prefix = '=' if formula.startswith('=') else ''; expression = formula[len(prefix):]
    try: tokens = Tokenizer('=' + expression).items
    except (TokenizerError, ValueError, IndexError) as exc: raise ValueError('Unsupported structured formula') from exc
    if ''.join(token.value for token in tokens) != expression: raise ValueError('Structured formula cannot be preserved exactly')
    updated = []; count = 0
    for token in tokens:
        value = token.value
        if '[' in value or ']' in value:
            match = re.fullmatch(r'([A-Za-z_][A-Za-z0-9_.]*)?(\[.*\])', value) if token.type == 'OPERAND' and token.subtype == 'RANGE' else None
            if match is None: raise ValueError('Unsupported structured reference')
            owner, selector = match.groups()
            if owner is None and implicit_table is None: raise ValueError('Implicit structured reference has ambiguous table context')
            target = (implicit_table if owner is None else owner.casefold() == table_name.casefold())
            # shortcut: escaped structured names need a full Excel grammar before they can be rewritten.
            nested = selector.startswith('[[') and selector.endswith(']')
            inner = selector[1:-1] if nested else selector
            # XlsxWriter emits the compact form [[#This Row],Column] for calculated columns.
            compact = re.fullmatch(r'\[(#This Row)\],([^\[\]#@,\':]+)', inner, re.IGNORECASE) if nested else None
            if compact:
                inner = '[' + compact[1] + '],[' + compact[2] + ']'
            pieces = list(re.finditer(r'\[([^\[\]]+)\]', inner))
            if not pieces or (not nested and (len(pieces) != 1 or pieces[0].span() != (0, len(inner)))): raise ValueError('Unsupported structured reference')
            if pieces[0].start() != 0 or pieces[-1].end() != len(inner) or any(inner[a.end():b.start()] not in (',', ':') for a,b in zip(pieces,pieces[1:])): raise ValueError('Unsupported structured reference')
            replacements = []
            for piece in pieces:
                column = piece.group(1); marker = '@' if column.startswith('@') else ''; column = column[len(marker):]
                if column.startswith('#'):
                    if marker or column.casefold() not in ('#all', '#data', '#headers', '#totals', '#this row'): raise ValueError('Unsupported structured reference selector')
                elif not column or re.search(r"[#@,':]", column): raise ValueError('Unsupported structured reference escaping')
                elif target and column.casefold() == old_name.casefold():
                    if re.search(r"[#@,':]", new_name): raise ValueError('New name requires unsupported structured-reference escaping')
                    replacements.append((piece.start(1), piece.end(1), marker + new_name)); count += 1
            for a,b,replacement in reversed(replacements): inner = inner[:a] + replacement + inner[b:]
            if compact: inner = inner[:inner.index(',')+1] + inner[inner.index(',')+2:-1]
            selector = '[' + inner + ']' if nested else inner
            value = (owner or '') + selector
        updated.append(value)
    return prefix + ''.join(updated), count


def rename_table_column(directory, base_revision, sheet, table_name, old_name, new_name):
    """Rename one table header and its column metadata in one revision."""
    if not isinstance(sheet, str) or not sheet or '/' in sheet or '\x00' in sheet: raise ValueError('Invalid sheet name')
    for value in (table_name, old_name, new_name):
        if not isinstance(value, str) or not value or value != value.strip() or len(value.encode('utf-16-le')) // 2 > 255 or re.search(r'[\x00-\x1f\ud800-\udfff\ufffe\uffff\[\]]|_x[0-9a-fA-F]{4}_', value): raise ValueError('Invalid table or column name')
    if old_name == new_name: raise ValueError('The new column name must differ')
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'xlsx': raise ValueError('Table column rename requires an XLSX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.xlsx'
    with tempfile.TemporaryDirectory(dir=state['directory'], prefix='.office-table-column-') as tmp:
        snapshot, digest = snapshots.copy(source, tmp)
        constraints = spreadsheet.worksheet_constraints(snapshot, sheet)
        if constraints['sheet_protected']: raise ValueError('Protected worksheets cannot be edited')
        matches = [table for table in constraints['table_parts'] if table['name'] == table_name and table['display_name'] == table_name]
        if len(matches) != 1 or matches[0]['header_rows'] != 1: raise ValueError('A unique table with a header row is required')
        table = matches[0]; c1, r1, c2, r2 = range_boundaries(table['ref'])
        if c2 - c1 + 1 > 200: raise ValueError('Table exceeds 200 columns')
        if any(a <= c2 and c1 <= c and b <= r1 <= d for a,b,c,d in constraints['merged_ranges'] + constraints['formula_ranges']): raise ValueError('Table header overlaps a merged or formula range')
        with zipfile.ZipFile(snapshot) as archive:
            parts = archive.namelist()
            if any(name.startswith(('customXml/', 'xl/pivotTables/', 'xl/pivotCache/', 'xl/slicers/', 'xl/slicerCaches/', 'xl/queryTables/', 'xl/externalLinks/', 'xl/embeddings/', 'xl/webextensions/', 'xl/activeX/', 'xl/model/', 'xl/richData/')) or name == 'xl/connections.xml' or name.startswith('xl/') and name.endswith('.bin') for name in parts): raise ValueError('Workbook contains unsupported dependent features')
            sheet_part = grid_styles.sheet_part(snapshot, sheet); table_owners = []; table_relationships = {}
            for rel_path in (name for name in parts if name.endswith('.rels')):
                if archive.getinfo(rel_path).file_size > 4 * 1024 * 1024: raise ValueError('Relationship part exceeds table rename inspection limit')
                rel_root = ooxml.parse_xml(archive.read(rel_path)).getroot()
                if rel_root.tag != grid_styles.REL+'Relationships': raise ValueError('Invalid package relationships')
                if rel_path != '_rels/.rels' and '/_rels/' not in rel_path: raise ValueError('Unsupported package relationship path')
                owner = posixpath.join(posixpath.dirname(posixpath.dirname(rel_path)), posixpath.basename(rel_path)[:-5])
                relationships = rel_root.findall(grid_styles.REL+'Relationship'); ids = [rel.get('Id') for rel in relationships]
                if len(relationships) != len(rel_root) or any(not value for value in ids) or len(ids) != len(set(ids)): raise ValueError('Invalid package relationship IDs')
                for rel in relationships:
                    mode = rel.get('TargetMode', 'Internal')
                    if mode not in ('Internal', 'External'): raise ValueError('Invalid package relationship mode')
                    if mode == 'External': continue
                    target = rel.get('Target', '')
                    if not target or any(char in target for char in ('%', '?', '#', chr(92), ':')): raise ValueError('Unsupported package relationship for table rename')
                    resolved = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(posixpath.dirname(owner), target))
                    table_relationships.setdefault(resolved, []).append((owner, rel.get('Type')))
                    if resolved == table['part']: table_owners.append((owner, rel.get('Type')))
            if table_owners != [(sheet_part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/table')]: raise ValueError('Table part has shared or unsupported relationships')
            book = load_workbook(snapshot, read_only=True, data_only=False, keep_links=False)
            try:
                worksheet_paths = {}
                for worksheet in book.worksheets:
                    path = getattr(worksheet, '_worksheet_path', None)
                    if isinstance(path, str): worksheet_paths.setdefault(path, []).append(worksheet.title)
            finally: book.close()
            sheet_constraints = {sheet_part: constraints}
            def table_context(path, address=None, part=None):
                titles = worksheet_paths.get(path, [])
                if len(titles) != 1: return None
                if path not in sheet_constraints: sheet_constraints[path] = spreadsheet.worksheet_constraints(snapshot, titles[0])
                tables = sheet_constraints[path]['table_parts']
                if part is not None:
                    matching = [item for item in tables if item['part'] == part]
                else:
                    if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', address or ''): return None
                    col, row = range_boundaries(address)[:2]
                    matching = [item for item in tables if (lambda bounds: bounds[0] <= col <= bounds[2] and bounds[1] + item['header_rows'] <= row <= bounds[3])(range_boundaries(item['ref']))]
                if len(matching) != 1 or table_relationships.get(matching[0]['part']) != [(path, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/table')]: return None
                return matching[0]['part'] == table['part']
            xml_parts = [name for name in parts if name.startswith('xl/') and name.endswith('.xml')]
            if len(xml_parts) > 1000: raise ValueError('Workbook has too many XML parts to inspect')
            table_declarations = []; rewrites = {}; rewrite_count = 0
            for name in xml_parts:
                if archive.getinfo(name).file_size > 8 * 1024 * 1024: raise ValueError('Workbook XML part exceeds table rename inspection limit')
                root = ooxml.parse_xml(archive.read(name)).getroot()
                if root.tag == grid_styles.X+'table' and any((root.get(key) or '').casefold() == table_name.casefold() for key in ('name','displayName')): table_declarations.append(name)
                formula_index = 0
                for node in root.iter():
                    if not isinstance(node.tag, str): continue
                    eligible = table_formula_node(name, node)
                    if eligible: formula_index += 1
                    text = node.text or ''; other = [node.tail or '', *node.attrib.values()]
                    if any('[' in value or ']' in value or re.search(r'_x005[bBdD]_', value) for value in other) or re.search(r'_x005[bBdD]_', text): raise ValueError('Workbook contains ambiguous structured-reference text')
                    if node.tag.rsplit('}',1)[-1].lower() in ('f','formula','formula1','formula2','definedname','calculatedcolumnformula','totalsrowformula') and re.search(r'(?i)\b(?:INDIRECT|EVALUATE)\s*\(', text): raise ValueError('Workbook contains indirect references; rename requires separate review')
                    if '[' in text or ']' in text:
                        if not eligible: raise ValueError('Workbook contains possible structured references outside supported formulas')
                        implicit_table = None
                        if name.startswith('xl/tables/') and node.getparent().tag == grid_styles.X+'tableColumn' and node.getparent().getparent() is not None and node.getparent().getparent().tag == grid_styles.X+'tableColumns' and node.getparent().getparent().getparent() is root:
                            owners = table_relationships.get(name, [])
                            if len(owners) == 1 and owners[0][1] == 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/table': implicit_table = table_context(owners[0][0], part=name)
                        elif name.startswith('xl/worksheets/') and node.tag == grid_styles.X+'f' and node.getparent().tag == grid_styles.X+'c':
                            implicit_table = table_context(name, address=node.getparent().get('r'))
                        updated, count = rewrite_table_formula(text, table_name, old_name, new_name, implicit_table)
                        if count:
                            if name.startswith('xl/tables/') and (node.attrib or len(node) or node.getparent().tag != grid_styles.X+'tableColumn'): raise ValueError('Unsupported table formula metadata')
                            if node.tag == grid_styles.X+'f' and (node.getparent().tag != grid_styles.X+'c' or node.get('t') not in (None, 'normal') or any(key != 't' for key in node.attrib)): raise ValueError('Shared or array formulas require separate review')
                            rewrites.setdefault(name, []).append((formula_index, text, updated)); rewrite_count += count
                            if rewrite_count > 500 or len(rewrites) > 100: raise ValueError('Structured-reference rewrite limit exceeded')
            if table_declarations != [table['part']]: raise ValueError('Table name is duplicated or ambiguous')
            tree = ooxml.parse_xml(archive.read(table['part'])); root = tree.getroot(); ns = grid_styles.X
            columns = root.findall(ns+'tableColumns')
            if len(columns) != 1 or root.get('tableType') not in (None, 'worksheet') or root.get('connectionId') not in (None, '0') or len(list(root.iter(ns+'extLst'))) or root.find('.//'+ns+'sortState') is not None: raise ValueError('Unsupported table structure')
            filter_node = root.find(ns+'autoFilter')
            if filter_node is not None and len(filter_node): raise ValueError('Table filter criteria require separate review')
            nodes = columns[0].findall(ns+'tableColumn')
            if len(nodes) != c2-c1+1 or len(nodes) != len(columns[0]) or columns[0].get('count') not in (None, str(len(nodes))): raise ValueError('Table columns are inconsistent')
            names = [node.get('name') for node in nodes]
            ids = [node.get('id') for node in nodes]
            if any(not name for name in names) or len({name.casefold() for name in names}) != len(names) or any(not value or not value.isdecimal() or len(value) > 8 or not 1 <= int(value) <= 2147483647 for value in ids) or len({int(value) for value in ids}) != len(ids): raise ValueError('Table column metadata is ambiguous')
            if sum(name == old_name for name in names) != 1 or any(name.casefold() == new_name.casefold() and name != old_name for name in names): raise ValueError('Column missing or new name conflicts')
            for node in nodes:
                if node.get('uniqueName') is not None or node.get('queryTableFieldId') is not None: raise ValueError('Linked table columns require separate review')
                children = list(node)
                if any(child.tag not in (ns+'calculatedColumnFormula', ns+'totalsRowFormula') or child.attrib or len(child) for child in children) or len({child.tag for child in children}) != len(children): raise ValueError('Unsupported table column metadata')
        first = f'{get_column_letter(c1)}{r1}'; last = f'{get_column_letter(c2)}{r1}'
        cells = spreadsheet.read_range(snapshot, sheet, f'{first}:{last}')['rows'][0]
        if len(cells) != len(names) or any(cell.get('value') != name or 'formula' in cell or cell.get('text_complete') is False for cell,name in zip(cells,names)): raise ValueError('Table headers and column metadata do not match')
        index = names.index(old_name); address = f'{get_column_letter(c1 + index)}{r1}'
        with zipfile.ZipFile(snapshot) as archive:
            worksheet = ooxml.parse_xml(archive.read(sheet_part)).getroot()
            matches = [cell for cell in worksheet.findall(ns+'sheetData/'+ns+'row/'+ns+'c') if cell.get('r') == address]
            if len(matches) != 1: raise ValueError('Table header cell is unavailable or ambiguous')
            cell = matches[0]; kind = cell.get('t')
            if kind == 'inlineStr':
                inline = cell.findall(ns+'is')
                if len(inline) != 1 or len(inline[0]) != 1 or inline[0][0].tag != ns+'t': raise ValueError('Rich or unsupported table header text requires separate review')
            elif kind == 's':
                values = cell.findall(ns+'v')
                if len(values) != 1 or not re.fullmatch(r'[0-9]{1,10}', values[0].text or ''): raise ValueError('Invalid table header shared string')
                manifest = ooxml.parse_xml(archive.read('[Content_Types].xml')).getroot()
                ct = '{http://schemas.openxmlformats.org/package/2006/content-types}'
                shared = [item.get('PartName', '') for item in manifest.findall(ct+'Override') if item.get('ContentType') == grid_values.SHARED]
                if len(shared) != 1 or not shared[0].startswith('/') or shared[0][1:] not in archive.namelist() or archive.getinfo(shared[0][1:]).file_size > 8 * 1024 * 1024: raise ValueError('Shared string table is unavailable')
                strings = ooxml.parse_xml(archive.read(shared[0][1:])).getroot()
                if strings.tag != ns+'sst': raise ValueError('Invalid shared string table')
                items = strings.findall(ns+'si'); number = int(values[0].text)
                if number >= len(items) or len(items[number]) != 1 or items[number][0].tag != ns+'t': raise ValueError('Rich or unsupported table header text requires separate review')
            elif kind != 'str' or len(cell.findall(ns+'v')) != 1:
                raise ValueError('Unsupported table header cell storage')
        candidate = Path(tmp) / 'cell.xlsx'
        change = {'cell': address, 'value': new_name, 'type': 'text'}
        operation = {'command': 'set', 'path': f'/{sheet}/{address}', 'props': {'value': 'office-use-literal-text'}}
        result = grid_values.apply(snapshot, candidate, sheet, [change], [operation])
        def edit(tree, archive):
            current = tree.getroot().findall(ns+'tableColumns/'+ns+'tableColumn')
            if len(current) != len(names) or [item.get('name') for item in current] != names: raise ValueError('Table column metadata changed')
            current[index].set('name', new_name)
            additional = {}
            for part, edits in rewrites.items():
                raw = archive.read(part); document = tree if part == table['part'] else ooxml.parse_xml(raw); root = document.getroot()
                formulas = [node for node in root.iter() if isinstance(node.tag, str) and table_formula_node(part, node)]
                sheet_changed = False
                for ordinal, old, replacement in edits:
                    if ordinal > len(formulas) or formulas[ordinal-1].text != old: raise ValueError('Formula changed during table rename')
                    node = formulas[ordinal-1]; node.text = replacement
                    if node.tag == ns+'f':
                        cell = node.getparent(); cached = cell.findall(ns+'v')
                        if cell.tag != ns+'c' or len(cached) > 1: raise ValueError('Unsupported formula cache')
                        for item in cached: cell.remove(item)
                        sheet_changed = True
                if sheet_changed: cell_text.request_recalculation(root)
                if part != table['part']: additional[part] = ooxml.serialize_xml(document, raw)
            if rewrites:
                raw = additional.get('xl/workbook.xml', archive.read('xl/workbook.xml')); document = ooxml.parse_xml(raw); root = document.getroot()
                if root.tag != ns+'workbook': raise ValueError('Invalid workbook XML')
                calculations = root.findall(ns+'calcPr')
                if len(calculations) > 1: raise ValueError('Ambiguous workbook calculation settings')
                if calculations: calculation = calculations[0]
                else:
                    from lxml import etree
                    calculation = etree.Element(ns+'calcPr')
                    later = {ns+key for key in ('oleSize','customWorkbookViews','pivotCaches','smartTagPr','smartTagTypes','webPublishing','fileRecoveryPr','webPublishObjects','extLst')}
                    position = next((i for i,item in enumerate(root) if item.tag in later),len(root)); root.insert(position,calculation)
                for key in ('fullCalcOnLoad','forceFullCalc'):
                    if calculation.get(key) not in (None,'0','1','true','false'): raise ValueError('Invalid calculation setting')
                    calculation.set(key,'1')
                additional['xl/workbook.xml'] = ooxml.serialize_xml(document, raw)
                manifest = ooxml.parse_xml(archive.read('[Content_Types].xml'))
                chain, _ = grid_values.remove_calculation_chain(archive, manifest)
                if set(chain) & set(additional): raise ValueError('Calculation chain conflicts with formula edits')
                additional.update(chain)
            return additional
        finished = Path(tmp) / 'finished.xlsx'
        ooxml.edit_xml(candidate, finished, lambda _: table['part'], edit, allow_remove=True)
        with zipfile.ZipFile(finished) as archive:
            entries = archive.infolist()
            if finished.stat().st_size > 128 * 1024 * 1024 or len(entries) > 5000 or len({item.filename for item in entries}) != len(entries) or sum(item.file_size for item in entries) > 128 * 1024 * 1024:
                raise ValueError('Renamed workbook exceeds package limits')
        try: os.link(finished, output)
        except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'sheet': sheet, 'table': table_name, 'cell': address,
            'old_name': old_name, 'new_name': new_name, 'snapshot_sha256': digest, 'rewritten_references': rewrite_count, 'formula_cache_effect': result['formula_cache_effect'], 'chart_cache_effect': 'may_be_stale'}


def edit_grid(directory, base_revision, sheet, cell_range, changes, expected_sha256=None):
    if not isinstance(changes, list) or not changes or len(changes) > 200:
        raise ValueError('Provide 1 to 200 changed cells')
    grid = read_grid(directory, base_revision, sheet, cell_range)
    if expected_sha256 is not None and grid['snapshot_sha256'] != expected_sha256: raise ValueError('Workbook snapshot changed; reload the grid')
    min_col, min_row, max_col, max_row = grid_bounds(cell_range)
    original = {cell['cell']: cell for row in grid['rows'] for cell in row}
    seen, operations, normalized = set(), [], []
    for change in changes:
        if not isinstance(change, dict) or set(change) not in ({'cell', 'value'}, {'cell', 'value', 'type'}):
            raise ValueError('Each change requires cell and value and may specify type')
        kind = change.get('type', 'auto')
        if kind not in ('auto', 'text'): raise ValueError('Cell type must be auto or text')
        address, value = change['cell'], change['value']
        match = re.fullmatch(r'([A-Za-z]{1,3})([1-9][0-9]{0,6})', address) if isinstance(address, str) else None
        if match is None:
            raise ValueError('Invalid cell address')
        col, row = column_index_from_string(match.group(1)), int(match.group(2))
        address = match.group(1).upper() + str(row)
        if not (min_col <= col <= max_col and min_row <= row <= max_row) or address in seen:
            raise ValueError('Cell is outside the range or repeated')
        if not original[address]['editable']:
            raise ValueError(f"Cell cannot be edited: {address} ({original[address]['read_only_reason']}); Array, data-table and shared formula ranges are protected")
        if not isinstance(value, str) or len(value) > 4000:
            raise ValueError('Cell values must be text of at most 4,000 characters')
        if kind == 'text' and (not value or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', value)):
            raise ValueError('Explicit text must be nonempty and contain valid XML characters')
        if kind == 'auto' and value.lstrip().startswith('=') and not value.startswith('='):
            raise ValueError('Formula expressions must begin with =')
        if kind == 'auto' and value == '=':
            raise ValueError('A formula requires an expression')
        seen.add(address)
        mode = 'formula' if kind == 'auto' and value.startswith('=') else 'value'
        if mode == 'value': spreadsheet.validate_cell_value(value)
        operations.append({'command': 'set', 'path': f'/{sheet}/{address}', 'props': {mode: 'office-use-literal-text' if kind == 'text' else value}})
        normalized.append({'cell': address, 'value': value, 'type': kind})
    state = info(directory); source = revision_file(state, base_revision)
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.xlsx'
    try: result = grid_values.apply(source, output, sheet, normalized, operations, expected_sha256=grid['snapshot_sha256'])
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': result['output'], 'result': result['result'], 'formula_cache_effect': result['formula_cache_effect'], 'recalculation_requested': True, 'calculation_chain': result['calculation_chain'], 'dynamic_formula_cells': result['dynamic_formula_cells'], 'data_validation': {'evaluated': False, 'coverage_complete': grid['data_validation']['complete'], 'affected_cells': sorted(seen & {address for rule in grid['data_validation']['rules'] for address in rule['cells']})}}


def format_grid(directory, base_revision, sheet, cell_range, style_range, properties, expected_sha256=None):
    allowed = {'bold', 'italic', 'underline', 'font_name', 'font_size', 'font_color', 'fill', 'horizontal', 'vertical', 'wrap', 'numberformat', 'border', 'border_color', 'border_top', 'border_bottom', 'border_left', 'border_right', 'border_scope', 'border_top_color', 'border_bottom_color', 'border_left_color', 'border_right_color'}
    if not isinstance(properties, dict) or not properties or not set(properties) <= allowed:
        raise ValueError('Provide supported cell formatting properties')
    props = {}
    for name, target in (('bold', 'font.bold'), ('italic', 'font.italic'), ('wrap', 'alignment.wrapText')):
        if name in properties:
            if not isinstance(properties[name], bool): raise ValueError(f'{name} must be true or false')
            props[target] = str(properties[name]).lower()
    for name, target in (('font_color', 'font.color'), ('fill', 'fill')):
        if name in properties:
            if not isinstance(properties[name], str) or not re.fullmatch(r'[0-9A-Fa-f]{6}', properties[name]): raise ValueError(f'{name} must be a six-digit RGB color')
            props[target] = properties[name]
    for name, target, choices in (('horizontal', 'alignment.horizontal', ('left', 'center', 'right', 'justify', 'fill', 'distributed')), ('vertical', 'alignment.vertical', ('top', 'center', 'bottom')), ('underline', 'underline', ('none', 'single', 'double', 'singleAccounting', 'doubleAccounting'))):
        if name in properties:
            if properties[name] not in choices: raise ValueError(f'Invalid {name}')
            props[target] = properties[name]
    border_names = ('border', 'border_top', 'border_bottom', 'border_left', 'border_right')
    if 'border' in properties and any(name in properties for name in border_names[1:]): raise ValueError('Choose either all borders or individual sides')
    for name in border_names:
        if name in properties:
            if properties[name] not in ('thin', 'medium', 'thick', 'double', 'dashed', 'dotted', 'none'): raise ValueError(f'Invalid {name}')
            props['border.all' if name == 'border' else 'border.' + name[7:]] = properties[name]
    if 'border_color' in properties:
        color = properties['border_color']
        if not isinstance(color, str) or not re.fullmatch(r'[0-9A-Fa-f]{6}', color): raise ValueError('border_color must be a six-digit RGB color')
        if properties.get('border') in (None, 'none'): raise ValueError('border_color requires an all-sides border style')
        props['border.color'] = color
    scope = properties.get('border_scope', 'all')
    if scope not in ('all', 'outline', 'inside') or ('border_scope' in properties and 'border' not in properties): raise ValueError('Border scope requires an all-sides border style')
    for side in ('top', 'bottom', 'left', 'right'):
        key = f'border_{side}_color'
        if key in properties:
            color = properties[key]
            if not isinstance(color, str) or not re.fullmatch(r'[0-9A-Fa-f]{6}', color): raise ValueError(f'{key} must be a six-digit RGB color')
            if properties.get(f'border_{side}') in (None, 'none'): raise ValueError(f'{key} requires a matching border style')
            props[f'border.{side}.color'] = color
    if 'font_name' in properties:
        value = properties['font_name']
        if not isinstance(value, str) or not value.strip() or len(value) > 100 or any(ord(c) < 32 for c in value): raise ValueError('Font name must contain 1 to 100 characters without control characters')
        props['font.name'] = value
    if 'font_size' in properties:
        value = properties['font_size']
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not 1 <= value <= 409 or not math.isfinite(value): raise ValueError('Font size must be between 1 and 409 points')
        props['font.size'] = f'{value}pt'
    if 'numberformat' in properties:
        value = properties['numberformat']
        if not isinstance(value, str) or not value or len(value) > 120:
            raise ValueError('Number format must contain 1 to 120 characters')
        props['numberformat'] = value
    grid = read_grid(directory, base_revision, sheet, cell_range)
    if expected_sha256 is not None and grid['snapshot_sha256'] != expected_sha256: raise ValueError('Workbook snapshot changed; reload the grid')
    if grid['sheet_protected']: raise ValueError('Grid formatting is unavailable for protected worksheets')
    c1, r1, c2, r2 = grid_bounds(cell_range)
    sc1, sr1, sc2, sr2 = grid_bounds(style_range)
    if not (c1 <= sc1 <= sc2 <= c2 and r1 <= sr1 <= sr2 <= r2):
        raise ValueError('Formatting range must be within the loaded range')
    operations = []
    for row in grid['rows']:
        for cell in row:
            match = re.fullmatch(r'([A-Z]+)([0-9]+)', cell['cell'])
            col, number = column_index_from_string(match.group(1)), int(match.group(2))
            if sc1 <= col <= sc2 and sr1 <= number <= sr2:
                cell_props = props.copy()
                if scope != 'all':
                    cell_props.pop('border.all'); cell_props.pop('border.color', None)
                    edges = {'top': number == sr1, 'bottom': number == sr2, 'left': col == sc1, 'right': col == sc2}
                    for side, outer in edges.items():
                        if outer == (scope == 'outline'):
                            cell_props[f'border.{side}'] = properties['border']
                            if 'border_color' in properties: cell_props[f'border.{side}.color'] = properties['border_color']
                if not cell_props: continue
                if not cell['editable'] or cell['dynamic_formula_anchor']:
                    raise ValueError(f"Cell cannot be formatted in the grid: {cell['cell']} ({cell['read_only_reason'] or 'formula_range'})")
                operations.append({'command': 'set', 'path': f"/{sheet}/{cell['cell']}", 'props': cell_props})
    if not operations: raise ValueError('Formatting selection contains no applicable changes')
    state = info(directory); source = revision_file(state, base_revision)
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.xlsx'
    try: result = grid_styles.apply(source, output, sheet, [op['path'].rsplit('/', 1)[1] for op in operations], operations, expected_sha256=grid['snapshot_sha256'])
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': result['output'], 'result': result['result'], 'formula_cache_effect': result['formula_cache_effect']}


def merge_cells(directory, base_revision, sheet, cell_range, merge=True, expected_sha256=None):
    if not isinstance(merge, bool): raise ValueError('merge must be true or false')
    c1, r1, c2, r2 = grid_bounds(cell_range)
    if (c1, r1) == (c2, r2): raise ValueError('A merged range requires at least two cells')
    address = f'{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}'
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'xlsx': raise ValueError('Merging cells requires an XLSX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    with tempfile.TemporaryDirectory(prefix='office-grid-merge-') as tmp:
        source, digest = snapshots.copy(source, tmp)
        if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Workbook snapshot changed; reload the grid')
        constraints = spreadsheet.worksheet_constraints(source, sheet)
        if constraints['sheet_protected']: raise ValueError('Protected worksheets cannot be merged')
        def overlap(a, b): return a[0] <= b[2] and b[0] <= a[2] and a[1] <= b[3] and b[1] <= a[3]
        bounds = (c1, r1, c2, r2)
        if merge and any(overlap(bounds, other) for other in constraints['formula_ranges']): raise ValueError('Merged range overlaps a multi-cell formula')
        output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.xlsx'; ns = grid_styles.X
        part = grid_styles.sheet_part(source, sheet)
        def edit(tree, archive):
            from lxml import etree
            root = tree.getroot(); data = root.findall(ns + 'sheetData'); containers = root.findall(ns + 'mergeCells')
            if root.tag != ns + 'worksheet' or len(data) != 1 or len(list(root.iter(ns + 'sheetData'))) != 1 or len(containers) > 1 or len(list(root.iter(ns + 'mergeCells'))) != len(containers): raise ValueError('Ambiguous worksheet structure')
            container = containers[0] if containers else None; matches = []
            if container is not None:
                count = container.get('count')
                if set(container.attrib) - {'count'} or (count is not None and (not re.fullmatch(r'[0-9]{1,8}', count) or int(count) != len(container))): raise ValueError('Invalid merge collection')
                for item in container:
                    if item.tag != ns + 'mergeCell' or set(item.attrib) != {'ref'} or len(item): raise ValueError('Unsupported merge record')
                    try: other = range_boundaries(item.get('ref'))
                    except (ValueError, TypeError): raise ValueError('Invalid merge record') from None
                    if any(v is None for v in other) or not (1 <= other[0] <= other[2] <= 16384 and 1 <= other[1] <= other[3] <= 1048576): raise ValueError('Invalid merge record')
                    if overlap(bounds, other): matches.append((item, other))
            if merge:
                if matches: raise ValueError('Range overlaps an existing merged range')
                rel_path = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
                relationships = []
                if rel_path in archive.namelist():
                    rels = ooxml.parse_xml(archive.read(rel_path)).getroot()
                    if rels.tag != grid_styles.REL + 'Relationships': raise ValueError('Invalid worksheet relationships')
                    relationships = rels.findall(grid_styles.REL + 'Relationship')
                    ids = [item.get('Id') for item in relationships]
                    if any(not rid for rid in ids) or len(ids) != len(set(ids)): raise ValueError('Ambiguous worksheet relationships')
                base = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
                def related(item, expected_type):
                    target = item.get('Target', '')
                    if item.get('Type') != expected_type or item.get('TargetMode', 'Internal') != 'Internal' or not target or any(c in target for c in ('%', '?', '#', chr(92), ':')): raise ValueError('Unsupported related worksheet part')
                    path = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(posixpath.dirname(part), target))
                    if not path.startswith('xl/') or path not in archive.namelist(): raise ValueError('Invalid related worksheet part')
                    return ooxml.parse_xml(archive.read(path)).getroot()
                table_parts = root.findall(ns + 'tableParts')
                if len(table_parts) > 1 or len(list(root.iter(ns + 'tableParts'))) != len(table_parts): raise ValueError('Ambiguous table parts')
                table_rels = [item for item in relationships if item.get('Type', '').lower().endswith('/table')]
                table_ids = []
                if table_parts:
                    collection = table_parts[0]; count = collection.get('count')
                    if set(collection.attrib) - {'count'} or (count is not None and (not re.fullmatch(r'[0-9]{1,8}', count) or int(count) != len(collection))): raise ValueError('Invalid table parts')
                    rid_key = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
                    for item in collection:
                        if item.tag != ns + 'tablePart' or set(item.attrib) != {rid_key} or len(item): raise ValueError('Unsupported table part')
                        table_ids.append(item.get(rid_key))
                if len(table_ids) != len(set(table_ids)) or set(table_ids) != {item.get('Id') for item in table_rels}: raise ValueError('Ambiguous table references')
                for rel in table_rels:
                    table = related(rel, base + 'table')
                    if table.tag != ns + 'table': raise ValueError('Invalid table part')
                    try: table_bounds = range_boundaries(table.get('ref'))
                    except (TypeError, ValueError): raise ValueError('Invalid table range') from None
                    if any(v is None for v in table_bounds) or not (1 <= table_bounds[0] <= table_bounds[2] <= 16384 and 1 <= table_bounds[1] <= table_bounds[3] <= 1048576): raise ValueError('Invalid table range')
                    if overlap(bounds, table_bounds): raise ValueError('Merged range overlaps a table')
                comment_rels = [item for item in relationships if item.get('Type', '').lower().endswith('/comments')]
                threaded_rels = [item for item in relationships if item.get('Type', '').lower().endswith('/threadedcomment')]
                thread_type = 'http://schemas.microsoft.com/office/2017/10/relationships/threadedComment'
                if len(comment_rels) > 1 or len(threaded_rels) > 1: raise ValueError('Ambiguous comment references')
                for rel in comment_rels:
                    comments = related(rel, base + 'comments'); lists = comments.findall(ns + 'commentList')
                    if comments.tag != ns + 'comments' or len(lists) != 1 or len(list(comments.iter(ns + 'commentList'))) != 1 or any(child.tag not in (ns + 'authors', ns + 'commentList') for child in comments): raise ValueError('Unsupported comment structure')
                    seen_comments = set()
                    for comment in lists[0]:
                        ref = comment.get('ref', '')
                        if comment.tag != ns + 'comment' or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', ref) or ref in seen_comments: raise ValueError('Invalid comment reference')
                        seen_comments.add(ref); col, number = range_boundaries(ref)[:2]
                        if col > 16384 or number > 1048576: raise ValueError('Invalid comment reference')
                        if c1 <= col <= c2 and r1 <= number <= r2 and (col, number) != (c1, r1): raise ValueError(f'Non-anchor cell {ref} has a comment')
                thread_ns = '{http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments}'
                for rel in threaded_rels:
                    thread = related(rel, thread_type)
                    if thread.tag != thread_ns + 'ThreadedComments' or any(child.tag != thread_ns + 'threadedComment' for child in thread): raise ValueError('Unsupported threaded comment structure')
                    for comment in thread:
                        ref = comment.get('ref', '')
                        if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', ref): raise ValueError('Invalid threaded comment reference')
                        col, number = range_boundaries(ref)[:2]
                        if col > 16384 or number > 1048576: raise ValueError('Invalid threaded comment reference')
                        if c1 <= col <= c2 and r1 <= number <= r2 and (col, number) != (c1, r1): raise ValueError(f'Non-anchor cell {ref} has a threaded comment')
                for link in root.iter(ns + 'hyperlink'):
                    try: linked = range_boundaries(link.get('ref'))
                    except (ValueError, TypeError): raise ValueError('Invalid hyperlink range') from None
                    if any(value is None for value in linked): raise ValueError('Invalid hyperlink range')
                    if overlap(bounds, linked) and (max(c1, linked[0]), max(r1, linked[1]), min(c2, linked[2]), min(r2, linked[3])) != (c1, r1, c1, r1): raise ValueError('Non-anchor cell contains a hyperlink')
                seen = set()
                for row in data[0].findall(ns + 'row'):
                    for cell in row.findall(ns + 'c'):
                        ref = cell.get('r', '')
                        if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', ref): raise ValueError('Invalid cell reference')
                        col, number = range_boundaries(ref)[:2]
                        if ref in seen: raise ValueError('Duplicate worksheet cell')
                        seen.add(ref)
                        if c1 <= col <= c2 and r1 <= number <= r2 and (col, number) != (c1, r1):
                            if len(cell) or set(cell.attrib) - {'r', 's', 't'} or cell.get('t') not in (None, 'n') or (cell.text and cell.text.strip()): raise ValueError(f'Non-anchor cell {ref} contains data or metadata')
                if container is None:
                    container = etree.Element(ns + 'mergeCells')
                    preceding = {ns + name for name in ('sheetData','sheetCalcPr','sheetProtection','protectedRanges','scenarios','autoFilter','sortState','dataConsolidate','customSheetViews')}
                    position = max(i for i, node in enumerate(root) if node.tag in preceding) + 1
                    root.insert(position, container)
                etree.SubElement(container, ns + 'mergeCell', ref=address)
            else:
                if len(matches) != 1 or matches[0][1] != bounds: raise ValueError('Select exactly one existing merged range')
                container.remove(matches[0][0])
            if len(container): container.set('count', str(len(container)))
            else: root.remove(container)
            dimensions = root.findall(ns + 'dimension')
            if len(dimensions) > 1: raise ValueError('Ambiguous worksheet dimension')
            if merge and dimensions:
                try: left, top, right, bottom = range_boundaries(dimensions[0].get('ref'))
                except (ValueError, TypeError): raise ValueError('Invalid worksheet dimension') from None
                if any(v is None for v in (left, top, right, bottom)) or not (1 <= left <= right <= 16384 and 1 <= top <= bottom <= 1048576): raise ValueError('Invalid worksheet dimension')
                dimensions[0].set('ref', f'{get_column_letter(min(left,c1))}{min(top,r1)}:{get_column_letter(max(right,c2))}{max(bottom,r2)}')
        try: ooxml.edit_xml(source, output, lambda _: part, edit)
        except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
        return {'revision': base_revision + 1, 'file': str(output), 'sheet': sheet, 'range': address, 'merged': merge, 'formula_cache_effect': 'preserved'}


def freeze_panes(directory, base_revision, sheet, cell, expected_sha256=None):
    if cell is not None:
        if not isinstance(cell, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', cell): raise ValueError('Freeze position must be an uppercase cell address or null')
        col, row = range_boundaries(cell)[:2]
        if col > 16384 or row > 1048576: raise ValueError('Freeze position is outside the worksheet')
        if cell == 'A1': cell = None
    state = info(directory); source = revision_file(state, base_revision)
    if state['format'] != 'xlsx': raise ValueError('Freeze panes require an XLSX session')
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    if base_revision >= 999999: raise ValueError('Session revision limit reached')
    output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.xlsx'
    ns = grid_styles.X
    def edit(tree, archive):
        root = tree.getroot(); containers = root.findall(ns + 'sheetViews')
        if root.tag != ns + 'worksheet' or len(containers) != 1 or len(list(root.iter(ns + 'sheetViews'))) != 1 or len(containers[0]) != 1 or containers[0][0].tag != ns + 'sheetView': raise ValueError('Unsupported worksheet views')
        view = containers[0][0]
        if view.get('workbookViewId') != '0' or any(child.tag not in (ns + 'pane', ns + 'selection') for child in view) or len(view.findall(ns + 'pane')) > 1: raise ValueError('Unsupported worksheet view')
        pane = view.find(ns + 'pane')
        if pane is not None and (pane.get('state') != 'frozen' or set(pane.attrib) - {'state', 'xSplit', 'ySplit', 'topLeftCell', 'activePane'} or len(pane)): raise ValueError('Unsupported existing pane')
        for selection in view.findall(ns + 'selection'):
            if len(selection) or set(selection.attrib) - {'pane', 'activeCell', 'sqref'}: raise ValueError('Unsupported existing selection')
        for child in list(view): view.remove(child)
        if cell is not None:
            attrs = {'topLeftCell':cell, 'state':'frozen'}
            if col > 1: attrs['xSplit'] = str(col - 1)
            if row > 1: attrs['ySplit'] = str(row - 1)
            active = 'bottomRight' if col > 1 and row > 1 else 'topRight' if col > 1 else 'bottomLeft'
            attrs['activePane'] = active
            from lxml import etree
            etree.SubElement(view, ns + 'pane', **attrs)
            if col > 1 and row > 1:
                etree.SubElement(view, ns + 'selection', pane='topRight')
                etree.SubElement(view, ns + 'selection', pane='bottomLeft')
            etree.SubElement(view, ns + 'selection', pane=active, activeCell=cell, sqref=cell)
        else:
            from lxml import etree
            etree.SubElement(view, ns + 'selection', activeCell='A1', sqref='A1')
    try: ooxml.edit_xml(source, output, lambda path: grid_styles.sheet_part(path, sheet), edit, expected_sha256=expected_sha256)
    except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
    return {'revision': base_revision + 1, 'file': str(output), 'sheet': sheet, 'freeze_panes': cell, 'formula_cache_effect': 'preserved'}


def resize_grid(directory, base_revision, sheet, cell_range, target_range, column_width=None, row_height=None, expected_sha256=None):
    for name, value, limit in (('column_width', column_width, 255), ('row_height', row_height, 409.5)):
        if value is not None and (isinstance(value, bool) or not isinstance(value, (int, float)) or not 0 < value <= limit or not math.isfinite(value)):
            raise ValueError(f'{name} must be greater than zero and at most {limit}')
    if column_width is None and row_height is None: raise ValueError('Provide a column width or row height')
    grid = read_grid(directory, base_revision, sheet, cell_range)
    if expected_sha256 is not None and grid['snapshot_sha256'] != expected_sha256: raise ValueError('Workbook snapshot changed; reload the grid')
    if grid['sheet_protected']: raise ValueError('Grid sizing is unavailable for protected worksheets')
    c1, r1, c2, r2 = grid_bounds(cell_range); sc1, sr1, sc2, sr2 = grid_bounds(target_range)
    if not (c1 <= sc1 <= sc2 <= c2 and r1 <= sr1 <= sr2 <= r2): raise ValueError('Sizing range must be within the loaded range')
    state = info(directory); source = revision_file(state, base_revision)
    if base_revision != state['latest_revision']: raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    with tempfile.TemporaryDirectory(prefix='office-grid-resize-') as tmp:
        source, digest = snapshots.copy(source, tmp)
        if digest != grid['snapshot_sha256']: raise ValueError('Workbook snapshot changed; reload the grid')
        constraints = spreadsheet.worksheet_constraints(source, sheet, (sc1, sr1, sc2, sr2), column_width is not None, row_height is not None)
        if constraints['dimension_issues']: raise ValueError('Ambiguous worksheet dimensions: ' + ', '.join(constraints['dimension_issues']))
        if base_revision >= 999999: raise ValueError('Session revision limit reached')
        output = Path(state['directory']) / f'revision-{base_revision + 1:06d}.xlsx'
        try: spreadsheet.resize_dimensions(source, output, sheet, (sc1, sr1, sc2, sr2), column_width, row_height)
        except FileExistsError as exc: raise ValueError('Revision conflict: another edit committed first') from exc
        return {'revision': base_revision + 1, 'file': str(output), 'sheet': sheet, 'target_range': target_range,
                'column_width': column_width, 'row_height': row_height, 'formula_cache_effect': 'preserved'}


def sheet_names(directory, base_revision):
    state = info(directory)
    source = revision_file(state, base_revision)
    if state['format'] != 'xlsx':
        raise ValueError('Worksheet listing requires an XLSX session')
    if base_revision != state['latest_revision']:
        raise ValueError(f"Revision conflict: latest revision is {state['latest_revision']}")
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        names = workbook.sheetnames
    finally:
        workbook.close()
    latest = info(directory)['latest_revision']
    if base_revision != latest:
        raise ValueError(f'Revision conflict: latest revision is {latest}')
    return {'revision': base_revision, 'sheets': names}

"""Check picture replacement, shared media isolation and revision conflicts."""
import io
import hashlib
import tempfile
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import session
import officecli_backend


def image(color):
    data = io.BytesIO(); Image.new('RGB', (32, 24), color).save(data, format='PNG'); return data.getvalue()

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); original_image = root / 'original.png'; original_image.write_bytes(image('red'))
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    first = slide.shapes.add_picture(str(original_image), Inches(1), Inches(1), Inches(2), Inches(2))
    first.crop_left = .2; first.rotation = 13
    second = slide.shapes.add_picture(str(original_image), Inches(4), Inches(1), Inches(2), Inches(2))
    source = root / 'source.pptx'; deck.save(source); original = source.read_bytes()
    draft = root / 'draft'; session.start(draft, source=source)
    replacement = image('blue'); path = f'/slide[1]/picture[@id={first.shape_id}]'
    expected = hashlib.sha256(original).hexdigest()
    altered=Presentation(source); altered.slides[0].shapes[0].rotation=21; alternate=root/'alternate.pptx';altered.save(alternate)
    revision_file=draft/'revision-000000.pptx';revision_file.write_bytes(alternate.read_bytes())
    try: session.replace_picture(draft, 0, path, replacement, expected_sha256=expected)
    except ValueError as exc: assert 'snapshot changed' in str(exc)
    else: raise AssertionError('Changed picture source accepted')
    assert session.info(draft)['latest_revision']==0
    revision_file.write_bytes(original)
    result = session.replace_picture(draft, 0, path, replacement, expected_sha256=expected)
    changed = Presentation(result['file']); a, b = changed.slides[0].shapes
    assert a.image.blob == replacement and b.image.blob == image('red')
    assert a.crop_left == first.crop_left and a.rotation == first.rotation and a.left == first.left and a.width == first.width
    validation = officecli_backend.audit(result['file'])['validation']
    assert validation['exit_code'] == 0, validation
    assert source.read_bytes() == original and (draft / 'revision-000000.pptx').read_bytes() == original
    try: session.replace_picture(draft, 0, path, replacement)
    except ValueError as exc: assert 'conflict' in str(exc).lower()
    else: raise AssertionError('Stale edit succeeded')
    try: session.replace_picture(draft, 1, path, b'bad image')
    except ValueError: pass
    else: raise AssertionError('Invalid image accepted')
    assert session.info(draft)['latest_revision'] == 1
    grouped = Presentation(); group_slide = grouped.slides.add_slide(grouped.slide_layouts[6])
    group = group_slide.shapes.add_group_shape(); nested = group.shapes.add_picture(str(original_image), Inches(1), Inches(1), Inches(2), Inches(2))
    grouped_source = root / 'grouped.pptx'; grouped.save(grouped_source)
    grouped_draft = root / 'grouped-draft'; session.start(grouped_draft, source=grouped_source)
    group_path = f'/slide[1]/group[@id={group.shape_id}]/picture[@id={nested.shape_id}]'
    grouped_result = session.replace_picture(grouped_draft, 0, group_path, replacement)
    assert Presentation(grouped_result['file']).slides[0].shapes[0].shapes[0].image.blob == replacement
    replacement_file = root / 'replacement.png'; replacement_file.write_bytes(replacement)
    cli_draft = root / 'cli-draft'; session.start(cli_draft, source=source)
    subprocess.run([sys.executable, str(Path(__file__).with_name('open_office_use.py')), 'session', 'replace-picture', str(cli_draft), '--base', '0', '--path', path, '--image', str(replacement_file), '--expected-sha256', expected], check=True, capture_output=True)
    assert Presentation(cli_draft / 'revision-000001.pptx').slides[0].shapes[0].image.blob == replacement
print('Picture smoke passed: shared images, geometry, revisions and invalid input')

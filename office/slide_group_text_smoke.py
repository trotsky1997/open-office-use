"""Check rendered text inside scaled and nested slide groups."""
import tempfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE
import open_office_use
import session
import slide_text_check

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'groups.pptx'; deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape(); normal = group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(.7)); normal.text_frame.text = 'Group heading'
    inner = group.shapes.add_group_shape(); overflow = inner.shapes.add_textbox(Inches(1), Inches(2), Inches(1), Inches(.3))
    overflow.text_frame.auto_size = MSO_AUTO_SIZE.NONE; overflow.text_frame.text = 'A long sentence that cannot fit in a small text box and requires a review'
    overflow.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    group.shapes._recalculate_extents(); group.left = Inches(2); group.top = Inches(1); group.width = Inches(5); deck.save(source)
    result = slide_text_check.scan(source, open_office_use.render)
    assert result['complete'] and result['text_shapes_checked'] == 2, result
    path = f'/slide[1]/group[@id={group.shape_id}]/group[@id={inner.shape_id}]/shape[@id={overflow.shape_id}]'
    assert [item['path_hint'] for item in result['candidates']] == [path], result
    draft = root / 'draft'; session.start(draft, source=source)
    edit = session.replace_text(draft, 0, path, overflow.text_frame.text, 'Fits')
    assert slide_text_check.scan(edit['file'], open_office_use.render)['candidate_count'] == 0
    table_source = root / 'group-table.pptx'; table_deck = Presentation(); table_slide = table_deck.slides.add_slide(table_deck.slide_layouts[6])
    table_group = table_slide.shapes.add_group_shape()
    table_shape = table_slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(1.5), Inches(.4))
    cell = table_shape.table.cell(0, 0); cell.text = 'A long grouped table cell needs more room for readable text'
    cell.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    table_slide.shapes._spTree.remove(table_shape._element); table_group._element.append(table_shape._element); table_group.shapes._recalculate_extents()
    table_deck.save(table_source)
    table_result = slide_text_check.scan(table_source, open_office_use.render)
    table_path = f'/slide[1]/group[@id={table_group.shape_id}]/table[@id={table_shape.shape_id}]/tr[1]/tc[1]'
    assert table_result['table_cells_checked'] == 1 and table_result['candidates'][0]['path_hint'] == table_path, table_result
    table_draft = root / 'table-draft'; session.start(table_draft, source=table_source)
    table_edit = session.replace_text(table_draft, 0, table_path, cell.text, 'Fits')
    assert slide_text_check.scan(table_edit['file'], open_office_use.render)['candidate_count'] == 0
    rotated = root / 'rotated.pptx'; group.rotation = 15; deck.save(rotated)
    check = slide_text_check.scan(rotated, open_office_use.render)
    assert not check['complete'] and any(item['reason'] == 'rotated_or_flipped_group' for item in check['skipped']), check
print('Rendered slide group text check passed')

"""Exercise single-snapshot review bundles for all three Office formats."""
import hashlib
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from openpyxl import Workbook
from pptx import Presentation

import open_office_use
import review_bundle
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); doc_path = root / 'report.docx'; doc = Document(); doc.add_paragraph('Original snapshot'); doc.save(doc_path)
    book_path = root / 'book.xlsx'; book = Workbook(); book.active['A1'] = '=1+1'; book.save(book_path)
    deck_path = root / 'deck.pptx'; deck = Presentation(); deck.slides.add_slide(deck.slide_layouts[1]); deck.save(deck_path)
    for source, check in ((doc_path, 'word_comments'), (book_path, 'formula_scan'), (deck_path, 'slide_layout')):
        out = root / (source.stem + '-review'); result = review_bundle.create(source, out, open_office_use.render)
        report = json.loads((out / 'report.json').read_text())
        assert result['snapshot_sha256'] == hashlib.sha256(source.read_bytes()).hexdigest() == report['snapshot_sha256']
        assert report['checks']['officecli_audit']['status'] == 'completed' and report['checks'][check]['status'] == 'completed', report
        assert report['checks']['overview']['status'] == 'completed' and (out / 'overview.png').read_bytes().startswith(b'\x89PNG')
        assert (out / 'preview.pdf').read_bytes().startswith(b'%PDF') and not report['check_errors']
        if source.suffix == '.xlsx': assert 'formula_cache_missing' in report['flags'], report
        try: review_bundle.create(source, out, open_office_use.render)
        except FileExistsError: pass
        else: raise AssertionError('Existing bundle overwritten')
    altered = root / 'changing.docx'; changing = Document(); changing.add_paragraph('Frozen original'); changing.save(altered)
    old_bytes = altered.read_bytes()
    def change_after_snapshot(snapshot, out):
        edited = Document(); edited.add_paragraph('Later external edit'); edited.save(altered)
        return open_office_use.render(snapshot, out)
    result = review_bundle.create(altered, root / 'frozen-review', change_after_snapshot)
    assert result['snapshot_sha256'] == hashlib.sha256(old_bytes).hexdigest()
    text = subprocess.run(['pdftotext', result['artifacts']['preview.pdf'], '-'], capture_output=True, text=True, check=True).stdout
    assert 'Frozen original' in text and 'Later external edit' not in text
    def failed_render(_snapshot, _out): raise RuntimeError('render unavailable')
    failed = review_bundle.create(doc_path, root / 'partial-review', failed_render)
    failed_report = json.loads(Path(failed['report']).read_text())
    assert failed['check_errors'] == ['render_pdf'] and failed_report['checks']['overview']['status'] == 'skipped'
    assert 'preview.pdf' not in failed['artifacts'] and 'overview.png' not in failed['artifacts']
    invalid = Document(); paragraph = invalid.add_paragraph('Invalid'); paragraph._p.append(OxmlElement('w:ins')); invalid_path = root / 'invalid.docx'; invalid.save(invalid_path)
    invalid_result = review_bundle.create(invalid_path, root / 'invalid-review', open_office_use.render)
    assert 'schema_validation_failed' in invalid_result['flags'], invalid_result
    broken = root / 'broken.xlsx'
    with zipfile.ZipFile(broken, 'w') as archive: archive.writestr('unrelated', 'content')
    broken_result = review_bundle.create(broken, root / 'broken-review', failed_render)
    broken_report = json.loads(Path(broken_result['report']).read_text())
    assert broken_report['checks']['formula_scan']['status'] == 'error' and 'formula_scan' in broken_result['check_errors']
    draft = root / 'draft'; session.start(draft, source=doc_path); session.add_comment(draft, 0, '/body/p[1]', 'Review', 'Editor')
    cli = Path(__file__).with_name('open_office_use.py'); out = root / 'revision-review'
    completed = subprocess.run([sys.executable, str(cli), 'session', 'review', str(draft), '--revision', '0', '--out-dir', str(out)], check=True, capture_output=True, text=True)
    assert json.loads(completed.stdout)['revision'] == 0 and json.loads((out / 'report.json').read_text())['checks']['word_comments']['result']['total_comments'] == 0
    rejected = subprocess.run([sys.executable, str(cli), 'session', 'review', str(draft), '--revision', '0', '--out-dir', str(draft / 'bad')], capture_output=True, text=True)
    assert rejected.returncode != 0 and not (draft / 'bad').exists()
print('Review bundle smoke passed: snapshots, three formats, partial failure and revision isolation')

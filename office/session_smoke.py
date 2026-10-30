"""Exercise local revision sessions and competing writers."""
import concurrent.futures
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

cli = Path(__file__).with_name('open_office_use.py')
def run(*args, ok=True):
    result = subprocess.run([sys.executable, str(cli), 'session', *map(str, args)], capture_output=True, text=True)
    assert (result.returncode == 0) == ok, (result.stdout, result.stderr)
    return json.loads(result.stdout) if ok else result.stderr

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); spec = root / 'ops.json'
    for suffix in ('.docx', '.xlsx', '.pptx'):
        source = root / ('source' + suffix); directory = root / suffix[1:]
        if suffix == '.docx':
            obj = Document(); obj.add_paragraph('Original'); obj.save(source)
            ops = [{'command': 'set', 'path': '/body/p[1]', 'props': {'text': 'Edited'}}]
        elif suffix == '.xlsx':
            obj = Workbook(); obj.active['A1'] = 'Original'; obj.save(source)
            ops = [{'command': 'set', 'path': '/Sheet/A1', 'props': {'value': 'Edited'}}]
        else:
            obj = Presentation(); slide = obj.slides.add_slide(obj.slide_layouts[1]); slide.shapes.title.text = 'Original'; obj.save(source)
            ops = [{'command': 'set', 'path': '/slide[1]/shape[1]', 'props': {'text': 'Edited'}}]
        source_bytes = source.read_bytes()
        state = run('start', directory, '--source', source)
        assert state['latest_revision'] == 0
        spec.write_text(json.dumps(ops))
        edited = run('apply', directory, '--base', '0', '--spec', spec)
        assert edited['revision'] == 1 and source.read_bytes() == source_bytes
        assert 'conflict' in run('apply', directory, '--base', '0', '--spec', spec, ok=False).lower()
        bad = root / 'bad.json'; bad.write_text(json.dumps([{'command': 'remove', 'path': '/missing[999]'}]))
        run('apply', directory, '--base', '1', '--spec', bad, ok=False)
        assert run('info', directory)['latest_revision'] == 1
        exported = root / ('export' + suffix)
        run('export', directory, '--revision', '1', '--out', exported)
        assert exported.exists() and source.read_bytes() == source_bytes
        assert 'already exists' in run('export', directory, '--revision', '1', '--out', exported, ok=False)
        assert 'outside' in run('export', directory, '--revision', '1', '--out', directory / ('x' + suffix), ok=False)
        if suffix == '.xlsx':
            assert load_workbook(exported).active['A1'].value == 'Edited'
        # Both writers target the same next revision; only one can publish it.
        def writer():
            return subprocess.run([sys.executable, str(cli), 'session', 'apply', str(directory), '--base', '1', '--spec', str(spec)], capture_output=True, text=True)
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: writer(), range(2)))
        assert sorted(r.returncode == 0 for r in results) == [False, True], [(r.stdout, r.stderr) for r in results]
        assert run('info', directory)['latest_revision'] == 2
    blank = run('start', root / 'blank', '--format', 'docx')
    assert blank['latest_revision'] == 0
print('Session smoke passed: three formats, failed edits, exports and concurrent revision conflict')

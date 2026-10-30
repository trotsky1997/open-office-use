"""Check exact Word comment anchors, run formatting and boundary rejection."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

import officecli_backend
import session


def styles(paragraph):
    return [(char, run.bold, run.italic) for run in paragraph.runs for char in run.text]


with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.docx'; draft = root / 'draft'; doc = Document()
    p = doc.add_paragraph(); p.add_run('Start ').bold = True; p.add_run('mixed').italic = True; p.add_run(' tail 🚀 end')
    doc.add_paragraph('same same')
    linked = doc.add_paragraph('before '); link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), doc.part.relate_to('https://example.com', RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'linked'; run.append(text); link.append(run); linked._p.append(link); linked.add_run(' after')
    tracked_run = doc.add_paragraph('Tracked format').runs[0]; change = OxmlElement('w:rPrChange'); change.set(qn('w:id'), '10'); change.set(qn('w:author'), 'Editor'); change.append(OxmlElement('w:rPr')); tracked_run._r.get_or_add_rPr().append(change)
    doc.save(source); original = source.read_bytes(); expected = styles(p); session.start(draft, source=source)
    first = session.add_comment(draft, 0, '/body/p[1]', 'Review phrase', 'Reviewer', 'art mix')
    result = Document(first['file']); assert result.paragraphs[0].text == p.text and styles(result.paragraphs[0]) == expected
    listing = session.list_comments(draft, 1); assert listing['comments'][0]['anchor_text'] == 'art mix', listing
    validation = officecli_backend.audit(first['file'])['validation']; assert validation['exit_code'] == 0, validation
    second = session.add_comment(draft, 1, '/body/p[1]', 'Nested', 'Reviewer', 'rt mi')
    second_listing = session.list_comments(draft, 2); assert [item['anchor_text'] for item in second_listing['comments']] == ['art mix', 'rt mi'], second_listing
    validation = officecli_backend.audit(second['file'])['validation']; assert validation['exit_code'] == 0, validation
    third = session.add_comment(draft, 2, '/body/p[1]', 'Emoji', 'Reviewer', ' 🚀 ')
    assert session.list_comments(draft, 3)['comments'][-1]['anchor_text'] == ' 🚀 '
    assert styles(Document(third['file']).paragraphs[0]) == expected and source.read_bytes() == original
    for path, anchor in [('/body/p[2]', 'same'), ('/body/p[3]', 'linked'), ('/body/p[3]', 'before linked'), ('/body/p[1]', 't mixed'), ('/body/p[1]', ''), ('/body/p[1]', 'missing'), ('/body/p[4]', 'rack')]:
        try: session.add_comment(draft, 3, path, 'Rejected', 'Reviewer', anchor)
        except ValueError: pass
        else: raise AssertionError(f'Invalid anchor accepted: {path} {anchor!r}')
    assert session.info(draft)['latest_revision'] == 3
    linked_result = session.add_comment(draft, 3, '/body/p[3]', 'After link', 'Reviewer', 'after')
    linked_read = Document(linked_result['file']).paragraphs[2]
    assert linked_read.hyperlinks[0].address == 'https://example.com' and linked_read.text == 'before linked after'
    assert session.list_comments(draft, 4)['comments'][-1]['anchor_text'] == 'after'
    reply_draft = root / 'reply'; session.start(reply_draft, source=third['file'])
    session.reply_comment(reply_draft, 0, '1', 'Nested reply', 'Editor')
    assert session.list_comments(reply_draft, 1)['comments'][-1]['anchor_text'] == 'rt mi'
    multi = Document(); paragraph = multi.add_paragraph(); run = paragraph.add_run('Alpha'); extra = OxmlElement('w:t'); extra.text = 'Beta'; run._r.append(extra)
    multi_file = root / 'multi.docx'; multi.save(multi_file); multi_draft = root / 'multi'; session.start(multi_draft, source=multi_file)
    output = session.add_comment(multi_draft, 0, '/body/p[1]', 'Multiple text nodes', 'Editor', 'haBe')
    assert Document(output['file']).paragraphs[0].text == 'AlphaBeta' and session.list_comments(multi_draft, 1)['comments'][0]['anchor_text'] == 'haBe'
    cli_draft = root / 'cli'; session.start(cli_draft, source=source)
    cli = Path(__file__).with_name('open_office_use.py')
    subprocess.run([sys.executable, str(cli), 'session', 'add-comment', str(cli_draft), '--base', '0', '--path', '/body/p[1]', '--text', 'CLI review', '--author', 'Editor', '--anchor', 'mixed'], check=True, capture_output=True)
    assert session.list_comments(cli_draft, 1)['comments'][0]['anchor_text'] == 'mixed'
print('Comment range smoke passed: run splits, nested anchors, Unicode and boundary rejection')

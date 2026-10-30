"""Create and render an ordered, styled Word document from blocks."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import open_office_use

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); output = root / 'brief.docx'
    spec = {'blocks': [
        {'type':'heading','level':1,'text':'Quarterly review'},
        {'type':'paragraph','runs':[{'text':'Revenue grew '},{'text':'18%','bold':True},{'text':' year over year.'}]},
        {'type':'bullet','text':'Retention was stable'},
        {'type':'numbered','text':'Prioritize onboarding'},
        {'type':'table','header':True,'style':'Light Shading Accent 1','rows':[['Metric','Value'],['Retention','92%']]},
        {'type':'page_break'}, {'type':'heading','level':2,'text':'Next steps'}]}
    path = root/'spec.json'; path.write_text(json.dumps(spec), encoding='utf-8')
    cli = Path(__file__).with_name('open_office_use.py')
    result = subprocess.run([sys.executable,str(cli),'author','--spec',str(path),'--out',str(output)], capture_output=True, text=True, check=True)
    assert json.loads(result.stdout)['output'] == str(output.resolve())
    doc = Document(output); assert [p.text for p in doc.paragraphs] == ['Quarterly review','Revenue grew 18% year over year.','Retention was stable','Prioritize onboarding','','Next steps']
    assert doc.paragraphs[0].style.name == 'Heading 1' and doc.paragraphs[2].style.name == 'List Bullet' and doc.paragraphs[3].style.name == 'List Number'
    assert doc.paragraphs[1].runs[1].bold and doc.tables[0].cell(1,1).text == '92%'
    assert doc.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None and doc.tables[0].cell(0,0).paragraphs[0].runs[0].bold
    assert doc.tables[0].style.name == 'Light Shading Accent 1'
    body = list(doc._element.body)
    assert [item.tag.rsplit('}', 1)[-1] for item in body] == ['p','p','p','p','tbl','p','p','sectPr']
    assert body[5].find('.//' + qn('w:br')).get(qn('w:type')) == 'page'
    pdf = open_office_use.render(output, root/'render'); assert Path(pdf).read_bytes().startswith(b'%PDF')
    for i, bad in enumerate([{'blocks':[{'type':'table','rows':[['a'],['a','b']]}]}, {'blocks':[{'type':'heading','level':10,'text':'bad'}]}, {'blocks':[{'type':'paragraph','text':'a','runs':[]}]}, {'blocks':[{'type':'paragraph','runs':[{'text':'a','bold':'yes'}]}]}]):
        target = root/f'bad{i}.docx'
        try: open_office_use.author(bad, target)
        except ValueError: pass
        else: raise AssertionError(f'Invalid block accepted: {bad}')
        assert not target.exists()
print('Word author blocks passed: structure, formatting, rendering and invalid input')

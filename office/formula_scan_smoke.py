"""Check workbook-wide formula cache scanning with real XLSX files."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
import spreadsheet

cli = Path(__file__).with_name('open_office_use.py')
def run(*args):
    result = subprocess.run([sys.executable, str(cli), *map(str, args)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    return json.loads(result.stdout)

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.xlsx'; calculated = root / 'calculated.xlsx'
    book = Workbook(); first = book.active; first.title = 'Forecast'; first['A1'] = '=1/0'; first['A2'] = '=1+2'
    second = book.create_sheet('Hidden'); second.sheet_state = 'hidden'; second['B1'] = '=4+5'; second['B2'] = '#N/A'
    book.save(source); original = source.read_bytes()
    result = run('scan-formulas', source)
    assert result['complete'] and result['total_sheets'] == 2 and result['formulas'] == 3, result
    assert result['empty_or_missing_cache'] == 3 and result['literal_errors'] == 1, result
    assert result['errors'][0]['sheet'] == 'Hidden' and result['errors'][0]['cell'] == 'B2'
    assert result['cache_freshness'] == 'unknown' and source.read_bytes() == original
    run('recalculate', source, '--out', calculated)
    result = run('scan-formulas', calculated)
    assert result['formula_errors'] >= 1 and result['empty_or_missing_cache'] == 0, result
    assert any(item['cell'] == 'A1' and item['error'] == '#DIV/0!' for item in result['errors']), result
    many = root / 'many.xlsx'; large = Workbook();
    for index in range(1, 204):
        large.active.cell(index, 1, '=1+1'); large.active.cell(index, 2, '#N/A')
    large.save(many); capped = run('scan-formulas', many)
    assert capped['complete'] and capped['empty_or_missing_cache'] == 203 and len(capped['empty_or_missing']) == 200 and capped['empty_or_missing_truncated'], capped
    assert capped['literal_errors'] == 203 and len(capped['errors']) == 200 and capped['errors_truncated'], capped
    spreadsheet.FORMULA_SCAN_MAX_CELLS = 2
    limited = spreadsheet.scan_formulas(source)
    assert not limited['complete'] and limited['scanned_cells'] == 2 and limited['total_sheets'] == 2, limited
    assert len(limited['sheets']) == 2 and not limited['sheets'][-1]['complete'], limited
print('Formula scan smoke passed: hidden sheets, errors, caches and scan limit')

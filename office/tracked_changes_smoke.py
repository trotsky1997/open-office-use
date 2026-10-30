"""Exercise tracked-change inventory, snapshot scope and review reporting."""
import hashlib
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import open_office_use
import review_bundle
import session
import word

def revision(kind, identifier, text=None):
    item = OxmlElement('w:' + kind); item.set(qn('w:id'), str(identifier)); item.set(qn('w:author'), 'Reviewer'); item.set(qn('w:date'), '2025-01-01T00:00:00Z')
    if text is not None:
        run = OxmlElement('w:r'); node = OxmlElement('w:delText' if kind == 'del' else 'w:t'); node.text = text; run.append(node); item.append(run)
    return item

with tempfile.TemporaryDirectory() as temp:
    root = Path(temp); source = root / 'changes.docx'; draft = root / 'draft'; doc = Document()
    paragraph = doc.add_paragraph('Before '); paragraph._p.append(revision('ins', 1, 'inserted')); paragraph._p.append(revision('del', 2, 'deleted'))
    ppr = paragraph._p.get_or_add_pPr(); change = revision('pPrChange', 3); change.append(OxmlElement('w:pPr')); ppr.append(change)
    rpr = OxmlElement('w:rPr'); rpr.append(revision('del', 4)); ppr.append(rpr)
    header = doc.sections[0].header.paragraphs[0]; header._p.append(revision('ins', 5, 'header addition'))
    setting = OxmlElement('w:trackRevisions'); doc.settings.element.append(setting); doc.save(source)
    original = source.read_bytes(); result = word.tracked_changes(source)
    assert result['snapshot_sha256'] == hashlib.sha256(original).hexdigest() and result['tracking_enabled'] is True
    assert result['scan_complete'] and result['counts']['ins'] == 2 and result['counts']['del'] == 2 and result['counts']['pPrChange'] == 1
    assert any(r['story'] == 'header' and r['text'] == 'header addition' for r in result['records'])
    assert any(r['kind'] == 'del' and r['parent'] == 'rPr' and r['text'] == '' for r in result['records'])
    assert any(r['kind'] == 'del' and r['text'] == 'deleted' for r in result['records'])
    footnote_file = root / 'footnote.docx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(footnote_file, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry)
            if entry.filename == 'word/_rels/document.xml.rels':
                rels = etree.fromstring(data); rel = etree.SubElement(rels, word.REL + 'Relationship')
                rel.set('Id', 'rIdFootnoteTest'); rel.set('Type', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes'); rel.set('Target', 'footnotes.xml'); data = etree.tostring(rels)
            elif entry.filename == '[Content_Types].xml':
                types = etree.fromstring(data); item = etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override')
                item.set('PartName', '/word/footnotes.xml'); item.set('ContentType', 'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml'); data = etree.tostring(types)
            after.writestr(entry, data)
        footnotes = etree.Element(word.W + 'footnotes'); note = etree.SubElement(footnotes, word.W + 'footnote'); note.set(word.W + 'id', '1'); p = etree.SubElement(note, word.W + 'p'); p.append(revision('ins', 10, 'footnote revision'))
        after.writestr('word/footnotes.xml', etree.tostring(footnotes))
    footnote_result = word.tracked_changes(footnote_file)
    assert footnote_result['scan_complete'] and any(item['story'] == 'footnotes' and item['text'] == 'footnote revision' for item in footnote_result['records'])
    session.start(draft, source=source); assert session.tracked_changes(draft, 0)['revision'] == 0
    cli = Path(__file__).with_name('open_office_use.py')
    out = subprocess.run([sys.executable, str(cli), 'session', 'tracked-changes', str(draft), '--revision', '0'], capture_output=True, text=True, check=True)
    assert json.loads(out.stdout)['records_scanned'] == 5 and source.read_bytes() == original
    bundle = review_bundle.create(source, root / 'bundle', open_office_use.render)
    assert 'word_tracked_changes_present' in bundle['flags']
    many = Document(); p = many.add_paragraph()
    for index in range(201): p._p.append(revision('ins', index, 'x'))
    limited = root / 'many.docx'; many.save(limited); report = word.tracked_changes(limited)
    assert report['records_scanned'] == 201 and len(report['records']) == 200 and report['records_truncated'] and report['scan_complete']
print('Tracked changes smoke passed: text, paragraph marks, formatting, header, limits and review bundle')

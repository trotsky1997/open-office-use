"""Exercise recalculation requests after a preserved-cache text edit."""
import tempfile
import json
import subprocess
import sys
import zipfile
from pathlib import Path
from lxml import etree
from openpyxl import Workbook, load_workbook
import officecli_backend
import open_office_use
import review_bundle
import session
import spreadsheet

X = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'manual.xlsx'; doc = Workbook(); doc.active['A1'] = 'short'; doc.active['B1'] = '=LEN(A1)'
    doc.calculation.calcMode = 'manual'; doc.calculation.fullCalcOnLoad = False; doc.calculation.forceFullCalc = False; doc.save(source)
    patched = root / 'cached.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(patched, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                tree = etree.fromstring(data); tree.xpath('./x:sheetData/x:row/x:c[@r="B1"]/x:v', namespaces={'x': X[1:-1]})[0].text = '5'; data = etree.tostring(tree)
            after.writestr(entry, data)
    patched.replace(source); original = source.read_bytes(); draft = root / 'draft'; session.start(draft, source=source)
    result = session.replace_text(draft, 0, '/Sheet/A1', 'short', 'longer input')
    assert result['formula_cache_effect'] == 'may_be_stale' and result['recalculation_requested']
    assert load_workbook(result['file'], data_only=True).active['B1'].value == 5 and source.read_bytes() == original
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(result['file']) as after:
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name != 'xl/worksheets/sheet1.xml')
        sheet = etree.fromstring(after.read('xl/worksheets/sheet1.xml')); marker = sheet.find(X + 'sheetCalcPr')
        assert marker.get('fullCalcOnLoad') == '1' and sheet.index(marker) == sheet.index(sheet.find(X + 'sheetData')) + 1
        calc = etree.fromstring(after.read('xl/workbook.xml')).find(X + 'calcPr'); assert calc.get('calcMode') == 'manual' and calc.get('fullCalcOnLoad') == '0'
    scan = spreadsheet.scan_formulas(result['file']); assert scan['recalculation_requests'] == 1 and scan['sheets'][0]['recalculation_requested'] and scan['workbook_calculation_mode'] == 'manual' and scan['calculation_metadata_complete'], scan
    saved_limit = spreadsheet.FORMULA_SCAN_MAX_CELLS; spreadsheet.FORMULA_SCAN_MAX_CELLS = 1
    try:
        partial = spreadsheet.scan_formulas(result['file'])
        assert not partial['calculation_metadata_complete'] and partial['sheets'][0]['recalculation_requested'] is None, partial
    finally: spreadsheet.FORMULA_SCAN_MAX_CELLS = saved_limit
    cli_draft = root / 'cli'; session.start(cli_draft, source=source)
    cli = Path(__file__).with_name('open_office_use.py')
    cli_result = subprocess.run([sys.executable, str(cli), 'session', 'replace-text', str(cli_draft), '--base', '0', '--path', '/Sheet/A1', '--find', 'short', '--replace', 'longer input'], capture_output=True, text=True, check=True)
    assert json.loads(cli_result.stdout)['recalculation_requested']
    audit = officecli_backend.audit(result['file']); assert audit['validation']['exit_code'] == 0 and audit['validation']['result']['success'], audit
    bundle = review_bundle.create(result['file'], root / 'review', open_office_use.render)
    assert 'sheet_recalculation_requested' in bundle['flags'], bundle
    recalculated = session.recalculate(draft, 1); assert load_workbook(recalculated['file'], data_only=True).active['B1'].value == 12
    assert load_workbook(recalculated['file'], read_only=True).calculation.calcMode == 'manual'
    final_audit = officecli_backend.audit(recalculated['file']); assert final_audit['validation']['exit_code'] == 0 and final_audit['validation']['result']['success'], final_audit
    changed_again = session.replace_text(draft, 2, '/Sheet/A1', 'longer input', 'hi')
    assert load_workbook(changed_again['file'], data_only=True).active['B1'].value == 12
    recalculated_again = session.recalculate(draft, 3)
    assert load_workbook(recalculated_again['file'], data_only=True).active['B1'].value == 2
    assert load_workbook(recalculated_again['file'], read_only=True).calculation.calcMode == 'manual'
    signed = root / 'signed.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(signed, 'w') as after:
        for entry in before.infolist(): after.writestr(entry, before.read(entry))
        after.writestr('_xmlsignatures/signature.xml', '<signature/>')
    try: spreadsheet.recalculate(signed, root / 'signed-output.xlsx')
    except ValueError: pass
    else: raise AssertionError('Signed workbook recalculation accepted')
    assert not (root / 'signed-output.xlsx').exists()
    invalid = root / 'duplicate.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(invalid, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                tree = etree.fromstring(data); sheet_data = tree.find(X + 'sheetData'); sheet_data.addnext(etree.Element(X + 'sheetCalcPr', fullCalcOnLoad='0')); sheet_data.addnext(etree.Element(X + 'sheetCalcPr', fullCalcOnLoad='1')); data = etree.tostring(tree)
            after.writestr(entry, data)
    bad_scan = spreadsheet.scan_formulas(invalid); assert not bad_scan['calculation_metadata_complete'] and bad_scan['sheets'][0]['calculation_metadata_issue'] == 'duplicate_sheet_calculation_properties'
    bad_draft = root / 'bad'; session.start(bad_draft, source=invalid)
    try: session.replace_text(bad_draft, 0, '/Sheet/A1', 'short', 'changed')
    except ValueError: pass
    else: raise AssertionError('Ambiguous calculation properties accepted')
    assert session.info(bad_draft)['latest_revision'] == 0
print('Worksheet recalculation request smoke passed')

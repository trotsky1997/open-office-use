"""Create editable PowerPoint charts with matching embedded data."""
import io
import tempfile
from pathlib import Path
from openpyxl import load_workbook
from pptx import Presentation
import open_office_use
import officecli_backend

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); out = root/'charts.pptx'
    charts = [{'type':'chart','chart_type':'column','x':1,'y':1,'width':5,'height':3,'title':'Revenue','x_axis_title':'Quarter','y_axis_title':'USD','categories':['Q1','Q2','Q3'],'series':[{'name':'Revenue','values':[100,120,135],'color':'236B8E'},{'name':'Target','values':[90,110,130],'color':'62A894'}]},
              {'type':'chart','chart_type':'line','x':1,'y':4.5,'width':5,'height':2,'categories':['Q1','Q2','Q3'],'series':[{'name':'Growth','values':[0.1,None,0.2],'number_format':'0.0%','color':'BA5A32'}]}]
    open_office_use.author({'slides':[{'layout':'blank','elements':charts}]},out)
    code, validation, _ = officecli_backend.invoke('validate',out); assert code==0 and validation['success'],validation
    deck = Presentation(out); first = deck.slides[0].shapes[0].chart; second = deck.slides[0].shapes[1].chart
    assert first.has_title and first.chart_title.text_frame.text == 'Revenue' and first.has_legend
    assert first.category_axis.axis_title.text_frame.text == 'Quarter' and first.value_axis.axis_title.text_frame.text == 'USD'
    assert tuple(first.series[0].values) == (100,120,135) and tuple(second.series[0].values) == (0.1,None,0.2)
    book = load_workbook(io.BytesIO(first.part.chart_workbook.xlsx_part.blob)); ws = book.active
    assert [ws.cell(i,1).value for i in range(2,5)] == ['Q1','Q2','Q3'] and [ws.cell(i,2).value for i in range(2,5)] == [100,120,135]
    assert first.series[0].format.fill.fore_color.rgb == (0x23,0x6b,0x8e)
    assert Path(open_office_use.render(out,root/'render')).read_bytes().startswith(b'%PDF')
    bad_cases = [dict(charts[0],categories=['Q1']),dict(charts[0],chart_type='pie'),dict(charts[0],series=[{'name':'Revenue','values':[1,float('nan'),3]}]),dict(charts[0],categories=['=1+1','Q2','Q3']),dict(charts[0],categories=['https://example.com','Q2','Q3'])]
    for index, chart in enumerate(bad_cases):
        bad = root/f'bad{index}.pptx'
        try: open_office_use.author({'slides':[{'layout':'blank','elements':[chart]}]},bad)
        except ValueError: pass
        else: raise AssertionError('Invalid chart data accepted')
        assert not bad.exists()
print('Slide chart author passed: editable data, cache, render and invalid input')

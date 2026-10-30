"""Create a formatted workbook with formulas and worksheet settings."""
import tempfile
import json
import sys
import zipfile
from lxml import etree
import subprocess
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, Reference
from PIL import Image
import xlsxwriter
import open_office_use
import officecli_backend
import chart_cache
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); output = root/'forecast.xlsx'
    spec = {'sheets':{'Forecast':{'A1':'Metric','B1':'Revenue','C1':'Growth','A2':'Q1','B2':100,'C2':0.18,'A3':'Q2','B3':'=B2*1.2','C3':'=B3/B2-1'}},
            'sheet_options':{'Forecast':{'freeze_panes':'B2','auto_filter':'A1:C3','fit_to_page':True,'column_widths':{'A':20,'B':16,'C':14},'row_heights':{'1':24},'formats':[
                {'range':'A1:C1','bold':True,'fill':'236B8E','font_color':'FFFFFF','horizontal':'center','wrap':True},
                {'range':'B2:B3','number_format':'#,##0.00'}, {'range':'C2:C3','number_format':'0.0%'}]}}}
    spec['sheet_options']['Forecast']['charts'] = [{'type':'column','data':'B1:B3','categories':'A2:A3','anchor':'E2','title':'Revenue','y_axis_title':'USD','width':12,'height':7}, {'type':'line','data':'C1:C3','categories':'A2:A3','anchor':'E18','title':'Growth'}]
    open_office_use.author(spec,output)
    sheet = load_workbook(output)['Forecast']
    assert sheet.freeze_panes == 'B2' and sheet.auto_filter.ref == 'A1:C3'
    assert sheet.sheet_properties.pageSetUpPr.fitToPage and sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
    assert sheet.column_dimensions['A'].width == 20 and sheet.row_dimensions[1].height == 24
    assert sheet['A1'].font.bold and sheet['A1'].font.color.rgb == 'FFFFFFFF' and sheet['A1'].fill.fgColor.rgb == 'FF236B8E'
    assert sheet['A1'].alignment.horizontal == 'center' and sheet['A1'].alignment.wrap_text
    assert sheet['B3'].value == '=B2*1.2' and sheet['B3'].number_format == '#,##0.00' and sheet['C3'].number_format == '0.0%'
    assert load_workbook(output,data_only=True)['Forecast']['B3'].value is None
    assert len(sheet._charts) == 2 and sheet._charts[0].__class__.__name__ == 'BarChart' and sheet._charts[1].__class__.__name__ == 'LineChart'
    assert sheet._charts[0].series[0].val.numRef.f == "'Forecast'!$B$2:$B$3" and sheet._charts[0].series[0].cat.strRef.f == "'Forecast'!$A$2:$A$3"
    report = chart_cache.scan(output)
    refs = report['charts'][0]['references']; assert [ref['status'] for ref in refs] == ['match','match','missing_cache']
    literal_spec = {'sheets':{'Sales':{'A1':'Quarter','B1':'Revenue','A2':'Q1','B2':100,'A3':'Q2','B3':120}},'sheet_options':{'Sales':{'formats':[{'range':'B2:B3','number_format':'0.00'}],'charts':[{'type':'column','data':'B1:B3','categories':'A2:A3'}]}}}
    literal = root/'literal.xlsx'; open_office_use.author(literal_spec, literal)
    assert load_workbook(literal)['Sales']._charts[0].series[0].val.numRef.numCache.formatCode == '0.00'
    checked = chart_cache.scan(literal); assert checked['complete'] and checked['missing_caches'] == 0 and checked['mismatch_references'] == 0
    draft = root/'draft'; session.start(draft, source=literal)
    edited = session.edit_grid(draft, 0, 'Sales', 'B2', [{'cell':'B2','value':'200'}])
    stale = chart_cache.scan(edited['file']); assert stale['mismatch_references'] == 1
    formula_spec = {'sheets':{'Data':{'A1':'=1','A2':'=2','B1':10,'B2':20}},'sheet_options':{'Data':{'charts':[{'type':'line','data':'B1:B2','categories':'A1:A2','category_type':'number','titles_from_data':False}]}}}
    formula_chart = root/'formula-categories.xlsx'; open_office_use.author(formula_spec, formula_chart)
    assert load_workbook(formula_chart)['Data']._charts[0].series[0].cat.numRef is not None
    del formula_spec['sheet_options']['Data']['charts'][0]['category_type']
    try: open_office_use.author(formula_spec, root/'unknown-categories.xlsx')
    except ValueError as exc: assert 'category_type' in str(exc)
    else: raise AssertionError('Formula category type was guessed')
    assert Path(open_office_use.render(output, root/'render')).read_bytes().startswith(b'%PDF')
    printing=root/'printing.xlsx'; cells={'A1':'Metric','B1':'Value','A100':'OUTSIDE PRINT AREA'}
    cells.update({f'A{i}':f'Row {i}' for i in range(2,100)}); cells.update({f'B{i}':i for i in range(2,100)})
    print_settings={'paper':'a4','orientation':'landscape','area':'A1:B99','repeat_rows':'1:1','repeat_columns':'A:A','header':{'text':'R&D REPORT'},'footer':{'text':'Page','page_number':True}}
    print_spec=root/'print-spec.json'; print_spec.write_text(json.dumps({'sheets':{'Data':cells},'sheet_options':{'Data':{'print_options':print_settings}}}))
    cli=Path(__file__).with_name('open_office_use.py')
    created=subprocess.run([sys.executable,str(cli),'author','--spec',str(print_spec),'--out',str(printing)],capture_output=True,text=True,check=True)
    assert json.loads(created.stdout)['output']==str(printing.resolve())
    printed=load_workbook(printing)['Data']
    assert str(printed.page_setup.paperSize)=='9' and printed.page_setup.orientation=='landscape' and printed.print_title_rows=='$1:$1' and printed.print_title_cols=='$A:$A'
    assert str(printed.print_area)=="'Data'!$A$1:$B$99" and printed.oddHeader.left.text=='R&&D REPORT' and printed.oddFooter.right.text=='Page &P'
    code, validation, _=officecli_backend.invoke('validate',printing); assert code==0 and validation['success'],validation
    pdf=open_office_use.render(printing,root/'print-render'); text=subprocess.run(['pdftotext','-layout',pdf,'-'],capture_output=True,text=True,check=True).stdout
    pages=[page for page in text.split('\f') if page.strip()]; assert len(pages)>=2
    assert all('Metric' in page and 'R&D REPORT' in page and f'Page {index}' in page for index,page in enumerate(pages,1)),pages
    assert 'OUTSIDE PRINT AREA' not in text
    image_path=root/'figure.png'; Image.new('RGB',(80,40),'#236B8E').save(image_path)
    image_source=image_path.read_bytes(); image_book=root/'images.xlsx'
    image_spec={'sheets':{'Data':{'A1':'Chart'}},'sheet_options':{'Data':{'images':[{'path':str(image_path),'anchor':'C3','width':200,'description':'Revenue graphic'}]}}}
    image_spec_path=root/'image-spec.json';image_spec_path.write_text(json.dumps(image_spec));
    result=subprocess.run([sys.executable,str(cli),'author','--spec',str(image_spec_path),'--out',str(image_book)],capture_output=True,text=True,check=True)
    assert json.loads(result.stdout)['output']==str(image_book.resolve())
    image_path.unlink(); loaded=load_workbook(image_book)['Data']; assert loaded['A1'].value=='Chart' and len(loaded._images)==1
    picture=loaded._images[0]; assert picture._data()==image_source and picture.anchor._from.col==2 and picture.anchor._from.row==2
    assert picture.anchor.ext.width==200*9525 and picture.anchor.ext.height==100*9525
    assert picture.anchor.pic.nvPicPr.cNvPr.descr=='Revenue graphic'
    code, validation, _=officecli_backend.invoke('validate',image_book); assert code==0 and validation['success'],validation
    image_pdf=open_office_use.render(image_book,root/'image-render')
    listed=subprocess.run(['pdfimages','-list',image_pdf],capture_output=True,text=True,check=True).stdout
    assert any('image' in line and '80' in line and '40' in line for line in listed.splitlines()[2:]),listed
    second_image=root/'second.jpg'; Image.new('RGB',(40,40),'#BA5A32').save(second_image)
    mixed_book=root/'mixed-images.xlsx'
    open_office_use.author({'sheets':{'Data':{'A2':1,'A3':2}},'sheet_options':{'Data':{'images':[{'path':str(second_image),'anchor':'D4','width':120,'description':'Second graphic'}],'charts':[{'type':'line','data':'A1:A3'}]}}},mixed_book,image_book)
    mixed=load_workbook(mixed_book)['Data']; assert len(mixed._images)==2 and len(mixed._charts)==1
    assert [image.anchor.pic.nvPicPr.cNvPr.descr for image in mixed._images]==['Revenue graphic','Second graphic']
    with zipfile.ZipFile(mixed_book) as archive:
        drawing=etree.fromstring(archive.read('xl/drawings/drawing1.xml'))
        identifiers=drawing.xpath('//*[local-name()="cNvPr"]/@id')
    assert len(identifiers)==3 and len(set(identifiers))==3, identifiers
    code, validation, _=officecli_backend.invoke('validate',mixed_book); assert code==0 and validation['success'],validation
    for index, image in enumerate([{'path':str(image_book),'anchor':'A1','width':100}, {'path':str(image_book),'anchor':'XFE1','width':100}, {'path':str(image_book),'anchor':'A1','width':True}]):
        invalid=root/f'bad-image-{index}.xlsx'
        try: open_office_use.author({'sheets':{'Data':{}},'sheet_options':{'Data':{'images':[image]}}},invalid)
        except (ValueError, OSError): pass
        else: raise AssertionError('Invalid worksheet image accepted')
        assert not invalid.exists()
    table_file=root/'table.xlsx'
    table_spec={'sheets':{'Sales':{'A1':'Quarter','B1':'Revenue','A2':'Q1','B2':120,'A3':'Q2','B3':140}},'sheet_options':{'Sales':{'tables':[{'name':'SalesTable','range':'A1:B3','style':'TableStyleMedium2','row_stripes':True}]}}}
    table_json=root/'table-spec.json';table_json.write_text(json.dumps(table_spec))
    result=subprocess.run([sys.executable,str(cli),'author','--spec',str(table_json),'--out',str(table_file)],capture_output=True,text=True,check=True)
    assert json.loads(result.stdout)['output']==str(table_file.resolve())
    table_sheet=load_workbook(table_file)['Sales']; table=table_sheet.tables['SalesTable']
    assert table.ref=='A1:B3' and [column.name for column in table.tableColumns]==['Quarter','Revenue'] and table.autoFilter.ref=='A1:B3' and table.tableStyleInfo.showRowStripes
    code, validation, _=officecli_backend.invoke('validate',table_file); assert code==0 and validation['success'],validation
    try: open_office_use.author({'sheets':{'Sales':{'A1':'Period'}}},root/'unsafe-table-header.xlsx',table_file)
    except ValueError as exc: assert 'table' in str(exc).lower()
    else: raise AssertionError('Table header authoring edit accepted')
    assert not (root/'unsafe-table-header.xlsx').exists()
    copied_table=root/'copied-table.xlsx'
    open_office_use.author({'sheets':{'Sales':{'B2':160}}},copied_table,table_file)
    copied_sheet=load_workbook(copied_table)['Sales']; assert copied_sheet['B2'].value==160 and copied_sheet.tables['SalesTable'].tableColumns[0].name=='Quarter'
    rename_draft=root/'rename-draft';session.start(rename_draft,source=table_file)
    renamed=session.rename_table_column(rename_draft,0,'Sales','SalesTable','Revenue','Net Revenue')
    renamed_sheet=load_workbook(renamed['file'])['Sales']
    assert renamed['revision']==1 and renamed_sheet['B1'].value=='Net Revenue' and renamed_sheet.tables['SalesTable'].tableColumns[1].name=='Net Revenue'
    assert load_workbook(table_file)['Sales']['B1'].value=='Revenue' and session.info(rename_draft)['latest_revision']==1
    with zipfile.ZipFile(table_file) as before,zipfile.ZipFile(renamed['file']) as after:
        changes={name for name in before.namelist() if before.read(name)!=after.read(name)}
        assert {'xl/worksheets/sheet1.xml','xl/tables/table1.xml'} <= changes
        assert changes <= {'xl/worksheets/sheet1.xml','xl/tables/table1.xml','xl/workbook.xml','xl/styles.xml','xl/sharedStrings.xml'}, changes

    code, validation, _=officecli_backend.invoke('validate',renamed['file']); assert code==0 and validation['success'],validation
    try: session.rename_table_column(rename_draft,0,'Sales','SalesTable','Revenue','Other')
    except ValueError as exc: assert 'conflict' in str(exc).lower()
    else: raise AssertionError('Stale table rename accepted')
    try: session.rename_table_column(rename_draft,1,'Sales','SalesTable','Net Revenue','Quarter')
    except ValueError: pass
    else: raise AssertionError('Duplicate table column name accepted')
    assert session.info(rename_draft)['latest_revision']==1
    cli_draft=root/'rename-cli';session.start(cli_draft,source=table_file)
    command=[sys.executable,str(cli),'session','rename-table-column',str(cli_draft),'--base','0','--sheet','Sales','--table','SalesTable','--old','Revenue','--new','Amount']
    assert json.loads(subprocess.run(command,capture_output=True,text=True,check=True).stdout)['new_name']=='Amount'
    for index, expression in enumerate(['=SUM(SalesTable[[#Unknown],[Revenue]])','=INDIRECT(A2)','=[@Revenue]']):
        dependent=root/f'dependent-{index}.xlsx'
        open_office_use.author({'sheets':{'Sales':{'A1':'Quarter','B1':'Revenue','A2':'Q1','B2':120,'C1':expression}},'sheet_options':{'Sales':{'tables':[{'name':'SalesTable','range':'A1:B2'}]}}},dependent)
        blocked=root/f'dependent-draft-{index}';session.start(blocked,source=dependent)
        try: session.rename_table_column(blocked,0,'Sales','SalesTable','Revenue','Amount')
        except ValueError: pass
        else: raise AssertionError('Unsafe table rename accepted')
        assert session.info(blocked)['latest_revision']==0
    encoded_source=root/'dependent-0.xlsx'; encoded=root/'encoded-reference.xlsx'
    with zipfile.ZipFile(encoded_source) as before, zipfile.ZipFile(encoded,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='xl/worksheets/sheet1.xml': data=data.replace(b'[',b'&#91;').replace(b']',b'&#93;')
            after.writestr(entry,data)
    encoded_draft=root/'encoded-draft';session.start(encoded_draft,source=encoded)
    try: session.rename_table_column(encoded_draft,0,'Sales','SalesTable','Revenue','Amount')
    except ValueError as exc: assert 'structured' in str(exc)
    else: raise AssertionError('Encoded structured reference accepted')
    implicit_source=root/'implicit-table.xlsx'
    implicit_spec={'sheets':{'Sales':{'A1':'Quarter','B1':'Revenue','C1':'Double','A2':'Q1','B2':120,'C2':'=[@Revenue]*2'}},'sheet_options':{'Sales':{'tables':[{'name':'SalesTable','range':'A1:C2'}]}}}
    open_office_use.author(implicit_spec,implicit_source)
    implicit_draft=root/'implicit-draft';session.start(implicit_draft,source=implicit_source)
    implicit_result=session.rename_table_column(implicit_draft,0,'Sales','SalesTable','Revenue','Amount')
    assert implicit_result['rewritten_references']==1 and load_workbook(implicit_result['file'])['Sales']['C2'].value=='=[@Amount]*2'
    assert session.rewrite_table_formula('=SUM(SalesTable[[Revenue]:[Quarter]])','SalesTable','Revenue','Amount')[0]=='=SUM(SalesTable[[Amount]:[Quarter]])'
    calculated=root/'calculated-table.xlsx'
    calculated_book=xlsxwriter.Workbook(calculated); calculated_sheet=calculated_book.add_worksheet('Sales')
    calculated_sheet.add_table('A1:C4',{'name':'SalesTable','total_row':True,'data':[['Q1',120],['Q2',140]],'columns':[{'header':'Quarter'},{'header':'Revenue','total_function':'sum'},{'header':'Double','formula':'=[@Revenue]*2','total_function':'sum'}]}); calculated_book.close()
    calculated_totals=root/'calculated-totals.xlsx'
    with zipfile.ZipFile(calculated) as before,zipfile.ZipFile(calculated_totals,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='xl/tables/table1.xml':
                table_root=etree.fromstring(data); ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                column=table_root.find(ns+'tableColumns')[1]; column.set('totalsRowFunction','custom')
                etree.SubElement(column,ns+'totalsRowFormula').text='SUBTOTAL(109,[Revenue])'; data=etree.tostring(table_root)
            after.writestr(entry,data)
    calculated=calculated_totals
    calculated_draft=root/'calculated-draft';session.start(calculated_draft,source=calculated)
    calculated_result=session.rename_table_column(calculated_draft,0,'Sales','SalesTable','Revenue','Net Revenue')
    calculated_after=load_workbook(calculated_result['file'])['Sales']
    assert calculated_result['rewritten_references']==5
    assert calculated_after.tables['SalesTable'].tableColumns[1].totalsRowFormula.text=='SUBTOTAL(109,[Net Revenue])'
    assert calculated_after.tables['SalesTable'].tableColumns[2].calculatedColumnFormula.text=='[[#This Row],Net Revenue]*2'
    assert calculated_after['C2'].value=='=[[#This Row],Net Revenue]*2' and calculated_after['C3'].value=='=[[#This Row],Net Revenue]*2'
    assert calculated_after['B4'].value=='=SUBTOTAL(109,[Net Revenue])'
    assert load_workbook(calculated_result['file'],data_only=True)['Sales']['C2'].value is None and load_workbook(calculated_result['file'],data_only=True)['Sales']['B4'].value is None
    assert load_workbook(calculated)['Sales']['C2'].value=='=[[#This Row],Revenue]*2'
    array_metadata=root/'array-table.xlsx'
    with zipfile.ZipFile(calculated) as before,zipfile.ZipFile(array_metadata,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='xl/tables/table1.xml':
                table_root=etree.fromstring(data); formula=next(table_root.iter(ns+'calculatedColumnFormula')); formula.set('array','1'); data=etree.tostring(table_root)
            after.writestr(entry,data)
    array_draft=root/'array-table-draft';session.start(array_draft,source=array_metadata)
    try: session.rename_table_column(array_draft,0,'Sales','SalesTable','Revenue','Amount')
    except ValueError: pass
    else: raise AssertionError('Array table formula accepted')
    assert session.info(array_draft)['latest_revision']==0
    multi=root/'multiple-tables.xlsx'
    multi_book=xlsxwriter.Workbook(multi); sales=multi_book.add_worksheet('Sales'); other=multi_book.add_worksheet('Other')
    sales.add_table('A1:B3',{'name':'SalesTable','data':[['Q1',120],['Q2',140]],'columns':[{'header':'Quarter'},{'header':'Revenue'}]})
    other.add_table('A1:B3',{'name':'OtherTable','data':[[30],[40]],'columns':[{'header':'Revenue'},{'header':'Adjusted','formula':'=[@Revenue]+SUM(SalesTable[Revenue])'}]}); multi_book.close()
    multi_draft=root/'multiple-tables-draft';session.start(multi_draft,source=multi)
    multi_result=session.rename_table_column(multi_draft,0,'Sales','SalesTable','Revenue','Amount')
    multi_after=load_workbook(multi_result['file']); other_after=multi_after['Other']
    assert multi_result['rewritten_references']==3
    assert other_after['B2'].value=='=[[#This Row],Revenue]+SUM(SalesTable[Amount])'
    assert other_after.tables['OtherTable'].tableColumns[1].calculatedColumnFormula.text=='[[#This Row],Revenue]+SUM(SalesTable[Amount])'
    assert other_after.tables['OtherTable'].tableColumns[0].name=='Revenue'
    assert load_workbook(multi_result['file'],data_only=True)['Other']['B2'].value is None
    ambiguous_multi=root/'ambiguous-multiple-tables.xlsx'
    with zipfile.ZipFile(multi) as before,zipfile.ZipFile(ambiguous_multi,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='xl/_rels/workbook.xml.rels':
                relationships=etree.fromstring(data); rel=etree.SubElement(relationships,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship')
                rel.set('Id','rIdAmbiguousOtherTable'); rel.set('Type','http://schemas.openxmlformats.org/officeDocument/2006/relationships/table'); rel.set('Target','/xl/tables/table2.xml'); data=etree.tostring(relationships)
            after.writestr(entry,data)
    ambiguous_draft=root/'ambiguous-multiple-tables-draft';session.start(ambiguous_draft,source=ambiguous_multi)
    try: session.rename_table_column(ambiguous_draft,0,'Sales','SalesTable','Revenue','Amount')
    except ValueError as exc: assert 'ambiguous' in str(exc).lower()
    else: raise AssertionError('Ambiguous foreign table context accepted')
    assert session.info(ambiguous_draft)['latest_revision']==0
    shared=root/'shared-table.xlsx'
    with zipfile.ZipFile(table_file) as before, zipfile.ZipFile(shared,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='xl/_rels/workbook.xml.rels':
                tree=etree.fromstring(data); rel=etree.SubElement(tree,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship')
                rel.set('Id','rIdSharedTable'); rel.set('Type','http://schemas.openxmlformats.org/officeDocument/2006/relationships/table'); rel.set('Target','/xl/tables/table1.xml'); data=etree.tostring(tree)
            after.writestr(entry,data)
    shared_draft=root/'shared-table-draft';session.start(shared_draft,source=shared)
    try: session.rename_table_column(shared_draft,0,'Sales','SalesTable','Revenue','Amount')
    except ValueError as exc: assert 'shared' in str(exc)
    else: raise AssertionError('Shared table part accepted')
    assert session.info(shared_draft)['latest_revision']==0
    rich_header=root/'rich-header.xlsx'
    with zipfile.ZipFile(table_file) as before,zipfile.ZipFile(rich_header,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='xl/worksheets/sheet1.xml':
                tree=etree.fromstring(data);ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                cell=next(node for node in tree.iter(ns+'c') if node.get('r')=='B1');inline=cell.find(ns+'is'); inline.clear()
                for text in ('Rev','enue'):
                    run=etree.SubElement(inline,ns+'r'); etree.SubElement(run,ns+'t').text=text
                data=etree.tostring(tree)
            after.writestr(entry,data)
    rich_draft=root/'rich-header-draft';session.start(rich_draft,source=rich_header)
    try: session.rename_table_column(rich_draft,0,'Sales','SalesTable','Revenue','Amount')
    except ValueError as exc: assert 'Rich' in str(exc)
    else: raise AssertionError('Rich table header rename accepted')
    assert session.info(rich_draft)['latest_revision']==0
    shared_text=root/'shared-text-table.xlsx'; book=xlsxwriter.Workbook(str(shared_text));ws=book.add_worksheet('Data')
    ws.add_table('A1:B3',{'name':'SharedTable','columns':[{'header':'Quarter'},{'header':'Revenue'}],'data':[['Q1',120],['Q2',140]]});book.close()
    shared_text_draft=root/'shared-text-draft';session.start(shared_text_draft,source=shared_text)
    shared_text_result=session.rename_table_column(shared_text_draft,0,'Data','SharedTable','Revenue','Amount')
    shared_text_sheet=load_workbook(shared_text_result['file'])['Data']
    assert shared_text_sheet['B1'].value=='Amount' and shared_text_sheet.tables['SharedTable'].tableColumns[1].name=='Amount'
    with zipfile.ZipFile(shared_text_result['file']) as archive:
        shared_root=etree.fromstring(archive.read('xl/sharedStrings.xml'))
        assert shared_root.get('count') is None and shared_root.get('uniqueCount') is None
    referenced=root/'referenced.xlsx'
    ref_spec={'sheets':{'Sales':{'A1':'Quarter','B1':'Revenue','A2':'Q1','B2':120,'C2':'=SUM(SalesTable[[#All],[Revenue]])'},'Other':{'A1':'=SUM(SalesTable[Revenue])+SalesTable[Revenue]'}},'sheet_options':{'Sales':{'tables':[{'name':'SalesTable','range':'A1:B2'}]}}}
    open_office_use.author(ref_spec,referenced)
    book=load_workbook(referenced);book.defined_names.add(DefinedName('RevenueRef',attr_text='SalesTable[Revenue]'))
    book.calculation.calcMode='manual';book.calculation.fullCalcOnLoad=False;book.calculation.forceFullCalc=False
    book['Other'].conditional_formatting.add('A1',FormulaRule(formula=['SUM(SalesTable[Revenue])>0']))
    validation_rule=DataValidation(type='custom',formula1='SUM(SalesTable[Revenue])>0');book['Other'].add_data_validation(validation_rule);validation_rule.add('B1')
    linked_chart=LineChart();linked_chart.add_data(Reference(book['Sales'],min_col=2,min_row=1,max_row=2),titles_from_data=True)
    linked_chart.series[0].val.numRef.f='SalesTable[Revenue]';book['Other'].add_chart(linked_chart,'D2');book.save(referenced)
    cached_chain=root/'referenced-cached.xlsx'
    with zipfile.ZipFile(referenced) as before,zipfile.ZipFile(cached_chain,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='xl/worksheets/sheet1.xml':
                tree=etree.fromstring(data); ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                cell=next(item for item in tree.iter(ns+'c') if item.get('r')=='C2'); value=cell.find(ns+'v')
                if value is None: value=etree.SubElement(cell,ns+'v')
                value.text='240'; data=etree.tostring(tree)
            elif entry.filename=='xl/_rels/workbook.xml.rels':
                tree=etree.fromstring(data); rel=etree.SubElement(tree,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship')
                rel.set('Id','rIdCalcChain');rel.set('Type','http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain');rel.set('Target','calcChain.xml');data=etree.tostring(tree)
            elif entry.filename=='[Content_Types].xml':
                tree=etree.fromstring(data); item=etree.SubElement(tree,'{http://schemas.openxmlformats.org/package/2006/content-types}Override')
                item.set('PartName','/xl/calcChain.xml');item.set('ContentType','application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml');data=etree.tostring(tree)
            after.writestr(entry,data)
        after.writestr('xl/calcChain.xml','<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="C2" i="1"/></calcChain>')
    assert load_workbook(cached_chain,data_only=True)['Sales']['C2'].value==240
    referenced=cached_chain
    referenced_draft=root/'referenced-draft';session.start(referenced_draft,source=referenced)
    ref_result=session.rename_table_column(referenced_draft,0,'Sales','SalesTable','Revenue','Net Revenue')
    renamed_book=load_workbook(ref_result['file'])
    assert ref_result['rewritten_references']==7 and renamed_book['Sales']['B1'].value=='Net Revenue'
    assert renamed_book['Sales']['C2'].value=='=SUM(SalesTable[[#All],[Net Revenue]])'
    assert renamed_book['Other']['A1'].value=='=SUM(SalesTable[Net Revenue])+SalesTable[Net Revenue]'
    assert renamed_book.defined_names['RevenueRef'].attr_text=='SalesTable[Net Revenue]'
    assert next(iter(renamed_book['Other'].conditional_formatting._cf_rules.values()))[0].formula==['SUM(SalesTable[Net Revenue])>0']
    assert renamed_book['Other'].data_validations.dataValidation[0].formula1=='SUM(SalesTable[Net Revenue])>0'
    assert renamed_book['Other']._charts[0].series[0].val.numRef.f=='SalesTable[Net Revenue]'

    assert renamed_book.calculation.calcMode=='manual' and renamed_book.calculation.fullCalcOnLoad and renamed_book.calculation.forceFullCalc
    assert load_workbook(ref_result['file'],data_only=True)['Sales']['C2'].value is None
    with zipfile.ZipFile(ref_result['file']) as archive: assert 'xl/calcChain.xml' not in archive.namelist()
    code, validation, _=officecli_backend.invoke('validate',ref_result['file']); assert code==0 and validation['success'],validation
    assert load_workbook(referenced)['Sales']['C2'].value=='=SUM(SalesTable[[#All],[Revenue]])'
    shared_formula=root/'shared-formula.xlsx'
    with zipfile.ZipFile(referenced) as before,zipfile.ZipFile(shared_formula,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='xl/worksheets/sheet1.xml':
                tree=etree.fromstring(data);ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                formula=next(item for item in tree.iter(ns+'c') if item.get('r')=='C2').find(ns+'f')
                formula.set('t','shared');formula.set('si','0');formula.set('ref','C2:D2');data=etree.tostring(tree)
            after.writestr(entry,data)
    shared_formula_draft=root/'shared-formula-draft';session.start(shared_formula_draft,source=shared_formula)
    try: session.rename_table_column(shared_formula_draft,0,'Sales','SalesTable','Revenue','Amount')
    except ValueError as exc: assert 'formula' in str(exc).lower()
    else: raise AssertionError('Shared formula rewrite accepted')
    assert session.info(shared_formula_draft)['latest_revision']==0
    special_draft=root/'special-name-draft';session.start(special_draft,source=referenced)
    try: session.rename_table_column(special_draft,0,'Sales','SalesTable','Revenue','X#Y')
    except ValueError: pass
    else: raise AssertionError('Unsupported structured reference name accepted')

    normal_formula=root/'normal-formula.xlsx'
    open_office_use.author({'sheets':{'Sales':{'A1':'Quarter','B1':'Revenue','A2':'Q1','B2':120,'C1':'=B2*2'}},'sheet_options':{'Sales':{'tables':[{'name':'SalesTable','range':'A1:B2'}]}}},normal_formula)
    formula_draft=root/'normal-formula-draft';session.start(formula_draft,source=normal_formula)
    formula_renamed=session.rename_table_column(formula_draft,0,'Sales','SalesTable','Revenue','Amount')
    assert load_workbook(formula_renamed['file'])['Sales']['C1'].value=='=B2*2' and formula_renamed['formula_cache_effect']=='may_be_stale'
    table_draft=root/'table-draft';session.start(table_draft,source=table_file)
    grid=session.read_grid(table_draft,0,'Sales','A1:B2');assert grid['rows'][0][0]['read_only_reason']=='table_header' and grid['rows'][1][1]['editable']
    try: session.replace_text(table_draft,0,'/Sales/A1','Quarter','Period')
    except ValueError as exc: assert 'table' in str(exc).lower()
    else: raise AssertionError('Table header replacement accepted')
    try: session.edit(table_draft,0,'/Sales/A1','literal','Period')
    except ValueError as exc: assert 'table_header' in str(exc)
    else: raise AssertionError('Table header element edit accepted')
    try: session.edit_grid(table_draft,0,'Sales','A1',[{'cell':'A1','value':'Period'}])
    except ValueError as exc: assert 'table_header' in str(exc)
    else: raise AssertionError('Table header grid edit accepted')
    changed=session.edit_grid(table_draft,0,'Sales','B2',[{'cell':'B2','value':'150'}]);updated=load_workbook(changed['file'])['Sales']
    assert updated['B2'].value==150 and updated.tables['SalesTable'].tableColumns[0].name=='Quarter'
    assert Path(open_office_use.render(table_file,root/'table-render')).read_bytes().startswith(b'%PDF')
    for index, options in enumerate([{'tables':[{'name':'A1','range':'A1:B3'}]}, {'tables':[{'name':'T','range':'A1:B1'}]}, {'tables':[{'name':'T','range':'A1:B3','style':'Invalid'}]}, {'tables':[{'name':'T','range':'A1:B3'},{'name':'t','range':'D1:E3'}]}, {'tables':[{'name':'T','range':'A1:B3'},{'name':'U','range':'B2:C3'}]}, {'auto_filter':'A1:B3','tables':[{'name':'T','range':'A1:B3'}]}]):
        target=root/f'bad-structured-table-{index}.xlsx'
        try: open_office_use.author({'sheets':{'Sales':{'A1':'Quarter','B1':'Revenue','A2':'Q1','B2':10,'D1':'X','E1':'Y'}},'sheet_options':{'Sales':options}},target)
        except ValueError: pass
        else: raise AssertionError(f'Invalid table accepted: {options}')
        assert not target.exists()
    for index, headers in enumerate([['Name','name'],['Name',''],['=1','Value']]):
        target=root/f'bad-table-headers-{index}.xlsx'
        try: open_office_use.author({'sheets':{'Sales':{'A1':headers[0],'B1':headers[1],'A2':'x','B2':1}},'sheet_options':{'Sales':{'tables':[{'name':'SalesTable','range':'A1:B2'}]}}},target)
        except ValueError: pass
        else: raise AssertionError('Invalid table headers accepted')
        assert not target.exists()
    totals_file=root/'totals.xlsx';book=xlsxwriter.Workbook(str(totals_file));ws=book.add_worksheet('Data')
    ws.add_table('A1:B4',{'name':'TotalsTable','total_row':True,'columns':[{'header':'Item','total_string':'Total'},{'header':'Amount','total_function':'sum'}],'data':[['A',10],['B',20]]});book.close()
    totals_draft=root/'totals-draft';session.start(totals_draft,source=totals_file)
    totals_grid=session.read_grid(totals_draft,0,'Data','A4:B4')
    assert all(cell['read_only_reason']=='table_totals' for cell in totals_grid['rows'][0]),totals_grid
    try: session.edit_grid(totals_draft,0,'Data','A4',[{'cell':'A4','value':'Other'}])
    except ValueError: pass
    else: raise AssertionError('Table total edit accepted')
    before = output.read_bytes(); revised = root/'revised.xlsx'
    open_office_use.author({'sheet_options':{'Forecast':{'formats':[{'range':'A1','italic':True}]}}}, revised, output)
    revised_sheet = load_workbook(revised)['Forecast']; assert revised_sheet['A1'].font.bold and revised_sheet['A1'].font.italic and output.read_bytes() == before
    filtered = root/'filtered.xlsx'; filtered_book = load_workbook(output); filtered_book['Forecast'].auto_filter.add_filter_column(0, ['Q1']); filtered_book.save(filtered)
    try: open_office_use.author({'sheet_options':{'Forecast':{'auto_filter':'A1:C2'}}}, root/'unsafe.xlsx', filtered)
    except ValueError as exc: assert 'filter criteria' in str(exc)
    else: raise AssertionError('Existing filter criteria were silently changed')
    assert not (root/'unsafe.xlsx').exists()
    for index, opts in enumerate([{'freeze_panes':'A1:B2'}, {'fit_to_page':'yes'}, {'print_options':{'paper':'bad'}}, {'print_options':{'repeat_rows':'2:1'}}, {'print_options':{'repeat_columns':'A:XFE'}}, {'print_options':{'paper':'letter','margins':{'left':5,'right':5}}}, {'print_options':{'footer':{'page_number':'yes'}}}, {'column_widths':{'XFE':10}}, {'row_heights':{'1':True}}, {'formats':[{'range':'A1:Z10000','bold':True}]}, {'formats':[{'range':'A1','font_color':'BAD'}]}, {'charts':[{'type':'line','data':'B1:B3','categories':'A2:A2'}]}, {'charts':[{'type':'pie','data':'B1:B3'}]}]):
        target = root/f'bad{index}.xlsx'
        try: open_office_use.author({'sheets':{'Forecast':{}},'sheet_options':{'Forecast':opts}}, target)
        except ValueError: pass
        else: raise AssertionError('Invalid sheet setting accepted')
        assert not target.exists()
print('Workbook author passed: settings, formulas, formatting, render and invalid input')

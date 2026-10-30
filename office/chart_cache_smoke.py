"""Exercise chart cache comparisons against a real workbook and revisions."""
import json
import copy
import subprocess
import sys
import tempfile
import zipfile
import urllib.request
import urllib.error
from urllib.parse import urlsplit
from pathlib import Path
from lxml import etree
from openpyxl import load_workbook
import xlsxwriter
import chart_cache
import officecli_backend
import review_bundle
import session

C=chart_cache.C
with tempfile.TemporaryDirectory() as tmp:
    root=Path(tmp); source=root/'chart.xlsx'; draft=root/'draft'
    book=xlsxwriter.Workbook(str(source)); sheet=book.add_worksheet("Sales Data")
    sheet.write_column('A1',['Month','Jan','Feb']); sheet.write_column('B1',['Revenue',10,20])
    chart=book.add_chart({'type':'column'}); chart.set_title({'name':'Revenue'}); chart.add_series({'name':"='Sales Data'!$B$1",'categories':"='Sales Data'!$A$2:$A$3",'values':"='Sales Data'!$B$2:$B$3"}); sheet.insert_chart('D2',chart); book.close()
    baseline=chart_cache.scan(source); assert baseline['chart_parts_total']==1 and baseline['complete'] and baseline['mismatch_references']==0,baseline
    pie_source=root/'pie-colors.xlsx'; pie_book=xlsxwriter.Workbook(str(pie_source)); pie_sheet=pie_book.add_worksheet('Data')
    pie_sheet.write_column('A1',['A','B','C']);pie_sheet.write_column('B1',[45,30,25]);pie=pie_book.add_chart({'type':'pie'});pie.add_series({'categories':'=Data!$A$1:$A$3','values':'=Data!$B$1:$B$3','points':[{'border':{'color':'black'}},{},{}]});pie_sheet.insert_chart('D2',pie);pie_book.close()
    plot=etree.Element(C+'pieChart'); etree.SubElement(plot,C+'ser')
    def literal(role,kind,values): return {'role':role,'kind':kind,'status':'literal','points':len(values),'stored_points':len(values),'examples':[{'point':i,'value':value,'value_truncated':False} for i,value in enumerate(values,1)]}
    labels=literal('cat','strLit',['Enterprise 45%','SMB 30%','Consumer 25%']); nums=literal('val','numLit',['45','30','25'])
    assert not chart_cache.pie_category_percentages([plot],[labels,nums])['mismatch_candidate']
    nums['examples'][0]['value']='50'; assert chart_cache.pie_category_percentages([plot],[labels,nums])['mismatch_points']==[1,2,3]
    labels['examples'][0]['value']='Enterprise'; assert chart_cache.pie_category_percentages([plot],[labels,nums])['status']=='not_applicable'
    pie_report=chart_cache.scan(pie_source);pie_part=pie_report['charts'][0]['part'];pie_draft=root/'pie-colors-draft';session.start(pie_draft,source=pie_source)
    assert pie_report['charts'][0]['pie_point_appearance']['default_fill_candidate']
    for colors in ([],['FFFFFF'],['GGGGGG','FFFFFF','FFFFFF']):
        try: session.set_chart_pie_point_colors(pie_draft,0,pie_part,colors)
        except ValueError: pass
        else: raise AssertionError('Invalid pie colors accepted')
    try: session.set_chart_pie_point_colors(pie_draft,0,pie_part,['FF0000','00FF00','0000FF'],'0'*64)
    except ValueError: pass
    else: raise AssertionError('Stale pie color snapshot accepted')
    pie_colors=session.set_chart_pie_point_colors(pie_draft,0,pie_part,['FF0000','00FF00','0000FF'],pie_report['snapshot_sha256'])
    appearance=chart_cache.scan(pie_colors['output'])['charts'][0]['pie_point_appearance']; assert not appearance['default_fill_candidate'] and [point['rgb'] for point in appearance['points']]==['FF0000','00FF00','0000FF']
    with zipfile.ZipFile(pie_source) as before,zipfile.ZipFile(pie_colors['output']) as after:
        assert all(before.read(name)==after.read(name) for name in before.namelist() if name!=pie_part)
        tree=etree.fromstring(after.read(pie_part)); colors=tree.findall('.//'+C+'dPt/'+C+'spPr/{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill/{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr'); assert [node.get('val') for node in colors]==['FF0000','00FF00','0000FF']
        assert tree.find('.//'+C+'dPt/'+C+'spPr/{http://schemas.openxmlformats.org/drawingml/2006/main}ln') is not None
    code,validation,_=officecli_backend.invoke('validate',pie_colors['output']);assert code==0 and validation['success'],validation
    cli=Path(__file__).with_name('open_office_use.py');cmd=[sys.executable,str(cli),'session','set-chart-pie-point-colors',str(pie_draft),'--base','1','--part',pie_part,'--colors',json.dumps(['112233','445566','778899'])]
    assert json.loads(subprocess.run(cmd,capture_output=True,text=True,check=True).stdout)['revision']==2
    with zipfile.ZipFile(pie_draft/'revision-000002.xlsx') as archive:
        tree=etree.fromstring(archive.read(pie_part)); fills=tree.findall('.//'+C+'dPt/'+C+'spPr/{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill/{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
        assert [node.get('val') for node in fills]==['112233','445566','778899']
    code,validation,_=officecli_backend.invoke('validate',pie_draft/'revision-000002.xlsx');assert code==0 and validation['success'],validation
    wrong_draft=root/'not-pie-draft';session.start(wrong_draft,source=source)
    try: session.set_chart_pie_point_colors(wrong_draft,0,baseline['charts'][0]['part'],['FF0000','00FF00'])
    except ValueError: pass
    else: raise AssertionError('Non-pie color edit accepted')
    shared_source=root/'shared-text-chart.xlsx'; shared_book=xlsxwriter.Workbook(str(shared_source)); shared_sheet=shared_book.add_worksheet('Data')
    shared_sheet.write('A1','prefix x005F_text'); shared_sheet.write('B1',10)
    shared_chart=shared_book.add_chart({'type':'column'}); shared_chart.add_series({'categories':'=Data!$A$1','values':'=Data!$B$1'}); shared_sheet.insert_chart('D1',shared_chart); shared_book.close()
    shared_scan=chart_cache.scan(shared_source); assert shared_scan['complete'] and shared_scan['mismatch_references']==0,shared_scan
    session.start(draft,source=source); changed=session.replace_text(draft,0,'/Sales Data/A2','Jan','March')
    result=session.scan_chart_caches(draft,1); assert result['mismatch_references']==1 and result['complete'],result
    mismatch=next(r for r in result['charts'][0]['references'] if r['status']=='mismatch')
    assert mismatch['role']=='cat' and mismatch['mismatches'][0]['cell']=='A2' and mismatch['mismatches'][0]['stored_value']=='March' and mismatch['mismatches'][0]['cached_value']=='Jan'
    assert session.scan_chart_caches(draft,0)['mismatch_references']==0
    assert load_workbook(changed['file'])['Sales Data']['A2'].value=='March'
    try: session.sync_chart_cache(draft,1,result['charts'][0]['part'],mismatch['index'],expected_sha256='0'*64)
    except ValueError as exc: assert 'snapshot changed' in str(exc)
    else: raise AssertionError('Wrong workbook chart snapshot accepted')
    assert session.info(draft)['latest_revision']==1
    synced=session.sync_chart_cache(draft,1,result['charts'][0]['part'],mismatch['index'],expected_sha256=result['snapshot_sha256'])
    assert synced['revision']==2 and synced['updated_points']==1 and synced['source_formula_cells']==0
    assert session.scan_chart_caches(draft,2)['mismatch_references']==0
    with zipfile.ZipFile(changed['file']) as before,zipfile.ZipFile(synced['output']) as after:
        changed_parts={name for name in before.namelist() if before.read(name)!=after.read(name)}
        assert changed_parts=={'xl/charts/chart1.xml'},changed_parts
    validation_code, validation, _=officecli_backend.invoke('validate',synced['output']); assert validation_code==0 and validation['success'],validation
    try: session.sync_chart_cache(draft,1,result['charts'][0]['part'],mismatch['index'])
    except ValueError as exc: assert 'conflict' in str(exc)
    else: raise AssertionError('Stale sync accepted')
    blank_draft=root/'blank-draft'; session.start(blank_draft,source=source)
    session.edit_grid(blank_draft,0,'Sales Data','B2',[{'cell':'B2','value':''}])
    blank_scan=session.scan_chart_caches(blank_draft,1); blank_ref=next(ref for ref in blank_scan['charts'][0]['references'] if ref['status']=='mismatch')
    assert blank_ref['mismatches'][0]['reason']=='extra_point' and blank_ref['unverified_points']==0
    blank_sync=session.sync_chart_cache(blank_draft,1,blank_scan['charts'][0]['part'],blank_ref['index'])
    assert blank_sync['updated_points']==1 and session.scan_chart_caches(blank_draft,2)['complete'] and session.scan_chart_caches(blank_draft,2)['mismatch_references']==0
    blank_fill=session.edit_grid(blank_draft,2,'Sales Data','B2',[{'cell':'B2','value':'15'}])
    fill_scan=session.scan_chart_caches(blank_draft,3); fill_ref=next(ref for ref in fill_scan['charts'][0]['references'] if ref['status']=='mismatch')
    assert fill_ref['mismatches'][0]['reason']=='missing_point'
    session.sync_chart_cache(blank_draft,3,fill_scan['charts'][0]['part'],fill_ref['index'])
    assert session.scan_chart_caches(blank_draft,4)['mismatch_references']==0
    signed=root/'signed.xlsx'
    with zipfile.ZipFile(changed['file']) as before,zipfile.ZipFile(signed,'w') as after:
        for entry in before.infolist(): after.writestr(entry,before.read(entry))
        after.writestr('_xmlsignatures/sig1.xml','<Signature/>')
    try: chart_cache.sync(signed,root/'signed-output.xlsx',result['charts'][0]['part'],mismatch['index'])
    except ValueError as exc: assert 'signed' in str(exc)
    else: raise AssertionError('Signed workbook sync accepted')
    assert not (root/'signed-output.xlsx').exists()
    def no_render(*_): raise RuntimeError('preview unavailable')
    pie_bundle=review_bundle.create(pie_source,root/'pie-review',no_render,1); assert 'chart_pie_default_fill_candidates' in pie_bundle['flags']
    bundle=review_bundle.create(changed['file'],root/'review',no_render,1)
    assert 'chart_cache_mismatch' in bundle['flags'] and 'render_pdf' in bundle['check_errors'],bundle
    report=json.loads(Path(bundle['report']).read_text()); assert report['checks']['chart_cache_scan']['result']['mismatch_references']==1
    cli=Path(__file__).with_name('open_office_use.py'); output=subprocess.run([sys.executable,str(cli),'session','scan-chart-caches',str(draft),'--revision','1'],capture_output=True,text=True,check=True)
    assert json.loads(output.stdout)['mismatch_references']==1
    cli_draft=root/'cli-draft'; session.start(cli_draft,source=changed['file'])
    cli_result=subprocess.run([sys.executable,str(cli),'session','sync-chart-cache',str(cli_draft),'--base','0','--part',result['charts'][0]['part'],'--reference',str(mismatch['index']),'--expected-sha256',result['snapshot_sha256']],capture_output=True,text=True,check=True)
    assert json.loads(cli_result.stdout)['updated_points']==1
    # Missing caches and unsupported external references must be visible as incomplete coverage.
    def rewrite(name, mutate):
        target=root/name
        with zipfile.ZipFile(source) as before,zipfile.ZipFile(target,'w') as after:
            for entry in before.infolist():
                data=before.read(entry)
                if entry.filename=='xl/charts/chart1.xml':
                    tree=etree.fromstring(data); mutate(tree); data=etree.tostring(tree,encoding='utf-8',xml_declaration=True)
                after.writestr(entry,data)
        return target
    def literal_values(tree):
        parent=tree.find('.//'+C+'val'); reference=parent.find(C+'numRef'); cache=reference.find(C+'numCache')
        reference.remove(cache); cache.tag=C+'numLit'; parent.replace(reference,cache)
    static_source=rewrite('static-values.xlsx',literal_values); static_scan=chart_cache.scan(static_source)
    static_ref=next(item for item in static_scan['charts'][0]['references'] if item['kind']=='numLit')
    assert static_scan['complete'] and static_scan['literal_sources']==1 and static_ref['status']=='literal' and static_ref['reason']=='no_worksheet_reference' and static_ref['examples'][0]['value']=='10', static_scan
    static_draft=root/'static-draft'; session.start(static_draft,source=static_source)
    session.edit_grid(static_draft,0,'Sales Data','B2',[{'cell':'B2','value':'99'}])
    after_static=session.scan_chart_caches(static_draft,1)
    assert after_static['literal_sources']==1 and after_static['mismatch_references']==0
    batch_draft=root/'static-batch-draft'; session.start(batch_draft,source=static_source)
    batch=[{'reference':static_ref['index'],'point':1,'value':11},{'reference':static_ref['index'],'point':2,'value':22}]
    for bad in (batch+[batch[0]], [batch[0],dict(batch[1],point=3)], [batch[0],dict(batch[1],value=9007199254740993)]):
        try: session.set_chart_literal_values(batch_draft,0,static_scan['charts'][0]['part'],bad)
        except ValueError: pass
        else: raise AssertionError('Invalid literal batch accepted')
        assert session.info(batch_draft)['latest_revision']==0
    try: session.set_chart_literal_values(batch_draft,0,static_scan['charts'][0]['part'],batch,'0'*64)
    except ValueError: pass
    else: raise AssertionError('Stale literal batch accepted')
    batched=session.set_chart_literal_values(batch_draft,0,static_scan['charts'][0]['part'],batch,static_scan['snapshot_sha256'])
    assert [item['old_value'] for item in batched['changes']]==[10,20] and [item['old_value_raw'] for item in batched['changes']]==['10','20'] and batched['revision']==1
    values=next(item for item in session.scan_chart_caches(batch_draft,1)['charts'][0]['references'] if item['kind']=='numLit')['examples']
    assert [item['value'] for item in values]==['11','22']
    with zipfile.ZipFile(static_source) as before,zipfile.ZipFile(batched['output']) as after:
        assert all(before.read(name)==after.read(name) for name in before.namelist() if name!=static_scan['charts'][0]['part'])
    command=[sys.executable,str(cli),'session','set-chart-literal-values',str(batch_draft),'--base','1','--part',static_scan['charts'][0]['part'],'--changes',json.dumps(batch)]
    assert json.loads(subprocess.run(command,capture_output=True,text=True,check=True).stdout)['revision']==2
    literal_draft=root/'static-point-draft'; session.start(literal_draft,source=static_source)
    for point,value in ((0,20),(3,20),(1,True),(1,float('inf'))):
        try: session.set_chart_literal_value(literal_draft,0,static_scan['charts'][0]['part'],static_ref['index'],point,value)
        except ValueError: pass
        else: raise AssertionError('Invalid static chart value accepted')
    try: session.set_chart_literal_value(literal_draft,0,static_scan['charts'][0]['part'],static_ref['index'],1,30,'0'*64)
    except ValueError: pass
    else: raise AssertionError('Stale static chart snapshot accepted')
    assert session.info(literal_draft)['latest_revision']==0
    changed_literal=session.set_chart_literal_value(literal_draft,0,static_scan['charts'][0]['part'],static_ref['index'],1,30,static_scan['snapshot_sha256'])
    assert changed_literal['old_value']==10
    with zipfile.ZipFile(static_source) as before,zipfile.ZipFile(changed_literal['output']) as after:
        assert all(before.read(name)==after.read(name) for name in before.namelist() if name!=static_scan['charts'][0]['part'])
        tree=etree.fromstring(after.read(static_scan['charts'][0]['part'])); assert tree.find('.//'+C+'val/'+C+'numLit/'+C+'pt/'+C+'v').text=='30'
    assert load_workbook(changed_literal['output'])['Sales Data']['B2'].value==10
    precision_draft=root/'literal-precision-draft'; session.start(precision_draft,source=static_source)
    precision_cli=[sys.executable,str(cli),'session','set-chart-literal-value',str(precision_draft),'--base','0','--part',static_scan['charts'][0]['part'],'--reference',str(static_ref['index']),'--point','1','--value']
    rejected=subprocess.run(precision_cli+['9007199254740993'],capture_output=True,text=True)
    assert rejected.returncode != 0 and session.info(precision_draft)['latest_revision']==0
    accepted=subprocess.run(precision_cli+['9007199254740994'],capture_output=True,text=True,check=True)
    assert json.loads(accepted.stdout)['revision']==1
    precise=session.scan_chart_caches(precision_draft,1); example=next(item for item in precise['charts'][0]['references'] if item['kind']=='numLit')['examples'][0]
    assert example['value']=='9007199254740994',example
    code,validation,_=officecli_backend.invoke('validate',changed_literal['output']); assert code==0 and validation['success'],validation
    command=[sys.executable,str(cli),'session','set-chart-literal-value',str(literal_draft),'--base','1','--part',static_scan['charts'][0]['part'],'--reference',str(static_ref['index']),'--point','2','--value','35']
    assert json.loads(subprocess.run(command,capture_output=True,text=True,check=True).stdout)['revision']==2
    referenced_draft=root/'referenced-point-draft';session.start(referenced_draft,source=source)
    ref=next(item for item in baseline['charts'][0]['references'] if item['role']=='val')
    try: session.set_chart_literal_value(referenced_draft,0,baseline['charts'][0]['part'],ref['index'],1,30)
    except ValueError: pass
    else: raise AssertionError('Referenced cache edited as literal')
    try: session.sync_chart_cache(static_draft,1,static_scan['charts'][0]['part'],static_ref['index'])
    except ValueError: pass
    else: raise AssertionError('Literal chart data accepted for cache sync')
    def invalid_literal(tree):
        literal_values(tree); points=tree.find('.//'+C+'val/'+C+'numLit').findall(C+'pt'); points[1].set('idx',points[0].get('idx'))
    bad_literal=chart_cache.scan(rewrite('invalid-literal.xlsx',invalid_literal))
    assert not bad_literal['complete'] and any(item.get('reason')=='invalid_literal_points' for item in bad_literal['charts'][0]['references'])
    def no_render(_source,_out): raise RuntimeError('Rendering intentionally unavailable')
    static_review=review_bundle.create(static_source,root/'static-review',no_render)
    assert 'chart_literal_data_present' in static_review['flags']
    def literal_categories(tree):
        ser=tree.find('.//'+C+'ser'); parent=ser.find(C+'cat'); ref=parent.find(C+'strRef'); cache=ref.find(C+'strCache')
        ref.remove(cache); cache.tag=C+'strLit'; parent.replace(ref,cache)
        duplicate=copy.deepcopy(ser); duplicate.find(C+'idx').set('val','1'); duplicate.find(C+'order').set('val','1'); ser.addnext(duplicate)
    category_source=rewrite('literal-categories.xlsx',literal_categories); category_report=chart_cache.scan(category_source)
    category_ref=next(item for item in category_report['charts'][0]['references'] if item['kind']=='strLit')
    category_draft=root/'category-draft';session.start(category_draft,source=category_source)
    for text in ('','x'*256,'bad\x01'):
        try: session.set_chart_literal_category(category_draft,0,category_report['charts'][0]['part'],category_ref['index'],1,text)
        except ValueError: pass
        else: raise AssertionError('Invalid category accepted')
    try: session.set_chart_literal_category(category_draft,0,category_report['charts'][0]['part'],category_ref['index'],1,'North','0'*64)
    except ValueError: pass
    else: raise AssertionError('Stale category snapshot accepted')
    changed_category=session.set_chart_literal_category(category_draft,0,category_report['charts'][0]['part'],category_ref['index'],1,'North & East',category_report['snapshot_sha256'])
    assert changed_category['series_updated']==2 and changed_category['old_text']=='Jan'
    with zipfile.ZipFile(category_source) as before,zipfile.ZipFile(changed_category['output']) as after:
        assert all(before.read(name)==after.read(name) for name in before.namelist() if name!=category_report['charts'][0]['part'])
        tree=etree.fromstring(after.read(category_report['charts'][0]['part'])); labels=tree.findall('.//'+C+'cat/'+C+'strLit/'+C+'pt[@idx="0"]/'+C+'v'); assert len(labels)==2 and all(node.text=='North & East' for node in labels)
    code,validation,_=officecli_backend.invoke('validate',changed_category['output']);assert code==0 and validation['success'],validation
    command=[sys.executable,str(cli),'session','set-chart-literal-category',str(category_draft),'--base','1','--part',category_report['charts'][0]['part'],'--reference',str(category_ref['index']),'--point','2','--text','South']
    assert json.loads(subprocess.run(command,capture_output=True,text=True,check=True).stdout)['revision']==2
    def different_categories(tree):
        literal_categories(tree); labels=tree.findall('.//'+C+'cat/'+C+'strLit/'+C+'pt[@idx="0"]/'+C+'v'); labels[-1].text='Different'
    conflict_source=rewrite('different-categories.xlsx',different_categories); conflict_draft=root/'category-conflict-draft';session.start(conflict_draft,source=conflict_source)
    try: session.set_chart_literal_category(conflict_draft,0,category_report['charts'][0]['part'],category_ref['index'],1,'North')
    except ValueError: pass
    else: raise AssertionError('Mismatched category arrays accepted')
    assert session.info(conflict_draft)['latest_revision']==0
    def literal_name(tree):
        title=tree.find('.//'+C+'ser/'+C+'tx'); title.clear(); etree.SubElement(title,C+'v').text='Revenue'
    names_source=rewrite('literal-series.xlsx',literal_name); names_report=chart_cache.scan(names_source)
    assert names_report['charts'][0]['series_complete'] and names_report['charts'][0]['series'][0]['name']=='Revenue' and names_report['charts'][0]['series'][0]['name_status']=='literal'
    names_draft=root/'literal-name-draft';session.start(names_draft,source=names_source)
    for text in ('','x'*256,'bad\x01'):
        try: session.set_chart_literal_series_name(names_draft,0,names_report['charts'][0]['part'],1,text)
        except ValueError: pass
        else: raise AssertionError('Invalid chart name accepted')
    try: session.set_chart_literal_series_name(names_draft,0,names_report['charts'][0]['part'],1,'New','0'*64)
    except ValueError: pass
    else: raise AssertionError('Stale chart name snapshot accepted')
    named=session.set_chart_literal_series_name(names_draft,0,names_report['charts'][0]['part'],1,'Sales & Service',names_report['snapshot_sha256'])
    assert named['old_text']=='Revenue' and session.scan_chart_caches(names_draft,1)['charts'][0]['series'][0]['name']=='Sales & Service'
    with zipfile.ZipFile(names_source) as before,zipfile.ZipFile(named['output']) as after:
        assert all(before.read(name)==after.read(name) for name in before.namelist() if name!=names_report['charts'][0]['part'])
    code,validation,_=officecli_backend.invoke('validate',named['output']); assert code==0 and validation['success'],validation
    cmd=[sys.executable,str(cli),'session','set-chart-literal-series-name',str(names_draft),'--base','1','--part',names_report['charts'][0]['part'],'--series','1','--text','Updated']
    assert json.loads(subprocess.run(cmd,capture_output=True,text=True,check=True).stdout)['revision']==2
    source_draft=root/'referenced-name-draft';session.start(source_draft,source=source)
    try: session.set_chart_literal_series_name(source_draft,0,baseline['charts'][0]['part'],1,'Wrong')
    except ValueError: pass
    else: raise AssertionError('Referenced chart name accepted')
    def literal_viewer_fixture(tree):
        literal_categories(tree); literal_values(tree); literal_name(tree)
    viewer_source=rewrite('literal-viewer.xlsx',literal_viewer_fixture); viewer_draft=root/'literal-viewer-draft';session.start(viewer_draft,source=viewer_source)
    server=subprocess.Popen([sys.executable,str(cli),'watch',str(viewer_draft)],stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True)
    try:
        url=json.loads(server.stdout.readline())['url']; origin='http://'+urlsplit(url).netloc
        def post(route,payload,request_origin=origin):
            request=urllib.request.Request(url+route,data=json.dumps(payload).encode(),headers={'Origin':request_origin,'Content-Type':'application/json'})
            with urllib.request.urlopen(request,timeout=30) as response: return json.load(response)
        report=post('chart-cache-scan',{'revision':0}); chart=report['charts'][0]
        value_ref=next(item for item in chart['references'] if item['kind']=='numLit' and item['role']=='val')
        payload={'base_revision':0,'snapshot_sha256':report['snapshot_sha256'],'kind':'value','part':chart['part'],'reference':value_ref['index'],'point':1,'series':None,'value':42}
        for bad,request_origin,code in [(dict(payload,snapshot_sha256='0'*64),origin,400),(dict(payload,value=True),origin,400),(payload,'http://evil.test',403)]:
            try: post('chart-literal-save',bad,request_origin)
            except urllib.error.HTTPError as error: assert error.code==code
            else: raise AssertionError('Invalid literal viewer edit accepted')
        assert post('chart-literal-save',payload)['revision']==1
        report=post('chart-cache-scan',{'revision':1}); category=next(item for item in report['charts'][0]['references'] if item['kind']=='strLit' and item['role']=='cat')
        category_payload={'base_revision':1,'snapshot_sha256':report['snapshot_sha256'],'kind':'category','part':chart['part'],'reference':category['index'],'point':1,'series':None,'value':'January'}
        assert post('chart-literal-save',category_payload)['series_updated']==2
        report=post('chart-cache-scan',{'revision':2})
        name_payload={'base_revision':2,'snapshot_sha256':report['snapshot_sha256'],'kind':'series_name','part':chart['part'],'reference':None,'point':None,'series':1,'value':'Sales'}
        assert post('chart-literal-save',name_payload)['revision']==3
        report=post('chart-cache-scan',{'revision':3}); assert report['charts'][0]['series'][0]['name']=='Sales'
        batch_payload={'base_revision':3,'snapshot_sha256':report['snapshot_sha256'],'part':chart['part'],'changes':[{'reference':value_ref['index'],'point':1,'value':11},{'reference':value_ref['index'],'point':2,'value':22}]}
        for bad,request_origin,code in [(dict(batch_payload,snapshot_sha256='0'*64),origin,400),(dict(batch_payload,changes=batch_payload['changes']*2),origin,400),(batch_payload,'http://evil.test',403)]:
            try: post('chart-batch-save',bad,request_origin)
            except urllib.error.HTTPError as error: assert error.code==code
            else: raise AssertionError('Invalid batch viewer edit accepted')
        assert post('chart-batch-save',batch_payload)['revision']==4
        updated=post('chart-cache-scan',{'revision':4}); values=next(item for item in updated['charts'][0]['references'] if item['index']==value_ref['index']); assert [item['value'] for item in values['examples']]==['11','22']
    finally:
        server.terminate(); server.wait(timeout=10)
    def missing(tree):
        ref=tree.find('.//'+C+'cat/'+C+'strRef'); ref.remove(ref.find(C+'strCache'))
    missing_result=chart_cache.scan(rewrite('missing.xlsx',missing)); assert missing_result['missing_caches']==1 and not missing_result['complete']
    def external(tree): tree.find('.//'+C+'val/'+C+'numRef/'+C+'f').text="'[other.xlsx]Data'!$B$2:$B$3"
    external_result=chart_cache.scan(rewrite('external.xlsx',external)); assert external_result['unverified_references']>=1 and not external_result['complete']
    def escaped(tree): tree.find('.//'+C+'cat/'+C+'strRef/'+C+'strCache/'+C+'pt/'+C+'v').text='_x0041_'
    escape_result=chart_cache.scan(rewrite('escaped.xlsx',escaped)); assert not escape_result['complete'] and escape_result['unverified_references']>=1
    def missing_point(tree):
        cache=tree.find('.//'+C+'cat/'+C+'strRef/'+C+'strCache'); cache.remove(cache.find(C+'pt'))
    gap_source=rewrite('missing-point.xlsx',missing_point); gap=chart_cache.scan(gap_source); gap_ref=next(x for x in gap['charts'][0]['references'] if x['status']=='mismatch')
    gap_result=chart_cache.sync(gap_source,root/'gap-output.xlsx',gap['charts'][0]['part'],gap_ref['index'])
    assert gap_result['updated_points']==1 and chart_cache.scan(gap_result['output'])['mismatch_references']==0
    validation_code, validation, _=officecli_backend.invoke('validate',gap_result['output']); assert validation_code==0 and validation['success'],validation
    def extension(tree): etree.SubElement(tree,C+'extLst')
    extension_result=chart_cache.scan(rewrite('extension.xlsx',extension)); assert extension_result['charts'][0]['extensions_present'] and not extension_result['complete']
    def extended_mismatch(tree):
        extension(tree); tree.find('.//'+C+'cat/'+C+'strRef/'+C+'strCache/'+C+'pt/'+C+'v').text='Wrong'
    extended_draft=root/'extension-draft'; session.start(extended_draft,source=rewrite('extended-mismatch.xlsx',extended_mismatch))
    extended=session.scan_chart_caches(extended_draft,0); extended_ref=next(r for r in extended['charts'][0]['references'] if r['status']=='mismatch')
    try: session.sync_chart_cache(extended_draft,0,extended['charts'][0]['part'],extended_ref['index'])
    except ValueError as exc: assert 'extensions' in str(exc)
    else: raise AssertionError('Extended chart sync accepted')
    assert chart_cache.parse_reference("'O''Brien!Data'!$A$1:$A$2")[0]=="O'Brien!Data"
    assert chart_cache.parse_reference('Sheet!A1:B2') is None
    mixed_source=root/'mixed-formula.xlsx'; mixed_book=xlsxwriter.Workbook(str(mixed_source)); mixed_sheet=mixed_book.add_worksheet('Calc')
    mixed_sheet.write_formula('A1','=1+1',None,''); mixed_sheet.write_number('A2',20)
    mixed_chart=mixed_book.add_chart({'type':'column'}); mixed_chart.add_series({'values':'=Calc!$A$1:$A$2'}); mixed_sheet.insert_chart('D2',mixed_chart); mixed_book.close()
    mixed_stale=root/'mixed-formula-stale.xlsx'
    with zipfile.ZipFile(mixed_source) as before,zipfile.ZipFile(mixed_stale,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='xl/charts/chart1.xml':
                tree=etree.fromstring(data); points=tree.findall('.//'+C+'numRef/'+C+'numCache/'+C+'pt'); next(p for p in points if p.get('idx')=='1').find(C+'v').text='9'; data=etree.tostring(tree)
            after.writestr(entry,data)
    mixed_scan=chart_cache.scan(mixed_stale)
    assert not mixed_scan['complete'] and mixed_scan['mismatch_references']==1 and mixed_scan['unverified_references']==1, mixed_scan
    formula_source=root/'formula.xlsx'; book=xlsxwriter.Workbook(str(formula_source)); sheet=book.add_worksheet('Calc'); sheet.write_formula('A1','=5+5',None,10); sheet.write_number('A2',20)
    chart=book.add_chart({'type':'column'}); chart.add_series({'values':'=Calc!$A$1:$A$2'}); sheet.insert_chart('D2',chart); book.close()
    stale=root/'formula-stale.xlsx'
    with zipfile.ZipFile(formula_source) as before,zipfile.ZipFile(stale,'w') as after:
        for entry in before.infolist():
            data=before.read(entry)
            if entry.filename=='xl/charts/chart1.xml':
                tree=etree.fromstring(data); tree.find('.//'+C+'numRef/'+C+'numCache/'+C+'pt/'+C+'v').text='7'; data=etree.tostring(tree,encoding='utf-8',xml_declaration=True)
            after.writestr(entry,data)
    formula_draft=root/'formula-draft'; session.start(formula_draft,source=stale); view=session.scan_chart_caches(formula_draft,0); ref=next(x for x in view['charts'][0]['references'] if x['status']=='mismatch')
    try: session.sync_chart_cache(formula_draft,0,view['charts'][0]['part'],ref['index'])
    except ValueError as exc: assert 'formula' in str(exc).lower()
    else: raise AssertionError('Formula cache sync without opt-in accepted')
    formula_result=session.sync_chart_cache(formula_draft,0,view['charts'][0]['part'],ref['index'],allow_formula_caches=True)
    assert formula_result['source_formula_cells']==1 and formula_result['updated_points']==1 and session.scan_chart_caches(formula_draft,1)['mismatch_references']==0
print('Chart cache smoke passed: mismatch, baseline, revision, unsupported sources and review bundle')

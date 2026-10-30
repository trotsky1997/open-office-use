"""Exercise tracked insertion and accepting or rejecting its revision."""
import tempfile
import json
import subprocess
import sys
import zipfile
from pathlib import Path
from docx import Document
from lxml import etree
import session
import word
import officecli_backend

with tempfile.TemporaryDirectory() as temp:
    root = Path(temp); source = root / 'source.docx'; doc = Document()
    p = doc.add_paragraph('Start '); p.add_run('an').bold = True; p.add_run('chor').italic = True; p.add_run(' end'); doc.save(source)
    original = source.read_bytes()
    cli_draft = root / 'cli'; session.start(cli_draft, source=source)
    cli = Path(__file__).with_name('open_office_use.py')
    output = subprocess.run([sys.executable, str(cli), 'session', 'track-insert-text', str(cli_draft), '--base', '0', '--path', '/body/p[1]', '--anchor', 'anchor', '--text', 'new ', '--position', 'before', '--author', 'Editor'], capture_output=True, text=True, check=True)
    assert json.loads(output.stdout)['id'] == session.tracked_changes(cli_draft, 1)['records'][0]['id']
    for position, insertion, expected, style in [('before', 'new ', 'Start new anchor end', 'b'), ('after', ' new', 'Start anchor new end', 'i')]:
        draft = root / position; session.start(draft, source=source)
        for args in [('anchor', '', position, 'Editor'), ('anchor', 'x', 'middle', 'Editor'), ('missing', 'x', position, 'Editor')]:
            try: session.track_insert_text(draft, 0, '/body/p[1]', *args)
            except ValueError: pass
            else: raise AssertionError('Invalid insertion accepted')
        assert session.info(draft)['latest_revision'] == 0
        result = session.track_insert_text(draft, 0, '/body/p[1]', 'anchor', insertion, position, 'Editor')
        changes = session.tracked_changes(draft, 1)['records']; assert len(changes) == 1
        assert (changes[0]['id'], changes[0]['kind'], changes[0]['text'], changes[0]['author']) == (result['id'], 'ins', insertion, 'Editor')
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(result['file']) as after:
            assert all(before.read(name) == after.read(name) for name in before.namelist() if name != 'word/document.xml')
            tree = etree.fromstring(after.read('word/document.xml')); ins = tree.find('.//' + word.W + 'ins')
            assert ins.find(word.W + 'r/' + word.W + 'rPr/' + word.W + style) is not None
        audit = officecli_backend.audit(result['file']); assert audit['validation']['exit_code'] == 0 and audit['validation']['result']['success'], audit
        for accept, wanted in [(True, expected), (False, 'Start anchor end')]:
            copy = root / (position + str(accept)); session.start(copy, source=result['file'])
            resolved = session.resolve_tracked_runs(copy, 0, 'word/document.xml', [result['id']], accept)
            assert Document(resolved['file']).paragraphs[0].text == wanted
        assert source.read_bytes() == original
print('tracked insertion smoke passed')

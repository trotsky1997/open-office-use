"""Reorder slides while retaining each slide's parts and notes."""
import tempfile
import zipfile
import subprocess
import sys
from lxml import etree
from pathlib import Path
from pptx import Presentation
import officecli_backend
import open_office_use
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root/'deck.pptx'; draft = root/'draft'
    open_office_use.author({'slides':[{'title':'First','notes':'Notes one'}, {'title':'Second','notes':'Notes two'}, {'title':'Third','notes':'Notes three'}]},source)
    session.start(draft, source=source)
    result = session.reorder_slides(draft, 0, [3,1,2]); result_file = Path(result['file'])
    deck = Presentation(result_file)
    assert [slide.shapes.title.text for slide in deck.slides] == ['Third','First','Second']
    assert [slide.notes_slide.notes_text_frame.text for slide in deck.slides] == ['Notes three','Notes one','Notes two']
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(result_file) as changed:
        assert original.namelist() == changed.namelist()
        assert all(original.read(name) == changed.read(name) for name in original.namelist() if name != 'ppt/presentation.xml')
    code, validation, _ = officecli_backend.invoke('validate', result_file); assert code == 0 and validation['success']
    pdf = open_office_use.render(result_file, root/'render')
    rendered = subprocess.run(['pdftotext', '-layout', str(pdf), '-'], capture_output=True, text=True, check=True).stdout.split('\f')
    assert ['Third','First','Second'] == [next(name for name in ('First','Second','Third') if name in page) for page in rendered if page.strip()]
    for order in ([1,1,2], [1,2], [True,2,3]):
        try: session.reorder_slides(draft, 1, order)
        except ValueError: pass
        else: raise AssertionError('Invalid order accepted')
    try: session.reorder_slides(draft, 0, [1,2,3])
    except ValueError as exc: assert 'conflict' in str(exc)
    else: raise AssertionError('Stale revision accepted')
    assert session.info(draft)['latest_revision'] == 1
    cli = subprocess.run([sys.executable, str(Path(open_office_use.__file__)), 'session', 'reorder-slides', str(draft), '--base', '1', '--order', '[2,3,1]'], capture_output=True, text=True)
    assert cli.returncode == 0, cli.stderr
    assert [slide.shapes.title.text for slide in Presentation(draft/'revision-000002.pptx').slides] == ['First','Second','Third']
    extended = root/'extended.pptx'
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(extended, 'w') as modified:
        for item in original.infolist():
            data = original.read(item)
            if item.filename == 'ppt/presentation.xml':
                tree = etree.fromstring(data); etree.SubElement(tree, '{http://schemas.openxmlformats.org/presentationml/2006/main}extLst'); data = etree.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True)
            modified.writestr(item, data)
    extended_draft = root/'extended-draft'; session.start(extended_draft, source=extended)
    try: session.reorder_slides(extended_draft, 0, [3,1,2])
    except ValueError as exc: assert 'extensions' in str(exc)
    else: raise AssertionError('Presentation extensions accepted')
    assert session.info(extended_draft)['latest_revision'] == 0
    section_source = root/'section-source.pptx'
    open_office_use.author({'slides':[{'title':name} for name in ('One','Two','Three','Four')]}, section_source)
    section_file = root/'sections.pptx'
    p = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
    p14 = '{http://schemas.microsoft.com/office/powerpoint/2010/main}'
    with zipfile.ZipFile(section_source) as original, zipfile.ZipFile(section_file, 'w') as modified:
        for item in original.infolist():
            data = original.read(item)
            if item.filename == 'ppt/presentation.xml':
                tree = etree.fromstring(data); ids = [node.get('id') for node in tree.find(p+'sldIdLst')]
                extension = etree.SubElement(etree.SubElement(tree,p+'extLst'),p+'ext',uri='{521415D9-36F7-43E2-AB2F-B90AF26B5E84}')
                sections = etree.SubElement(extension,p14+'sectionLst',nsmap={'p14':p14[1:-1]})
                for index in range(2):
                    section = etree.SubElement(sections,p14+'section',name=f'Section {index+1}',id=f'{{00000000-0000-0000-0000-00000000000{index+1}}}')
                    members = etree.SubElement(section,p14+'sldIdLst')
                    for value in ids[index*2:index*2+2]: etree.SubElement(members,p14+'sldId',id=value)
                data = etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
            modified.writestr(item,data)
    section_draft = root/'sections-draft'; session.start(section_draft,source=section_file)
    try: session.reorder_slides(section_draft,0,[1,3,2,4])
    except ValueError as exc: assert 'contiguous' in str(exc)
    else: raise AssertionError('Interleaved sections accepted')
    assert session.info(section_draft)['latest_revision'] == 0
    section_result = session.reorder_slides(section_draft,0,[4,3,2,1])
    with zipfile.ZipFile(section_result['file']) as archive:
        presentation_xml = etree.fromstring(archive.read('ppt/presentation.xml'))
        section_nodes = presentation_xml.find(p+'extLst/'+p+'ext/'+p14+'sectionLst')
        assert [item.get('name') for item in section_nodes] == ['Section 2','Section 1']
        assert [[item.get('id') for item in section[0]] for section in section_nodes] == [ids[2:4][::-1],ids[0:2][::-1]]
    assert [slide.shapes.title.text for slide in Presentation(section_result['file']).slides] == ['Four','Three','Two','One']
    code, validation, _ = officecli_backend.invoke('validate',section_result['file']); assert code == 0 and validation['success'], validation
    section_pdf = open_office_use.render(section_result['file'],root/'section-render')
    pages = subprocess.run(['pdftotext','-layout',str(section_pdf),'-'],capture_output=True,text=True,check=True).stdout.split('\f')
    assert [next(name for name in ('One','Two','Three','Four') if name in page) for page in pages if page.strip()] == ['Four','Three','Two','One']
print('Slide reorder passed: notes, package preservation, validation and conflicts')

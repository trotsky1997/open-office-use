"""Verify merge and unmerge preserve stored workbook data."""
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.formula import ArrayFormula
import officecli_backend
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.xlsx'; book = Workbook(); sheet = book.active
    sheet['A1'] = 'Header'; sheet['A1'].font = Font(bold=True); sheet['B1'].fill = PatternFill('solid', fgColor='FFCC00')
    sheet['C3'] = '=1+1'; sheet['A5'] = 'Occupied'; book.create_sheet('Other')['A1'] = 42; book.save(source)
    patched = root / 'patched.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(patched, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                tree = etree.fromstring(data); ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                tree.xpath('./x:sheetData/x:row/x:c[@r="C3"]/x:v', namespaces=ns)[0].text = '2'; data = etree.tostring(tree)
            after.writestr(entry, data)
    patched.replace(source); draft = root / 'draft'; session.start(draft, source=source)
    merged = session.merge_cells(draft, 0, 'Sheet', 'A1:B2')
    assert merged['merged'] and load_workbook(merged['file']).active.merged_cells.ranges
    grid = session.read_grid(draft, 1, 'Sheet', 'A1:B2'); assert grid['rows'][0][1]['read_only_reason'] == 'merged_placeholder'
    assert load_workbook(merged['file'], data_only=True).active['C3'].value == 2
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(merged['file']) as after:
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name != 'xl/worksheets/sheet1.xml')
    assert officecli_backend.audit(merged['file'])['validation']['exit_code'] == 0
    unmerged = session.merge_cells(draft, 1, 'Sheet', 'A1:B2', False)
    workbook = load_workbook(unmerged['file']); assert not workbook.active.merged_cells.ranges
    assert workbook.active['A1'].value == 'Header' and workbook.active['A1'].font.bold and workbook.active['B1'].fill.fgColor.rgb[-6:] == 'FFCC00'
    assert load_workbook(unmerged['file'], data_only=True).active['C3'].value == 2
    for selected, operation in [('A4:A5', True), ('A1', True), ('A1:C2', False)]:
        try: session.merge_cells(draft, 2, 'Sheet', selected, operation)
        except ValueError: pass
        else: raise AssertionError(f'Invalid merge accepted: {selected}')
    cli = Path(__file__).with_name('open_office_use.py')
    result = subprocess.run([sys.executable, str(cli), 'session', 'merge-cells', str(draft), '--base', '2', '--sheet', 'Sheet', '--range', 'E1:F1'], capture_output=True, text=True, check=True)
    assert json.loads(result.stdout)['merged'] and 'E1:F1' in str(load_workbook(draft / 'revision-000003.xlsx').active.merged_cells)
    for kind in ('hyperlink', 'comment', 'table', 'array', 'overlap'):
        file = root / f'{kind}.xlsx'; book = Workbook(); sheet = book.active
        if kind == 'hyperlink': sheet['B1'].hyperlink = 'https://example.com'
        if kind == 'comment': sheet['C3'].comment = Comment('note', 'Reviewer')
        if kind == 'table':
            sheet['A1'] = 'Column'; sheet['A2'] = 1; sheet.add_table(Table(displayName='Example', ref='A1:A2'))
        if kind == 'array': sheet['A1'] = ArrayFormula(ref='A1:B2', text='=1')
        if kind == 'overlap': sheet.merge_cells('A1:C1')
        book.save(file); trial = root / f'{kind}-draft'; session.start(trial, source=file)
        try: session.merge_cells(trial, 0, 'Sheet', 'A1:B1')
        except ValueError: pass
        else: raise AssertionError(f'{kind} merge accepted')
        assert session.info(trial)['latest_revision'] == 0
    print('Merge cells smoke passed: merge, unmerge, formatting, caches and content protection')

"""Compare saved Office files using bounded text views and OOXML parts."""
import difflib
import hashlib
import json
import shutil
import tempfile
import zipfile
from pathlib import Path

from openpyxl.utils.cell import get_column_letter, range_boundaries

import officecli_backend
import spreadsheet
import notes


def part_hashes(file):
    with zipfile.ZipFile(file) as archive:
        parts = [entry for entry in archive.infolist() if not entry.is_dir()]
        names = [entry.filename for entry in parts]
        if len(parts) > 5000 or sum(entry.file_size for entry in parts) > 128 * 1024 * 1024:
            raise ValueError('Office package exceeds the review limit of 5,000 parts or 128 MiB uncompressed')
        if len(names) != len(set(names)):
            raise ValueError('Office package contains duplicate part names')
        hashes = {}
        for entry in parts:
            digest = hashlib.sha256()
            with archive.open(entry) as part:
                for chunk in iter(lambda: part.read(1024 * 1024), b''):
                    digest.update(chunk)
            hashes[entry.filename] = digest.hexdigest()
        return hashes


def compare(before, after):
    before, after = Path(before).resolve(strict=True), Path(after).resolve(strict=True)
    suffix = before.suffix.lower()
    if suffix not in ('.docx', '.xlsx', '.pptx') or after.suffix.lower() != suffix:
        raise ValueError('Both files must have the same supported Office format')
    with tempfile.TemporaryDirectory(prefix='office-review-') as tmp:
        left, right = Path(tmp) / ('before' + suffix), Path(tmp) / ('after' + suffix)
        shutil.copyfile(before, left)
        shutil.copyfile(after, right)
        try:
            left_parts, right_parts = part_hashes(left), part_hashes(right)
        except zipfile.BadZipFile as exc:
            raise ValueError('Invalid Office package') from exc
        before_view = officecli_backend.call('view', left, 'text', '--max-lines', '2000')
        after_view = officecli_backend.call('view', right, 'text', '--max-lines', '2000')
    before_text = json.dumps(before_view, ensure_ascii=False, sort_keys=True, indent=2)
    after_text = json.dumps(after_view, ensure_ascii=False, sort_keys=True, indent=2)
    # shortcut: views beyond 100,000 characters are clipped, use range-specific comparison for large files.
    view_clipped = len(before_text) > 100000 or len(after_text) > 100000
    lines = difflib.unified_diff(before_text[:100000].splitlines(), after_text[:100000].splitlines(),
                                 fromfile='before', tofile='after', lineterm='')
    diff = []
    for line in lines:
        if len(diff) == 500:
            break
        diff.append(line)
    else:
        return {'parts': changes(left_parts, right_parts), 'text_diff': diff,
                'text_diff_truncated': view_clipped, 'text_view_max_lines': 2000,
                'text_view_may_be_truncated': True}
    return {'parts': changes(left_parts, right_parts), 'text_diff': diff,
            'text_diff_truncated': True, 'text_view_max_lines': 2000,
            'text_view_may_be_truncated': True}


def changes(before, after):
    return {'added': sorted(after.keys() - before.keys()),
            'removed': sorted(before.keys() - after.keys()),
            'modified': sorted(name for name in before.keys() & after.keys() if before[name] != after[name])}


def compare_notes(before, after, before_slide, after_slide):
    """Compare speaker notes on two explicitly corresponding slides."""
    before, after = Path(before).resolve(strict=True), Path(after).resolve(strict=True)
    if before.suffix.lower() != '.pptx' or after.suffix.lower() != '.pptx':
        raise ValueError('Notes comparison requires two PPTX files')
    if any(file.stat().st_size > 128 * 1024 * 1024 for file in (before, after)):
        raise ValueError('Document exceeds the 128 MiB size limit')
    with tempfile.TemporaryDirectory(prefix='office-notes-review-') as tmp:
        left, right = Path(tmp) / 'before.pptx', Path(tmp) / 'after.pptx'
        shutil.copyfile(before, left); shutil.copyfile(after, right)
        old, new = notes.read(left, before_slide), notes.read(right, after_slide)
    text_changed = old['text'] != new['text'] or (None if old['truncated'] or new['truncated'] else False)
    old_lines, new_lines = old['text'].split('\n'), new['text'].split('\n')
    # shortcut: diff at most 2,000 lines and 500 output lines; use individual notes reads for larger changes.
    diff = difflib.unified_diff(old_lines[:2000], new_lines[:2000], fromfile=f'before slide {before_slide}', tofile=f'after slide {after_slide}', lineterm='')
    lines = []
    for line in diff:
        if len(lines) == 500: break
        lines.append(line)
    else:
        return {'before_slide': before_slide, 'after_slide': after_slide, 'before_has_notes': old['has_notes_slide'], 'after_has_notes': new['has_notes_slide'],
                'text_changed': text_changed, 'diff': lines, 'truncated': old['truncated'] or new['truncated'] or len(old_lines) > 2000 or len(new_lines) > 2000}
    return {'before_slide': before_slide, 'after_slide': after_slide, 'before_has_notes': old['has_notes_slide'], 'after_has_notes': new['has_notes_slide'],
            'text_changed': text_changed, 'diff': lines, 'truncated': True}



def compare_dimensions(before, after, sheet, cell_range):
    """Compare saved row, column and sheet-default attributes in a bounded range."""
    bounds = range_boundaries(cell_range)
    old = spreadsheet.worksheet_constraints(before, sheet, bounds, True, True)['dimensions']
    new = spreadsheet.worksheet_constraints(after, sheet, bounds, True, True)['dimensions']
    changes = []; count = 0
    def record(kind, target, left, right):
        nonlocal count
        fields = sorted(key for key in left.keys() | right.keys() if left.get(key) != right.get(key))
        if not fields: return
        count += 1
        if len(changes) < 500: changes.append({'kind': kind, 'target': target, 'changed_fields': fields, 'before': left, 'after': right})
    record('defaults', sheet, old['defaults'], new['defaults'])
    c1, r1, c2, r2 = bounds
    for col in range(c1, c2 + 1):
        key = get_column_letter(col)
        left = {k: v for k, v in old['columns'].get(key, {}).items() if k not in ('min', 'max')}
        right = {k: v for k, v in new['columns'].get(key, {}).items() if k not in ('min', 'max')}
        record('column', key, left, right)
    for row in range(r1, r2 + 1):
        key = str(row); record('row', key, old['rows'].get(key, {}), new['rows'].get(key, {}))
    return {'status': 'completed', 'changed_count': count, 'changes': changes, 'truncated': count > len(changes),
            'complete': not old['issues'] and not new['issues'], 'before_issues': old['issues'], 'after_issues': new['issues'],
            'notes': ['Comparison uses stored XML attributes; equivalent numeric encodings and explicit defaults can be reported as changes.',
                      'Column span boundaries are excluded; row and column changes affect the entire worksheet dimension.']}

def compare_cells(before, after, sheet, cell_range):
    """Compare formulas, values and direct cell formatting in a bounded XLSX range."""
    before, after = Path(before).resolve(strict=True), Path(after).resolve(strict=True)
    if before.suffix.lower() != '.xlsx' or after.suffix.lower() != '.xlsx':
        raise ValueError('Cell comparison requires two .xlsx files')
    with tempfile.TemporaryDirectory(prefix='office-cell-review-') as tmp:
        left, right = Path(tmp) / 'before.xlsx', Path(tmp) / 'after.xlsx'
        shutil.copyfile(before, left)
        shutil.copyfile(after, right)
        left_rows = spreadsheet.read_range(left, sheet, cell_range, include_style='review')['rows']
        right_rows = spreadsheet.read_range(right, sheet, cell_range, include_style='review')['rows']
        try: dimensions = compare_dimensions(left, right, sheet, cell_range)
        except Exception as exc:
            dimensions = {'status': 'error', 'complete': False, 'error_type': type(exc).__name__, 'error': str(exc)[:1000]}
    min_col, min_row, _, _ = range_boundaries(cell_range)
    if min_col is None or min_row is None:
        raise ValueError('A bounded cell range is required')
    text_comparison_complete = all(cell.get('text_complete') is not False for rows in (left_rows, right_rows) for row in rows for cell in row)
    result = []
    total = 0
    for row_index, (left_row, right_row) in enumerate(zip(left_rows, right_rows)):
        for col_index, (old, new) in enumerate(zip(left_row, right_row)):
            fields = [field for field in ('formula', 'value', 'type') if old.get(field) != new.get(field)]
            fields.extend('format.' + key for key in old['format'] if old['format'][key] != new['format'][key])
            if fields:
                total += 1
                if len(result) < 500:
                    address = f'{get_column_letter(min_col + col_index)}{min_row + row_index}'
                    result.append({'cell': address, 'changed_fields': fields, 'before': old, 'after': new})
    return {'sheet': sheet, 'range': cell_range, 'changed_count': total, 'changes': result, 'text_comparison_complete': text_comparison_complete,
            'truncated': total > len(result), 'cache_freshness': 'unknown', 'dimensions': dimensions}

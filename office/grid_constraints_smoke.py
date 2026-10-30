"""Check merged and protected worksheet behavior in editing workflows."""
import tempfile
import zipfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection
import cell_text
import session

with tempfile.TemporaryDirectory() as tmp:
    root=Path(tmp); source=root/'merged.xlsx'; draft=root/'draft'; wb=Workbook(); sheet=wb.active
    sheet['A1']='Label'; sheet.merge_cells('A1:B2'); sheet['C1']='Outside'; wb.save(source); session.start(draft,source=source)
    grid=session.read_grid(draft,0,'Sheet','A1:C2'); cells={c['cell']:c for row in grid['rows'] for c in row}
    assert cells['A1']['editable'] and cells['A1']['merged_anchor']=='A1'
    assert not cells['B1']['editable'] and cells['B1']['read_only_reason']=='merged_placeholder' and cells['B1']['merged_anchor']=='A1'
    assert session.read_grid(draft,0,'Sheet','B2')['rows'][0][0]['merged_anchor']=='A1'
    for fn in (lambda:session.edit_grid(draft,0,'Sheet','A1:C2',[{'cell':'C1','value':'new'},{'cell':'B1','value':'hidden'}]),
               lambda:session.format_grid(draft,0,'Sheet','A1:C2','B1',{'bold':True}),
               lambda:session.edit(draft,0,'/Sheet/B1','value','hidden')):
        try:fn()
        except (ValueError,RuntimeError):pass
        else:raise AssertionError('Merged placeholder accepted')
    assert session.info(draft)['latest_revision']==0
    out=session.edit_grid(draft,0,'Sheet','A1:C2',[{'cell':'A1','value':'New title'}])
    assert load_workbook(out['file']).active['A1'].value=='New title'
    protected=root/'protected.xlsx'; protected_draft=root/'protected-draft'; wb=Workbook(); sheet=wb.active
    sheet['A1']='Locked'; sheet['B1']='Open'; sheet['B1'].protection=Protection(locked=False); sheet.protection.sheet=True; wb.save(protected)
    session.start(protected_draft,source=protected); grid=session.read_grid(protected_draft,0,'Sheet','A1:C1')
    assert grid['sheet_protected'] and not grid['rows'][0][0]['editable'] and grid['rows'][0][1]['editable'] and not grid['rows'][0][2]['editable']
    target=session.edit_target(protected_draft,0,'/Sheet/A1'); assert target['read_only_reason']=='locked_cell'
    for fn in (lambda:session.edit(protected_draft,0,'/Sheet/A1','value','no'),
               lambda:session.replace_text(protected_draft,0,'/Sheet/A1','Locked','no'),
               lambda:session.edit_grid(protected_draft,0,'Sheet','A1:B1',[{'cell':'A1','value':'no'},{'cell':'B1','value':'no'}]),
               lambda:session.format_grid(protected_draft,0,'Sheet','A1:B1','B1',{'bold':True})):
        try:fn()
        except ValueError:pass
        else:raise AssertionError('Protected edit accepted')
    assert session.info(protected_draft)['latest_revision']==0
    result=session.edit_grid(protected_draft,0,'Sheet','B1',[{'cell':'B1','value':'Allowed'}]); changed=load_workbook(result['file']).active
    assert changed['A1'].value=='Locked' and changed['B1'].value=='Allowed' and changed.protection.sheet
    session.replace_text(protected_draft,1,'/Sheet/B1','Allowed','OK')
    assert load_workbook(protected_draft/'revision-000002.xlsx').active['B1'].value=='OK'
    overlap=root/'overlap.xlsx'
    with zipfile.ZipFile(source) as a,zipfile.ZipFile(overlap,'w') as b:
        for entry in a.infolist():
            data=a.read(entry)
            if entry.filename=='xl/worksheets/sheet1.xml':data=data.replace(b'<mergeCell ref="A1:B2"/>',b'<mergeCell ref="A1:B2"/><mergeCell ref="A1:C1"/>')
            b.writestr(entry,data)
    overlap_draft=root/'overlap-draft';session.start(overlap_draft,source=overlap)
    assert session.read_grid(overlap_draft,0,'Sheet','A1')['rows'][0][0]['read_only_reason']=='overlapping_merges'
print('Grid constraints passed: merged cells, overlap, protection and unlocked cells')

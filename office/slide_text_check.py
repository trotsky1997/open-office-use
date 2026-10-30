"""Compare slide text boxes with text extracted from a rendered PDF."""
import hashlib
import math
import shutil
import subprocess
import tempfile
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import layout
import visual


def normalize(text):
    return ''.join(c for c in unicodedata.normalize('NFKC', text).casefold() if c.isalnum())



def validate_package(source):
    with zipfile.ZipFile(source) as archive:
        entries = archive.infolist(); names = [item.filename for item in entries]
        if len(entries) > 5000 or len(set(names)) != len(names) or sum(item.file_size for item in entries) > 128 * 1024 * 1024:
            raise ValueError('Presentation exceeds package limits or contains duplicate parts')


def table_cells(shape, prefix, matrix, width, height):
    table = shape.table; rows = len(table.rows); columns = len(table.columns)
    if not rows or not columns or rows * columns > 10000: raise ValueError('table_size')
    col_sizes = [column.width for column in table.columns]; row_sizes = [row.height for row in table.rows]
    if any(value is None or value <= 0 for value in col_sizes + row_sizes): raise ValueError('table_grid')
    x_edges = [shape.left]; y_edges = [shape.top]; total_width = sum(col_sizes); total_height = sum(row_sizes)
    for value in col_sizes: x_edges.append(x_edges[-1] + value * shape.width / total_width)
    for value in row_sizes: y_edges.append(y_edges[-1] + value * shape.height / total_height)
    covered = set(); entries = []
    for row in range(rows):
        for column in range(columns):
            cell = table.cell(row, column); path = f'{prefix}/tr[{row + 1}]/tc[{column + 1}]'
            if (row, column) in covered:
                if not cell.is_spanned or cell.is_merge_origin or cell.text.strip(): raise ValueError('unsupported_merged_cell')
                continue
            if cell.is_spanned: raise ValueError('invalid_merged_cell')
            row_span = cell.span_height if cell.is_merge_origin else 1
            col_span = cell.span_width if cell.is_merge_origin else 1
            if row_span < 1 or col_span < 1 or row + row_span > rows or column + col_span > columns:
                raise ValueError('invalid_merged_cell')
            for r in range(row, row + row_span):
                for c in range(column, column + col_span):
                    if (r, c) == (row, column): continue
                    if (r, c) in covered: raise ValueError('overlapping_merged_cells')
                    covered.add((r, c))
            x1, y1 = layout.point(matrix, x_edges[column], y_edges[row])
            x2, y2 = layout.point(matrix, x_edges[column + col_span], y_edges[row + row_span])
            bounds = [x1 / width, y1 / height, x2 / width, y2 / height]
            entries.append((path, cell.text, bounds))
    return entries, rows * columns

def inspect(source, pdf):
    source = Path(source)
    if source.suffix.lower() != '.pptx': raise ValueError('Slide text inspection requires PPTX')
    if source.stat().st_size > 128 * 1024 * 1024: raise ValueError('Presentation exceeds the 128 MiB limit')
    validate_package(source)
    presentation = Presentation(source); slides = presentation.slides
    width, height = presentation.slide_width, presentation.slide_height
    if not width or not height or width <= 0 or height <= 0: raise ValueError('Invalid slide size')
    page_count = visual.pdf_page_count(pdf)
    visibility = [slide._element.get('show') for slide in slides]
    reasons = []
    if any(value not in (None, '1', 'true') for value in visibility): reasons.append('hidden_or_unknown_slide_visibility')
    if page_count != len(slides): reasons.append('page_count_mismatch')
    if len(slides) > 100: reasons.append('slide_limit')
    result = {'slides_total': len(slides), 'pdf_pages': page_count, 'mapping_valid': not reasons,
              'mapping_issues': reasons, 'shapes_scanned': 0, 'text_shapes_checked': 0, 'table_cells_scanned': 0, 'table_cells_checked': 0,
              'skipped_count': 0, 'skipped': [], 'skipped_truncated': False, 'candidate_count': 0, 'candidates': [], 'candidates_truncated': False, 'complete': False,
              'notes': ['Candidates compare source text with extracted PDF text within the text box or table cell bounds; inspect the rendered page.',
                        'Font encoding, overlapping objects, line order, clipping and renderer differences affect results. Masters, layouts, rotated or flipped groups and OCR are outside this check.']}
    if reasons: return result
    executable = shutil.which('pdftotext')
    if not executable: raise RuntimeError('Poppler pdftotext is required for slide text inspection')
    with tempfile.TemporaryDirectory(prefix='office-slide-text-') as tmp:
        output = Path(tmp) / 'text.xhtml'
        command = subprocess.run([executable, '-bbox-layout', str(pdf), str(output)], capture_output=True, text=True, timeout=120)
        if command.returncode or not output.is_file(): raise RuntimeError(f'PDF text extraction failed: {command.stderr[:1000]}')
        if output.stat().st_size > 32 * 1024 * 1024: raise ValueError('PDF text extraction exceeds 32 MiB')
        try: pages = [node for node in ET.parse(output).getroot().iter() if node.tag.rsplit('}', 1)[-1] == 'page']
        except ET.ParseError as exc: raise ValueError('Invalid PDF text extraction') from exc
    if len(pages) != len(slides): raise ValueError('PDF text extraction has an unexpected page count')
    checks = 0
    def skip(slide, path, reason):
        result['skipped_count'] += 1
        if len(result['skipped']) < 200: result['skipped'].append({'slide': slide, 'path_hint': path, 'reason': reason})
    stop = False
    for index, (slide, page) in enumerate(zip(slides, pages), 1):
        try: pw, ph = float(page.attrib['width']), float(page.attrib['height'])
        except (ValueError, KeyError) as exc: raise ValueError('Invalid PDF page dimensions') from exc
        if not math.isfinite(pw) or not math.isfinite(ph) or pw <= 0 or ph <= 0 or abs(pw / ph - width / height) > 0.01:
            result['mapping_valid'] = False; result['mapping_issues'].append('page_aspect_mismatch'); result['candidates'] = []; result['candidate_count'] = 0; result['text_shapes_checked'] = 0; result['table_cells_checked'] = 0; return result
        words = []
        for node in page.iter():
            if node.tag.rsplit('}', 1)[-1] != 'word': continue
            try: box = [float(node.attrib[name]) for name in ('xMin', 'yMin', 'xMax', 'yMax')]
            except (KeyError, ValueError) as exc: raise ValueError('Invalid PDF word bounds') from exc
            if not all(math.isfinite(value) for value in box): raise ValueError('Invalid PDF word bounds')
            words.append(((box[0] + box[2]) / (2 * pw), (box[1] + box[3]) / (2 * ph), ''.join(node.itertext())))
        def walk(shapes, matrix, prefix, depth):
            nonlocal checks, stop
            for shape in shapes:
                if result['shapes_scanned'] >= 10000: skip(index, None, 'shape_limit'); stop = True; break
                result['shapes_scanned'] += 1
                metadata = next(shape._element.iter(layout.P + 'cNvPr'), None)
                if metadata is not None and metadata.get('hidden') in ('1', 'true'): continue
                path = f'{prefix}/shape[@id={shape.shape_id}]'
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    path = f'{prefix}/group[@id={shape.shape_id}]'
                    if depth >= 8: skip(index, path, 'group_depth'); continue
                    try:
                        transform = shape._element.find(layout.P + 'grpSpPr/' + layout.A + 'xfrm')
                        if transform is None or any(transform.get(attr) not in (None, '0', 'false', '1', 'true') for attr in ('flipH', 'flipV')):
                            raise ValueError('invalid_group_transform')
                        child_matrix = layout.group_transform(shape, matrix)
                        if not all(math.isfinite(v) for v in child_matrix) or child_matrix[0] <= 0 or child_matrix[3] <= 0 or abs(child_matrix[1]) > 1e-9 or abs(child_matrix[2]) > 1e-9:
                            raise ValueError('rotated_or_flipped_group')
                        walk(shape.shapes, child_matrix, path, depth + 1)
                        if stop: return
                    except (ValueError, TypeError, AttributeError, OverflowError) as exc:
                        skip(index, path, str(exc)[:100] or 'invalid_group')
                    continue
                is_table = shape.has_table
                if is_table: path = f'{prefix}/table[@id={shape.shape_id}]'
                if not is_table and (not shape.has_text_frame or not shape.text_frame.text.strip()): continue
                if shape.rotation or shape.left is None or shape.top is None or shape.width is None or shape.height is None or shape.width <= 0 or shape.height <= 0:
                    skip(index, path, 'rotated_or_empty_bounds'); continue
                if is_table:
                    transform = shape._element.find(layout.P + 'xfrm')
                    if transform is not None and any(transform.get(attr) not in (None, '0', 'false') for attr in ('flipH', 'flipV')):
                        skip(index, path, 'flipped_table'); continue
                    try: entries, cell_count = table_cells(shape, path, matrix, width, height)
                    except (ValueError, TypeError, IndexError, AttributeError, OverflowError) as exc:
                        skip(index, path, str(exc)[:100] or 'invalid_table'); continue
                    if result['table_cells_scanned'] + cell_count > 10000:
                        skip(index, path, 'table_cell_limit'); stop = True; break
                    result['table_cells_scanned'] += cell_count
                else:
                    x1, y1 = layout.point(matrix, shape.left, shape.top)
                    x2, y2 = layout.point(matrix, shape.left + shape.width, shape.top + shape.height)
                    x1, y1, x2, y2 = x1 / width, y1 / height, x2 / width, y2 / height
                    entries = [(path, shape.text_frame.text, [x1, y1, x2, y2])]
                for path, text, bounds in entries:
                    if not text.strip(): continue
                    if len(text) > 10000: skip(index, path, 'text_limit'); continue
                    expected = normalize(text)
                    if not expected: skip(index, path, 'no_comparable_characters'); continue
                    if checks + len(words) > 2000000: skip(index, path, 'comparison_limit'); stop = True; break
                    checks += len(words)
                    x1, y1, x2, y2 = bounds
                    extracted = normalize(' '.join(text for x, y, text in words if x1 <= x <= x2 and y1 <= y <= y2))
                    result['table_cells_checked' if is_table else 'text_shapes_checked'] += 1
                    if expected not in extracted:
                        result['candidate_count'] += 1
                        if len(result['candidates']) < 200:
                            result['candidates'].append({'slide': index, 'path_hint': path, 'name': str(shape.name or '')[:200], 'source_text': text[:200], 'source_text_truncated': len(text) > 200,
                                                         'bounds': bounds, 'extracted_text': extracted[:200]})
                if stop: return
        walk(slide.shapes, (1, 0, 0, 1, 0, 0), f'/slide[{index}]', 0)
        if stop: break
    result['complete'] = not result['skipped_count'] and result['mapping_valid']
    result['candidates_truncated'] = result['candidate_count'] > len(result['candidates'])
    result['skipped_truncated'] = result['skipped_count'] > len(result['skipped'])
    return result


def scan(source, render):
    source = Path(source).resolve(strict=True)
    if source.suffix.lower() != '.pptx' or source.stat().st_size > 128 * 1024 * 1024: raise ValueError('Provide a PPTX within the 128 MiB limit')
    with tempfile.TemporaryDirectory(prefix='office-slide-text-snapshot-') as tmp:
        snapshot = Path(tmp) / 'snapshot.pptx'; shutil.copyfile(source, snapshot)
        if snapshot.stat().st_size > 128 * 1024 * 1024: raise ValueError('Presentation exceeds the 128 MiB limit')
        digest = hashlib.sha256()
        with snapshot.open('rb') as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b''): digest.update(chunk)
        digest = digest.hexdigest()
        validate_package(snapshot)
        pdf = render(snapshot, Path(tmp) / 'render')
        return {'snapshot_sha256': digest, **inspect(snapshot, pdf)}

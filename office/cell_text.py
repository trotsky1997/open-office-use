"""Version-safe replacement inside an XLSX text cell."""
import copy
import re
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, range_boundaries

import ooxml
import spreadsheet

X = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
CT = '{http://schemas.openxmlformats.org/package/2006/content-types}'
SHARED_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml'
PATH = re.compile(r'/([^/\x00]{1,31})/([A-Za-z]{1,3})([1-9][0-9]{0,6})\Z')
ESCAPED = re.compile(r'_x[0-9a-fA-F]{4}_')


def cell_path(path):
    match = PATH.fullmatch(path) if isinstance(path, str) and len(path) <= 512 else None
    return bool(match and column_index_from_string(match.group(2).upper()) <= 16384 and int(match.group(3)) <= 1048576)


def content_key(item):
    normalized = copy.deepcopy(item)
    space = '{http://www.w3.org/XML/1998/namespace}space'
    for node in normalized.iter(X + 't'):
        value = node.text or ''
        if node.get(space) == 'preserve' and not any(char.isspace() for char in value):
            del node.attrib[space]
    return etree.tostring(normalized, method='c14n', exclusive=True, with_tail=False)


def request_recalculation(root):
    sheet_data = root.findall(X + 'sheetData'); calc = root.findall(X + 'sheetCalcPr')
    if len(sheet_data) != 1 or len(list(root.iter(X + 'sheetData'))) != 1 or len(calc) > 1 or len(list(root.iter(X + 'sheetCalcPr'))) != len(calc):
        raise ValueError('Worksheet calculation metadata is ambiguous')
    if calc:
        marker = calc[0]
        if root.index(marker) <= root.index(sheet_data[0]) or any(isinstance(item.tag, str) for item in root[root.index(sheet_data[0]) + 1:root.index(marker)]):
            raise ValueError('Worksheet calculation metadata is out of order')
        if len(marker): raise ValueError('Unsupported worksheet calculation metadata')
        if marker.get('fullCalcOnLoad') not in (None, '0', '1', 'false', 'true'): raise ValueError('Invalid worksheet calculation flag')
    else:
        marker = etree.Element(X + 'sheetCalcPr'); sheet_data[0].addnext(marker)
    marker.set('fullCalcOnLoad', '1')


def replace_text(source, output, path, find, replacement):
    if not cell_path(path):
        raise ValueError('Target must be a worksheet cell path such as /Sheet/A1')
    if Path(source).suffix.lower() != '.xlsx' or Path(output).suffix.lower() != '.xlsx':
        raise ValueError('Text replacement requires XLSX input and output')
    if not isinstance(find, str) or not isinstance(replacement, str) or ESCAPED.search(find) or ESCAPED.search(replacement):
        raise ValueError('Excel escaped text sequences are not supported')
    match = PATH.fullmatch(path); sheet = match.group(1); col = column_index_from_string(match.group(2).upper()); row = int(match.group(3))
    address = match.group(2).upper() + str(row)

    def select_part(file):
        workbook = load_workbook(file, read_only=True, data_only=False, keep_links=False)
        try:
            if sheet not in workbook.sheetnames:
                raise ValueError(f'Worksheet not found: {sheet}')
            part = getattr(workbook[sheet], '_worksheet_path', None)
            if not isinstance(part, str):
                raise ValueError('Cannot locate worksheet XML')
            if sum(getattr(item, '_worksheet_path', None) == part for item in workbook.worksheets) != 1:
                raise ValueError('Worksheet part is referenced by multiple sheets')
        finally:
            workbook.close()
        constraints = spreadsheet.worksheet_constraints(file, sheet)
        locked = spreadsheet.read_range(file, sheet, address, include_style=True)['rows'][0][0]['locked'] if constraints['sheet_protected'] else False
        reason, _ = spreadsheet.cell_constraint(constraints, col, row, locked)
        if reason == 'formula_range': raise ValueError('Array, data-table or shared formula ranges cannot be edited')
        if reason in ('merged_placeholder', 'overlapping_merges'): raise ValueError('Merged-cell placeholders or overlapping merges cannot be edited')
        if reason == 'locked_cell': raise ValueError('Locked cells on protected worksheets cannot be edited')
        return part

    def edit(tree, archive):
        root = tree.getroot()
        if root.tag != X + 'worksheet':
            raise ValueError('Worksheet XML has an invalid root')
        for merge in root.findall(X + 'mergeCells/' + X + 'mergeCell'):
            try:
                bounds = range_boundaries(merge.get('ref'))
                if any(value is None for value in bounds):
                    raise ValueError
                c1, r1, c2, r2 = bounds
            except (TypeError, ValueError):
                raise ValueError('Invalid merged-cell range') from None
            if c1 <= col <= c2 and r1 <= row <= r2 and (col, row) != (c1, r1):
                raise ValueError('Merged-cell placeholders cannot be edited; use the merged cell origin')
        cells = [cell for data in root.findall(X + 'sheetData') for row_node in data.findall(X + 'row')
                 for cell in row_node.findall(X + 'c') if cell.get('r', '').upper() == address]
        if len(cells) != 1:
            raise ValueError('Target cell must exist exactly once')
        request_recalculation(root)
        cell = cells[0]
        if cell.get('cm') is not None or cell.get('vm') is not None or cell.get('ph') not in (None, '0', 'false'):
            raise ValueError('Cell metadata or phonetic content is not supported')
        kind = cell.get('t'); additional = {}
        if kind == 'str':
            if len(cell) != 1 or cell[0].tag != X + 'v' or len(cell[0]) or any(key != '{http://www.w3.org/XML/1998/namespace}space' for key in cell[0].attrib):
                raise ValueError('String cell has unsupported content or a formula')
            node = cell[0]; value = node.text or ''
            if ESCAPED.search(value): raise ValueError('Excel escaped text sequences are not supported')
            start = ooxml.match_index(value, find, replacement)
            value = value[:start] + replacement + value[start + len(find):]
            if ESCAPED.search(value): raise ValueError('Replacement creates an unsupported Excel escaped text sequence')
            if len(value.encode('utf-16-le')) // 2 > 32767: raise ValueError('Cell text exceeds Excel’s 32,767 UTF-16-unit limit')
            node.text = value
            return additional
        elif kind == 'inlineStr':
            if len(cell) != 1 or cell[0].tag != X + 'is':
                raise ValueError('Inline text cell has unsupported content')
            content = cell[0]
        elif kind == 's':
            if len(cell) != 1 or cell[0].tag != X + 'v' or not re.fullmatch(r'[0-9]{1,10}', cell[0].text or ''):
                raise ValueError('Shared text cell has unsupported content')
            manifest = ooxml.parse_xml(archive.read('[Content_Types].xml')).getroot()
            if manifest.tag != CT + 'Types':
                raise ValueError('Content-type manifest has an invalid root')
            declarations = [node.get('PartName', '') for node in manifest.findall(CT + 'Override') if node.get('ContentType') == SHARED_CONTENT_TYPE]
            if len(declarations) != 1 or not declarations[0].startswith('/'):
                raise ValueError('Shared string table must have one content-type declaration')
            shared_path = declarations[0][1:]; shared_xml = archive.read(shared_path); shared_tree = ooxml.parse_xml(shared_xml); root_sst = shared_tree.getroot()
            if root_sst.tag != X + 'sst' or any(child.tag != X + 'si' for child in root_sst):
                raise ValueError('Shared string table contains unsupported content')
            entries = root_sst.findall(X + 'si'); index = int(cell[0].text)
            if index >= len(entries):
                raise ValueError('Shared string index does not exist')
            content = copy.deepcopy(entries[index])
        else:
            raise ValueError('Target must be a text cell with string, inline or shared storage')
        nodes = []
        for child in content:
            if child.tag == X + 't':
                nodes.append(child)
            elif child.tag == X + 'r':
                if any(item.tag not in (X + 'rPr', X + 't') for item in child):
                    raise ValueError('Rich text contains unsupported structures')
                nodes.extend(child.findall(X + 't'))
            else:
                raise ValueError('Text contains unsupported phonetic or extension content')
        if content.find(X + 't') is not None and content.find(X + 'r') is not None:
            raise ValueError('Text mixes unsupported rich and plain structures')
        if ESCAPED.search(''.join(node.text or '' for node in nodes)):
            raise ValueError('Excel escaped text sequences are not supported')
        ooxml.replace_nodes(nodes, find, replacement, preserve_space=True)
        text = ''.join(node.text or '' for node in nodes)
        if ESCAPED.search(text):
            raise ValueError('Replacement creates an unsupported Excel escaped text sequence')
        if len(text.encode('utf-16-le')) // 2 > 32767:
            raise ValueError('Cell text exceeds Excel’s 32,767 UTF-16-unit limit')
        if kind == 's':
            candidate = content_key(content)
            same = next((i for i, item in enumerate(entries)
                         if ''.join(node.text or '' for node in item.iter(X + 't')) == text and content_key(item) == candidate), None)
            if same is None:
                unique = root_sst.get('uniqueCount')
                if unique is not None and (not re.fullmatch(r'[0-9]{1,10}', unique) or int(unique) != len(entries)):
                    raise ValueError('Shared string unique count is inconsistent')
                root_sst.append(content); same = len(entries)
                if unique is not None:
                    root_sst.set('uniqueCount', str(len(entries) + 1))
                additional[shared_path] = ooxml.serialize_xml(shared_tree, shared_xml)
            cell[0].text = str(same)
        return additional

    return ooxml.edit_xml(source, output, select_part, edit)

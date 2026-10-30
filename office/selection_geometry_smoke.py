"""Exercise revision-bound PowerPoint selection against real rendered PDFs."""
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import open_office_use
import selection_candidates
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    image = root / 'image.png'
    Image.new('RGB', (80, 80), 'red').save(image)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(image), Inches(1), Inches(1), Inches(2), Inches(2))
    for x in (4, 7):
        shape = slide.shapes.add_textbox(Inches(x), Inches(1), Inches(2), Inches(1))
        shape.text = 'Repeated text'
    source = root / 'slides.pptx'; deck.save(source)
    draft = root / 'draft'; session.start(draft, source=source)
    session.set_selection(draft, {'revision': 0, 'page': 1, 'box': [.09, .12, .31, .41]})
    result = selection_candidates.candidates(draft, open_office_use.render)
    assert result['geometry_available'] and result['geometry_match_count'] == 1, result
    assert result['candidates'][0]['type'] == 'picture' and result['candidates'][0]['geometry_match'], result
    session.set_selection(draft, {'revision': 0, 'page': 1, 'box': [.39, .12, .61, .28]})
    result = selection_candidates.candidates(draft, open_office_use.render)
    assert result['geometry_match_count'] == 1 and result['candidates'][0]['geometry_match'], result
    assert any(not item.get('geometry_match') for item in result['candidates']), result
    hidden = deck.slides.add_slide(deck.slide_layouts[6]); hidden._element.set('show', '0')
    hidden_source = root / 'hidden.pptx'; deck.save(hidden_source)
    hidden_draft = root / 'hidden-draft'; session.start(hidden_draft, source=hidden_source)
    session.set_selection(hidden_draft, {'revision': 0, 'page': 1, 'box': [.09, .12, .31, .41]})
    fallback = selection_candidates.candidates(hidden_draft, open_office_use.render)
    assert not fallback['geometry_available'] and fallback['method'] == 'text_search', fallback
print('Selection geometry smoke passed: images, repeated text and hidden-slide fallback')

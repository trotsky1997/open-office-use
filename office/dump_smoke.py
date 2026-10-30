"""Verify bounded OfficeCLI extraction and replay with fidelity checks."""
import hashlib
import json
import subprocess
import sys
import tempfile
from types import SimpleNamespace
from unittest.mock import patch
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
import xlsxwriter
import officecli_backend
import session

with tempfile.TemporaryDirectory() as tmp:
    root=Path(tmp); doc=root/'source.docx'; docx=Document(); p=docx.add_paragraph(); p.add_run('Hello ').bold=True; p.add_run('world').italic=True; docx.save(doc)
    result=officecli_backend.dump(doc,'/body/p[1]'); assert result['snapshot_sha256']==hashlib.sha256(doc.read_bytes()).hexdigest()
    assert result['operations'][0]['command']=='meta' and len(result['operations'])>=3
    doc_draft=root/'doc-draft'; session.start(doc_draft,format='docx'); edited=session.apply(doc_draft,0,result['operations']); out=Document(edited['file'])
    paragraph=next(p for p in out.paragraphs if p.text=='Hello world'); assert paragraph.runs[0].bold and paragraph.runs[1].italic
    assert session.dump(doc_draft,0)['revision']==0
    cli=Path(__file__).with_name('open_office_use.py'); output=subprocess.run([sys.executable,str(cli),'dump',str(doc),'/body/p[1]'],capture_output=True,text=True,check=True)
    assert json.loads(output.stdout)['operations']==result['operations']
    image=root/'image.png'; Image.new('RGB',(30,30),'red').save(image); ppt=root/'source.pptx'; deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6]); slide.shapes.add_picture(str(image),Inches(1),Inches(1)); deck.save(ppt)
    ppt_dump=officecli_backend.dump(ppt,'/slide[1]'); assert any(op.get('type')=='picture' and op.get('props',{}).get('src','').startswith('data:image/png;base64,') for op in ppt_dump['operations'])
    compatible=root/'compatible.pptx'; Presentation().save(compatible)
    ppt_draft=root/'ppt-draft'; session.start(ppt_draft,source=compatible); ppt_out=session.apply(ppt_draft,0,ppt_dump['operations'])
    assert len(Presentation(ppt_out['file']).slides)==1 and Presentation(ppt_out['file']).slides[0].shapes[0].image.blob==image.read_bytes()
    book_path=root/'source.xlsx'; book=xlsxwriter.Workbook(str(book_path)); sheet=book.add_worksheet('Sales'); sheet.write_column('A1',['Jan','Feb']); sheet.write_column('B1',[10,20]); chart=book.add_chart({'type':'column'}); chart.add_series({'categories':'=Sales!$A$1:$A$2','values':'=Sales!$B$1:$B$2'}); sheet.insert_chart('D1',chart); book.close()
    dumped=officecli_backend.dump(book_path,'/Sales'); xlsx_draft=root/'xlsx-draft'; session.start(xlsx_draft,format='xlsx'); replay=session.apply(xlsx_draft,0,dumped['operations'])
    assert load_workbook(replay['file'])['Sales']['B2'].value==20
    assert len(load_workbook(book_path)['Sales']._charts)==1
    if not load_workbook(replay['file'])['Sales']._charts:
        assert any('charts' in note.lower() for note in dumped['notes'])
    for bad in ('bad-path','/bad\x00path'):
        try: officecli_backend.dump(doc,bad)
        except ValueError:pass
        else:raise AssertionError('Invalid dump path accepted')
    invalid=subprocess.run([sys.executable,str(cli),'dump',str(book_path),'/Sales/chart[1]'],capture_output=True,text=True)
    assert invalid.returncode!=0 and 'unsupported' in invalid.stderr.lower()
    def fake_output(content):
        def run(*args,stdout,**kwargs): stdout.write(content); return SimpleNamespace(returncode=0)
        return run
    for data,expected in [(b'x'*(1024*1024+1),'1 MiB'),(json.dumps({'success':True,'data':[{'command':'get'}]*1001}).encode(),'1,000 operations')]:
        with patch.object(officecli_backend.subprocess,'run',fake_output(data)):
            try: officecli_backend.dump(doc,'/body')
            except ValueError as exc: assert expected in str(exc)
            else: raise AssertionError('Dump limit accepted')
print('Office dump smoke passed: Word, image slide, worksheet and explicit fidelity limit')

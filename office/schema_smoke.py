"""Discover the installed schema and use it to author a real chart."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path
from openpyxl import Workbook,load_workbook
import chart_cache
import officecli_backend
import session

catalog=officecli_backend.schema('xlsx',operation='add',query='chart')
assert catalog['officecli_version'] and any(item['element']=='chart' for item in catalog['elements'])
chart=officecli_backend.schema('xlsx','chart',operation='add')
for property in ('dataRange','chartType','title','anchor'): assert property in chart['schema']['properties'],property
assert chart['property_operation_filter_applied'] and all(item.get('add') is True for item in chart['schema']['properties'].values())
filtered=officecli_backend.schema('xlsx','chart',operation='add',query='datarange')
assert 'dataRange' in filtered['schema']['properties'] and filtered['returned_properties'] < filtered['total_properties']
for format,element in [('docx','paragraph'),('pptx','shape')]: assert officecli_backend.schema(format,element)['schema']['element']==element
for args in [('xlsx','--help'),('xlsx','doesnotexist'),('xlsx','chart','remove','\n')]:
    try: officecli_backend.schema(*args)
    except ValueError: pass
    else: raise AssertionError('Invalid schema request accepted')
cli=Path(__file__).with_name('open_office_use.py'); response=subprocess.run([sys.executable,str(cli),'schema','xlsx','--element','chart','--operation','add','--query','datarange'],capture_output=True,text=True,check=True)
assert json.loads(response.stdout)['returned_properties']>=1
with tempfile.TemporaryDirectory() as tmp:
    root=Path(tmp); source=root/'source.xlsx'; draft=root/'draft'; book=Workbook(); sheet=book.active; sheet.title='Sales'
    for row in [('Month','Revenue'),('Jan',10),('Feb',20)]: sheet.append(row)
    book.save(source); session.start(draft,source=source)
    created=session.apply(draft,0,[{'command':'add','parent':'/Sales','type':'chart','props':{'chartType':'column','dataRange':'Sales!A1:B3','title':'Revenue','anchor':'D2:J15'}}])
    chart_data=officecli_backend.read(created['file'],'query','chart'); result=chart_data['results'][0]
    assert result['format']['title']=='Revenue' and result['children'][0]['format']['valuesRef']=='Sales!$B$2:$B$3',chart_data
    code,validation,_=officecli_backend.invoke('validate',created['file']); assert code==0 and validation['success']
    assert load_workbook(created['file']).worksheets[0]['B3'].value==20
    assert chart_cache.scan(created['file'])['chart_parts_total']==1
print('Office schema smoke passed: discovery, filtering and chart authoring')

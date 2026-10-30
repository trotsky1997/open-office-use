"""Exercise frozen rows and columns as versioned worksheet view edits."""
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from lxml import etree
from openpyxl import Workbook, load_workbook
import officecli_backend
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'book.xlsx'; book = Workbook(); book.active['A1'] = 'Header'; book.active['B2'] = '=1+1'; book.create_sheet('Other')['A1'] = 5; book.save(source)
    patch = root / 'cached.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(patch, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                tree = etree.fromstring(data); ns = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                tree.xpath('./x:sheetData/x:row/x:c[@r="B2"]/x:v', namespaces=ns)[0].text = '2'; data = etree.tostring(tree)
            after.writestr(entry, data)
    patch.replace(source); draft = root / 'draft'; session.start(draft, source=source)
    for index, cell in enumerate(('A2','B1','B2',None)):
        result = session.freeze_panes(draft, index, 'Sheet', cell)
        assert result['revision'] == index + 1 and result['freeze_panes'] == cell and result['formula_cache_effect'] == 'preserved'
        saved = Path(result['file']); workbook = load_workbook(saved)
        assert workbook.active.freeze_panes == cell and workbook.active['B2'].value == '=1+1'
        assert load_workbook(saved, data_only=True).active['B2'].value == 2
        view = session.read_grid(draft, index + 1, 'Sheet', 'A1')['worksheet_view']
        assert view['issue'] is None and (view['pane']['topLeftCell'] if view['pane'] else None) == cell
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(saved) as after:
            assert all(before.read(name) == after.read(name) for name in before.namelist() if name != 'xl/worksheets/sheet1.xml')
        assert officecli_backend.audit(saved)['validation']['exit_code'] == 0
    cli = Path(__file__).with_name('open_office_use.py')
    result = subprocess.run([sys.executable, str(cli), 'session', 'freeze-panes', str(draft), '--base', '4', '--sheet', 'Sheet', '--cell', 'C3'], text=True, capture_output=True, check=True)
    assert json.loads(result.stdout)['freeze_panes'] == 'C3' and load_workbook(draft / 'revision-000005.xlsx').active.freeze_panes == 'C3'
    for cell in ('XFE1', 'A1048577', 'a2', 'A0', 'A1:B2'):
        try: session.freeze_panes(draft, 5, 'Sheet', cell)
        except ValueError: pass
        else: raise AssertionError(f'Invalid freeze position accepted: {cell}')
    assert session.info(draft)['latest_revision'] == 5
    try: session.freeze_panes(draft, 4, 'Sheet', 'A2')
    except ValueError as exc: assert 'conflict' in str(exc).lower()
    else: raise AssertionError('Stale revision accepted')
    split = root / 'split.xlsx'; book = Workbook(); book.active.freeze_panes = 'B2'; book.save(split)
    patched = root / 'split-patched.xlsx'
    with zipfile.ZipFile(split) as before, zipfile.ZipFile(patched, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry)
            if entry.filename == 'xl/worksheets/sheet1.xml': data = data.replace(b'state="frozen"', b'state="split"')
            after.writestr(entry, data)
    split_draft = root / 'split-draft'; session.start(split_draft, source=patched)
    try: session.freeze_panes(split_draft, 0, 'Sheet', 'A2')
    except ValueError: pass
    else: raise AssertionError('Split view overwritten')
    assert session.info(split_draft)['latest_revision'] == 0
print('Freeze panes smoke passed: axes, unfreeze, preserved caches and protected split views')

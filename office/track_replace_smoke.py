"""Exercise tracked Word replacement, formatting and revision safety."""
import copy
import concurrent.futures
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import officecli_backend
import session
import word

with tempfile.TemporaryDirectory() as temp:
    root = Path(temp); source = root / 'source.docx'; draft = root / 'draft'; doc = Document()
    p = doc.add_paragraph('Start '); p.add_run('old').bold = True; p.add_run(' text').italic = True; p.add_run(' end')
    table = doc.add_table(rows=1, cols=1); table.cell(0, 0).text = 'table old'
    doc.save(source); original = source.read_bytes(); session.start(draft, source=source)
    result = session.track_replace_text(draft, 0, '/body/p[1]', 'old text', 'new', 'Reviewer')
    output = Path(result['file']); changes = session.tracked_changes(draft, 1)
    assert [(x['kind'], x['text'], x['author']) for x in changes['records']] == [('del', 'old text', 'Reviewer'), ('ins', 'new', 'Reviewer')]
    assert source.read_bytes() == original and len({x['id'] for x in changes['records']}) == 2
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        assert set(before.namelist()) == set(after.namelist())
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name != 'word/document.xml')
        tree = etree.fromstring(after.read('word/document.xml')); paragraph = tree.find('.//' + word.W + 'body/' + word.W + 'p')
        deleted = paragraph.find(word.W + 'del'); inserted = paragraph.find(word.W + 'ins')
        assert deleted[0].find(word.W + 'rPr/' + word.W + 'b') is not None and deleted[1].find(word.W + 'rPr/' + word.W + 'i') is not None
        assert inserted[0].find(word.W + 'rPr/' + word.W + 'b') is not None
        def text_as(accept):
            node = copy.deepcopy(paragraph)
            for element in list(node):
                if element.tag == word.W + ('del' if accept else 'ins'): node.remove(element)
                elif element.tag == word.W + ('ins' if accept else 'del'):
                    index = node.index(element)
                    for run in list(element):
                        for text in run.findall(word.W + 'delText'): text.tag = word.W + 't'
                        node.insert(index, run); index += 1
                    node.remove(element)
            return ''.join(t.text or '' for t in node.iter(word.W + 't'))
        assert text_as(True) == 'Start new end' and text_as(False) == 'Start old text end'
    audit = officecli_backend.audit(output)
    assert audit['validation']['exit_code'] == 0 and audit['validation']['result']['success'], audit
    for base, path, find, repl, author in [(0, '/body/p[1]', 'old', 'new', 'A'), (1, '/body/p[1]', 'new', 'x', 'A'), (1, '/body/p[1]', 'x', 'y', '')]:
        try: session.track_replace_text(draft, base, path, find, repl, author)
        except ValueError: pass
        else: raise AssertionError('invalid tracked edit accepted')
    assert session.info(draft)['latest_revision'] == 1
    cli = Path(__file__).with_name('open_office_use.py')
    command = subprocess.run([sys.executable, str(cli), 'session', 'track-replace-text', str(draft), '--base', '1', '--path', '/body/tbl[1]/tr[1]/tc[1]/p[1]', '--find', 'old', '--replace', '', '--author', 'Editor'], capture_output=True, text=True, check=True)
    assert json.loads(command.stdout)['revision'] == 2
    updated = session.tracked_changes(draft, 2); assert updated['counts']['del'] == 2 and updated['counts']['ins'] == 1
    collision_source = root / 'collision.docx'; collision_doc = Document(); collision_doc.add_paragraph('needle'); other = collision_doc.add_paragraph()
    existing = OxmlElement('w:ins'); existing.set(qn('w:id'), '00000000000000000007'); other._p.append(existing); collision_doc.save(collision_source)
    collision_draft = root / 'collision-draft'; session.start(collision_draft, source=collision_source)
    with patch('secrets.randbelow', side_effect=[7, 8, 8, 9]):
        session.track_replace_text(collision_draft, 0, '/body/p[1]', 'needle', 'replacement', 'Editor')
    ids = [record['id'] for record in session.tracked_changes(collision_draft, 1)['records']]
    assert sorted(int(value) for value in ids) == [7, 8, 9]
    boundary_source = root / 'boundary.docx'; boundary_doc = Document(); boundary = boundary_doc.add_paragraph('left')
    marker = OxmlElement('w:bookmarkStart'); marker.set(qn('w:id'), '0'); marker.set(qn('w:name'), 'mark'); boundary._p.append(marker); boundary.add_run('right'); boundary_doc.save(boundary_source)
    boundary_draft = root / 'boundary-draft'; session.start(boundary_draft, source=boundary_source)
    try: session.track_replace_text(boundary_draft, 0, '/body/p[1]', 'leftright', 'x', 'Editor')
    except ValueError: pass
    else: raise AssertionError('annotation boundary was crossed')
    assert session.info(boundary_draft)['latest_revision'] == 0
    range_source = root / 'range.docx'; range_doc = Document(); range_doc.add_paragraph('outside target'); inside = range_doc.add_paragraph('inside target')
    start = OxmlElement('w:moveFromRangeStart'); start.set(qn('w:id'), '40'); inside._p.addprevious(start)
    end = OxmlElement('w:moveFromRangeEnd'); end.set(qn('w:id'), '40'); inside._p.addnext(end); range_doc.save(range_source)
    range_draft = root / 'range'; session.start(range_draft, source=range_source)
    try: session.track_replace_text(range_draft, 0, '/body/p[2]', 'target', 'new', 'Editor')
    except ValueError as exc: assert 'inside' in str(exc)
    else: raise AssertionError('tracked replacement was created inside a move range')
    assert session.info(range_draft)['latest_revision'] == 0
    range_edit = session.track_replace_text(range_draft, 0, '/body/p[1]', 'target', 'new', 'Editor')
    range_records = session.tracked_changes(range_draft, 1)
    ids = [r['id'] for r in range_records['records'] if r['kind'] in ('ins', 'del')]
    assert session.resolve_tracked_runs(range_draft, 1, 'word/document.xml', ids, True)['revision'] == 2
    assert Document(range_edit['file']).paragraphs[1].text == 'inside target'
    row_source = root / 'row.docx'; row_doc = Document(); table = row_doc.add_table(rows=1, cols=1); table.cell(0, 0).text = 'target'
    row_props = table.rows[0]._tr.get_or_add_trPr(); row_mark = OxmlElement('w:ins'); row_mark.set(qn('w:id'), '41'); row_props.append(row_mark); row_doc.save(row_source)
    row_draft = root / 'row'; session.start(row_draft, source=row_source)
    try: session.track_replace_text(row_draft, 0, '/body/tbl[1]/tr[1]/tc[1]/p[1]', 'target', 'new', 'Editor')
    except ValueError as exc: assert 'properties' in str(exc)
    else: raise AssertionError('tracked replacement inside a tracked row was accepted')
    assert session.info(row_draft)['latest_revision'] == 0
    for kind in ('tblGridChange', 'tblPrExChange'):
        structural_doc = Document(); structural_table = structural_doc.add_table(rows=1, cols=1); structural_table.cell(0, 0).text = 'target'
        if kind == 'tblGridChange': container = structural_table._tbl.tblGrid
        else:
            container = OxmlElement('w:tblPrEx'); structural_table.rows[0]._tr.insert(0, container)
        marker = OxmlElement('w:' + kind); marker.set(qn('w:id'), '42'); container.append(marker)
        structural_source = root / (kind + '.docx'); structural_doc.save(structural_source)
        structural_draft = root / kind; session.start(structural_draft, source=structural_source)
        try: session.track_replace_text(structural_draft, 0, '/body/tbl[1]/tr[1]/tc[1]/p[1]', 'target', 'new', 'Editor')
        except ValueError as exc: assert 'properties' in str(exc)
        else: raise AssertionError('tracked table structure was accepted')
        assert session.info(structural_draft)['latest_revision'] == 0 and session.tracked_changes(structural_draft, 0)['counts'][kind] == 1
    conflict_source = root / 'run-conflict.docx'; conflict_doc = Document(); run = conflict_doc.add_paragraph().add_run('target')
    run._r.get_or_add_rPr().append(etree.Element(word.W14 + 'conflictIns')); conflict_doc.save(conflict_source)
    conflict_draft = root / 'run-conflict'; session.start(conflict_draft, source=conflict_source)
    try: session.track_replace_text(conflict_draft, 0, '/body/p[1]', 'target', 'new', 'Editor')
    except ValueError: pass
    else: raise AssertionError('run with tracked conflict formatting was accepted')
    assert session.info(conflict_draft)['latest_revision'] == 0
    concurrent_draft = root / 'concurrent-draft'; session.start(concurrent_draft, source=source)
    def concurrent_edit():
        try: return session.track_replace_text(concurrent_draft, 0, '/body/p[1]', 'old text', 'new', 'Editor')
        except ValueError: return None
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: concurrent_edit(), range(2)))
    assert sum(item is not None for item in results) == 1 and session.info(concurrent_draft)['latest_revision'] == 1
print('Tracked replacement smoke passed: formatting, accept/reject text, audit, tables and conflicts')

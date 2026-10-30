"""Exercise revision-bound formula scans and recalculation."""
import concurrent.futures
import json
import signal
import subprocess
import sys
import tempfile
import urllib.error
import urllib.request
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from docx import Document
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.xlsx'; draft = root / 'draft'
    book = Workbook(); book.active['A1'] = 4; book.active['A2'] = '=A1*3'; book.active['A3'] = '=1/0'; book.save(source)
    original = source.read_bytes(); session.start(draft, source=source)
    before = session.scan_formulas(draft, 0)
    assert before['revision'] == 0 and before['empty_or_missing_cache'] == 2 and before['complete'], before
    with patch('spreadsheet.subprocess.run', return_value=SimpleNamespace(returncode=1, stderr='failed', stdout='')):
        try:
            session.recalculate(draft, 0)
            raise AssertionError('Failed conversion was published')
        except RuntimeError:
            pass
    assert session.info(draft)['latest_revision'] == 0
    result = session.recalculate(draft, 0)
    assert result['revision'] == 1 and result['engine'] == 'LibreOffice' and source.read_bytes() == original
    assert load_workbook(result['file'], read_only=True).calculation.calcMode == book.calculation.calcMode
    values = load_workbook(result['file'], data_only=True).active
    assert values['A2'].value == 12 and values['A3'].value == '#DIV/0!'
    after = session.scan_formulas(draft, 1)
    assert after['revision'] == 1 and after['formula_errors'] == 1 and after['empty_or_missing_cache'] == 0, after
    for operation in (session.recalculate, session.scan_formulas):
        try:
            operation(draft, 0)
            raise AssertionError('Stale operation accepted')
        except ValueError as error:
            assert 'conflict' in str(error)
    assert (draft / 'revision-000000.xlsx').read_bytes() == original

    doc_source = root / 'source.docx'; Document().save(doc_source); doc_draft = root / 'doc'; session.start(doc_draft, source=doc_source)
    for operation in (session.recalculate, session.scan_formulas):
        try:
            operation(doc_draft, 0)
            raise AssertionError('Non-XLSX operation accepted')
        except ValueError as error:
            assert 'XLSX' in str(error)
    assert session.info(doc_draft)['latest_revision'] == 0

    cli = Path(__file__).with_name('open_office_use.py'); cli_draft = root / 'cli'; session.start(cli_draft, source=source)
    scanned = subprocess.run([sys.executable, str(cli), 'session', 'scan-formulas', str(cli_draft), '--base', '0'], check=True, capture_output=True, text=True)
    assert json.loads(scanned.stdout)['empty_or_missing_cache'] == 2
    calculated = subprocess.run([sys.executable, str(cli), 'session', 'recalculate', str(cli_draft), '--base', '0'], check=True, capture_output=True, text=True)
    assert json.loads(calculated.stdout)['revision'] == 1 and load_workbook(cli_draft / 'revision-000001.xlsx', data_only=True).active['A2'].value == 12

    concurrent_draft = root / 'concurrent'; session.start(concurrent_draft, source=source)
    def convert():
        try:
            return session.recalculate(concurrent_draft, 0)['revision']
        except ValueError as error:
            assert 'conflict' in str(error)
            return 'conflict'
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
        outcomes = list(pool.map(lambda _: convert(), range(2)))
    assert sorted(map(str, outcomes)) == ['1', 'conflict'] and session.info(concurrent_draft)['latest_revision'] == 1, outcomes

    server_draft = root / 'server'; session.start(server_draft, source=source)
    cli = Path(__file__).with_name('open_office_use.py')
    server = subprocess.Popen([sys.executable, str(cli), 'watch', str(server_draft)], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    try:
        url = json.loads(server.stdout.readline())['url']; origin = 'http://' + urlsplit(url).netloc
        def post(route, base, request_origin=origin):
            request = urllib.request.Request(url + route, data=json.dumps({'base_revision': base}).encode(), headers={'Origin': request_origin, 'Content-Type': 'application/json'})
            with urllib.request.urlopen(request, timeout=60) as response: return json.loads(response.read())
        assert post('formula-scan', 0)['empty_or_missing_cache'] == 2
        try:
            post('formula-recalculate', 0, 'http://evil.test')
            raise AssertionError('Cross-origin recalculation accepted')
        except urllib.error.HTTPError as error:
            assert error.code == 403
        assert post('formula-recalculate', 0)['revision'] == 1
        assert post('formula-scan', 1)['formula_errors'] == 1
        try:
            post('formula-recalculate', 0)
            raise AssertionError('Stale recalculation accepted')
        except urllib.error.HTTPError as error:
            assert error.code == 400
    finally:
        server.send_signal(signal.SIGINT); server.wait(timeout=10)
print('Session calculation smoke passed: scans, caches, failures, concurrency and origin checks')

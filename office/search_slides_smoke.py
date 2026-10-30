"""Exercise slide search paths, groups, tables, pagination and editing."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Inches
import session
import slides

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.pptx'; draft = root / 'draft'; deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.paragraphs[0]
    text.add_run().text = 'find '; text.add_run().text = 'needle'
    group = slide.shapes.add_group_shape(); group.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1)).text = 'group needle'
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(2), Inches(1)).table; table.cell(0, 0).text = 'table needle'
    hidden = deck.slides.add_slide(deck.slide_layouts[6]); hidden._element.set('show', '0')
    for _ in range(101): hidden.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = 'needle'
    deck.save(source)
    result = slides.search_text(source, 'needle')
    assert result['matches'][0]['path'] == '/slide[1]/shape[1]/paragraph[1]'
    assert result['matches'][1]['path'] == '/slide[1]/group[1]/shape[1]/paragraph[1]'
    assert result['matches'][2]['path'] == '/slide[1]/table[1]/tr[1]/tc[1]/paragraph[1]'
    assert result['stop_reason'] == 'match_limit' and result['matches'][3]['hidden'] is True
    rest = slides.search_text(source, 'needle', offset=result['next_offset'], snapshot_sha256=result['snapshot_sha256'])
    assert rest['scan_complete'] and len(rest['matches']) == 4
    selected = slides.search_text(source, 'needle', slide=1)
    assert selected['complete'] and len(selected['matches']) == 3
    alternate = root / 'alternate.pptx'; slide.shapes._spTree.append(etree.Element(slides.MC + 'AlternateContent')); deck.save(alternate)
    alternate_result = slides.search_text(alternate, 'needle', slide=1)
    assert not alternate_result['complete'] and 'alternate_content' in alternate_result['unsupported_structure']
    session.start(draft, source=source); hit = session.search_slides(draft, 0, 'group needle')['matches'][0]
    session.replace_text(draft, 0, hit['path'], 'group needle', 'edited')
    assert session.search_slides(draft, 1, 'edited')['matches'][0]['path'] == hit['path']
    cli = Path(__file__).with_name('open_office_use.py')
    out = subprocess.run([sys.executable, str(cli), 'session', 'search-slides', str(draft), '--revision', '0', '--query', 'table needle', '--slide', '1'], capture_output=True, text=True, check=True)
    assert json.loads(out.stdout)['matches'][0]['path'] == result['matches'][2]['path']
    for offset, digest in ((1, None), (1, '0' * 64)):
        try: slides.search_text(source, 'needle', offset=offset, snapshot_sha256=digest)
        except ValueError: pass
        else: raise AssertionError('invalid continuation accepted')
print('Slide search smoke passed: shapes, groups, tables, hidden slides and targeted edits')

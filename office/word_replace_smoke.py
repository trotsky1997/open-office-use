"""Exercise versioned, run-preserving Word text replacement."""
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from PIL import Image
from lxml import etree

import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.docx'; draft = root / 'draft'
    image = root / 'image.png'; Image.new('RGB', (10, 10), 'red').save(image)
    doc = Document(); para = doc.add_paragraph(); para.add_run('Hello ').bold = True; para.add_run('world!').italic = True
    doc.add_paragraph('red red'); doc.add_paragraph('before\tafter'); doc.add_picture(str(image))
    marked = doc.add_paragraph('marked'); marker = OxmlElement('w:bookmarkStart'); marker.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id', '0'); marker.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name', 'test'); marked._p.insert(0, marker)
    doc.save(source); original = source.read_bytes(); session.start(draft, source=source)
    for path, find in [('/body/p[2]', 'red'), ('/body/p[3]', 'before'), ('/body/p[5]', 'marked'), ('/body/p[1]', 'missing')]:
        try:
            session.replace_text(draft, 0, path, find, 'new')
            raise AssertionError('Unsafe or ambiguous replacement succeeded')
        except ValueError:
            pass
    assert session.info(draft)['latest_revision'] == 0
    result = session.replace_text(draft, 0, '/body/p[1]', 'llo wor', 'i earth')
    assert result['revision'] == 1 and source.read_bytes() == original
    revised = Document(result['file']); runs = revised.paragraphs[0].runs
    assert revised.paragraphs[0].text == 'Hei earthld!' and runs[0].bold and runs[1].italic
    assert runs[0].text == 'Hei earth' and runs[1].text == 'ld!'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(result['file']) as after:
        assert set(before.namelist()) == set(after.namelist())
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name != 'word/document.xml')
        assert before.read('word/document.xml') != after.read('word/document.xml')
        old_tree = etree.fromstring(before.read('word/document.xml')); new_tree = etree.fromstring(after.read('word/document.xml'))
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        for tree in (old_tree, new_tree):
            for node in tree.xpath('/w:document/w:body/w:p[1]/w:r/w:t', namespaces=ns): node.text = 'MATCH'
        assert etree.tostring(old_tree, method='c14n') == etree.tostring(new_tree, method='c14n')

    signed = root / 'signed.docx'
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(signed, 'w') as dest:
        for item in src.infolist(): dest.writestr(item, src.read(item))
        dest.writestr('_xmlsignatures/sig1.xml', '<Signature/>')
    signed_draft = root / 'signed-draft'; session.start(signed_draft, source=signed)
    try:
        session.replace_text(signed_draft, 0, '/body/p[1]', 'Hello', 'Hi')
        raise AssertionError('Signed document accepted')
    except ValueError as error:
        assert 'signed' in str(error)
    assert session.info(signed_draft)['latest_revision'] == 0
    try:
        session.replace_text(draft, 0, '/body/p[1]', 'Hello', 'Bad')
        raise AssertionError('Stale revision accepted')
    except ValueError as error:
        assert 'conflict' in str(error)
    assert session.info(draft)['latest_revision'] == 1
print('Word replacement smoke passed: run formatting, package preservation and rejected edits')

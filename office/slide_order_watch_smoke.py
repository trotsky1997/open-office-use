"""Check slide order viewer endpoints and revision protection."""
import json
import subprocess
import sys
import tempfile
import urllib.error
import urllib.request
from pathlib import Path
from urllib.parse import urlsplit
from pptx import Presentation
import open_office_use
import session

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root/'deck.pptx'; draft = root/'draft'
    open_office_use.author({'slides':[{'title':'One'}, {'title':'Two'}]}, source)
    session.start(draft, source=source)
    server = subprocess.Popen([sys.executable, str(Path(open_office_use.__file__)), 'watch', str(draft)], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    try:
        url = json.loads(server.stdout.readline())['url']; origin = 'http://' + urlsplit(url).netloc
        def post(route, data, request_origin=origin):
            request = urllib.request.Request(url+route,data=json.dumps(data).encode(),headers={'Origin':request_origin,'Content-Type':'application/json'})
            with urllib.request.urlopen(request,timeout=30) as response: return json.load(response)
        inventory = post('slide-order-read',{'revision':0})
        assert [item['title'] for item in inventory['slides']] == ['One','Two']
        payload = {'base_revision':0,'snapshot_sha256':inventory['snapshot_sha256'],'order':[2,1]}
        for bad, request_origin, code in [(dict(payload,snapshot_sha256='0'*64),origin,400),(dict(payload,snapshot_sha256=None),origin,400),(payload,'http://evil.test',403)]:
            try: post('slide-order-save',bad,request_origin)
            except urllib.error.HTTPError as exc: assert exc.code == code
            else: raise AssertionError('Invalid slide order request accepted')
        result = post('slide-order-save',payload); assert result['revision'] == 1
        assert [item.shapes.title.text for item in Presentation(result['file']).slides] == ['Two','One']
        try: post('slide-order-save',payload)
        except urllib.error.HTTPError as exc: assert exc.code == 400
        else: raise AssertionError('Stale revision accepted')
        assert session.info(draft)['latest_revision'] == 1
    finally:
        server.terminate(); server.wait(timeout=10)
print('Slide order viewer routes passed: inventory, reorder, source, origin and revision checks')

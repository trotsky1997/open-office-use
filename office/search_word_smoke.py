"""Exercise Word search paths, coverage and safe targeted editing."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
import session
import word

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.docx'; draft = root / 'draft'; doc = Document()
    paragraph = doc.add_paragraph(); paragraph.add_run('find ').bold = True; paragraph.add_run('needle')
    cell = doc.add_table(rows=1, cols=1).cell(0, 0); cell.paragraphs[0].text = 'table needle'
    nested = cell.add_table(rows=1, cols=1); nested.cell(0, 0).text = 'nested needle'
    special = doc.add_paragraph('tab needle'); special.add_run().add_tab()
    doc.add_paragraph('aaa')
    for _ in range(101): doc.add_paragraph('needle')
    doc.save(source)
    result = word.search_body(source, 'needle')
    assert result['matches'][0]['path'] == '/body/p[1]' and result['matches'][1]['path'] == '/body/tbl[1]/tr[1]/tc[1]/p[1]'
    assert result['matches'][2]['path'] == '/body/tbl[1]/tr[1]/tc[1]/tbl[1]/tr[1]/tc[1]/p[1]'
    assert not result['matches'][3]['text_complete'] and result['unverified_paragraphs'] == 1
    assert result['stop_reason'] == 'match_limit' and len(result['matches']) == 100
    rest = word.search_body(source, 'needle', result['next_offset'], result['snapshot_sha256'])
    assert rest['scan_complete'] and len(rest['matches']) == 5
    assert word.search_body(source, 'aa')['matches'][0]['occurrences'] == 2
    wrapper_doc = Document(); wrapper_doc.add_paragraph('plain needle'); wrapper_doc.element.body.append(OxmlElement('w:sdt'))
    wrapper = root / 'wrapper.docx'; wrapper_doc.save(wrapper)
    wrapper_result = word.search_body(wrapper, 'needle')
    assert not wrapper_result['complete'] and 'container_structure' in wrapper_result['unsupported_structure']
    session.start(draft, source=source)
    candidate = session.search_body(draft, 0, 'nested needle')['matches'][0]
    edited = session.replace_text(draft, 0, candidate['path'], 'nested needle', 'replaced')
    assert session.search_body(draft, 1, 'replaced')['matches'][0]['path'] == candidate['path']
    assert word.search_body(source, 'nested needle')['matches'] and Path(edited['file']).exists()
    cli = Path(__file__).with_name('open_office_use.py')
    out = subprocess.run([sys.executable, str(cli), 'session', 'search-word', str(draft), '--revision', '0', '--query', 'nested needle'], capture_output=True, text=True, check=True)
    assert json.loads(out.stdout)['matches'][0]['path'] == candidate['path']
    for offset, digest in ((1, None), (1, '0' * 64)):
        try: word.search_body(source, 'needle', offset, digest)
        except ValueError: pass
        else: raise AssertionError('invalid continuation accepted')
print('Word search smoke passed: body, nested tables, pagination and targeted edits')

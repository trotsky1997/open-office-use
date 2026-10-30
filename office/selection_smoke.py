"""Exercise word geometry and revision binding on rendered Office files."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from session import set_selection

cli = Path(__file__).with_name('open_office_use.py')
def run(*args):
    result = subprocess.run([sys.executable, str(cli), *map(str, args)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    return json.loads(result.stdout)

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    for suffix in ('.docx', '.xlsx', '.pptx'):
        source = root / ('source' + suffix); directory = root / suffix[1:]
        if suffix == '.docx':
            obj = Document(); obj.add_heading('First title', 0); obj.add_paragraph('Hello highlighted world'); obj.add_page_break(); obj.add_paragraph('Second page marker'); obj.save(source)
        elif suffix == '.xlsx':
            obj = Workbook(); obj.active['A1'] = 'Highlighted'; obj.save(source)
        else:
            obj = Presentation(); slide = obj.slides.add_slide(obj.slide_layouts[1]); slide.shapes.title.text = 'Highlighted'; obj.save(source)
        run('session', 'start', directory, '--source', source)
        set_selection(directory, {'revision': 0, 'page': 1, 'box': [0, 0, 1, 1]})
        text = run('session', 'selection-text', directory)
        assert 'highlighted' in text['text'].lower() and text['revision'] == 0 and not text['stale'], text
        assert text['words'] and not text['truncated']
        if suffix == '.docx':
            set_selection(directory, {'revision': 0, 'page': 1, 'box': [0.13, 0.08, 0.31, 0.13]})
            narrow = run('session', 'selection-text', directory)
            assert 'First title' in narrow['text'] and 'highlighted' not in narrow['text'], narrow
            set_selection(directory, {'revision': 0, 'page': 2, 'box': [0, 0, 1, 1]})
            second = run('session', 'selection-text', directory)
            assert 'Second page marker' in second['text'] and 'First title' not in second['text'], second
            set_selection(directory, {'revision': 0, 'page': 1, 'box': [0.13, 0.08, 0.31, 0.13]})
            spec = root / 'ops.json'; spec.write_text(json.dumps([{'command': 'set', 'path': '/body/p[1]', 'props': {'text': 'Changed'}}]))
            run('session', 'apply', directory, '--base', '0', '--spec', spec)
            stale = run('session', 'selection-text', directory)
            assert stale['stale'] and stale['revision'] == 0 and 'First title' in stale['text'], stale
print('Selection smoke passed: three formats, bounded text and exact stale revision')

"""Create a positioned slide with editable text, shape and image."""
import tempfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import open_office_use

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); image = root/'picture.png'; Image.new('RGB', (32, 20), '#2b8a91').save(image)
    spec = {'slide_size':{'width':13.333,'height':7.5}, 'slides':[{'layout':'blank','background':'F5F8FC','elements':[
        {'type':'rectangle','x':0.5,'y':0.5,'width':0.2,'height':1.0,'fill':'236B8E'},
        {'type':'textbox','x':1,'y':0.5,'width':9,'height':1,'text':'Quarterly review','font_size':32,'font':'Arial','bold':True,'color':'14283F'},
        {'type':'textbox','x':1,'y':2,'width':6,'height':2,'text':'Revenue +18%\nRetention 92%','font_size':20,'color':'31485F'},
        {'type':'image','path':str(image),'x':9,'y':2,'width':3,'height':2,'description':'Teal illustration'}]}]}
    out = root/'slide.pptx'; open_office_use.author(spec, out)
    deck = Presentation(out); slide = deck.slides[0]
    assert abs(deck.slide_width-Inches(13.333)) < 5 and slide.background.fill.fore_color.rgb == (0xf5,0xf8,0xfc)
    assert len(slide.shapes) == 4 and slide.shapes[1].text == 'Quarterly review'
    assert slide.shapes[1].text_frame.paragraphs[0].runs[0].font.size.pt == 32
    assert slide.shapes[3]._element.nvPicPr.cNvPr.get('descr') == 'Teal illustration'
    pdf = open_office_use.render(out, root/'render'); assert Path(pdf).read_bytes().startswith(b'%PDF')
    original = out.read_bytes(); appended = root/'appended.pptx'
    open_office_use.author({'slides':[{'layout':'blank','elements':[{'type':'textbox','x':1,'y':1,'width':4,'height':1,'text':'Append'}]}]}, appended, out)
    assert out.read_bytes() == original and len(Presentation(appended).slides) == 2
    try: open_office_use.author({'slide_size':{'width':10,'height':7.5}}, root/'resize.pptx', out)
    except ValueError: pass
    else: raise AssertionError('Existing presentation was resized')
    for i, bad in enumerate([
        {'slides':[{'layout':'blank','elements':[{'type':'textbox','x':0,'y':0,'width':0,'height':1,'text':'x'}]}]},
        {'slides':[{'layout':'blank','elements':[{'type':'rectangle','x':0,'y':0,'width':1,'height':1,'fill':'INVALID'}]}]},
        {'slides':[{'layout':'blank','elements':[{'type':'image','x':0,'y':0,'width':1,'height':1,'path':str(root/'missing')}]}]},
        {'slide_size':{'width':float('nan'),'height':7.5}}]):
        target = root/f'bad{i}.pptx'
        try: open_office_use.author(bad, target)
        except (ValueError, OSError): pass
        else: raise AssertionError('Invalid slide specification accepted')
        assert not target.exists()
print('Slide author passed: layout, formatting, image, render and invalid input')

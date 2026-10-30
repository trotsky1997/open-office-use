"""Verify text replacement after OfficeCLI writes a plain string cell."""
import tempfile
import zipfile
from pathlib import Path
from lxml import etree
from openpyxl import load_workbook
import xlsxwriter
import officecli_backend
import session
import spreadsheet

with tempfile.TemporaryDirectory() as tmp:
    root=Path(tmp); source=root/'source.xlsx'; draft=root/'draft'
    wb=xlsxwriter.Workbook(str(source)); sheet=wb.add_worksheet('Sheet'); style=wb.add_format({'bold':True,'bg_color':'#FFCC00'})
    sheet.write_string('A1','before',style); sheet.write_comment('A1','Keep'); sheet.write_url('A2','https://example.com',string='Link'); sheet.write_formula('B1','=CONCAT("alpha"," beta")',None,'alpha beta'); wb.close()
    session.start(draft,source=source); grid=session.edit_grid(draft,0,'Sheet','A1',[{'cell':'A1','value':'alpha beta'}]); written=Path(grid['file'])
    target=session.edit_target(draft,1,'/Sheet/A1'); assert target['replace_text_path_supported'] and target['value']=='alpha beta', target
    formula_target=session.edit_target(draft,1,'/Sheet/B1'); assert formula_target['mode']=='formula' and not formula_target['replace_text_path_supported'], formula_target
    for path,find in [('/Sheet/B1','alpha'),('/Sheet/A1','a'),('/Sheet/A1','_x0041_')]:
        try: session.replace_text(draft,1,path,find,'bad')
        except ValueError: pass
        else: raise AssertionError('Unsafe replacement accepted')
    assert session.info(draft)['latest_revision']==1
    edited=session.replace_text(draft,1,'/Sheet/A1','alpha',' α '); result=Path(edited['file'])
    book=load_workbook(result); cell=book['Sheet']['A1']; assert cell.value==' α  beta' and cell.font.bold and cell.fill.fgColor.rgb.endswith('FFCC00') and cell.comment.text=='Keep'
    assert book['Sheet']['B1'].value=='=CONCAT("alpha"," beta")' and book['Sheet']['A2'].hyperlink.target=='https://example.com'
    with zipfile.ZipFile(written) as before,zipfile.ZipFile(result) as after:
        assert set(before.namelist())==set(after.namelist())
        assert all(before.read(n)==after.read(n) for n in before.namelist() if n!='xl/worksheets/sheet1.xml')
        sheet_xml=etree.fromstring(after.read('xl/worksheets/sheet1.xml')); ns={'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        assert sheet_xml.xpath('string(//x:c[@r="A1"]/@t)',namespaces=ns)=='str' and sheet_xml.xpath('string(//x:c[@r="A1"]/x:v)',namespaces=ns)==' α  beta'
    code, validation, _=officecli_backend.invoke('validate',result); assert code==0 and validation['success'],validation
    converted=root/'roundtrip.xlsx'; spreadsheet.recalculate(result,converted)
    assert load_workbook(converted)['Sheet']['A1'].value==' α  beta'
print('String cell smoke passed: grid to replacement, formula guard and whitespace')

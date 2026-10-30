"""Exercise concurrent-write rejection at search and review entry points."""
import hashlib
import tempfile
from pathlib import Path
from unittest.mock import patch
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
import review_bundle
import slides
import snapshot
import spreadsheet
import word

real_sha256 = hashlib.sha256
with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    class MutatingDigest:
        def __init__(self, source): self.inner = real_sha256(); self.source = source
        def update(self, block):
            self.inner.update(block)
            if self.source:
                with self.source.open('r+b') as stream:
                    stream.seek(-1, 2); value = stream.read(1); stream.seek(-1, 2); stream.write(bytes([value[0] ^ 1]))
                self.source = None
        def hexdigest(self): return self.inner.hexdigest()
    document = root / 'source.docx'; doc = Document(); doc.add_paragraph('needle'); doc.save(document)
    presentation = root / 'source.pptx'; deck = Presentation(); deck.slides.add_slide(deck.slide_layouts[0]); deck.save(presentation)
    workbook = root / 'source.xlsx'; book = Workbook(); book.active['A1'] = 'needle'; book.save(workbook)
    cases = [(document, lambda p: word.search_body(p, 'needle')),
             (document, word.tracked_changes), (presentation, lambda p: slides.search_text(p, 'needle')),
             (workbook, lambda p: spreadsheet.search_cells(p, 'needle')),
             (document, lambda p: review_bundle.create(p, root / 'report', lambda *_: None))]
    for index, (original, call) in enumerate(cases):
        source = root / f'case-{index}{original.suffix}'; source.write_bytes(original.read_bytes())
        with patch.object(snapshot.hashlib, 'sha256', side_effect=lambda: MutatingDigest(source)):
            try: call(source)
            except ValueError as exc: assert 'changed while creating snapshot' in str(exc)
            else: raise AssertionError(f'Concurrent write accepted in case {index}')
    assert not (root / 'report').exists()
print('Snapshot integration smoke passed: searches, tracked changes and review bundle')

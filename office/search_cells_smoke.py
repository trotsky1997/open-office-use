"""Exercise sparse search, result limits and revision-specific access."""
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from openpyxl import Workbook
import xlsxwriter
import session
import spreadsheet
import review
import open_office_use

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp); source = root / 'source.xlsx'; book = Workbook(); sheet = book.active
    sheet['XFD1048576'] = 'distant needle'; sheet['A1'] = '=SUM(needle)'
    hidden = book.create_sheet('Hidden'); hidden.sheet_state = 'hidden'
    for i in range(1, 203): hidden.cell(i, 1, 'needle')
    book.save(source)
    result = spreadsheet.search_cells(source, 'needle')
    assert not result['complete'] and result['match_count'] == 200 and result['matches_truncated'] and result['stop_reason'] == 'match_limit'
    assert result['matches'][0]['kind'] == 'formula' and result['matches'][1]['cell'] == 'XFD1048576'
    assert result['matches'][2]['sheet_state'] == 'hidden'
    rest = spreadsheet.search_cells(source, 'needle', offset=result['next_offset'], snapshot_sha256=result['snapshot_sha256'])
    assert rest['scan_complete'] and rest['match_count'] == 4 and rest['next_offset'] is None
    hidden_only = spreadsheet.search_cells(source, 'needle', sheet='Hidden')
    assert hidden_only['matches'][0]['sheet'] == 'Hidden' and hidden_only['match_count'] == 200
    try: spreadsheet.search_cells(source, 'needle', offset=1)
    except ValueError: pass
    else: raise AssertionError('continuation without snapshot hash accepted')
    try: spreadsheet.search_cells(source, 'needle', offset=1, snapshot_sha256='0' * 64)
    except ValueError: pass
    else: raise AssertionError('changed snapshot accepted')
    assert spreadsheet.search_cells(source, 'Needle')['match_count'] == 0
    draft = root / 'draft'; session.start(draft, source=source)
    assert session.search_cells(draft, 0, 'distant')['matches'][0]['cell'] == 'XFD1048576'
    cli = Path(__file__).with_name('open_office_use.py')
    output = subprocess.run([sys.executable, str(cli), 'session', 'search-cells', str(draft), '--revision', '0', '--query', 'distant'], capture_output=True, text=True, check=True)
    assert json.loads(output.stdout)['revision'] == 0
    output = subprocess.run([sys.executable, str(cli), 'search-cells', str(source), '--query', 'needle', '--sheet', 'Hidden', '--offset', str(hidden_only['next_offset']), '--snapshot-sha256', hidden_only['snapshot_sha256']], capture_output=True, text=True, check=True)
    assert [item['cell'] for item in json.loads(output.stdout)['matches']] == ['A201', 'A202']
    shared = root / 'shared.xlsx'; writer = xlsxwriter.Workbook(str(shared)); writer.add_worksheet().write('A1', '共享文本'); writer.close()
    shared_result = spreadsheet.search_cells(shared, '共享')
    assert shared_result['matches'][0]['cell'] == 'A1' and shared_result['complete'] and shared_result['unverified_text_cells'] == 0
    mixed = root / 'mixed.xlsx'; writer = xlsxwriter.Workbook(str(mixed)); tab = writer.add_worksheet()
    tab.write('A1', 'prefix x005F_text'); tab.write('A2', 'literal _x0041_')
    tab.write_rich_string('A3', 'rich ', writer.add_format({'bold': True}), 'needle'); writer.close()
    preserved = spreadsheet.search_cells(mixed, 'x005F_text')
    assert preserved['matches'][0]['cell'] == 'A1' and preserved['matches'][0]['text_complete']
    assert spreadsheet.search_cells(mixed, 'needle')['matches'][0]['cell'] == 'A3'
    escaped = spreadsheet.search_cells(mixed, '_x0041_')
    assert escaped['matches'][0]['cell'] == 'A2' and not escaped['matches'][0]['text_complete'] and escaped['unverified_text_cells'] == 1
    cells = spreadsheet.read_range(mixed, 'Sheet1', 'A1:A3')['rows']
    assert cells[0][0]['value'] == 'prefix x005F_text' and cells[0][0]['text_complete']
    assert not cells[1][0]['text_complete'] and cells[2][0]['value'] == 'rich needle'
    mixed_draft = root / 'mixed-draft'; session.start(mixed_draft, source=mixed)
    grid = session.read_grid(mixed_draft, 0, 'Sheet1', 'A1:A3')['rows']
    assert grid[0][0]['value'] == 'prefix x005F_text' and grid[0][0]['editable']
    assert not grid[1][0]['editable'] and grid[1][0]['read_only_reason'] == 'unverified_text'
    edited = session.replace_text(mixed_draft, 0, '/Sheet1/A1', 'x005F_text', 'fixed')
    assert spreadsheet.read_range(edited['file'], 'Sheet1', 'A1')['rows'][0][0]['value'] == 'prefix fixed'
    try: session.edit_grid(mixed_draft, 1, 'Sheet1', 'A2', [{'cell': 'A2', 'value': 'overwrite'}])
    except ValueError: pass
    else: raise AssertionError('unverified text was editable')
    assert not review.compare_cells(mixed, mixed, 'Sheet1', 'A1:A3')['text_comparison_complete']
    refused = root / 'unverified-author.xlsx'
    try: open_office_use.author({'sheets': {'Sheet1': {'B1': 'new'}}}, refused, mixed)
    except ValueError as exc: assert 'unverified shared strings' in str(exc)
    else: raise AssertionError('author accepted unverified shared strings')
    assert not refused.exists()
    safe = root / 'safe.xlsx'; writer = xlsxwriter.Workbook(str(safe)); writer.add_worksheet().write('A1', 'prefix x005F_text'); writer.close()
    authored = root / 'authored.xlsx'; open_office_use.author({'sheets': {'Sheet1': {'B1': 'new'}}}, authored, safe)
    assert spreadsheet.read_range(authored, 'Sheet1', 'A1')['rows'][0][0]['value'] == 'prefix x005F_text'
    phonetic = root / 'phonetic.xlsx'
    with zipfile.ZipFile(shared) as original, zipfile.ZipFile(phonetic, 'w') as output:
        for entry in original.infolist():
            data = original.read(entry)
            if entry.filename == 'xl/sharedStrings.xml': data = data.replace(b'</si>', b'<rPh sb="0" eb="1"><t>phonetic</t></rPh></si>', 1)
            output.writestr(entry, data)
    uncertain = spreadsheet.search_cells(phonetic, '共享')
    assert uncertain['matches'][0]['cell'] == 'A1' and not uncertain['complete'] and uncertain['unverified_text_cells'] == 1
    large = root / 'large.xlsx'; big = Workbook(write_only=True); page = big.create_sheet()
    for _ in range(100000): page.append(['other'])
    page.append(['needle']); big.save(large)
    limited = spreadsheet.search_cells(large, 'needle')
    assert limited['stored_cells_scanned'] == 100000 and not limited['scan_complete'] and not limited['matches']
    continued = spreadsheet.search_cells(large, 'needle', offset=limited['next_offset'], snapshot_sha256=limited['snapshot_sha256'])
    assert continued['scan_complete'] and continued['matches'][0]['cell'] == 'A100001'
    for query in ('', 'a' * 201, '\x00'):
        try: spreadsheet.search_cells(source, query)
        except ValueError: pass
        else: raise AssertionError('invalid search query accepted')
print('Cell search smoke passed: sparse, formulas, hidden sheets, bounds and saved revisions')

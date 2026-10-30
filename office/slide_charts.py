"""Inspect bounded cached chart data in saved PowerPoint presentations."""
import hashlib
import copy
import io
import itertools
import math
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE

import snapshot as snapshots
import chart_cache
import ooxml

C = '{http://schemas.openxmlformats.org/drawingml/2006/chart}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


SYNTHETIC_PART = 'xl/charts/office-use-inspection.xml'


def prepare_workbook(blob, chart_blob, candidate):
    if len(blob) > 16 * 1024 * 1024: raise ValueError('Embedded workbook exceeds size limit')
    with zipfile.ZipFile(io.BytesIO(blob)) as source:
        entries = source.infolist(); names = [entry.filename for entry in entries]
        if len(entries) > 5000 or len(names) != len(set(names)) or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024:
            raise ValueError('Embedded workbook exceeds package limits or contains duplicates')
        manifest = ooxml.parse_xml(source.read('[Content_Types].xml')).getroot()
        if manifest.tag != chart_cache.CT + 'Types': raise ValueError('Invalid content types')
        if SYNTHETIC_PART in names or any(item.get('PartName') == '/' + SYNTHETIC_PART for item in manifest):
            raise ValueError('Synthetic chart part already exists')
        entry = etree.Element(chart_cache.CT + 'Override', PartName='/' + SYNTHETIC_PART, ContentType=chart_cache.CHART_TYPE)
        overrides = [index for index, item in enumerate(manifest) if item.tag == chart_cache.CT + 'Override']
        manifest.insert(overrides[0] if overrides else len(manifest), entry)
        chart_root = ooxml.parse_xml(chart_blob).getroot()
        for node in chart_root.findall(C + 'externalData'): chart_root.remove(node)
        chart_xml = etree.tostring(chart_root, encoding='utf-8', xml_declaration=True)
        if len(chart_xml) > 4 * 1024 * 1024: raise ValueError('Chart exceeds size limit')
        with zipfile.ZipFile(candidate, 'w') as destination:
            for item in entries:
                data = etree.tostring(manifest, encoding='utf-8', xml_declaration=True) if item.filename == '[Content_Types].xml' else source.read(item)
                destination.writestr(item, data)
            destination.writestr(SYNTHETIC_PART, chart_xml)


def reconcile(chart, budget):
    """Compare chart caches against a bounded embedded workbook copy."""
    if budget['charts'] >= 20: return {'status': 'incomplete', 'reason': 'chart_limit', 'complete': False}
    budget['charts'] += 1
    try:
        part = chart.part.chart_workbook.xlsx_part
        if part is None: return {'status': 'unavailable', 'reason': 'missing_workbook', 'complete': False}
        blob = part.blob
        if len(blob) > 16 * 1024 * 1024 or budget['bytes'] + len(blob) > 64 * 1024 * 1024:
            return {'status': 'incomplete', 'reason': 'workbook_size_limit', 'complete': False}
        budget['bytes'] += len(blob)
        with tempfile.TemporaryDirectory(prefix='office-slide-chart-reconcile-') as tmp:
            candidate = Path(tmp) / 'embedded.xlsx'
            prepare_workbook(blob, chart.part.blob, candidate)
            report = chart_cache.scan(candidate)
        synthetic_part = SYNTHETIC_PART
        found = [item for item in report['charts'] if item['part'] == synthetic_part]
        if len(found) != 1 or found[0]['status'] != 'scanned': return {'status': 'incomplete', 'reason': 'unsupported_chart', 'complete': False}
        chart_report = found[0]; refs = chart_report['references']
        incomplete = chart_report['extensions_present'] or chart_report['external_data'] or any(ref['status'] not in ('match', 'mismatch') or ref.get('unverified_points', 0) for ref in refs)
        mismatch_count = sum(ref.get('mismatch_count', 0) for ref in refs)
        visible = []
        for ref in refs[:30]:
            item = dict(ref); mismatches = item.get('mismatches', [])
            item['mismatches'] = mismatches[:3]
            item['mismatches_truncated'] = item.get('mismatches_truncated', False) or len(mismatches) > 3
            visible.append(item)
        return {'status': 'mismatch' if mismatch_count else 'incomplete' if incomplete or len(refs) > 30 else 'match',
                'complete': not incomplete and len(refs) <= 30, 'mismatch_count': mismatch_count,
                'references': visible, 'references_truncated': len(refs) > 30,
                'formula_freshness': 'unknown'}
    except (ValueError, TypeError, KeyError, OSError, EOFError, IndexError, AttributeError, OverflowError, RuntimeError, NotImplementedError, zipfile.BadZipFile, etree.XMLSyntaxError) as exc:
        return {'status': 'incomplete', 'reason': 'embedded_workbook_unreadable', 'error_type': type(exc).__name__, 'complete': False}


def inspect(file):
    source = Path(file).resolve(strict=True)
    if source.suffix.lower() != '.pptx': raise ValueError('Slide chart inspection requires PPTX')
    with tempfile.TemporaryDirectory(prefix='office-slide-charts-') as tmp:
        snapshot, digest = snapshots.copy(source, tmp)
        with zipfile.ZipFile(snapshot) as archive:
            entries = archive.infolist()
            if len(entries) > 5000 or len({entry.filename for entry in entries}) != len(entries) or sum(entry.file_size for entry in entries) > 128 * 1024 * 1024:
                raise ValueError('Presentation exceeds package limits or contains duplicate parts')
        presentation = Presentation(snapshot); slides = presentation.slides
        result = {'snapshot_sha256': digest, 'slides_total': len(slides), 'slides_scanned': 0, 'shapes_scanned': 0,
                  'charts': [], 'issues': [], 'issue_count': 0, 'complete': True,
                  'scope': 'Chart values are saved caches; reconciliation compares them with saved embedded workbook values. Formula cache freshness remains unknown.'}
        reconcile_budget = {'charts': 0, 'bytes': 0}
        points = 0
        def issue(slide, path, reason):
            result['complete'] = False; result['issue_count'] += 1
            if len(result['issues']) < 200: result['issues'].append({'slide': slide, 'path': path, 'reason': reason})
        def reference(series, role):
            parent = series._element.find(C + role)
            if parent is None: return {'status': 'absent', 'formula': None}
            nodes = [child for child in parent if child.tag in (C + 'numRef', C + 'strRef')]
            if len(nodes) != 1: return {'status': 'unsupported', 'formula': None}
            node = nodes[0]; formulas = node.findall(C + 'f'); caches = node.findall(C + ('numCache' if node.tag == C + 'numRef' else 'strCache'))
            formula = formulas[0].text if len(formulas) == 1 else None
            status = 'present' if len(formulas) == 1 and bool(formula) and len(caches) == 1 else 'missing_or_ambiguous'
            count = None
            if status == 'present':
                counts = caches[0].findall(C + 'ptCount'); raw = counts[0].get('val') if len(counts) == 1 else None
                if raw is None or not raw.isdecimal() or len(raw) > 8: status = 'invalid_cache'
                else:
                    count = int(raw); seen = set()
                    for point in caches[0].findall(C + 'pt'):
                        raw_index = point.get('idx'); values = point.findall(C + 'v')
                        if raw_index is None or not raw_index.isdecimal() or len(raw_index) > 8 or len(values) != 1 or len(values[0]): status = 'invalid_cache'; break
                        index = int(raw_index)
                        if index >= count or index in seen: status = 'invalid_cache'; break
                        seen.add(index)
            if formula is not None and len(formula) > 500: status = 'truncated'
            return {'status': status, 'formula': formula[:500] if formula else None, 'count': count}
        def workbook(chart):
            nodes = chart.part._element.findall(C + 'externalData')
            if not nodes: return {'status': 'absent'}
            if len(nodes) != 1 or not nodes[0].get(R + 'id'): return {'status': 'ambiguous'}
            try: relation = chart.part.rels[nodes[0].get(R + 'id')]
            except KeyError: return {'status': 'invalid_relationship'}
            if relation.is_external: return {'status': 'external'}
            if relation.reltype != RELATIONSHIP_TYPE.PACKAGE: return {'status': 'unsupported_relationship'}
            part = relation.target_part; blob = part.blob
            metadata = {'part': str(part.partname), 'size': len(blob), 'sha256': hashlib.sha256(blob).hexdigest(), 'content_type': part.content_type}
            if part.content_type != 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': return {'status': 'unsupported_embedded_part', **metadata}
            return {'status': 'embedded', **metadata}
        def walk(shapes, slide_no, prefix, depth, hidden):
            nonlocal points
            for shape in shapes:
                if result['shapes_scanned'] >= 10000: issue(slide_no, prefix, 'shape_limit'); return False
                result['shapes_scanned'] += 1
                path = f'{prefix}/shape[@id={shape.shape_id}]'
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    path = f'{prefix}/group[@id={shape.shape_id}]'
                    if depth >= 8: issue(slide_no, path, 'group_depth'); continue
                    if not walk(shape.shapes, slide_no, path, depth + 1, hidden): return False
                    continue
                if not shape.has_chart: continue
                path = f'{prefix}/chart[@id={shape.shape_id}]'
                if len(result['charts']) >= 100: issue(slide_no, path, 'chart_limit'); return False
                try:
                    chart = shape.chart; series = chart.series; plots = chart.plots
                    metadata = next(shape._element.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr'), None)
                    hidden_shape = metadata.get('hidden') if metadata is not None else None
                    if hidden_shape not in (None, '0', 'false', '1', 'true'): issue(slide_no, path, 'unknown_shape_visibility')
                    record = {'slide': slide_no, 'hidden_slide': hidden, 'hidden_shape': hidden_shape in ('1', 'true'), 'shape_id': shape.shape_id, 'name': str(shape.name or '')[:200],
                              'path': path, 'chart_part': str(chart.part.partname), 'chart_type': str(chart.chart_type),
                              'plot_count': len(plots), 'series_total': len(series), 'series': [], 'workbook': workbook(chart),
                              'description': (metadata.get('descr', '') if metadata is not None else '')[:1000],
                              'description_truncated': metadata is not None and len(metadata.get('descr', '')) > 1000}
                    title = chart.part._element.find('.//' + C + 'chart/' + C + 'title')
                    if title is not None:
                        rich = title.find(C + 'tx/' + C + 'rich')
                        if rich is None: issue(slide_no, path, 'unsupported_chart_title'); record['title'] = None
                        else:
                            title_text = ''.join(node.text or '' for node in rich.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'))
                            record['title'] = title_text[:500]
                            if len(title_text) > 500: issue(slide_no, path, 'chart_title_truncated')
                    else: record['title'] = None
                    if record['workbook']['status'] not in ('embedded', 'absent'): issue(slide_no, path, 'workbook_' + record['workbook']['status'])
                    record['reconciliation'] = reconcile(chart, reconcile_budget) if record['workbook']['status'] == 'embedded' else {'status': 'unavailable', 'reason': record['workbook']['status'], 'complete': False}
                    if len(plots) != 1: issue(slide_no, path, 'multiple_plots')
                    if len(series) > 30: issue(slide_no, path, 'series_limit')
                    category_safe = True
                    for index, item in enumerate(itertools.islice(series, 30), 1):
                        refs = {role: reference(item, role) for role in ('tx', 'cat', 'val')}
                        cache_points = sum(1 for node in item._element.iter(C + 'pt'))
                        if cache_points > 1000 or points + cache_points > 5000:
                            issue(slide_no, path, 'point_limit'); category_safe = False; break
                        points += cache_points
                        value_count = refs['val'].get('count')
                        category_count = refs['cat'].get('count')
                        if refs['val']['status'] != 'present' or refs['cat']['status'] not in ('present', 'absent') or refs['tx']['status'] not in ('present', 'absent'):
                            issue(slide_no, path, 'unsupported_or_missing_cache'); category_safe = False
                        if category_count is not None and category_count > 200: issue(slide_no, path, 'category_limit'); category_safe = False
                        if value_count is None or value_count > 200:
                            issue(slide_no, path, 'series_value_limit' if value_count is not None else 'unavailable_values'); values = ()
                        else: values = item.values
                        name = item.name
                        if name is not None and len(name) > 500: issue(slide_no, path, 'series_name_truncated')
                        clipped = []
                        for value in values[:200]:
                            if value is None or isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value): clipped.append(value)
                            else: clipped.append(None); issue(slide_no, path, 'unverified_cached_value')
                        record['series'].append({'index': index, 'name': name[:500] if name else name, 'values': clipped,
                                                 'values_total': value_count, 'values_truncated': value_count is not None and value_count > 200, 'references': refs})
                    if len(plots) == 1 and category_safe:
                        category_view = plots[0].categories; count = len(category_view)
                        if count > 200: issue(slide_no, path, 'category_limit')
                        labels = []
                        for category in itertools.islice(category_view, 200):
                            value = category.label
                            label = str(value) if value is not None else None
                            if label is not None and len(label) > 500: issue(slide_no, path, 'category_truncated'); label = label[:500]
                            labels.append(label)
                        record['categories'] = labels; record['categories_total'] = count
                    result['charts'].append(record)
                except (ValueError, TypeError, AttributeError, KeyError, IndexError, OverflowError) as exc:
                    issue(slide_no, path, 'unsupported_chart'); result['charts'].append({'slide': slide_no, 'path': path, 'error': type(exc).__name__})
            return True
        # shortcut: bounded chart inventory and reconciliation; add pagination for larger presentations.
        for slide_no in range(1, min(len(slides), 200) + 1):
            slide = slides[slide_no - 1]; visibility = slide._element.get('show')
            if visibility not in (None, '0', 'false', '1', 'true'): issue(slide_no, None, 'unknown_visibility')
            result['slides_scanned'] = slide_no
            if not walk(slide.shapes, slide_no, f'/slide[{slide_no}]', 0, visibility in ('0', 'false')): break
        if len(slides) > 200: issue(None, None, 'slide_limit')
        result['issues_truncated'] = result['issue_count'] > len(result['issues'])
        return result


def sync(source, output, part, reference, allow_formula_caches=False, expected_sha256=None):
    """Synchronize one verified chart reference from its saved embedded workbook."""
    if not isinstance(part, str) or not part.startswith('/'): raise ValueError('Provide a chart part from slide chart inspection')
    if isinstance(reference, bool) or not isinstance(reference, int) or reference < 1: raise ValueError('Provide a positive reference index')
    if not isinstance(allow_formula_caches, bool): raise ValueError('allow_formula_caches must be a boolean')
    source = Path(source).resolve(strict=True); output = Path(output).absolute()
    if output.exists() or output.is_symlink(): raise FileExistsError(f'Output already exists: {output}')
    with tempfile.TemporaryDirectory(dir=output.parent, prefix='.office-slide-chart-sync-') as tmp:
        snapshot, digest = snapshots.copy(source, tmp)
        if expected_sha256 is not None and digest != expected_sha256: raise ValueError('Presentation snapshot changed; rescan chart caches')
        inventory = inspect(snapshot)
        charts = [chart for chart in inventory['charts'] if chart.get('chart_part') == part]
        if not charts or any(chart['workbook']['status'] != 'embedded' for chart in charts): raise ValueError('Chart with embedded workbook was not available in the bounded inventory')
        if any(chart['workbook']['part'] != charts[0]['workbook']['part'] for chart in charts): raise ValueError('Ambiguous embedded workbook')
        workbook_part = charts[0]['workbook']['part'].lstrip('/')
        chart_part = part.lstrip('/')
        with zipfile.ZipFile(snapshot) as archive:
            workbook_blob = archive.read(workbook_part); chart_blob = archive.read(chart_part)
        candidate = Path(tmp) / 'embedded.xlsx'; updated = Path(tmp) / 'updated.xlsx'
        prepare_workbook(workbook_blob, chart_blob, candidate)
        synced = chart_cache.sync(candidate, updated, SYNTHETIC_PART, reference, allow_formula_caches)
        with zipfile.ZipFile(updated) as archive: new_root = ooxml.parse_xml(archive.read(SYNTHETIC_PART)).getroot()
        def edit(tree, _archive):
            refs = [node for node in tree.getroot().iter() if node.tag in chart_cache.REF_TYPES]
            new_refs = [node for node in new_root.iter() if node.tag in chart_cache.REF_TYPES]
            if reference > len(refs) or reference > len(new_refs): raise ValueError('Chart reference changed')
            old_ref, new_ref = refs[reference - 1], new_refs[reference - 1]
            cache_tag = C + ('numCache' if old_ref.tag == C + 'numRef' else 'strCache')
            old_caches, new_caches = old_ref.findall(cache_tag), new_ref.findall(cache_tag)
            if old_ref.tag != new_ref.tag or len(old_caches) != 1 or len(new_caches) != 1: raise ValueError('Chart cache changed')
            old_ref.replace(old_caches[0], copy.deepcopy(new_caches[0]))
        ooxml.edit_xml(snapshot, output, lambda _: chart_part, edit)
    return {'output': str(output), 'chart_part': part, 'reference': reference, 'updated_points': synced['updated_points'],
            'source_formula_cells': synced['source_formula_cells'], 'formula_cache_freshness': 'unknown', 'snapshot_sha256': digest}
